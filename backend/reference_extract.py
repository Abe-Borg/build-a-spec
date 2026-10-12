"""Turning an attached reference file into the text the model reads.

This is the one place a reference upload's bytes become text. ``.docx`` was
the first supported type and its extraction stays with the importer (it shares
the Accept-All tracked-changes resolution and table flattening a master import
needed); every other supported type is read here.

Supported kinds, and why each is read the way it is
--------------------------------------------------
``docx``  Word body text in document order with tracked changes resolved to
          the Accept-All view — delegated to
          :func:`backend.spec_doc.importer.extract_reference_text`, behind the
          same bounded ZIP/OPC inspection a master import gets, because an
          attachment is the same attack surface. Nothing is retained but the
          text.
``pdf``   Per-page text through ``pypdf`` (already a dependency — the research
          resend sanitizer counts pages with it). Pages carry a ``[page N]``
          marker so the model can tell the user where something came from. A
          PDF with no text layer at all is refused *with the reason* rather
          than attached as an empty document: silently accepting a scan the
          model cannot read is the one outcome worse than a clear error.
``txt``   Decoded and normalized, nothing else.
``csv`` / ``xml``
          Also decoded verbatim: for these two the structure *is* the content
          — a CSV's rows and an XML's tags are exactly what makes the file
          worth reading — so nothing is parsed away or flattened.
``md``    Markdown (``.md`` or ``.markdown``), decoded verbatim like the three
          above and never rendered. Its headings, lists, tables and emphasis
          markers are the document's outline, and the model reads Markdown as
          written, so rendering it to plain text would only throw the outline
          away. Its lines are kept exactly, which the other text types' are
          not: their trailing whitespace is dropped and their blank-line runs
          collapse to one, but in Markdown two trailing spaces are a hard
          line break, and trailing spaces and blank lines inside a code block
          are the code. Only line endings (to ``\n``), control characters and
          blank lines at the very start and end are normalized.
``xlsx``  An Excel workbook (``.xlsx`` or the macro-enabled ``.xlsm``; the
          macros are never run, nor even read). Read through ``openpyxl`` in
          read-only mode, behind a bounded ZIP check of its own and with
          ``defusedxml`` installed, which makes openpyxl refuse any entity
          declaration in the XML (an attachment is untrusted, and Word's
          reader never expands entities either). One sheet at a time in
          workbook order: each sheet opens with a ``[sheet N of M:
          Name]`` marker (``[hidden sheet …]`` for one the author hid) and its
          rows follow as CSV — the same reading the ``csv`` kind gets, since
          an equipment schedule's rows and columns are its content. A cell is
          read as Excel last showed it: a formula's saved result, a date as
          ``YYYY-MM-DD``, a percentage with its ``%``, TRUE/FALSE. A formula
          with no saved result (a workbook written by a script, never opened
          in Excel) is read as its formula and reported, never left blank.
          Columns empty in every row of a sheet (spacer columns) and runs of
          empty rows are dropped, because a reference's tokens are budgeted.
          Legacy ``.xls``, binary ``.xlsb`` and OpenDocument ``.ods`` are
          refused with the way out (save as ``.xlsx`` or export ``.csv``), as
          is a password-protected workbook, which Excel encrypts whole.

Everything here is pure CPU over the uploaded bytes and must run on a worker
thread; see the event-loop rule in ``backend/app.py``.

Token posture is unchanged and set by ``reference_docs``: extracted text is
bounded (``MAX_TEXT_CHARS``), never rides the per-turn PROJECT CONTEXT, and is
elided from committed history after the turn that read it. This module only
decides what the text *is*.
"""
from __future__ import annotations

import codecs
import csv
import datetime
import io
import math
import re
import tempfile
import zipfile
from itertools import zip_longest
from pathlib import Path
from typing import Any

from .spec_doc.importer import ReferenceExtraction, extract_reference_text
from .spec_doc.source_package import (
    MAX_ZIP_MEMBER_BYTES,
    MAX_ZIP_MEMBERS,
    MAX_ZIP_UNCOMPRESSED_BYTES,
    inspect_docx_package,
    sanitize_source_filename,
)

# Extension → kind. Insertion order is the order the supported types are
# listed to the user, so the rejection message can never drift from the set
# actually accepted.
REFERENCE_KINDS: dict[str, str] = {
    ".docx": "docx",
    ".pdf": "pdf",
    ".txt": "txt",
    ".xml": "xml",
    ".csv": "csv",
    # Two spellings, one kind. The first listed is the canonical extension
    # (``reference_extension_for_kind``); a file keeps whichever it came with
    # (``sanitize_reference_filename``).
    ".md": "md",
    ".markdown": "md",
    # A macro-enabled workbook is the same package with a VBA part beside
    # the sheets; the reader never loads that part.
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
}

# Display names for the panel list and the model's context stub. The stub says
# what kind of file a reference is because it changes how the model should read
# it — a CSV is data, a PDF carries page markers, a Word standard is prose.
REFERENCE_KIND_LABELS: dict[str, str] = {
    "docx": "Word",
    "pdf": "PDF",
    "txt": "text",
    "xml": "XML",
    "csv": "CSV",
    "md": "Markdown",
    "xlsx": "Excel",
}

# Spreadsheet formats a user plausibly has but no extractor reads. Picking one
# gets the way out instead of the bare list of what is accepted: Excel, Numbers
# and LibreOffice all save or export an ``.xlsx`` or a ``.csv``.
_UNREADABLE_SPREADSHEETS: dict[str, str] = {
    ".xls": "an Excel 97-2003 workbook (.xls)",
    ".xlsb": "an Excel binary workbook (.xlsb)",
    ".ods": "an OpenDocument spreadsheet (.ods)",
    ".numbers": "a Numbers spreadsheet (.numbers)",
}

# Runaway breaker for a workbook (the MAX_PDF_PAGES posture): a 25 MiB upload
# can hold millions of cells, and reading cells is the slow part. Sized past
# what a reference stores (``reference_docs.MAX_TEXT_CHARS``, 400,000) so the
# stored truncation, which tells the model the tail was not read, is what
# normally cuts; hitting this is reported, never silent.
MAX_WORKBOOK_TEXT_CHARS = 1_000_000

# Every OLE2 compound file starts with these bytes. An ``.xlsx`` that does is
# either password-protected (Excel encrypts the whole package into one) or a
# legacy ``.xls`` renamed — neither is a ZIP, and both deserve the reason.
_OLE_COMPOUND_FILE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_ZIP_ENCRYPTED_FLAG = 0x1

_NOT_A_WORKBOOK = (
    "That file could not be read as an Excel workbook. Save it from Excel as "
    "an Excel Workbook (.xlsx), or export the sheet as .csv, and attach that."
)

# Runaway breaker, not a quality limit (the MAX_TOOL_ROUNDS posture): the
# upload is already bounded at 25 MiB, but a text-only PDF that size can run
# to thousands of pages and text extraction is the slow part. Sized well past
# any real design standard; hitting it is reported, never silent.
MAX_PDF_PAGES = 1_000

# Runs of blank lines longer than this collapse. Blank lines are structure in a
# text file (paragraph breaks, record separators), so they are kept — but a
# file with fifty blank lines between sections would spend real tokens on
# nothing.
MAX_CONSECUTIVE_BLANK_LINES = 1

# Kinds whose lines are kept exactly as written: no trailing whitespace is
# stripped and no run of blank lines is collapsed. Markdown only — there both
# can be meaning (a hard line break, the inside of a code block).
_VERBATIM_LINE_KINDS = frozenset({"md"})

# C0 controls have no meaning in extracted reference text; tabs and newlines
# do. DEL goes too. Anything left would only confuse the model or the panel.
_CONTROL_TRANSLATION = {
    code: None
    for code in list(range(0x20)) + [0x7F]
    if code not in (0x09, 0x0A)
}


class ReferenceExtractError(ValueError):
    """The attachment could not be read as text.

    A ``ValueError``, like ``MasterImportError`` and ``SourcePackageError``, so
    the upload endpoint's existing 400 mapping covers every supported type.
    """


def supported_extensions_phrase() -> str:
    """``.docx, .pdf, …, .xlsx, or .xlsm`` — for messages."""
    extensions = list(REFERENCE_KINDS)
    return ", ".join(extensions[:-1]) + f", or {extensions[-1]}"


def unsupported_reference_message(filename: str) -> str:
    """Why ``filename`` cannot be attached, and what to attach instead.

    A spreadsheet format no extractor reads names the save-as that fixes it;
    anything else gets the list of what is accepted.
    """
    lowered = (filename or "").lower()
    for extension, description in _UNREADABLE_SPREADSHEETS.items():
        if lowered.endswith(extension):
            return (
                f"That is {description}, which cannot be read as a "
                "reference. Save it as an Excel Workbook (.xlsx), or export "
                "the sheet as .csv, and attach that."
            )
    return (
        f"That file type is not supported as a reference. Attach a "
        f"{supported_extensions_phrase()} file."
    )


def reference_kind_for_filename(filename: str) -> str | None:
    """The kind an upload will be read as, or ``None`` when unsupported.

    Extension-based on purpose: the browser's reported content type is not
    trustworthy, and every extractor validates the bytes it actually gets
    (a ``.pdf`` that is not a PDF is refused by the PDF reader, a ``.docx``
    that is not a package by the OPC inspection).
    """
    lowered = (filename or "").lower()
    for extension, kind in REFERENCE_KINDS.items():
        if lowered.endswith(extension):
            return kind
    return None


def reference_extension_for_kind(kind: str) -> str:
    """``"pdf"`` → ``".pdf"``. Empty for an unknown kind."""
    for extension, known in REFERENCE_KINDS.items():
        if known == kind:
            return extension
    return ""


def reference_extension_for_filename(filename: str, *, kind: str) -> str:
    """The extension ``filename`` already carries for ``kind``, else the
    kind's canonical one.

    A kind can have more than one spelling (``.md`` and ``.markdown``), and a
    file keeps the one it arrived with: ``notes.markdown`` must not become
    ``notes.markdown.md``.
    """
    lowered = (filename or "").lower()
    for extension, known in REFERENCE_KINDS.items():
        if known == kind and lowered.endswith(extension):
            return extension
    return reference_extension_for_kind(kind)


def sanitize_reference_filename(filename: str, *, kind: str) -> str:
    """A path-free display name that keeps the attachment's own type.

    The shared sanitizer defaults to the imported-master case and appends
    ``.docx`` to anything else, which would rename ``acme.pdf`` and then send
    it to the wrong extractor.
    """
    extension = reference_extension_for_filename(filename, kind=kind) or ".txt"
    return sanitize_source_filename(
        filename, extension=extension, fallback=f"attachment{extension}"
    )


def extract_reference_document(
    source_bytes: bytes, *, filename: str
) -> ReferenceExtraction:
    """Read an attached file's text. Raises :class:`ValueError` subclasses.

    Blocking: seconds of CPU on a long PDF, Word document or workbook.
    Callers must be on a worker thread.
    """
    kind = reference_kind_for_filename(filename)
    if kind is None:
        raise ReferenceExtractError(unsupported_reference_message(filename))
    if kind == "docx":
        return _extract_docx(source_bytes)
    if kind == "pdf":
        return _extract_pdf(source_bytes)
    if kind == "xlsx":
        return _extract_workbook(source_bytes)
    return _extract_plain_text(
        source_bytes,
        kind=kind,
        extension=reference_extension_for_filename(filename, kind=kind),
    )


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------


def _extract_docx(source_bytes: bytes) -> ReferenceExtraction:
    """The same package inspection a master import gets, then body text.

    Nothing is retained: a reference document has no source package, no
    preservation map, and no export path, so the bytes are dropped as soon as
    the text is out.
    """
    inspect_docx_package(source_bytes)
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as handle:
        handle.write(source_bytes)
        temp_path = Path(handle.name)
    try:
        return extract_reference_text(temp_path)
    finally:
        temp_path.unlink(missing_ok=True)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _pdf_reader(source_bytes: bytes):
    """Open a PDF, unlocking an owner-password-only file if possible."""
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(source_bytes), strict=False)
    except Exception as exc:  # noqa: BLE001 — pypdf raises broadly on junk
        raise ReferenceExtractError(
            "That file could not be read as a PDF."
        ) from exc
    try:
        encrypted = bool(reader.is_encrypted)
    except Exception:  # noqa: BLE001
        encrypted = False
    if encrypted:
        # A PDF restricted only by an owner password (print/copy permissions)
        # opens with an empty user password, which is the common case for a
        # circulated standard. A real user password cannot be guessed.
        try:
            unlocked = bool(reader.decrypt(""))
        except Exception:  # noqa: BLE001
            unlocked = False
        if not unlocked:
            raise ReferenceExtractError(
                "That PDF is password-protected, so its text cannot be read. "
                "Attach an unprotected copy."
            )
    return reader


def _extract_pdf(source_bytes: bytes) -> ReferenceExtraction:
    reader = _pdf_reader(source_bytes)
    try:
        total_pages = len(reader.pages)
    except Exception as exc:  # noqa: BLE001
        raise ReferenceExtractError(
            "That PDF's page structure could not be read."
        ) from exc

    read_pages = min(total_pages, MAX_PDF_PAGES)
    lines: list[str] = []
    block_count = 0
    empty_pages = 0
    failed_pages = 0
    for index in range(read_pages):
        try:
            page_text = reader.pages[index].extract_text() or ""
        except Exception:  # noqa: BLE001 — one bad page is not a bad file
            failed_pages += 1
            continue
        page_lines = [
            collapsed
            for collapsed in (" ".join(raw.split()) for raw in page_text.splitlines())
            if collapsed
        ]
        if not page_lines:
            empty_pages += 1
            continue
        # The page marker is why a PDF reference can be cited back to the user
        # as "page 12 of the owner's standard". It is not a text block.
        lines.append(f"[page {index + 1}]")
        lines.extend(page_lines)
        block_count += len(page_lines)

    if not block_count:
        raise ReferenceExtractError(
            "No text could be read from that PDF. A scanned PDF is images of "
            "pages with no text layer — run OCR on it, or attach the content "
            "in another format."
        )

    warnings: list[str] = []
    if total_pages > read_pages:
        skipped = total_pages - read_pages
        note = (
            f"[... {skipped:,} further page(s) of this PDF were not read "
            f"(limit {MAX_PDF_PAGES:,} pages).]"
        )
        lines.append(note)
        warnings.append(
            f"Only the first {read_pages:,} of {total_pages:,} pages were "
            f"read. The model is told the rest was not."
        )
    if failed_pages:
        warnings.append(
            f"{failed_pages} page(s) could not be decoded and were skipped."
        )
    if empty_pages:
        warnings.append(
            f"{empty_pages} of {read_pages:,} page(s) held no extractable "
            "text (scanned images or drawings)."
        )
    return ReferenceExtraction(
        text="\n".join(lines),
        block_count=block_count,
        tracked_changes=False,
        kind="pdf",
        warnings=tuple(warnings),
    )


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def _inspect_workbook_package(source_bytes: bytes) -> None:
    """The ZIP limits a master import enforces, before ``openpyxl`` sees it.

    Declared sizes are enough here: ``zipfile`` never decompresses a member
    past its declared size, so bounding the declarations bounds what the
    reader can inflate. Nothing is retained, so the raw-layout checks a
    preserved Word package needs do not apply.
    """
    if source_bytes.startswith(_OLE_COMPOUND_FILE_SIGNATURE):
        raise ReferenceExtractError(
            "That workbook is password-protected (Excel encrypts a workbook "
            "that needs a password to open), or it is an older .xls file "
            "renamed. Either way its cells cannot be read: attach an "
            "unprotected copy saved as an Excel Workbook (.xlsx), or export "
            "the sheet as .csv."
        )
    try:
        archive = zipfile.ZipFile(io.BytesIO(source_bytes), "r")
    except (zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise ReferenceExtractError(_NOT_A_WORKBOOK) from exc
    with archive:
        infos = archive.infolist()
        if len(infos) > MAX_ZIP_MEMBERS:
            raise ReferenceExtractError(
                f"That workbook has too many parts to read (more than "
                f"{MAX_ZIP_MEMBERS:,})."
            )
        declared_total = 0
        for info in infos:
            if info.flag_bits & _ZIP_ENCRYPTED_FLAG:
                raise ReferenceExtractError(
                    "That workbook's contents are encrypted, so its cells "
                    "cannot be read. Attach an unprotected copy."
                )
            declared_total += info.file_size
            if (
                info.file_size > MAX_ZIP_MEMBER_BYTES
                or declared_total > MAX_ZIP_UNCOMPRESSED_BYTES
            ):
                raise ReferenceExtractError(
                    "That workbook unpacks to more data than a reference "
                    f"can hold (limit {MAX_ZIP_UNCOMPRESSED_BYTES // (1024 * 1024)} "
                    "MiB). Attach the sheets you need as a smaller workbook "
                    "or as .csv."
                )


def _extract_workbook(source_bytes: bytes) -> ReferenceExtraction:
    """Read every sheet's cells, as Excel last showed them, into CSV rows.

    Two read-only passes walk the same rows in step: one reads the saved
    values, the other the formulas, which is the only way ``openpyxl`` tells
    a formula Excel never calculated (no saved value) from an empty cell.
    """
    _inspect_workbook_package(source_bytes)
    from openpyxl import DEFUSEDXML, load_workbook

    if not DEFUSEDXML:
        # defusedxml is what makes openpyxl refuse entity declarations in a
        # sheet's XML; the standard parser it falls back to expands them
        # (bounded, but up to a hundredfold). Fail closed rather than read an
        # untrusted file without the guard.
        raise ReferenceExtractError(
            "This copy of the app cannot read Excel workbooks safely (its XML "
            "guard is missing). Export the sheet as .csv and attach that."
        )
    books: list[Any] = []
    try:
        for data_only in (True, False):
            books.append(
                load_workbook(
                    io.BytesIO(source_bytes),
                    read_only=True,
                    data_only=data_only,
                    keep_links=False,
                )
            )
        return _read_workbook(*books)
    except ReferenceExtractError:
        raise
    except Exception as exc:  # noqa: BLE001 — openpyxl raises broadly on junk
        raise ReferenceExtractError(_NOT_A_WORKBOOK) from exc
    finally:
        for book in books:
            try:
                book.close()
            except Exception:  # noqa: BLE001 — closing is best effort
                pass


def _read_workbook(values_book: Any, formulas_book: Any) -> ReferenceExtraction:
    sheets = list(zip(values_book.worksheets, formulas_book.worksheets))
    chart_sheets = len(getattr(values_book, "chartsheets", None) or ())
    total = len(sheets)
    titles = [" ".join(str(sheet.title).split()) for sheet, _ in sheets]
    hidden = [
        getattr(sheet, "sheet_state", "visible") != "visible"
        for sheet, _ in sheets
    ]

    lines: list[str] = []
    if total > 1:
        # The contents up front: when the stored copy is cut short, the model
        # still knows which sheets the cut left out.
        lines.append(
            "[workbook sheets: "
            + "; ".join(
                f"{position}. {title}" + (" (hidden)" if is_hidden else "")
                for position, (title, is_hidden) in enumerate(
                    zip(titles, hidden), start=1
                )
            )
            + "]"
        )

    block_count = 0
    chars = 0
    unsaved_formulas = 0
    stopped: tuple[str, int, int] | None = None
    for position, (values_sheet, formulas_sheet) in enumerate(sheets, start=1):
        title = titles[position - 1]
        prefix = "hidden sheet" if hidden[position - 1] else "sheet"
        lines.append(f"[{prefix} {position} of {total}: {title}]")
        sheet = _read_sheet(
            values_sheet,
            formulas_sheet,
            budget=MAX_WORKBOOK_TEXT_CHARS - chars,
        )
        lines.extend(sheet.lines)
        block_count += sheet.row_count
        chars += sheet.chars
        unsaved_formulas += sheet.unsaved_formulas
        if sheet.stopped_at_row is not None:
            stopped = (title, sheet.stopped_at_row, total - position)
            break

    if not block_count:
        raise ReferenceExtractError(
            "No cell values could be read from that workbook. Attach a "
            "workbook whose sheets hold data, or export the sheet as .csv."
        )

    warnings: list[str] = []
    if stopped is not None:
        title, row, later = stopped
        rest = f" and {later:,} later sheet(s)" if later else ""
        lines.append(
            f"[... reading stopped at row {row:,} of sheet {title!r} (limit "
            f"{MAX_WORKBOOK_TEXT_CHARS:,} characters); the rest of that "
            f"sheet{rest} were not read.]"
        )
        warnings.append(
            f"Reading stopped at row {row:,} of sheet {title!r}; the rest of "
            f"that sheet{rest} were not read. The model is told the rest was "
            "not read."
        )
    if unsaved_formulas:
        warnings.append(
            f"{unsaved_formulas:,} formula cell(s) had no saved result — the "
            "workbook was last saved by a program that does not calculate "
            "formulas — so the formulas were read instead. Open and save it "
            "in Excel to attach the results."
        )
    hidden_count = sum(hidden)
    if hidden_count:
        warnings.append(
            f"{hidden_count} hidden sheet(s) were read too; the model is told "
            "which ones are hidden."
        )
    if chart_sheets:
        warnings.append(
            f"{chart_sheets} chart sheet(s) hold no cell values and were "
            "skipped."
        )
    return ReferenceExtraction(
        text="\n".join(lines),
        block_count=block_count,
        tracked_changes=False,
        kind="xlsx",
        warnings=tuple(warnings),
    )


class _SheetRead:
    """One sheet's CSV lines and what reading it found."""

    __slots__ = ("chars", "lines", "row_count", "stopped_at_row", "unsaved_formulas")

    def __init__(self) -> None:
        self.lines: list[str] = []
        self.row_count = 0
        self.chars = 0
        self.unsaved_formulas = 0
        self.stopped_at_row: int | None = None


def _read_sheet(values_sheet: Any, formulas_sheet: Any, *, budget: int) -> _SheetRead:
    """Read one sheet's rows within ``budget`` characters.

    ``reset_dimensions`` first: read-only mode otherwise trusts the
    ``<dimension>`` the writing program recorded, and some record ``A1`` for
    a full sheet, which would silently read one row.
    """
    values_sheet.reset_dimensions()
    formulas_sheet.reset_dimensions()
    read = _SheetRead()
    # Empty rows are kept only as one separator between rows with values: a
    # blank row often divides a schedule's groups, a run of them is nothing.
    rows: list[list[str]] = []
    blank_pending = False
    used_columns: set[int] = set()
    row_pairs = zip_longest(
        values_sheet.iter_rows(), formulas_sheet.iter_rows(), fillvalue=()
    )
    for row_number, (values_row, formulas_row) in enumerate(row_pairs, start=1):
        cells: list[str] = []
        unsaved = 0
        for value_cell, formula_cell in zip_longest(values_row, formulas_row):
            text, from_formula = _cell_text(value_cell, formula_cell)
            cells.append(text)
            unsaved += from_formula
        while cells and not cells[-1]:
            cells.pop()
        if not cells:
            blank_pending = bool(rows)
            continue
        row_chars = sum(len(text) + 1 for text in cells)
        if read.chars + row_chars > budget:
            read.stopped_at_row = row_number
            break
        if blank_pending:
            rows.append([])
            blank_pending = False
        rows.append(cells)
        used_columns.update(index for index, text in enumerate(cells) if text)
        read.chars += row_chars
        read.row_count += 1
        read.unsaved_formulas += unsaved

    # A column empty in every row is a spacer, not data; dropping it keeps a
    # formatted schedule from spending tokens on runs of empty fields.
    kept = sorted(used_columns)
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    for cells in rows:
        if not cells:
            read.lines.append("")
            continue
        fields = [cells[index] if index < len(cells) else "" for index in kept]
        while fields and not fields[-1]:
            fields.pop()
        buffer.seek(0)
        buffer.truncate()
        writer.writerow(fields)
        read.lines.append(buffer.getvalue().rstrip("\n"))
    return read


def _cell_text(value_cell: Any, formula_cell: Any) -> tuple[str, bool]:
    """``(text, read_from_formula)`` for one cell of the two passes.

    A formula cell with no saved value comes back empty from the values pass;
    the formulas pass is what says it is a formula. A formula whose saved
    result is an empty string (``=IF(A1="","",…)``) is stored as an empty
    string value (type ``str``), which openpyxl also reads as empty — that
    cell IS blank in Excel, so only a numeric-typed empty one is unsaved.
    """
    value = getattr(value_cell, "value", None)
    if value is not None:
        number_format = getattr(value_cell, "number_format", None) or "General"
        return _format_cell_value(value, number_format), False
    if (
        getattr(formula_cell, "data_type", None) == "f"
        and getattr(value_cell, "data_type", "n") == "n"
    ):
        formula = getattr(formula_cell, "value", None)
        if not isinstance(formula, str):
            # An array formula's text rides an object.
            formula = getattr(formula, "text", None) or ""
        formula = " ".join(str(formula).translate(_CONTROL_TRANSLATION).split())
        if formula:
            return (formula if formula.startswith("=") else f"={formula}"), True
    return "", False


# Quoted literals, escaped characters and bracketed sections ([Red], [$-409],
# [h]) never make a number a percentage, even when they contain a ``%``.
_NUMBER_FORMAT_LITERALS = re.compile(r'"[^"]*"|\\.|\[[^\]]*\]')


def _format_cell_value(value: Any, number_format: str) -> str:
    """A cell's value as Excel would show it, without its decoration."""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime.datetime):
        if not (value.hour or value.minute or value.second or value.microsecond):
            return value.date().isoformat()
        spec = "seconds" if value.second or value.microsecond else "minutes"
        return value.isoformat(sep=" ", timespec=spec)
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, datetime.time):
        spec = "seconds" if value.second or value.microsecond else "minutes"
        return value.isoformat(timespec=spec)
    if isinstance(value, datetime.timedelta):
        seconds = round(value.total_seconds())
        sign = "-" if seconds < 0 else ""
        hours, rest = divmod(abs(seconds), 3600)
        minutes, seconds = divmod(rest, 60)
        return f"{sign}{hours}:{minutes:02d}:{seconds:02d}"
    if isinstance(value, (int, float)):
        if "%" in _NUMBER_FORMAT_LITERALS.sub("", number_format):
            return _format_number(value * 100) + "%"
        return _format_number(value)
    return " ".join(str(value).translate(_CONTROL_TRANSLATION).split())


def _format_number(value: int | float) -> str:
    """Up to Excel's 15 significant digits: ``2500``, not ``2500.0``, and
    ``0.3``, not ``0.30000000000000004``."""
    if isinstance(value, int):
        return str(value)
    if not math.isfinite(value):
        return str(value)
    if value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    return format(value, ".15g")


# ---------------------------------------------------------------------------
# Plain text, CSV, XML, Markdown
# ---------------------------------------------------------------------------


def _decode(source_bytes: bytes) -> tuple[str, str]:
    """Decode uploaded text, returning ``(text, encoding label)``.

    A ladder rather than a single attempt: Windows is the primary platform,
    where a hand-saved ``.txt`` or a spreadsheet's ``.csv`` export is often
    UTF-16 (Excel) or Windows-1252 rather than UTF-8. Latin-1 is the last
    resort because it cannot fail — a mojibake character beats refusing a
    file the user can plainly read.
    """
    if source_bytes.startswith(
        (codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)
    ):
        try:
            return source_bytes.decode("utf-16"), "UTF-16"
        except UnicodeDecodeError:
            pass
    for encoding, label in (("utf-8-sig", "UTF-8"), ("cp1252", "Windows-1252")):
        try:
            return source_bytes.decode(encoding), label
        except UnicodeDecodeError:
            continue
    return source_bytes.decode("latin-1"), "Latin-1"


def _extract_plain_text(
    source_bytes: bytes, *, kind: str, extension: str = ""
) -> ReferenceExtraction:
    text, encoding = _decode(source_bytes)
    if "\x00" in text:
        # Reached by a binary file renamed to .txt/.csv/.xml/.md. The decode
        # ladder ends in Latin-1, which never raises, so this is the only
        # place binary content can be caught.
        raise ReferenceExtractError(
            f"That file does not contain readable text (it looks like "
            f"binary data with a {extension or '.' + kind} name)."
        )
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.translate(_CONTROL_TRANSLATION)

    # Markdown keeps its lines exactly: two trailing spaces are a hard line
    # break, and trailing spaces and blank-line runs inside a fenced or
    # indented code block are the code. Every other text type is tidied.
    verbatim = kind in _VERBATIM_LINE_KINDS
    lines: list[str] = []
    block_count = 0
    blank_run = 0
    for raw in text.split("\n"):
        line = raw if verbatim else raw.rstrip()
        if line.strip():
            blank_run = 0
            block_count += 1
            lines.append(line)
            continue
        blank_run += 1
        if lines and (verbatim or blank_run <= MAX_CONSECUTIVE_BLANK_LINES):
            lines.append(line if verbatim else "")
    while lines and not lines[-1].strip():
        lines.pop()

    warnings: list[str] = []
    if encoding != "UTF-8":
        warnings.append(
            f"The file was not UTF-8; it was read as {encoding}. Check the "
            "excerpt if it contains unusual characters."
        )
    return ReferenceExtraction(
        text="\n".join(lines),
        block_count=block_count,
        tracked_changes=False,
        kind=kind,
        warnings=tuple(warnings),
    )
