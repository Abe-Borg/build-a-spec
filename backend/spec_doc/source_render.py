"""Emit an imported master back as Word, with its formatting intact.

The contract (decided with Abraham, 2026-08-21):

    Every part of the uploaded package except the body of
    ``word/document.xml`` is carried through byte-for-byte — headers,
    footers, styles, theme, fonts, numbering definitions, page setup,
    section properties. Inside the body:

    * a provision you did not touch is emitted as a byte-identical clone of
      its source element;
    * a provision you edited keeps its paragraph properties and its own
      runs: the new words are spliced into the runs that were already there
      (``source_splice``, D-2 of the Redline-on-your-original plan), so a
      bold phrase you did not change stays bold and the tab after a typed
      letter stays a tab — relettering included;
    * a preserved block (table, picture, embedded object, content control)
      is emitted verbatim;
    * a provision you added is cloned from the nearest kin of its own kind
      (a provision from a provision at its depth, an article heading from an
      article heading), taking its formatting, its label convention and its
      separator — never its identity, bookmarks, comment anchors or section
      break;
    * body content the tree never modelled travels with the element below
      it — blank spacers, a cover page, a picture-only paragraph, a page
      break — except a SECTION BREAK, which belongs to the content above it
      (see "Section breaks" below).

What is deliberately NOT promised is the older mode's byte-exact whole-file
clone. That promise bought so little editing surface (three of twenty-seven
body operations on a clean master) that it was the feature's real defect.

Two limits remain, and are disclosed rather than worked around:

* **New words inherit their formatting from a neighbour.** A word typed
  over others takes the formatting of the first character it replaced, and
  an inserted word the formatting of the character before it — what Word
  itself does. Unchanged words keep their own. A paragraph the splice cannot
  map (a hyperlink, field, content control, comment or note reference,
  symbol or drawing inside it) is rebuilt from its first run's properties
  instead; the export's diagnostics count those fallbacks by reason.
* **A revision-bearing paragraph is rewritten, never cloned.** The importer
  showed the Accept-All view, so cloning the original markup would export
  text the user never saw. Those paragraphs take the rewrite path even when
  their text is unchanged.

Section breaks (Phase 0 of the Redline-on-your-original plan): a Word
section break is the END of a section — the content above it — so it never
moves or disappears because of an edit to what sits below it. An empty
paragraph holding a break stays after the content above it; a provision that
holds one in its own ``w:pPr`` keeps it while it stays in place, and when it
is deleted or moved away an empty paragraph holding the break is left where
it was; a clone never copies one. No edit loses a break. Where "in place"
is ambiguous after a reorder, the break goes where it best separates the
content that was above it from the content that was below it.
"""

from __future__ import annotations

import copy
import re
import zipfile
from dataclasses import dataclass, field
from io import BytesIO

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph as DocxParagraph
from lxml import etree

from .importer import (
    _accept_all_paragraph_text,
    _default_paragraph_style_id,
    _effective_numbering,
    _element_has_tracked_changes,
    _load_numbering_catalog,
    _load_style_numbering,
    _promoted_heading_kind,
)
from .model import (
    Article,
    Paragraph,
    SpecSection,
    _paragraph_label,
    labelled_paragraphs,
)
from .raw_zip import replace_document_xml_raw
from .source_format import (
    HEADER_SOURCE_CHROME,
    HEADER_SOURCE_FRONT_MATTER,
    LABEL_AUTO,
    LABEL_MANUAL,
    NO_ORIGIN,
    SECTION_TITLE_UID,
    SourceFormatMap,
)
from .revision_marks import (
    delete_content,
    insert_content,
    mark_paragraph,
    mark_table,
    neutralize_last_paragraph,
    record_paragraph_properties,
)
from .source_splice import (
    FALLBACK_REVISIONS,
    ParagraphMap,
    SpliceOp,
    append_text,
    map_paragraph,
    plan_splice,
    render_clean,
    render_redline,
)

_DOCUMENT_PART = "word/document.xml"

_W_P = qn("w:p")
_W_PPR = qn("w:pPr")
_W_R = qn("w:r")
_W_RPR = qn("w:rPr")
_W_SECTPR = qn("w:sectPr")
_W_NUMPR = qn("w:numPr")
_W_NUMID = qn("w:numId")
_W_ILVL = qn("w:ilvl")
_W_VAL = qn("w:val")
_W14_PARA_ID = "{http://schemas.microsoft.com/office/word/2010/wordml}paraId"
_W14_TEXT_ID = "{http://schemas.microsoft.com/office/word/2010/wordml}textId"
# A paragraph carrying any of these is layout, not a spacer: a picture, a
# text box (most cover pages), an embedded object, a content control. It is
# never dropped as a blank, and it travels with the element below it.
_LAYOUT_TAGS = (
    qn("w:drawing"),
    qn("w:pict"),
    qn("w:object"),
    qn("w:txbxContent"),
    qn("w:sdt"),
)
# ``w:pPr`` children in schema order (CT_PPr), for inserting ``w:numPr``
# where Word expects it.
_PPR_ORDER = tuple(
    qn(f"w:{name}")
    for name in (
        "pStyle",
        "keepNext",
        "keepLines",
        "pageBreakBefore",
        "framePr",
        "widowControl",
        "numPr",
        "suppressLineNumbers",
        "pBdr",
        "shd",
        "tabs",
        "suppressAutoHyphens",
        "kinsoku",
        "wordWrap",
        "overflowPunct",
        "topLinePunct",
        "autoSpaceDE",
        "autoSpaceDN",
        "bidi",
        "adjustRightInd",
        "snapToGrid",
        "spacing",
        "ind",
        "contextualSpacing",
        "mirrorIndents",
        "suppressOverlap",
        "jc",
        "textDirection",
        "textAlignment",
        "textboxTightWrap",
        "outlineLvl",
        "divId",
        "cnfStyle",
        "rPr",
        "sectPr",
        "pPrChange",
    )
)
# An article heading as the importer's ``_ARTICLE_RE`` reads it, with the
# pieces the export must reproduce: the ordinal's width ("1.01" vs "1.1"), a
# trailing dot, and the raw separator (a tab, a space, " - ").
_ARTICLE_HEADING_RE = re.compile(r"^([1-5])\.(\d{1,2})(\.?)(\s+[-–—]?\s*)(\S.*)$")
_LEADING_TOKEN_RE = re.compile(r"\S+(\s+)")
# The line the importer skips under a PART with no articles.
_NOT_USED = "(not used.)"


class SourceRenderError(ValueError):
    """The retained package cannot back an appearance-preserving export."""


def _normalized(text: str) -> str:
    """The importer's whitespace fold, for comparing like with like."""
    return " ".join((text or "").split())


def _is_blank_paragraph(element) -> bool:
    """A spacer: a paragraph with no text and nothing drawn in it."""
    if element.tag != _W_P:
        return False
    if _accept_all_paragraph_text(element).strip():
        return False
    return all(element.find(f".//{tag}") is None for tag in _LAYOUT_TAGS)


def _holds_break(element) -> bool:
    """A paragraph whose mark ends a Word section."""
    return (
        element.tag == _W_P
        and element.find(f"{_W_PPR}/{_W_SECTPR}") is not None
    )


def _strip_break(element) -> None:
    properties = element.find(_W_PPR)
    if properties is None:
        return
    for sect_pr in properties.findall(_W_SECTPR):
        properties.remove(sect_pr)


def _strip_identity(element) -> None:
    """Drop the ``w14`` ids Word expects to be unique; it regenerates them."""
    for attribute in (_W14_PARA_ID, _W14_TEXT_ID):
        if attribute in element.attrib:
            del element.attrib[attribute]


def _numbering_properties(element):
    """``element``'s own ``w:numPr``, created at its schema position."""
    properties = element.find(_W_PPR)
    if properties is None:
        properties = etree.Element(_W_PPR)
        element.insert(0, properties)
    numbering = properties.find(_W_NUMPR)
    if numbering is None:
        numbering = etree.Element(_W_NUMPR)
        after = _PPR_ORDER[_PPR_ORDER.index(_W_NUMPR) + 1 :]
        successor = next(
            (child for child in properties if child.tag in after), None
        )
        if successor is None:
            properties.append(numbering)
        else:
            successor.addprevious(numbering)
    return numbering


def _cancel_numbering(element) -> None:
    """``w:numId 0`` — "no numbering", whatever the paragraph style says.

    Word prints the number of an empty numbered paragraph, so an empty
    paragraph left holding a section break must not stay in the list.
    """
    numbering = _numbering_properties(element)
    for child in list(numbering):
        numbering.remove(child)
    etree.SubElement(numbering, _W_NUMID).set(_W_VAL, "0")


def _set_numbering_level(element, num_id: int, ilvl: int) -> None:
    """Number ``element`` at level ``ilvl`` of numbering instance ``num_id``.

    Its own ``w:numPr`` takes the level (``w:ilvl`` first, then ``w:numId``:
    ``CT_NumPr``'s order) and names the instance explicitly — even when the
    paragraph style names it too — so Word and the importer read one answer.
    """
    numbering = _numbering_properties(element)
    level = numbering.find(_W_ILVL)
    if level is None:
        level = etree.Element(_W_ILVL)
        numbering.insert(0, level)
    level.set(_W_VAL, str(ilvl))
    instance = numbering.find(_W_NUMID)
    if instance is None:
        instance = etree.Element(_W_NUMID)
        level.addnext(instance)
    instance.set(_W_VAL, str(num_id))


def _write_paragraph_text(paragraph_element, text: str) -> None:
    """Replace a cloned paragraph's inline content with ``text``.

    The fallback for a paragraph the splice cannot map, and the writer for a
    new provision. ``w:pPr`` survives untouched — that is the whole of the
    paragraph-level promise (style, numbering, indent, spacing, and the run
    defaults the style carries). Everything else is rebuilt from the first
    run's properties.
    """
    run_properties = None
    for run in paragraph_element.iterchildren(_W_R):
        found = run.find(_W_RPR)
        run_properties = copy.deepcopy(found) if found is not None else None
        break
    for child in list(paragraph_element):
        if child.tag != _W_PPR:
            paragraph_element.remove(child)
    run = etree.SubElement(paragraph_element, _W_R)
    if run_properties is not None:
        run.append(run_properties)
    append_text(run, text)


def _blank_template(numbered: bool):
    """A last-resort paragraph for a document that anchors nothing."""
    paragraph = etree.Element(_W_P)
    if numbered:  # pragma: no cover - defensive; templates carry their own
        etree.SubElement(etree.SubElement(paragraph, _W_PPR), _W_NUMPR)
    return paragraph


# ---------------------------------------------------------------------------
# The walk: the current tree, in document order
# ---------------------------------------------------------------------------


@dataclass
class _Item:
    """One element the export emits, in current document order."""

    uid: str
    kind: str  # "section" | "part" | "article" | "paragraph"
    origin: int  # NO_ORIGIN when the element is new
    mode: str  # "text" | "verbatim" | "locked" | "new"
    text: str = ""
    template: int | None = None
    #: A new provision's nesting depth (0 = directly under its article).
    depth: int | None = None


def _classify(uid: str) -> tuple[str, int | None]:
    """``(kind, depth)`` from the stable id's shape."""
    if uid in ("sec", SECTION_TITLE_UID):
        return "section", None
    if re.fullmatch(r"pt\d+", uid):
        return "part", None
    if re.fullmatch(r"pt\d+\.a\d+", uid):
        return "article", None
    return "paragraph", uid.count(".p") - 1


@dataclass(frozen=True)
class _ArticleStyle:
    padded: bool = False
    dot: str = ""
    separator: str = " "


def _article_style(element) -> _ArticleStyle:
    """How a master writes its article numbers: ``1.01`` or ``1.1``, a
    trailing dot, and what separates the number from the title."""
    if element is None or element.tag != _W_P:
        return _ArticleStyle()
    match = _ARTICLE_HEADING_RE.match(_accept_all_paragraph_text(element).strip())
    if not match:
        return _ArticleStyle()
    _part, ordinal, dot, separator, _title = match.groups()
    return _ArticleStyle(
        padded=len(ordinal) == 2 and ordinal.startswith("0"),
        dot=dot,
        separator=separator,
    )


def _article_heading(
    part_number: int, index: int, title: str, style: _ArticleStyle
) -> str:
    ordinal = f"{index + 1:02d}" if style.padded else str(index + 1)
    return f"{part_number}.{ordinal}{style.dot}{style.separator}{title}".rstrip()


def _label_separator(element) -> str:
    """The whitespace a master writes after a typed label — usually a tab."""
    if element is None or element.tag != _W_P:
        return " "
    match = _LEADING_TOKEN_RE.match(_accept_all_paragraph_text(element).lstrip())
    return match.group(1) if match else " "


class _Walker:
    """Collects the export's items and chooses templates for new ones.

    A new element is cloned from its nearest kin OF ITS OWN KIND — a
    provision from a provision at its depth, an article heading from an
    article heading — preferring the last one emitted before it, then the
    first one in the upload. Cloning from "whatever came last" made a new
    article look like the provision above it.
    """

    def __init__(self, children: list, format_map: SourceFormatMap):
        self._children = children
        self._map = format_map
        self.items: list[_Item] = []
        self._last: dict[tuple, int] = {}
        self._last_any: int | None = None
        self._first: dict[tuple, int] = {}
        self._kind_at: dict[int, str] = {}
        #: The nesting depth of every provision a new one may be cloned from.
        self._depth_at: dict[int, int] = {}
        for anchor in sorted(format_map.anchors, key=lambda a: a.origin_index):
            index = anchor.origin_index
            if anchor.label_kind and 0 <= index < len(children):
                self._kind_at.setdefault(index, anchor.label_kind)
            if anchor.locked or not 0 <= index < len(children):
                continue
            if children[index].tag != _W_P:
                continue
            kind, depth = _classify(anchor.uid)
            if kind == "paragraph":
                self._depth_at.setdefault(index, depth)
            for key in self._keys(kind, depth):
                self._first.setdefault(key, index)

    @staticmethod
    def _keys(kind: str, depth: int | None) -> list[tuple]:
        if kind == "paragraph":
            return [("paragraph", depth), ("paragraph",)]
        if kind in ("part", "article"):
            return [(kind,)]
        return []

    # -- source access ---------------------------------------------------
    def origin(self, uid: str) -> int:
        anchor = self._map.anchor(uid)
        if anchor is None or anchor.origin_index == NO_ORIGIN:
            return NO_ORIGIN
        if not 0 <= anchor.origin_index < len(self._children):
            return NO_ORIGIN
        return anchor.origin_index

    def label_kind(self, uid: str) -> str:
        anchor = self._map.anchor(uid)
        return anchor.label_kind if anchor is not None else ""

    def depth_at(self, index: int) -> int | None:
        """The nesting depth of the provision at upload index ``index``
        (``None``: not a provision a new one is cloned from)."""
        return self._depth_at.get(index)

    def _label_kind_at(self, index: int) -> str:
        recorded = self._kind_at.get(index)
        if recorded:
            return recorded
        element = self._children[index]
        if element.tag == _W_P and element.find(f"{_W_PPR}/{_W_NUMPR}") is not None:
            return LABEL_AUTO
        return LABEL_MANUAL

    # -- templates -------------------------------------------------------
    def template_for(self, kind: str, depth: int | None = None) -> int | None:
        if kind == "paragraph":
            order = [
                self._last.get(("paragraph", depth)),
                self._first.get(("paragraph", depth)),
                self._last.get(("paragraph",)),
                self._first.get(("paragraph",)),
            ]
        else:
            order = [self._last.get((kind,)), self._first.get((kind,))]
        for index in order:
            if index is not None:
                return index
        return self._last_any

    def template_label_kind(self, kind: str, depth: int | None = None) -> str:
        """The label convention a NEW element inherits from its template.

        A clone of an auto-numbered sibling is auto-numbered too, and Word
        will render its label — so writing one into the text as well would
        print it twice. Absence of ``w:numPr`` does not mean "manual": an
        unstructured import records no label at all.
        """
        index = self.template_for(kind, depth)
        if index is None:
            return LABEL_MANUAL
        return self._label_kind_at(index)

    def template_separator(self, kind: str, depth: int | None = None) -> str:
        index = self.template_for(kind, depth)
        if index is None or self._label_kind_at(index) != LABEL_MANUAL:
            return " "
        return _label_separator(self._children[index])

    def template_article_style(self) -> _ArticleStyle:
        index = self.template_for("article")
        if index is None or self._label_kind_at(index) != LABEL_MANUAL:
            return _ArticleStyle()
        return _article_style(self._children[index])

    def article_style_of(self, uid: str) -> _ArticleStyle:
        index = self.origin(uid)
        if index == NO_ORIGIN:
            return _ArticleStyle()
        return _article_style(self._children[index])

    def _remember(self, kind: str, depth: int | None, index: int) -> None:
        for key in self._keys(kind, depth):
            self._last[key] = index
        self._last_any = index

    # -- items -----------------------------------------------------------
    def add_verbatim(self, uid: str, kind: str) -> bool:
        """An anchored element exactly as it arrived, whatever it says.

        Used where the CALLER knows the content is unchanged but cannot
        reconstruct the original wording — the section header, whose line
        has several legitimate forms the parse folds into one number and
        title.
        """
        index = self.origin(uid)
        if index == NO_ORIGIN:
            return False
        self.items.append(_Item(uid, kind, index, "verbatim"))
        return True

    def add_locked(self, uid: str) -> bool:
        """A preserved block exactly as it arrived."""
        index = self.origin(uid)
        if index == NO_ORIGIN:
            return False
        self.items.append(_Item(uid, "paragraph", index, "locked"))
        return True

    def add_text(
        self, uid: str, kind: str, text: str, *, depth: int | None = None
    ) -> None:
        index = self.origin(uid)
        if index != NO_ORIGIN:
            if self._children[index].tag == _W_P:
                self._remember(kind, depth, index)
                self.items.append(_Item(uid, kind, index, "text", text))
            else:
                # An anchored non-paragraph (a table claimed by an editable
                # element) can only be emitted as itself.
                self.items.append(_Item(uid, kind, index, "verbatim"))
            return
        self.items.append(
            _Item(
                uid,
                kind,
                NO_ORIGIN,
                "new",
                text,
                template=self.template_for(kind, depth),
                depth=depth if kind == "paragraph" else None,
            )
        )


def _render_body(
    section: SpecSection,
    walker: _Walker,
    format_map: SourceFormatMap,
) -> None:
    if section.number or section.title:
        # While the section identity is exactly what was imported, reproduce
        # the header ELEMENT rather than rebuild its text. The parse folds
        # several legitimate header forms into one number and title — a
        # "SECTION 23 05 48" line, a keyword-less "23 05 48 — TITLE", a
        # number and title on separate lines — so rebuilding a canonical
        # form rewrites the firm's header on a NO-OP export. Caught in
        # review on PR #141 (Codex).
        unchanged_identity = (
            section.number == format_map.section_number
            and section.title == format_map.section_title
        )
        if unchanged_identity and walker.add_verbatim("sec", "section"):
            walker.add_verbatim(SECTION_TITLE_UID, "section")
            return _render_parts(section, walker)
        if format_map.header_source in (
            HEADER_SOURCE_FRONT_MATTER,
            HEADER_SOURCE_CHROME,
        ):
            # The identity lives in content this export carries through
            # verbatim anyway — a cover page's "Section Number:" field, the
            # page header/footer — so there is no header element to rewrite
            # and inventing one would print the section twice. A changed
            # identity is reported by the stale-identifier lint instead.
            return _render_parts(section, walker)
        header = " ".join(part for part in ("SECTION", section.number) if part)
        if walker.origin(SECTION_TITLE_UID) != NO_ORIGIN:
            walker.add_text("sec", "section", header.strip())
            walker.add_text(SECTION_TITLE_UID, "section", section.title)
        else:
            combined = f"{header} {section.title}".strip() if section.title else header
            walker.add_text("sec", "section", combined.strip())
    _render_parts(section, walker)


def _render_parts(section: SpecSection, walker: _Walker) -> None:
    for part in section.parts:
        if not part.articles and walker.origin(part.uid) == NO_ORIGIN:
            # SectionFormat always has three parts; the master may not have
            # written all three. Emitting a heading the upload never carried
            # would ADD content to the user's file.
            continue
        # ``Part.title`` already carries the whole heading line
        # ("PART 1 - GENERAL"), so an auto-numbered master is the only case
        # that needs the number stripped back off.
        label_kind = walker.label_kind(part.uid) or walker.template_label_kind(
            "part"
        )
        if label_kind == LABEL_AUTO:
            walker.add_text(part.uid, "part", _part_title_only(part))
        else:
            walker.add_text(part.uid, "part", part.title)
        for article_index, article in enumerate(part.articles):
            _render_article(part.number, article_index, article, walker)


def _part_title_only(part) -> str:
    """``"GENERAL"`` from ``"PART 1 - GENERAL"`` — Word renders the number."""
    _prefix, separator, remainder = part.title.partition(" - ")
    return remainder.strip() if separator else part.title


def _render_article(
    part_number: int, index: int, article: Article, walker: _Walker
) -> None:
    label_kind = walker.label_kind(article.uid)
    if label_kind:
        style = walker.article_style_of(article.uid)
    else:
        label_kind = walker.template_label_kind("article")
        style = walker.template_article_style()
    if label_kind == LABEL_AUTO:
        walker.add_text(article.uid, "article", article.title)
    else:
        # The master's own number format — "1.01" stays "1.01", "1.1 -
        # TITLE" keeps its dash — or an untouched heading is rewritten on
        # an export with no edits at all.
        walker.add_text(
            article.uid,
            "article",
            _article_heading(part_number, index, article.title, style),
        )
    for paragraph, label_index in labelled_paragraphs(article.paragraphs):
        _render_paragraph(paragraph, 0, label_index, walker)


def _render_paragraph(
    paragraph: Paragraph, depth: int, index: int, walker: _Walker
) -> None:
    if paragraph.locked:
        # A locked block with no source behind it cannot be invented: writing
        # a flattened table back as a paragraph would replace a grid with
        # pipe characters. Its children are still rendered either way —
        # the importer no longer lets a preserved block become a parent, but
        # a project saved before that fix can still carry one, and dropping
        # provisions on the floor is exactly the failure this guards.
        walker.add_locked(paragraph.uid)
        for child, child_label in labelled_paragraphs(paragraph.children):
            _render_paragraph(child, depth, child_label, walker)
        return
    label_kind = walker.label_kind(paragraph.uid)
    separator = " "
    if not label_kind:
        label_kind = walker.template_label_kind("paragraph", depth)
        # A new provision writes its kin's separator — the master's tab, not
        # a space. (An edited one keeps its own through the splice.)
        separator = walker.template_separator("paragraph", depth)
    if label_kind == LABEL_MANUAL:
        text = f"{_paragraph_label(depth, index)}{separator}{paragraph.text}"
    else:
        text = paragraph.text
    walker.add_text(paragraph.uid, "paragraph", text, depth=depth)
    for child, child_label in labelled_paragraphs(paragraph.children):
        _render_paragraph(child, depth + 1, child_label, walker)


# ---------------------------------------------------------------------------
# Assembly: unmodelled content, section breaks, and the render
# ---------------------------------------------------------------------------


@dataclass
class _Group:
    """Content bound to a POSITION in the upload rather than to an element.

    ``position`` is a source body index (a half-step between two indexes is
    "right after" the lower one). Every emitted element whose origin sits
    before it is "above" the group; every one after it is "below".
    """

    position: float
    members: list[int] = field(default_factory=list)
    holder: int | None = None  # a provision whose own w:pPr holds the break
    leftover: bool = False  # emit an empty paragraph holding the holder's break
    gap: int = 0
    #: The group carries a section break (a break paragraph, or a holder's):
    #: ``_place`` alone decides where it goes, and no tracked change may move
    #: it (Redline on your original, D-3).
    breaks: bool = False


# ---------------------------------------------------------------------------
# The emission plan (D-1 of the Redline-on-your-original plan)
# ---------------------------------------------------------------------------

#: A source element emitted as a byte-identical clone: an untouched
#: provision, a heading reproduced verbatim, a preserved block.
RECORD_KEPT = "kept"
#: A source paragraph whose words changed: the word-level splice into its
#: own runs, or — when the splice cannot map it — the first-run fallback.
RECORD_SPLICED = "spliced"
#: A new element, cloned from its nearest kin of its own kind.
RECORD_INSERTED = "inserted"
#: An anchored source element the user removed.
RECORD_DELETED = "deleted"
#: Unmodelled source content, carried through verbatim.
RECORD_CARRIED = "carried"
#: Unmodelled source content the clean export omits: the blank spacers of a
#: deleted element, a stale "(Not used.)" line.
RECORD_DROPPED = "dropped"
#: The empty paragraph a displaced section-break holder leaves where its
#: break belongs — the holder's own paragraph properties, nothing else.
RECORD_LEFTOVER = "leftover"

RECORD_KINDS = frozenset(
    {
        RECORD_KEPT,
        RECORD_SPLICED,
        RECORD_INSERTED,
        RECORD_DELETED,
        RECORD_CARRIED,
        RECORD_DROPPED,
        RECORD_LEFTOVER,
    }
)
#: The kinds the clean export renders nothing for.
_REMOVED_KINDS = frozenset({RECORD_DELETED, RECORD_DROPPED})


@dataclass
class Record:
    """One body child of the emission plan.

    ``source`` is the upload body index the record reproduces, removes or
    derives from (``NO_ORIGIN`` for an inserted element). The clean export
    renders every record except ``deleted`` and ``dropped``; the redline
    renders every record, marking what Accept All and Reject All disagree
    about. One plan feeding both is what makes "Accept All equals the
    formatted export" true by construction.
    """

    kind: str
    source: int = NO_ORIGIN
    item: _Item | None = None
    #: The clean export emits the source WITHOUT its section break: the
    #: break was displaced and is left behind in a ``leftover`` record.
    strip_break: bool = False
    #: ``spliced``: the character map and edit script, or — when the splice
    #: could not map the paragraph — the fallback reason instead.
    pmap: ParagraphMap | None = None
    ops: list[SpliceOp] | None = None
    fallback: str = ""
    #: ``inserted``: the kin the new element is cloned from.
    template: int | None = None
    #: ``inserted``: ``(numId, ilvl)`` — the Word numbering level the clone
    #: takes instead of its kin's, when the kin sits at another depth
    #: (:meth:`_Assembler._nesting_level`); ``None`` keeps the kin's.
    level: tuple[int, int] | None = None
    #: Redline only: this record is one copy of a MOVED element — an
    #: ``inserted`` copy at its new position or a ``deleted`` copy at its
    #: old one (Phase 1, D-1/D-3).
    moved: bool = False
    #: ``carried`` leading content: the uid of the element it travels with.
    owner: str = ""
    #: Where ``_place`` put this record around a section break — content the
    #: redline must keep exactly there, because Word cannot track a moved
    #: section break.
    pinned: bool = False


class _NumberingTables:
    """The upload's Word numbering, read exactly the way the importer reads
    it (its own helpers, over its own python-docx view of the package).

    Loaded lazily — only a new provision cloned from kin at another depth
    ever asks — and degrading like the importer: a package whose numbering
    or styles cannot be read has none, and nothing is renumbered.
    """

    def __init__(self, source_bytes: bytes):
        self._source_bytes = source_bytes
        self._loaded = False
        self.catalog: dict[tuple[int, int], tuple[str, str]] = {}
        self.style_numbering: dict[str, tuple[int, int]] = {}
        self.default_style_id = ""

    def _load(self) -> None:
        if self._loaded:
            return
        self._loaded = True
        try:
            document = Document(BytesIO(self._source_bytes))
            self.catalog = _load_numbering_catalog(document)
            self.style_numbering = _load_style_numbering(document)
            self.default_style_id = _default_paragraph_style_id(document)
        except Exception:  # noqa: BLE001 - unreadable numbering is no numbering
            self.catalog, self.style_numbering, self.default_style_id = {}, {}, ""

    def numbering(self, element) -> tuple[int, int] | None:
        """``(numId, ilvl)`` of ``element`` the way Word resolves it: its
        own ``w:numPr``, else its style's (the importer's
        ``_effective_numbering``)."""
        self._load()
        return _effective_numbering(
            DocxParagraph(element, None), self.style_numbering, self.default_style_id
        )

    def draws_a_provision(self, num_id: int, ilvl: int) -> bool:
        """The instance defines level ``ilvl`` with a label Word draws — not
        ``numFmt="none"``, not a blank ``lvlText``, either of which would
        print the provision with no number at all — and not as a PART or
        article heading (whose label grammar the importer promotes to
        structure)."""
        self._load()
        entry = self.catalog.get((num_id, ilvl))
        if entry is None:
            return False
        num_fmt, lvl_text = entry
        if num_fmt == "none" or not lvl_text.strip():
            return False
        return not _promoted_heading_kind(self.catalog, num_id, ilvl)


class _Assembler:
    def __init__(
        self,
        children: list,
        format_map: SourceFormatMap,
        walker: _Walker,
        section: SpecSection,
        numbering: _NumberingTables,
    ):
        self._children = children
        self._walker = walker
        self._numbering = numbering
        self.items = walker.items
        self.stats: dict = {
            "cloned": 0,
            "spliced": 0,
            "fallback": {},
            "inserted": 0,
            "preserved": 0,
            "break_leftovers": 0,
            "not_used_dropped": 0,
            # New provisions cloned from Word-numbered kin at another depth:
            # renumbered to their own level, or — the master's numbering
            # defines no provision level there — left at their kin's.
            "level_offset": 0,
            "level_kept": 0,
        }
        anchored: dict[int, str] = {}
        for anchor in format_map.anchors:
            if 0 <= anchor.origin_index < len(children):
                anchored.setdefault(anchor.origin_index, anchor.uid)
        self._anchored = anchored
        self._anchored_order = sorted(anchored)
        self._position: dict[int, int] = {}
        for index, item in enumerate(self.items):
            if item.origin != NO_ORIGIN:
                self._position.setdefault(item.origin, index)
        self._parts_with_articles = {
            part.uid for part in section.parts if part.articles
        }
        self._dropped: set[int] = set()
        # The blank spacers of a deleted element: the clean export drops
        # them with it, the redline marks them deleted with it.
        self._dropped_spacers: list[int] = []
        self._tails: dict[int, list[int]] = {}
        self._trailing: list[int] = []
        self._strip_break: set[int] = set()
        self._groups: list[_Group] = []

    # -- classification ----------------------------------------------------
    def _is_stale_not_used(self, index: int, above: int) -> bool:
        """A PART's "(Not used.)" line once that PART has an article.

        The importer skips the line, so it is unmodelled content; left alone
        the export printed "2.1 ISOLATORS" and then "(Not used.)" under it.
        A line that also holds a section break is never dropped — no edit
        loses a break.
        """
        element = self._children[index]
        if element.tag != _W_P or _holds_break(element):
            return False
        if _normalized(_accept_all_paragraph_text(element)).casefold() != _NOT_USED:
            return False
        uid = self._anchored.get(above, "")
        return _classify(uid)[0] == "part" and uid in self._parts_with_articles

    def _classify_runs(self) -> None:
        """Assign every unmodelled body child to exactly one place.

        The children between two anchored elements lead the later one (they
        travel with it), except that everything up to and including the LAST
        section break among them is bound to its position: a break is the
        end of the section above it, not the start of the element below.
        Everything after the last anchored element is trailing content —
        ``END OF SECTION``, an appendix — carried verbatim, blanks, page
        breaks and section breaks included.
        """
        previous = -1
        boundaries = self._anchored_order + [len(self._children)]
        for origin in boundaries:
            between = list(range(previous + 1, origin))
            for index in between:
                if self._is_stale_not_used(index, previous):
                    self._dropped.add(index)
            between = [i for i in between if i not in self._dropped]
            if origin == len(self._children):
                self._trailing = between
                break
            breaks = [i for i in between if _holds_break(self._children[i])]
            if breaks:
                last = breaks[-1]
                cut = between.index(last) + 1
                self._groups.append(
                    _Group(position=last, members=between[:cut], breaks=True)
                )
                between = between[cut:]
            if origin in self._position:
                self._tails[origin] = between
            else:
                # The element below was deleted. Its spacers go with it; what
                # the tree never modelled (a picture-only paragraph, a body
                # bookmark) stays where it was instead of being swept to the
                # end of the file.
                kept = []
                for i in between:
                    if _is_blank_paragraph(self._children[i]):
                        self._dropped_spacers.append(i)
                    else:
                        kept.append(i)
                if kept:
                    self._groups.append(_Group(position=origin - 0.5, members=kept))
            previous = origin
        for origin in self._anchored_order:
            if _holds_break(self._children[origin]):
                self._groups.append(
                    _Group(
                        position=origin + 0.5,
                        holder=origin,
                        leftover=origin not in self._position,
                        breaks=True,
                    )
                )

    # -- placement ---------------------------------------------------------
    def _nearest_emitted_above(self, position: float) -> int | None:
        found = None
        for origin in self._anchored_order:
            if origin >= position:
                break
            if origin in self._position:
                found = origin
        return found

    def _place(self) -> None:
        """Choose each group's gap in the output.

        A gap is scored by how many emitted elements it keeps on the right
        side: those that were above the group's position before it, those
        that were below it after it. The best-scoring gap wins; among ties,
        the one right after the nearest surviving element that was above it
        in the upload ("a break belongs to the content above it"), then any
        gap right after an element from above, then the earliest. Groups are
        placed in upload order and never cross each other, so sections keep
        their order.
        """
        count = len(self.items)
        origins = [item.origin for item in self.items]
        floor = 0
        for group in sorted(self._groups, key=lambda g: g.position):
            above = [
                origin != NO_ORIGIN and origin < group.position for origin in origins
            ]
            below = [
                origin != NO_ORIGIN and origin > group.position for origin in origins
            ]
            before = [0] * (count + 1)
            for index in range(count):
                before[index + 1] = before[index] + above[index]
            after = [0] * (count + 1)
            for index in range(count - 1, -1, -1):
                after[index] = after[index + 1] + below[index]
            scores = [before[gap] + after[gap] for gap in range(count + 1)]
            best = max(scores[floor:])
            candidates = [
                gap for gap in range(floor, count + 1) if scores[gap] == best
            ]
            chosen = None
            nearest = self._nearest_emitted_above(group.position)
            if nearest is not None and self._position[nearest] + 1 in candidates:
                chosen = self._position[nearest] + 1
            if chosen is None:
                after_above = [gap for gap in candidates if gap and above[gap - 1]]
                chosen = after_above[0] if after_above else candidates[0]
            group.gap = chosen
            floor = chosen
            if group.holder is not None and not group.leftover:
                if chosen == self._position[group.holder] + 1:
                    # The holder is still where its break belongs: it keeps
                    # it, in its own w:pPr, and nothing else is emitted.
                    group.holder = None
                else:
                    self._strip_break.add(group.holder)
                    group.leftover = True
            if group.leftover:
                self.stats["break_leftovers"] += 1

    # -- the plan ----------------------------------------------------------
    def _item_record(self, item: _Item) -> Record:
        """What the export does with one tree element, decided once.

        The decision the renderers used to make while rendering — clone,
        splice or rebuild — is made here, so the clean export and the
        redline read the same answer.
        """
        strip_break = item.origin in self._strip_break
        if item.mode in ("verbatim", "locked"):
            self.stats["preserved"] += 1
            return Record(
                RECORD_KEPT, item.origin, item, strip_break=strip_break
            )
        if item.mode == "new":
            self.stats["inserted"] += 1
            return Record(
                RECORD_INSERTED,
                NO_ORIGIN,
                item,
                template=item.template,
                level=self._nesting_level(item),
            )
        source = self._children[item.origin]
        source_text = _accept_all_paragraph_text(source)
        revised = _element_has_tracked_changes(source)
        if not revised and _normalized(source_text) == _normalized(item.text):
            # Compare the NORMALIZED forms. The importer folds runs of
            # whitespace, so a source paragraph containing a double space
            # after a period — which is most office masters — would otherwise
            # never match its own semantic text. Caught in review on PR #141
            # (Codex). This is the untouched-provision guarantee: a
            # byte-identical clone, markup and all.
            self.stats["cloned"] += 1
            return Record(
                RECORD_KEPT, item.origin, item, strip_break=strip_break
            )
        if revised:
            pmap, reason = None, FALLBACK_REVISIONS
        else:
            pmap, reason = map_paragraph(source, expected_text=source_text)
        if pmap is not None:
            self.stats["spliced"] += 1
            return Record(
                RECORD_SPLICED,
                item.origin,
                item,
                strip_break=strip_break,
                pmap=pmap,
                ops=plan_splice(pmap.text, item.text),
            )
        fallbacks = self.stats["fallback"]
        fallbacks[reason] = fallbacks.get(reason, 0) + 1
        return Record(
            RECORD_SPLICED,
            item.origin,
            item,
            strip_break=strip_break,
            fallback=reason,
        )

    def _nesting_level(self, item: _Item) -> tuple[int, int] | None:
        """The Word numbering level a NEW provision takes, when its kin sits
        at another depth: ``(numId, ilvl)``, or ``None`` to keep the kin's.

        A new provision is cloned from kin of its own kind, preferring kin
        at its own depth — but a master may have no provision at that depth
        (the first sub-provision anywhere under an "A."), and then the kin
        sits shallower. Cloning its ``w:ilvl`` drew the new provision one
        level up, and the importer read it back as its parent's sibling.

        ``ilvl`` is a level RELATIVE to the article's list — that is how the
        importer reads it (``_TreeBuilder.numbered_paragraph``) — so the
        clone's level is the kin's offset by the depth difference, in the
        kin's own numbering instance, resolved the way Word resolves it.
        Only a level the instance defines with a visible label, and draws as
        a provision rather than a PART or article heading, is taken
        (:meth:`_NumberingTables.draws_a_provision`); otherwise the clone
        keeps its kin's level — one level up in Word, but never a number the
        master's numbering cannot draw, nor no number at all — and the
        export event counts it (``level_kept``).
        """
        if item.kind != "paragraph" or item.template is None or item.depth is None:
            return None
        template_depth = self._walker.depth_at(item.template)
        if template_depth is None or template_depth == item.depth:
            return None
        if self._walker._label_kind_at(item.template) != LABEL_AUTO:
            return None  # a typed label carries its own level
        numbering = self._numbering.numbering(self._children[item.template])
        if numbering is not None:
            num_id, ilvl = numbering
            target = ilvl + item.depth - template_depth
            if target >= 0 and self._numbering.draws_a_provision(num_id, target):
                self.stats["level_offset"] += 1
                return num_id, target
        self.stats["level_kept"] += 1
        return None

    def _removed_records(self) -> list[Record]:
        """What the upload holds and the clean export does not, in upload
        order: removed anchored elements, their blank spacers, and stale
        "(Not used.)" lines. The clean export renders none of them; the
        redline marks each one deleted."""
        holders = {
            group.holder
            for group in self._groups
            if group.holder is not None and group.leftover
        }
        removed = [
            Record(RECORD_DELETED, origin)
            for origin in self._anchored_order
            if origin not in self._position and origin not in holders
        ]
        removed.extend(
            Record(RECORD_DROPPED, index)
            for index in sorted(self._dropped) + self._dropped_spacers
        )
        removed.sort(key=lambda record: record.source)
        return removed

    def plan(self) -> tuple[list[Record], list[Record]]:
        """``(records, removed)``: the clean export's body children in
        output order, and what it leaves out of the upload."""
        self._classify_runs()
        self._place()
        self.stats["not_used_dropped"] = len(self._dropped)
        by_gap: dict[int, list[_Group]] = {}
        for group in sorted(self._groups, key=lambda g: g.position):
            if group.members or group.leftover:
                by_gap.setdefault(group.gap, []).append(group)
        records: list[Record] = []
        emitted_tails: set[int] = set()
        for gap in range(len(self.items) + 1):
            for group in by_gap.get(gap, ()):
                records.extend(
                    Record(RECORD_CARRIED, i, pinned=group.breaks)
                    for i in group.members
                )
                if group.leftover and group.holder is not None:
                    records.append(
                        Record(RECORD_LEFTOVER, group.holder, pinned=True)
                    )
            if gap == len(self.items):
                break
            item = self.items[gap]
            if item.origin != NO_ORIGIN and item.origin not in emitted_tails:
                emitted_tails.add(item.origin)
                records.extend(
                    Record(RECORD_CARRIED, i, owner=item.uid)
                    for i in self._tails.get(item.origin, ())
                )
            record = self._item_record(item)
            if (
                record.source != NO_ORIGIN
                and not record.strip_break
                and _holds_break(self._children[record.source])
            ):
                # A holder still where its break belongs keeps the break.
                record.pinned = True
            records.append(record)
        records.extend(Record(RECORD_CARRIED, i) for i in self._trailing)
        return records, self._removed_records()

    # -- the clean rendering ----------------------------------------------
    def render_leftover(self, holder: int):
        """An empty paragraph holding a displaced provision's section break.

        The provision's own paragraph properties, so the break keeps its
        place and spacing — minus its text, its identity (the provision may
        still exist elsewhere) and its list numbering.
        """
        paragraph = copy.deepcopy(self._children[holder])
        for child in list(paragraph):
            if child.tag != _W_PPR:
                paragraph.remove(child)
        _strip_identity(paragraph)
        if self.holder_auto_numbered(holder):
            _cancel_numbering(paragraph)
        return paragraph

    def holder_auto_numbered(self, holder: int) -> bool:
        uid = self._anchored.get(holder, "")
        return self._walker.label_kind(uid) == LABEL_AUTO

    def render_inserted(self, record: Record):
        """A new element: its kin's formatting, never its identity — and,
        when its kin sits at another depth, its own numbering level."""
        if record.template is None:
            element = _blank_template(False)
        else:
            element = copy.deepcopy(self._children[record.template])
            # Clone hygiene: the kin's formatting, never its identity.
            _strip_break(element)
            _strip_identity(element)
            if record.level is not None:
                _set_numbering_level(element, *record.level)
        _write_paragraph_text(element, record.item.text)
        return element

    def render_accepted(self, record: Record):
        """``record`` as the clean export emits it (``None``: nothing)."""
        kind = record.kind
        if kind in _REMOVED_KINDS:
            return None
        if kind == RECORD_INSERTED:
            return self.render_inserted(record)
        if kind == RECORD_LEFTOVER:
            return self.render_leftover(record.source)
        if kind == RECORD_SPLICED:
            if record.pmap is not None:
                element = render_clean(record.pmap, record.ops or [])
            else:
                element = copy.deepcopy(self._children[record.source])
                _write_paragraph_text(element, record.item.text)
        else:  # kept, carried
            element = copy.deepcopy(self._children[record.source])
        if record.strip_break:
            _strip_break(element)
        return element

    def assemble(self) -> list:
        records, _removed = self.plan()
        rendered = (self.render_accepted(record) for record in records)
        return [element for element in rendered if element is not None]


@dataclass
class _LoadedBody:
    root: object
    body: object
    content: list
    trailing_sect_pr: object | None
    #: The upload the body was read from (its numbering is read from it
    #: only when a new provision needs a level of its own).
    source_bytes: bytes = b""


def _load_body(source_bytes: bytes, format_map: SourceFormatMap) -> _LoadedBody:
    """Validate the map against the bytes and parse the upload's body."""
    if not isinstance(source_bytes, bytes):
        raise TypeError("source_bytes must be bytes")
    if format_map is None:
        raise SourceRenderError(
            "This document has no retained formatting map, so its original "
            "layout cannot be reproduced."
        )
    if not format_map.matches(source_bytes):
        # Origin indexes are positions in a specific package. Read beside
        # different bytes they address whatever now sits at that index, so
        # refusing is the only safe answer.
        raise SourceRenderError(
            "The retained formatting map does not describe the retained "
            "source document."
        )
    try:
        with zipfile.ZipFile(BytesIO(source_bytes), "r") as archive:
            document_xml = archive.read(_DOCUMENT_PART)
    except (KeyError, zipfile.BadZipFile, RuntimeError) as exc:
        raise SourceRenderError(
            "The retained source document could not be opened."
        ) from exc
    try:
        # python-docx's parser, NOT etree.fromstring: the Accept-All text
        # reader depends on python-docx's registered oxml element classes,
        # and a plain lxml tree reports every paragraph as empty — which
        # silently sends untouched provisions down the rewrite path and
        # loses the byte-identical-clone guarantee.
        root = parse_xml(document_xml)
    except (etree.XMLSyntaxError, ValueError) as exc:
        raise SourceRenderError(
            "The retained source document XML could not be parsed."
        ) from exc
    body = root.find(qn("w:body"))
    if body is None:
        raise SourceRenderError("The retained source document has no body.")

    children = [child for child in body.iterchildren() if isinstance(child.tag, str)]
    # Word requires the body's section properties to remain its last child;
    # they carry page size, margins and the header/footer references, which
    # is most of what "keep my formatting" means.
    trailing_sect_pr = (
        children[-1] if children and children[-1].tag == _W_SECTPR else None
    )
    content = children[:-1] if trailing_sect_pr is not None else children
    return _LoadedBody(root, body, content, trailing_sect_pr, source_bytes)


def _serialize(loaded: _LoadedBody, elements: list) -> bytes:
    """``word/document.xml`` with ``elements`` as the body."""
    body = loaded.body
    for child in list(body):
        body.remove(child)
    for element in elements:
        body.append(element)
    if loaded.trailing_sect_pr is not None:
        body.append(copy.deepcopy(loaded.trailing_sect_pr))
    return etree.tostring(
        loaded.root, xml_declaration=True, encoding="UTF-8", standalone=True
    )


def _plan_for(loaded: _LoadedBody, format_map: SourceFormatMap, current: SpecSection):
    walker = _Walker(loaded.content, format_map)
    _render_body(current, walker, format_map)
    return _Assembler(
        loaded.content,
        format_map,
        walker,
        current,
        _NumberingTables(loaded.source_bytes),
    )


def render_preserving_docx(
    *,
    source_bytes: bytes,
    format_map: SourceFormatMap,
    current: SpecSection,
    stats: dict | None = None,
) -> bytes:
    """Return the upload with a rebuilt body carrying ``current``.

    ``stats``, when given, is filled with counts for the export's
    diagnostics: elements cloned, spliced, rebuilt by the fallback (by
    reason), inserted and preserved; empty paragraphs left holding a
    displaced provision's section break (``break_leftovers``); stale
    "(Not used.)" lines dropped. Counts only — never provision text.
    """
    loaded = _load_body(source_bytes, format_map)
    assembler = _plan_for(loaded, format_map, current)
    rendered = assembler.assemble()
    if stats is not None:
        stats.update(assembler.stats)
    return replace_document_xml_raw(source_bytes, _serialize(loaded, rendered))


# ---------------------------------------------------------------------------
# Redline on your original (Phase 1)
# ---------------------------------------------------------------------------

#: Why a redline on the original could not be produced — a closed
#: vocabulary the route answers with and the ``export`` event records.
#: The first two are availability answers (the doc payload says so before
#: anyone clicks); the rest are refusals only a render can reach.
REDLINE_NO_BASELINE = "no_baseline"
REDLINE_NO_ORIGINAL = "no_original"
REDLINE_PENDING_REVISIONS = "pending_revisions"
REDLINE_REVISION_SCAN = "revision_scan_unavailable"
REDLINE_SECTION_BREAK = "section_break_reorder"
REDLINE_MOVED_ANNOTATION = "moved_annotation"
REDLINE_UNACCOUNTED = "unaccounted_content"
REDLINE_ACCEPT_CHECK = "accept_check_failed"
REDLINE_REJECT_CHECK = "reject_check_failed"
REDLINE_BOOKMARK_CHECK = "duplicate_bookmarks"
REDLINE_PACKAGE_CHECK = "package_check_failed"

_REDLINE_UNAVAILABLE = (
    "The redline of extracted provisions still works."
)
_REDLINE_MESSAGES = {
    REDLINE_NO_BASELINE: (
        "There is no imported master in this document's history to compare "
        "against. Redline vs version still works."
    ),
    REDLINE_NO_ORIGINAL: (
        "A redline on your original needs the Word file this project was "
        "imported from and the map Build-a-Spec builds from it at import, "
        "and this project does not keep both (a project imported before "
        "1.14.0 has no map; importing the file again gives it one). "
        + _REDLINE_UNAVAILABLE
    ),
    REDLINE_PENDING_REVISIONS: (
        "This Word file already carries tracked changes, so a redline on it "
        "could not promise that Reject All gives your original back. Accept "
        "or reject those changes in Word, save, and import the file again. "
        + _REDLINE_UNAVAILABLE
    ),
    REDLINE_REVISION_SCAN: (
        "Build-a-Spec could not confirm that this Word file carries no "
        "tracked changes, so it cannot promise that Reject All gives your "
        "original back. " + _REDLINE_UNAVAILABLE
    ),
    REDLINE_SECTION_BREAK: (
        "This version reorders content across a Word section break in a way "
        "a tracked change cannot show (Word cannot track a moved section "
        "break). " + _REDLINE_UNAVAILABLE
    ),
    REDLINE_MOVED_ANNOTATION: (
        "A provision you moved carries a comment or a footnote reference, "
        "and Word cannot show one annotation in two places. "
        + _REDLINE_UNAVAILABLE
    ),
    REDLINE_UNACCOUNTED: (
        "Part of your original could not be placed in the redline, so it "
        "was refused rather than risk losing it. " + _REDLINE_UNAVAILABLE
    ),
    "simple_field": (
        "A provision you changed contains a Word field that cannot be shown "
        "as a tracked change. " + _REDLINE_UNAVAILABLE
    ),
    "block_content_control": (
        "A content control you deleted or moved cannot be shown as a Word "
        "tracked change. " + _REDLINE_UNAVAILABLE
    ),
    "field_block": (
        "Part of a table of contents (a Word field) was deleted or moved, "
        "and Word cannot show that as a tracked change. "
        + _REDLINE_UNAVAILABLE
    ),
    "pending_revisions_in_body": (
        "This Word file already carries tracked changes. "
        + _REDLINE_UNAVAILABLE
    ),
    "untrackable_markup": (
        "Your original contains Word markup around a change that cannot be "
        "shown as a tracked change. " + _REDLINE_UNAVAILABLE
    ),
}
_SELF_CHECK_MESSAGE = (
    "The redline failed its own check ({check}): it could not prove that "
    "Accept All gives the formatted export and Reject All gives your "
    "original back, so it was not handed over. " + _REDLINE_UNAVAILABLE
)


class SourceRedlineError(SourceRenderError):
    """No redline on the original, for a named ``reason``.

    ``detail`` carries positions and element names only — never document
    text — for the ``export`` diagnostics event.
    """

    def __init__(self, reason: str, *, detail: dict | None = None):
        message = _REDLINE_MESSAGES.get(reason) or _SELF_CHECK_MESSAGE.format(
            check=reason.replace("_", " ")
        )
        super().__init__(message)
        self.reason = reason
        self.detail = dict(detail or {})


def redline_refusal_message(reason: str) -> str:
    """The user-facing sentence for a refusal ``reason``."""
    return SourceRedlineError(reason).args[0]


_W_TBL = qn("w:tbl")
_W_SDT = qn("w:sdt")
_W_FLD_CHAR = qn("w:fldChar")
_W_FLD_CHAR_TYPE = qn("w:fldCharType")
_W_BOOKMARK_START = qn("w:bookmarkStart")
_W_BOOKMARK_END = qn("w:bookmarkEnd")
_W_NAME = qn("w:name")
_ANNOTATION_TAGS = (
    qn("w:commentRangeStart"),
    qn("w:commentRangeEnd"),
    qn("w:commentReference"),
    qn("w:footnoteReference"),
    qn("w:endnoteReference"),
)


def _crosses_field_boundary(paragraph) -> bool:
    """A complex field that opens or closes outside this paragraph — a table
    of contents spans many. Tracking part of one would break the field."""
    depth = 0
    for char in paragraph.iter(_W_FLD_CHAR):
        kind = char.get(_W_FLD_CHAR_TYPE)
        if kind == "begin":
            depth += 1
        elif kind == "end":
            depth -= 1
            if depth < 0:
                return True
    return depth != 0


def _give_up_identity(element) -> None:
    """The old copy of a moved element: its bookmarks and ``w14`` ids go
    with the new copy, because a file must never carry a bookmark name (or
    a paragraph id) twice."""
    for marker in list(element.iter(_W_BOOKMARK_START, _W_BOOKMARK_END)):
        marker.getparent().remove(marker)
    for node in element.iter():
        if isinstance(node.tag, str):
            _strip_identity(node)


def _bookmark_names(elements) -> set[str]:
    return {
        start.get(_W_NAME, "")
        for element in elements
        for start in element.iter(_W_BOOKMARK_START)
    }


def _first_run_properties(paragraph):
    """What ``_write_paragraph_text`` borrows: the first direct run's rPr."""
    for run in paragraph.iterchildren(_W_R):
        found = run.find(_W_RPR)
        return copy.deepcopy(found) if found is not None else None
    return None


def _subtree_uids(section: SpecSection, roots) -> set[str]:
    """``roots`` and every descendant of them in ``section``."""
    roots = set(roots)
    found: set[str] = set()

    def paragraphs(nodes, inside):
        for node in nodes:
            here = inside or node.uid in roots
            if here:
                found.add(node.uid)
            paragraphs(node.children, here)

    for part in section.parts:
        for article in part.articles:
            inside = article.uid in roots
            if inside:
                found.add(article.uid)
            paragraphs(article.paragraphs, inside)
    return found


class _RedlineBuilder:
    """The body of the redline: the upload and the clean export, merged.

    Accept All must give the clean export (C) and Reject All the upload (U),
    so the redline is an interleaving of the two in which the content they
    SHARE — untouched and edited elements, carried content, section breaks
    — appears once, in an order both agree on. Everything only U has is a
    tracked deletion; everything only C has, a tracked insertion.

    Which content is shared is a heaviest increasing subsequence over the
    clean records' upload positions, weighted so that (1) every record
    ``_place`` pinned around a section break is shared — Word cannot track
    a moved break, and the clean export already decided where it goes;
    (2) then as much as possible of what the diff did NOT report moved; (3)
    then anything else. A record left out is a move: an inserted copy where
    C has it, a deleted copy where U had it. The diff's per-sibling moves
    are therefore honoured wherever a section break allows, and only
    extended — never contradicted without cause — where one does not.
    """

    def __init__(
        self,
        loaded: _LoadedBody,
        format_map: SourceFormatMap,
        assembler: _Assembler,
        moved_uids: set[str],
        marks,
    ):
        self.content = loaded.content
        self.format_map = format_map
        self.assembler = assembler
        self.moved_uids = moved_uids
        self.marks = marks
        self.stats: dict = {
            "kept": 0,
            "spliced": 0,
            "fallback": 0,
            "inserted": 0,
            "deleted": 0,
            "dropped": 0,
            "moved": 0,
            "moves_added": 0,
            "leftovers": 0,
            "last_mark_untracked": 0,
        }

    # -- helpers -------------------------------------------------------------
    def _uid_at(self, source: int) -> str:
        return self.assembler._anchored.get(source, "")

    def _locked_reason(self, source: int) -> str:
        anchor = self.format_map.anchor(self._uid_at(source))
        return anchor.locked if anchor is not None else ""

    def _prefers_move(self, record: Record) -> bool:
        if record.item is not None and record.item.uid in self.moved_uids:
            return True
        return bool(record.owner) and record.owner in self.moved_uids

    def _check_trackable(self, element, source: int, *, moved: bool) -> None:
        tag = element.tag
        if tag == _W_SDT:
            raise SourceRedlineError(
                "block_content_control", detail={"index": source, "tag": "sdt"}
            )
        if tag == _W_P:
            if self._locked_reason(source) == "field" or _crosses_field_boundary(
                element
            ):
                raise SourceRedlineError(
                    "field_block", detail={"index": source, "tag": "p"}
                )
            if moved and any(True for _ in element.iter(*_ANNOTATION_TAGS)):
                raise SourceRedlineError(
                    REDLINE_MOVED_ANNOTATION, detail={"index": source, "tag": "p"}
                )
            return
        if tag == _W_TBL:
            if moved and any(True for _ in element.iter(*_ANNOTATION_TAGS)):
                raise SourceRedlineError(
                    REDLINE_MOVED_ANNOTATION, detail={"index": source, "tag": "tbl"}
                )
            return
        raise SourceRedlineError(
            "untrackable_markup",
            detail={"index": source, "tag": etree.QName(tag).localname},
        )

    def _tracked(self, action):
        from .revision_marks import UntrackableContent

        try:
            return action()
        except UntrackableContent as exc:
            reason = (
                "pending_revisions_in_body"
                if exc.reason == "pending_revisions"
                else exc.reason
            )
            raise SourceRedlineError(reason, detail={"tag": exc.tag}) from exc

    # -- the merge -----------------------------------------------------------
    def _shared(self, records: list[Record]) -> set[int]:
        candidates = [i for i, r in enumerate(records) if r.source != NO_ORIGIN]
        movers = [self._prefers_move(records[i]) for i in candidates]
        kept = _keep_in_place(
            [records[i].source for i in candidates],
            [records[i].pinned for i in candidates],
            movers,
        )
        shared = {candidates[i] for i in kept}
        for position, index in enumerate(candidates):
            record = records[index]
            if record.pinned and index not in shared:
                raise SourceRedlineError(
                    REDLINE_SECTION_BREAK, detail={"index": record.source}
                )
            if index not in shared and not movers[position]:
                self.stats["moves_added"] += 1
        return shared

    def plan(self, records: list[Record], removed: list[Record]) -> list[Record]:
        """The redline's records in output order."""
        shared = self._shared(records)
        claimed = {records[i].source for i in shared}
        referenced = {r.source for r in records if r.source != NO_ORIGIN}
        removed_by_source = {r.source: r for r in removed}
        u_only: list[Record] = []
        for source in range(len(self.content)):
            if source in claimed:
                continue
            if source in removed_by_source:
                u_only.append(removed_by_source[source])
            elif source in referenced:
                u_only.append(Record(RECORD_DELETED, source, moved=True))
            else:
                raise SourceRedlineError(
                    REDLINE_UNACCOUNTED, detail={"index": source}
                )
        ordered: list[Record] = []
        next_shared = [0] * (len(records) + 1)
        upcoming = len(self.content)
        next_shared[len(records)] = upcoming
        for index in range(len(records) - 1, -1, -1):
            if index in shared:
                upcoming = records[index].source
            next_shared[index] = upcoming
        cursor = 0

        def flush(below: int) -> None:
            nonlocal cursor
            while cursor < len(u_only) and u_only[cursor].source < below:
                ordered.append(u_only[cursor])
                cursor += 1

        for index, record in enumerate(records):
            if index in shared:
                flush(record.source)
                ordered.append(record)
                continue
            # A copy only the clean export has. Deletions in the same gap go
            # first, the way a reviewer reads a replacement.
            flush(next_shared[index])
            if record.source == NO_ORIGIN:
                ordered.append(record)
            else:
                moved = copy.copy(record)
                moved.moved = True
                ordered.append(moved)
        flush(len(self.content) + 1)
        return ordered

    # -- rendering -----------------------------------------------------------
    def _render_shared(self, record: Record):
        """``(element, mark_flag)`` for content both views share."""
        if record.strip_break:  # pragma: no cover - a pinned leftover owns it
            raise SourceRedlineError(REDLINE_SECTION_BREAK)
        source = self.content[record.source]
        if record.kind in (RECORD_KEPT, RECORD_CARRIED):
            self.stats["kept"] += 1
            return copy.deepcopy(source), None
        if record.kind == RECORD_LEFTOVER:
            return self._render_emptied_holder(record.source), None
        # spliced
        if record.pmap is not None:
            self.stats["spliced"] += 1
            return render_redline(record.pmap, record.ops or [], self.marks), None
        self.stats["fallback"] += 1
        return self._render_fallback(source, record.item.text), None

    def _render_fallback(self, source, text: str):
        """Everything the paragraph held, deleted; one run of the new text,
        inserted — Accept All is exactly ``_write_paragraph_text``."""
        element = copy.deepcopy(source)
        properties = _first_run_properties(source)
        self._tracked(lambda: delete_content(element, self.marks))
        run = etree.Element(_W_R)
        if properties is not None:
            run.append(properties)
        append_text(run, text)
        wrapper = self.marks.wrapper("w:ins")
        wrapper.append(run)
        element.append(wrapper)
        return element

    def _render_emptied_holder(self, holder: int):
        """The paragraph that held a displaced section break: its words go,
        its mark (and the break it holds) stays. Accept All is the clean
        export's leftover; Reject All is the provision."""
        self.stats["leftovers"] += 1
        source = self.content[holder]
        element = copy.deepcopy(source)
        moved = self._uid_at(holder) in self.moved_uids or any(
            item.origin == holder for item in self.assembler.items
        )
        if moved:
            _give_up_identity(element)
        self._check_trackable(element, holder, moved=moved)
        self._tracked(lambda: delete_content(element, self.marks))
        if self.assembler.holder_auto_numbered(holder):
            original = copy.deepcopy(element.find(_W_PPR))
            _cancel_numbering(element)
            record_paragraph_properties(element, original, self.marks)
        return element

    def _render_inserted(self, record: Record):
        """``(element, mark_flag)`` for a copy only the clean export has."""
        if record.source == NO_ORIGIN:
            self.stats["inserted"] += 1
            element = self.assembler.render_inserted(record)
            source_index = -1
        else:
            self.stats["moved"] += 1
            element = self.assembler.render_accepted(record)
            source_index = record.source
            self._check_trackable(element, source_index, moved=True)
        if element.tag == _W_TBL:
            self._tracked(lambda: mark_table(element, "w:ins", self.marks))
            return element, None
        if element.tag != _W_P:  # pragma: no cover - _check_trackable refused it
            raise SourceRedlineError("untrackable_markup")
        self._tracked(lambda: insert_content(element, self.marks))
        return element, "w:ins"

    def _render_deleted(self, record: Record):
        """``(element, mark_flag)`` for content only the upload has."""
        self.stats[
            "deleted" if record.kind == RECORD_DELETED else "dropped"
        ] += 1
        element = copy.deepcopy(self.content[record.source])
        if _holds_break(element):  # pragma: no cover - pinned records own them
            raise SourceRedlineError(
                REDLINE_SECTION_BREAK, detail={"index": record.source}
            )
        self._check_trackable(element, record.source, moved=record.moved)
        if record.moved:
            _give_up_identity(element)
        if element.tag == _W_TBL:
            self._tracked(lambda: mark_table(element, "w:del", self.marks))
            return element, None
        self._tracked(lambda: delete_content(element, self.marks))
        return element, "w:del"

    def render(self, ordered: list[Record]) -> tuple[list, set[str]]:
        """The redline's body elements, and the bookmark names that moved
        with a new copy (the one documented Reject-All limit)."""
        rendered: list[tuple] = []
        moved_names: set[str] = set()
        for record in ordered:
            if record.kind in _REMOVED_KINDS:
                rendered.append(self._render_deleted(record))
            elif record.moved:
                element, flag = self._render_inserted(record)
                moved_names |= _bookmark_names([element])
                rendered.append((element, flag))
            elif record.kind == RECORD_INSERTED:
                rendered.append(self._render_inserted(record))
            else:
                rendered.append(self._render_shared(record))
        last = len(rendered) - 1
        for index, (element, flag) in enumerate(rendered):
            if flag is None:
                continue
            if index != last:
                mark_paragraph(element, flag, self.marks)
                continue
            # Word cannot track a document's last paragraph mark: its words
            # are marked, and the resolution that should remove it leaves an
            # empty paragraph behind — which must set nothing Word shows (a
            # number, a page break), so its formatting becomes a tracked
            # change whose side there is empty.
            neutralize_last_paragraph(element, flag, self.marks)
            self.stats["last_mark_untracked"] += 1
        return [element for element, _flag in rendered], moved_names


def _keep_in_place(sources: list[int], pinned: list[bool], movers: list[bool]) -> list[int]:
    """Which of the clean export's records stay where they are (positions
    into ``sources``, each a record's upload index, in clean order).

    A heaviest increasing subsequence of the upload positions, weighted so
    that (1) a pinned record — placed by ``_place`` around a section break —
    outweighs any number of others, because Word cannot track a moved break;
    (2) then a record the diff did not report moved (a "stayer") outweighs
    any number of reported movers; (3) a mover counts least. So the chain
    keeps every break where the clean export put it whenever any increasing
    chain can, honours the diff's choice wherever the breaks allow, and
    otherwise moves the fewest elements it can.
    """
    from .diffing import heaviest_increasing_subsequence

    count = len(sources) + 1
    stay = count
    pin = count * count + 1
    weights = [
        pin if pinned[i] else (1 if movers[i] else stay) for i in range(len(sources))
    ]
    return heaviest_increasing_subsequence(list(sources), weights)


def _body_of(loaded: _LoadedBody, elements: list):
    body = etree.Element(qn("w:body"))
    for element in elements:
        body.append(element)
    if loaded.trailing_sect_pr is not None:
        body.append(copy.deepcopy(loaded.trailing_sect_pr))
    return body


def _self_check(loaded: _LoadedBody, redline: list, clean: list, moved_names) -> None:
    """D-7: the file must keep its own promise before anyone sees it."""
    from .revisions import (
        accept_all,
        duplicate_bookmark_names,
        first_difference,
        reject_all,
    )

    redline_body = _body_of(loaded, redline)
    upload_body = _body_of(loaded, [copy.deepcopy(c) for c in loaded.content])
    accepted = first_difference(accept_all(redline_body), _body_of(loaded, clean))
    if accepted is not None:
        raise SourceRedlineError(REDLINE_ACCEPT_CHECK, detail=accepted.to_dict())
    rejected = first_difference(
        reject_all(redline_body), upload_body, exclude_bookmarks=moved_names
    )
    if rejected is not None:
        raise SourceRedlineError(REDLINE_REJECT_CHECK, detail=rejected.to_dict())
    extra = duplicate_bookmark_names(redline_body) - duplicate_bookmark_names(
        upload_body
    )
    if extra:
        raise SourceRedlineError(
            REDLINE_BOOKMARK_CHECK, detail={"duplicates": len(extra)}
        )


def render_preserving_redline(
    *,
    source_bytes: bytes,
    format_map: SourceFormatMap,
    baseline: SpecSection,
    current: SpecSection,
    author: str,
    date: str,
    stats: dict | None = None,
) -> bytes:
    """The upload with every change since ``baseline`` as a Word tracked
    change: Accept All gives :func:`render_preserving_docx`'s output and
    Reject All gives the upload, and both are checked before the bytes are
    returned (:class:`SourceRedlineError` names why when they cannot be).

    Every part other than ``word/document.xml`` is the upload's, byte for
    byte — ``word/settings.xml`` included: Track Changes is not switched on
    in the file (Decision 4).
    """
    from .diffing import diff_sections
    from .revision_marks import RevisionMarks, highest_annotation_id
    from .source_audit import SourceAuditError, audit_package_preservation_streaming
    from .source_mapping import (
        PENDING_REVISIONS,
        detect_pending_revisions,
    )

    loaded = _load_body(source_bytes, format_map)
    pending = detect_pending_revisions(source_bytes)
    if pending:
        raise SourceRedlineError(
            REDLINE_PENDING_REVISIONS
            if pending == PENDING_REVISIONS
            else REDLINE_REVISION_SCAN
        )
    assembler = _plan_for(loaded, format_map, current)
    records, removed = assembler.plan()
    clean = [
        element
        for element in (assembler.render_accepted(r) for r in records)
        if element is not None
    ]
    diff = diff_sections(baseline, current, detect_moves=True)
    marks = RevisionMarks(
        author=author,
        date=date,
        first_id=highest_annotation_id(source_bytes) + 1,
    )
    builder = _RedlineBuilder(
        loaded, format_map, assembler, _subtree_uids(current, diff.moved or []), marks
    )
    try:
        ordered = builder.plan(records, removed)
        redline, moved_names = builder.render(ordered)
    except AssertionError as exc:  # a mark holding a section break, deleted
        raise SourceRedlineError(REDLINE_SECTION_BREAK) from exc
    if stats is not None:
        stats.update(assembler.stats)
        stats["redline"] = {**builder.stats, "revisions": marks.count}
    _self_check(loaded, redline, clean, moved_names)
    rebuilt = _serialize(loaded, redline)
    payload = replace_document_xml_raw(source_bytes, rebuilt)
    try:
        audit_package_preservation_streaming(
            source_bytes, payload, expected_document_xml=rebuilt
        )
    except SourceAuditError as exc:
        raise SourceRedlineError(
            REDLINE_PACKAGE_CHECK, detail={"blocker": exc.blocker}
        ) from exc
    return payload


__all__ = [
    "RECORD_KINDS",
    "SECTION_TITLE_UID",
    "SourceRedlineError",
    "SourceRenderError",
    "redline_refusal_message",
    "render_preserving_docx",
    "render_preserving_redline",
]
