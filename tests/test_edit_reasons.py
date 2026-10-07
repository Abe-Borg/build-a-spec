"""Every edit the software makes carries a brief reason (owner rule, 2026-10-07).

The reason rides each ``apply_spec_edits`` operation — a single added word
included — and the engine refuses a model batch without one on every
operation, the ``check_drafted_edits`` posture. It is kept per element on the
tree (``SpecSection.edit_reasons``: transactional, undoable, versioned,
persisted), shown by the panel's "why" chip and inline under a block changed
this turn, and said in a Word comment by the redline on the original — which
used to comment only a change with a research, attached-document or Final QC
basis, so the user saw comments on some changes and none on the rest.
"""
from __future__ import annotations

import json
from types import SimpleNamespace

import pytest
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from backend import sessions, templates
from backend.app import create_app
from backend.llm.prompts import render_system_prompt
from backend.qc.apply import ops_with_fix_reasons, qc_fix_reason
from backend.qc.engine import _op_vocabulary
from backend.redline_basis import redline_comment_basis
from backend.spec_doc.model import (
    APPLY_SPEC_EDITS_TOOL,
    EDIT_REASON_MAX_CHARS,
    MAX_EDIT_REASONS_PER_ELEMENT,
    DocumentStore,
    SpecEditError,
    SpecSection,
    apply_edits,
    check_edit_reasons,
    edit_reason_problems,
    fold_reason,
)
from backend.spec_doc.redline_comments import CommentBasis
from backend.spec_doc.source_format import SECTION_TITLE_UID
from backend.spec_modules.generic import GENERIC
from backend.spec_modules.hyperscale_fire import HYPERSCALE_FIRE
from tests.fakes import (
    FAKE_EDIT_REASON,
    FakeClient,
    audit_grade_qc_result,
    text_turn,
    tool_turn,
    with_edit_reasons,
)
from tests.test_preserving_export import _master_bytes, _parse
from tests.test_qc_chat_apply import _SEED_OPS, _finding
from tests.test_redline_comments import _comments_root, _fix_entry, _render
from tests.test_redline_original import AUTHOR, _edit

# ---------------------------------------------------------------------------
# The tree
# ---------------------------------------------------------------------------

_ARTICLE = {
    "action": "add_article",
    "target_id": "pt1",
    "text": "SUMMARY",
    "reason": "The user asked for a SUMMARY article first.",
}
_PROVISION = {
    "action": "add_paragraph",
    "target_id": "pt1.a1",
    "text": "Provide a complete wet-pipe sprinkler system.",
    "status": "confirmed",
    "reason": "Scope the user stated.",
}


def _store(*ops: dict) -> DocumentStore:
    store = DocumentStore()
    store.begin_turn()
    store.apply_edits(list(ops))
    store.commit_turn()
    return store


def test_every_content_op_records_its_reason_on_its_element():
    store = _store(_ARTICLE, _PROVISION)
    store.begin_turn()
    applied = store.apply_edits(
        [
            {
                "action": "add_paragraph",
                "target_id": "pt1.a1.p1",
                "text": "Include the fire pump.",
                "reason": "  Pump   is in the user's   scope  ",
            },
            {
                "action": "replace",
                "target_id": "pt1.a1.p1",
                "text": "Provide a complete wet-pipe sprinkler system throughout.",
                "reason": "Added 'throughout': the user confirmed full coverage.",
            },
            {
                "action": "set_status",
                "target_id": "pt1.a1.p1.p1",
                "status": "confirmed",
                "reason": "User confirmed the pump.",
            },
            {
                "action": "add_article",
                "target_id": "pt1",
                "text": "REFERENCES",
                "position": 0,
                "reason": "REFERENCES leads PART 1.",
            },
            {
                "action": "move",
                "target_id": "pt1.a1",
                "position": 0,
                "reason": "SUMMARY reads first in SectionFormat.",
            },
            {
                "action": "replace",
                "target_id": "sec",
                "text": "WET-PIPE SPRINKLER SYSTEMS",
                "numbering": "21 13 13",
                "reason": "Section chosen by the user.",
            },
        ]
    )
    store.commit_turn()
    # Every applied record echoes its reason, folded to one line.
    assert [op["reason"] for op in applied] == [
        "Pump is in the user's scope",
        "Added 'throughout': the user confirmed full coverage.",
        "User confirmed the pump.",
        "REFERENCES leads PART 1.",
        "SUMMARY reads first in SectionFormat.",
        "Section chosen by the user.",
    ]
    assert store.doc.edit_reasons == {
        "pt1.a1": [
            "The user asked for a SUMMARY article first.",
            "SUMMARY reads first in SectionFormat.",
        ],
        "pt1.a1.p1": [
            "Scope the user stated.",
            "Added 'throughout': the user confirmed full coverage.",
        ],
        "pt1.a1.p1.p1": ["Pump is in the user's scope", "User confirmed the pump."],
        "pt1.a2": ["REFERENCES leads PART 1."],
        "sec": ["Section chosen by the user."],
    }
    # The reason never reaches the document's words.
    assert "scope" not in store.doc.parts[0].articles[0].title.lower()
    assert all(
        "reason" not in p.text.lower()
        for p in store.doc.parts[0].articles[0].paragraphs
    )


def test_a_deletion_keeps_the_reason_for_the_element_and_everything_under_it():
    store = _store(
        _ARTICLE,
        _PROVISION,
        {
            "action": "add_paragraph",
            "target_id": "pt1.a1.p1",
            "text": "Nested detail.",
            "reason": "Detail.",
        },
    )
    store.begin_turn()
    (applied,) = store.apply_edits(
        [
            {
                "action": "delete",
                "target_id": "pt1.a1",
                "reason": "This project has no SUMMARY article in its office format.",
            }
        ]
    )
    store.commit_turn()
    assert applied == {
        "action": "delete",
        "id": "pt1.a1",
        "reason": "This project has no SUMMARY article in its office format.",
    }
    # uids are never reused, so the entries can never name a later element;
    # the redline reads them for its comment on each deleted provision.
    assert store.doc.parts[0].articles == []
    for uid in ("pt1.a1", "pt1.a1.p1", "pt1.a1.p1.p1"):
        assert store.doc.edit_reasons[uid][-1] == (
            "This project has no SUMMARY article in its office format."
        )


def test_the_trail_drops_a_repeat_and_keeps_the_newest_few():
    store = _store(_ARTICLE, _PROVISION)
    store.begin_turn()
    for index in range(MAX_EDIT_REASONS_PER_ELEMENT + 3):
        store.apply_edits(
            [
                {
                    "action": "set_status",
                    "target_id": "pt1.a1.p1",
                    "status": "assumed" if index % 2 else "confirmed",
                    "reason": f"Round {index}.",
                },
                # The same reason twice in a row is one entry.
                {
                    "action": "set_status",
                    "target_id": "pt1.a1.p1",
                    "status": "assumed" if index % 2 else "confirmed",
                    "reason": f"Round {index}.",
                },
            ]
        )
    store.commit_turn()
    trail = store.doc.edit_reasons["pt1.a1.p1"]
    assert len(trail) == MAX_EDIT_REASONS_PER_ELEMENT
    assert trail[-1] == f"Round {MAX_EDIT_REASONS_PER_ELEMENT + 2}."
    assert "Scope the user stated." not in trail  # the oldest went first


def test_a_long_reason_is_cut_at_a_word_and_never_refused():
    long = "because " * 80
    folded = fold_reason(long)
    assert len(folded) <= EDIT_REASON_MAX_CHARS
    assert folded.endswith("…") and not folded.endswith(" …")
    assert fold_reason("  one   line  ") == "one line"
    assert fold_reason("") == ""
    section, (applied,) = apply_edits(
        SpecSection.empty(),
        [{**_ARTICLE, "reason": long}],
    )
    assert applied["reason"] == folded
    assert section.edit_reasons["pt1.a1"] == [folded]


def test_reasons_ride_undo_redo_and_serialization_and_legacy_bytes_are_untouched():
    # No model edit: the snapshot has no edit_reasons key at all.
    assert "edit_reasons" not in SpecSection.empty().to_dict()
    assert "edit_reasons" not in _store(
        {"action": "add_article", "target_id": "pt1", "text": "SUMMARY"}
    ).doc.to_dict()

    store = _store(_ARTICLE, _PROVISION)
    store.begin_turn()
    store.apply_edits(
        [{"action": "set_status", "target_id": "pt1.a1.p1", "status": "assumed",
          "reason": "Confirmed later."}]
    )
    store.commit_turn()
    assert store.doc.edit_reasons["pt1.a1.p1"][-1] == "Confirmed later."
    store.undo()
    assert store.doc.edit_reasons["pt1.a1.p1"] == ["Scope the user stated."]
    store.redo()
    assert store.doc.edit_reasons["pt1.a1.p1"][-1] == "Confirmed later."

    data = store.doc.to_dict()
    assert data["edit_reasons"] == store.doc.edit_reasons
    assert SpecSection.from_dict(json.loads(json.dumps(data))).edit_reasons == (
        store.doc.edit_reasons
    )
    restored = DocumentStore()
    restored.load(json.loads(json.dumps(store.to_dict())))
    assert restored.doc.edit_reasons == store.doc.edit_reasons
    assert "edit_reasons" in store.snapshot()


@pytest.mark.parametrize(
    "bad",
    [
        ["pt1.a1"],
        {"": ["x"]},
        {"pt1.a1": "not a list"},
        {"pt1.a1": []},
        {"pt1.a1": ["ok", ""]},
        {"pt1.a1": [1]},
    ],
)
def test_load_rejects_malformed_edit_reasons(bad):
    snapshot = SpecSection.empty().to_dict()
    snapshot["edit_reasons"] = bad
    with pytest.raises(ValueError, match="edit_reasons"):
        SpecSection.from_dict(snapshot)
    assert SpecSection.from_dict({**snapshot, "edit_reasons": None}).edit_reasons == {}


def test_metadata_ops_echo_the_reason_and_store_it_only_where_a_basis_lives():
    section, applied = apply_edits(
        SpecSection.empty(),
        [
            {
                "action": "set_project_profile",
                "target_id": "sec",
                "city": "Ashburn",
                "reason": "The user named the city.",
            },
            {
                "action": "set_project_identity",
                "target_id": "sec",
                "project_type": "Data Center",
                "reason": "The user described a hyperscale campus.",
            },
            # set_standard_edition: the required basis is the reason.
            {
                "action": "set_standard_edition",
                "target_id": "sec",
                "standard": "NFPA 13",
                "edition": "2019",
                "basis": "2021 VCC per user",
            },
            # set_standard_suppressed: no basis given, so the reason is kept
            # as the exclusion's basis — the context block says why.
            {
                "action": "set_standard_suppressed",
                "target_id": "sec",
                "standard": "NFPA 2001",
                "suppressed": True,
                "reason": "No clean-agent system on this project.",
            },
        ],
    )
    assert applied[0]["reason"] == "The user named the city."
    assert applied[1]["reason"] == "The user described a hyperscale campus."
    assert "reason" not in applied[2]
    assert applied[3]["reason"] == "No clean-agent system on this project."
    assert section.suppressed_standards == {
        "NFPA 2001": "No clean-agent system on this project."
    }
    # Section metadata is not an element: nothing to pin a chip to.
    assert section.edit_reasons == {}
    assert "edit_reasons" not in section.to_dict()


def test_a_non_string_reason_is_refused_and_move_accepts_one():
    with pytest.raises(SpecEditError, match="'reason' must be a string"):
        apply_edits(SpecSection.empty(), [{**_ARTICLE, "reason": 7}])
    section, _ = apply_edits(
        SpecSection.empty(),
        [_ARTICLE, {**_ARTICLE, "text": "REFERENCES", "reason": "Second."}],
    )
    moved, (applied,) = apply_edits(
        section,
        [{"action": "move", "target_id": "pt1.a2", "position": 0, "reason": "Order."}],
    )
    assert applied["reason"] == "Order."
    assert moved.edit_reasons["pt1.a2"] == ["Second.", "Order."]
    with pytest.raises(SpecEditError, match="unsupported field"):
        apply_edits(
            section,
            [{"action": "move", "target_id": "pt1.a2", "position": 0, "text": "x",
              "reason": "Order."}],
        )


def test_a_hand_rewrite_move_or_delete_drops_the_trail_but_a_confirm_keeps_it():
    """The user's own panel edit carries no reason. Where it rewrites,
    moves or deletes what the redline shows, the assistant's reasons would
    be stale — in the chip and in the comment — so they go; a status or
    source change keeps them, because the words are still the assistant's."""
    section, _ = apply_edits(
        SpecSection.empty(),
        [
            _ARTICLE,
            _PROVISION,
            {**_ARTICLE, "text": "REFERENCES", "reason": "Second article."},
            {
                "action": "replace",
                "target_id": "sec",
                "text": "WET-PIPE SPRINKLER SYSTEMS",
                "reason": "Section chosen.",
            },
        ],
    )
    confirmed, _ = apply_edits(
        section,
        [
            {"action": "set_status", "target_id": "pt1.a1.p1", "status": "confirmed"},
            {"action": "replace", "target_id": "pt1.a1.p1", "source_item_id": "pf-1"},
        ],
    )
    assert confirmed.edit_reasons["pt1.a1.p1"] == ["Scope the user stated."]
    retyped, _ = apply_edits(
        confirmed,
        [
            {"action": "replace", "target_id": "pt1.a1.p1", "text": "My own words."},
            {"action": "replace", "target_id": "pt1.a1", "text": "GENERAL"},
            {"action": "move", "target_id": "pt1.a2", "position": 0},
            {"action": "replace", "target_id": "sec", "numbering": "21 13 16"},
        ],
    )
    assert retyped.edit_reasons == {}
    deleted, _ = apply_edits(section, [{"action": "delete", "target_id": "pt1.a1"}])
    assert set(deleted.edit_reasons) == {"pt1.a2", "sec"}
    # The redline then has no reason to say for the user's own change (the
    # retyped provision still names a source, so its basis entry survives
    # as an unresolved source with no reasons).
    assert redline_comment_basis(retyped).get("pt1.a1.p1", CommentBasis()).reasons == ()
    assert "pt1.a1" not in redline_comment_basis(deleted)


# ---------------------------------------------------------------------------
# The contract
# ---------------------------------------------------------------------------


def test_check_edit_reasons_names_each_operation_without_one():
    edits = [
        {"action": "add_article", "target_id": "pt1", "text": "SUMMARY"},
        {**_PROVISION},
        {"action": "move", "target_id": "pt1.a1", "position": 0, "reason": "   "},
        # The adoption basis IS the reason.
        {
            "action": "set_standard_edition",
            "target_id": "sec",
            "standard": "NFPA 13",
            "edition": "2019",
            "basis": "2021 VCC per user",
        },
        # Without a basis, the op needs its reason like any other.
        {"action": "set_standard_edition", "target_id": "sec", "standard": "NFPA 72"},
        # Shapes apply_edits refuses are left to it.
        "not an op",
        {"action": "explode", "target_id": "pt1"},
        {"action": "delete", "reason": "Gone."},
    ]
    assert edit_reason_problems(edits) == [
        "#1 add_article pt1",
        "#3 move pt1.a1",
        "#5 set_standard_edition sec",
    ]
    with pytest.raises(SpecEditError) as excinfo:
        check_edit_reasons(edits)
    message = str(excinfo.value)
    assert "Every operation needs a 'reason'" in message
    assert "#1 add_article pt1; #3 move pt1.a1; #5 set_standard_edition sec" in message
    assert "single added word" in message
    # Nothing to report: a complete batch, or a shape that is not a batch.
    check_edit_reasons([_ARTICLE, _PROVISION])
    assert edit_reason_problems(None) == [] and edit_reason_problems({}) == []


@pytest.mark.parametrize("module", [GENERIC, HYPERSCALE_FIRE])
def test_the_tool_declares_the_reason_and_the_prompt_teaches_it(module):
    item = APPLY_SPEC_EDITS_TOOL["input_schema"]["properties"]["edits"]["items"]
    reason = item["properties"]["reason"]
    assert reason["type"] == "string"
    assert reason["description"].startswith("Required on every operation")
    assert "single added word" in reason["description"]
    assert "never goes into the specification text" in reason["description"]
    # Stated, not in 'required': saved histories carry inputs without it
    # (the status-enum precedent in the tool definition's comment).
    assert item["required"] == ["action", "target_id"]
    # The top-level description is what Final QC's lenses read as their op
    # vocabulary, and their strict schema has no reason field — so the rule
    # lives in the property and the prompt, and retained reports stay current.
    assert _op_vocabulary() == APPLY_SPEC_EDITS_TOOL["description"]
    assert "Required on every operation" not in APPLY_SPEC_EDITS_TOOL["description"]
    prompt = render_system_prompt(module)
    assert "Every operation carries a reason" in prompt
    assert "a single added word included" in prompt
    assert "rejected whole" in prompt


# ---------------------------------------------------------------------------
# The chat
# ---------------------------------------------------------------------------


def _client() -> TestClient:
    return TestClient(create_app())


def _parse_sse(body: str) -> list[dict]:
    return [
        json.loads(line[len("data: "):])
        for line in body.splitlines()
        if line.startswith("data: ")
    ]


def _patch_client(monkeypatch, fake: FakeClient) -> None:
    monkeypatch.setattr("backend.llm.conversation.get_client", lambda: fake)


def test_a_model_batch_without_reasons_is_refused_and_the_model_resends_with_them(
    monkeypatch,
):
    bare = {
        "edits": [
            {"action": "add_article", "target_id": "pt1", "text": "SUMMARY"},
            {
                "action": "add_paragraph",
                "target_id": "pt1.a1",
                "text": "Provide a complete wet-pipe sprinkler system.",
                "reason": "Scope the user stated.",
            },
        ]
    }
    fake = FakeClient(
        [
            tool_turn([], bare, tool_id="toolu_bare", reasons=False),
            tool_turn([], with_edit_reasons(bare), tool_id="toolu_good"),
            text_turn(["Drafted the summary. What hazard classification governs?"]),
        ]
    )
    _patch_client(monkeypatch, fake)
    client = _client()

    events = _parse_sse(client.post("/api/chat", json={"message": "go"}).text)
    assert events[-1]["type"] == "turn_complete"
    patches = [e for e in events if e["type"] == "doc_patch"]
    assert len(patches) == 1
    # The panel's copy of the ops carries each reason …
    assert [op["reason"] for op in patches[0]["ops"]] == [
        FAKE_EDIT_REASON,
        "Scope the user stated.",
    ]
    # … and so does every snapshot of the document.
    assert patches[0]["doc"]["edit_reasons"] == {
        "pt1.a1": [FAKE_EDIT_REASON],
        "pt1.a1.p1": ["Scope the user stated."],
    }
    snapshot = next(e for e in events if e["type"] == "doc_snapshot")
    assert snapshot["doc"]["edit_reasons"] == patches[0]["doc"]["edit_reasons"]
    assert client.get("/api/doc").json()["doc"]["edit_reasons"] == (
        patches[0]["doc"]["edit_reasons"]
    )

    history = sessions.get_session().history
    refusal = history[2]["content"][0]
    assert refusal["is_error"] is True
    assert "rejected (nothing was applied)" in refusal["content"]
    assert "Every operation needs a 'reason'" in refusal["content"]
    assert "#1 add_article pt1" in refusal["content"]
    assert "#2" not in refusal["content"].split("Add the reasons")[0]
    # The accepted batch's result echoes what was done, never the reasons
    # (the model's own input already carries them; history would hold
    # every reason twice and re-send it on every turn).
    accepted = json.loads(history[4]["content"][0]["content"])
    assert [op["id"] for op in accepted["applied"]] == ["pt1.a1", "pt1.a1.p1"]
    assert all("reason" not in op for op in accepted["applied"])
    assert history[3]["content"][-1]["input"]["edits"][1]["reason"] == (
        "Scope the user stated."
    )


def test_the_panels_own_edit_needs_no_reason_and_records_none():
    client = _client()
    response = client.post("/api/doc/edit", json={"ops": _SEED_OPS})
    assert response.status_code == 200, response.text
    data = response.json()
    assert all("reason" not in op for op in data["applied"])
    assert "edit_reasons" not in data["doc"]
    # A reason the panel chooses to send is kept like any other.
    response = client.post(
        "/api/doc/edit",
        json={
            "ops": [
                {
                    "action": "set_status",
                    "target_id": "pt1.a1.p1",
                    "status": "assumed",
                    "reason": "Reviewer set it back.",
                }
            ]
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["doc"]["edit_reasons"] == {
        "pt1.a1.p1": ["Reviewer set it back."]
    }


def test_reasons_survive_a_project_save_and_resume(monkeypatch):
    fake = FakeClient(
        [
            tool_turn([], {"edits": [_ARTICLE, _PROVISION]}, tool_id="toolu_seed"),
            text_turn(["Drafted."]),
        ]
    )
    _patch_client(monkeypatch, fake)
    client = _client()
    assert _parse_sse(client.post("/api/chat", json={"message": "go"}).text)[-1][
        "type"
    ] == "turn_complete"
    project = json.loads(json.dumps(sessions.project_payload(sessions.get_session())))
    client.post("/api/session/reset")
    assert sessions.get_session().doc.doc.edit_reasons == {}
    loaded = client.post("/api/project/load", json=project)
    assert loaded.status_code == 200, loaded.text
    assert loaded.json()["doc"]["edit_reasons"] == {
        "pt1.a1": ["The user asked for a SUMMARY article first."],
        "pt1.a1.p1": ["Scope the user stated."],
    }


# ---------------------------------------------------------------------------
# Final QC fixes
# ---------------------------------------------------------------------------


def test_ops_with_fix_reasons_names_the_finding_keeps_order_and_an_own_reason():
    first = {"action": "replace", "target_id": "pt1.a1.p1", "text": "Fixed."}
    second = {"action": "delete", "target_id": "pt1.a1.p2"}
    own = {"action": "set_status", "target_id": "pt1.a1.p3", "status": "confirmed",
           "reason": "Already stated."}
    result = SimpleNamespace(
        findings=[
            SimpleNamespace(finding_id="qc-1", title="Density  states\n0.20 gpm/sq ft"),
            SimpleNamespace(finding_id="qc-2", title="Dead provision"),
            SimpleNamespace(finding_id="qc-3", title=""),
        ]
    )
    eligible = [("qc-1", [first]), ("qc-2", [second, first]), ("qc-3", [own])]
    combined = [first, second, own]
    reasoned = ops_with_fix_reasons(combined, eligible, result)
    assert [op["reason"] for op in reasoned] == [
        "Final QC fix: Density states 0.20 gpm/sq ft",  # the first owner speaks
        "Final QC fix: Dead provision",
        "Already stated.",
    ]
    strip = lambda ops: [{k: v for k, v in op.items() if k != "reason"} for op in ops]  # noqa: E731
    assert strip(reasoned) == strip(combined)  # same ops, same order
    assert all("reason" not in op for op in (first, second))  # copies, not mutations
    # An op no finding owns, or whose finding has no title, still says what it is.
    assert ops_with_fix_reasons([{"action": "delete", "target_id": "x"}], [], result)[
        0
    ]["reason"] == "Final QC fix"


def test_a_final_qc_fix_applied_from_the_panel_carries_its_finding_as_the_reason():
    client = _client()
    assert client.post("/api/doc/edit", json={"ops": _SEED_OPS}).json()["ok"]
    session = sessions.get_session()
    finding = _finding(
        "qc-reason00001",
        [
            {
                "action": "replace",
                "target_id": "pt1.a1.p1",
                "text": "Provide a complete, hydraulically calculated wet-pipe system.",
                "status": "confirmed",
            }
        ],
    )
    result = audit_grade_qc_result(session, [finding])
    session.qc.result = result
    session.qc.status = "complete"
    session.qc.latest_attempt_run_id = result.run_id
    session.qc.latest_attempt_status = "complete"
    session.qc.latest_attempt_result = result
    applied = client.post("/api/qc/apply", json={"finding_ids": [finding.finding_id]})
    assert applied.status_code == 200, applied.text
    assert applied.json()["outcomes"] == {finding.finding_id: "applied"}
    doc = client.get("/api/doc").json()["doc"]
    assert doc["edit_reasons"] == {"pt1.a1.p1": ["Final QC fix: Finding qc-reason00001"]}
    # The fix's survival record still keys on action + id, untouched by the
    # reason riding the echo.
    assert session.qc_fix_log and session.qc_fix_log[-1]["finding_id"] == finding.finding_id


# ---------------------------------------------------------------------------
# The redline on the original
# ---------------------------------------------------------------------------


def _comment_lines(payload: bytes) -> list[list[str]]:
    """Each comment the render helper's test author wrote, as plain lines."""
    root = _comments_root(payload)
    if root is None:
        return []
    return [
        ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in c.iter(qn("w:p"))]
        for c in root.iter(qn("w:comment"))
        if c.get(qn("w:author")) == AUTHOR
    ]


def test_the_wording_is_one_line_or_a_numbered_trail_and_the_header_speaks_twice():
    section, _ = apply_edits(
        SpecSection.empty(),
        [
            _ARTICLE,
            _PROVISION,
            {
                "action": "replace",
                "target_id": "pt1.a1.p1",
                "text": "Provide a complete wet-pipe sprinkler system throughout.",
                "reason": "Added 'throughout' at the user's request.",
            },
            {
                "action": "replace",
                "target_id": "sec",
                "text": "WET-PIPE SPRINKLER SYSTEMS",
                "reason": "Section chosen.",
            },
        ],
    )
    bases = redline_comment_basis(section)
    assert bases["pt1.a1"] == CommentBasis(
        reasons=(("Reason: The user asked for a SUMMARY article first.",),)
    )
    assert bases["pt1.a1.p1"].reasons == (
        ("Reasons, oldest first:",),
        ("1. Scope the user stated.",),
        ("2. Added 'throughout' at the user's request.",),
    )
    assert bases["sec"].reasons == bases[SECTION_TITLE_UID].reasons == (
        ("Reason: Section chosen.",),
    )
    # No reasons: nothing changes for a tree the model never touched.
    assert redline_comment_basis(SpecSection.empty()) == {}


def test_a_qc_fix_reason_is_not_said_beside_the_record_of_that_fix():
    section, _ = apply_edits(
        SpecSection.empty(),
        [
            _ARTICLE,
            _PROVISION,
            {
                "action": "replace",
                "target_id": "pt1.a1.p1",
                "text": "Provide a complete, hydraulically calculated wet-pipe system.",
                "reason": qc_fix_reason("Density  not\nstated"),
            },
        ],
    )
    entry = _fix_entry(
        "pt1.a1.p1",
        "Provide a complete, hydraulically calculated wet-pipe system.",
        finding_id="qc-density",
        title="Density not stated",
    )
    basis = redline_comment_basis(section, fix_log=[entry])["pt1.a1.p1"]
    # The record speaks for the fix; the trail's own line for it is not
    # repeated, while the drafting reason before it still is.
    assert basis.qc
    assert basis.reasons == (("Reason: Scope the user stated.",),)
    # Without the record (undo, a hand edit), the fix reason is said.
    assert redline_comment_basis(section)["pt1.a1.p1"].reasons == (
        ("Reasons, oldest first:",),
        ("1. Scope the user stated.",),
        ("2. Final QC fix: Density not stated",),
    )


def test_a_reworded_provision_with_a_reason_gets_a_comment_where_it_had_none(tmp_path):
    source = _master_bytes()
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    uid = article.paragraphs[0].uid
    reason = "Reworded so the provision directs the Contractor instead of describing."
    section = _edit(
        imported.section,
        {"action": "replace", "target_id": uid, "text": "Reworded.", "reason": reason},
    )
    # The same edit without a reason is still the no-comment case the old
    # tests pin; with one, the change says why.
    payload, stats = _render(source, imported, section, redline_comment_basis(section))
    assert stats["redline"]["comments"]["added"] == 1
    assert stats["redline"]["comments"]["skipped"] == {}
    assert _comment_lines(payload) == [[f"Reason: {reason}"]]


def test_a_deleted_article_speaks_once_for_every_provision_it_took(tmp_path):
    source = _master_bytes()
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    assert len(article.paragraphs) >= 1
    reason = "This office format carries the scope statement in Division 01."
    section = _edit(
        imported.section,
        {"action": "delete", "target_id": article.uid, "reason": reason},
    )
    bases = redline_comment_basis(section)
    assert bases[article.uid].reasons == ((f"Reason: {reason}",),)
    for paragraph in article.paragraphs:
        assert bases[paragraph.uid].reasons == ((f"Reason: {reason}",),)
    payload, stats = _render(source, imported, section, bases)
    assert stats["redline"]["comments"]["added"] >= 1
    assert stats["redline"]["comments"]["skipped"] == {}
    lines = _comment_lines(payload)
    assert lines and all(comment == [f"Reason: {reason}"] for comment in lines)


def test_a_pure_move_with_a_reason_is_commented(tmp_path):
    source = _master_bytes()
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    if len(article.paragraphs) < 2:
        pytest.skip("the fixture master needs two provisions to reorder")
    reason = "Submittals read before quality assurance in the office order."
    section = _edit(
        imported.section,
        {
            "action": "move",
            "target_id": article.paragraphs[-1].uid,
            "position": 0,
            "reason": reason,
        },
    )
    payload, stats = _render(
        source, imported, section, redline_comment_basis(section), native_moves=True
    )
    assert stats["redline"]["comments"]["added"] >= 1
    assert [f"Reason: {reason}"] in _comment_lines(payload)


# ---------------------------------------------------------------------------
# Templates
# ---------------------------------------------------------------------------


def test_a_template_starter_carries_no_other_projects_reasons():
    section, _ = apply_edits(SpecSection.empty(), [_ARTICLE, _PROVISION])
    assert section.edit_reasons
    canonical = templates._canonical_document(section.to_dict(), rebase_statuses=True)
    assert "edit_reasons" not in canonical
    assert SpecSection.from_dict(canonical).edit_reasons == {}
