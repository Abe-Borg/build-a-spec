"""The fifth SectionFormat paragraph level: ``a)`` under ``1)``.

CSI SectionFormat nests a provision five levels deep under its article —
``A.`` / ``1.`` / ``a.`` / ``1)`` / ``a)`` — and a MasterSpec-derived office
master carries the fifth on its ``PR5`` paragraph style. The document model
stopped at four, and a real master showed what that cost (reported
2026-09-23 on a WATER BASED FIRE PROTECTION SYSTEMS master):

* Word-numbered, the fifth level was CLAMPED into the fourth. Every ``a)``
  line warned "nesting deeper than 4 levels — clamped to level 4" and became
  a sibling of the ``1)`` it belonged to, numbered ``2)``, ``3)`` … in the
  panel and in everything the model read.
* Typed, the ``a)`` label was not recognized at all. The lines fell through
  to the unlabelled branch: top-level provisions with "a)" left in their
  text, which relettered every provision after them (B. read as E.) and hung
  the next ``2)`` under the wrong parent — and the formatted export of an
  UNTOUCHED import wrote those new letters into the user's file.

These tests pin the fifth level end to end: the model's labels and limit, the
importer's three ways of reading it (a typed label, Word numbering on the
paragraph or its style, the ``PR5`` style name alone), both Word exports, the
redline on the original, and the import route. A sixth level is still beyond
SectionFormat and still clamps, with the warning saying where.
"""
from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from lxml import etree

from backend import sessions
from backend.app import create_app
from backend.spec_doc.diffing import diff_sections
from backend.spec_doc.docx_export import build_docx
from backend.spec_doc.importer import _accept_all_paragraph_text
from backend.spec_doc.model import (
    APPLY_SPEC_EDITS_TOOL,
    MAX_PARAGRAPH_DEPTH,
    DocumentStore,
    SpecEditError,
    SpecSection,
    _paragraph_label,
    apply_edits,
    iter_paragraphs,
    outline,
)
from backend.spec_doc.source_render import render_preserving_docx
from tests.docx_fidelity_helpers import DOCX_MEDIA_TYPE
from tests.test_import_office_master import _add_style
from tests.test_importer import _define_numbering, _numbered
from tests.test_preserving_export import _parse, _save
from tests.test_redline_original import _edit, _shape, _verify

# ilvl -> (numFmt, lvlText): the MasterSpec outline, every level of it. PART
# renders "PART 1 - ", the article "1.01", then the five provision levels.
_MASTERSPEC_LEVELS = {
    0: ("decimal", "PART %1 - "),
    1: ("decimalZero", "%1.%2"),
    2: ("upperLetter", "%3."),
    3: ("decimal", "%4."),
    4: ("lowerLetter", "%5."),
    5: ("decimal", "%6)"),
    6: ("lowerLetter", "%7)"),
}
_MASTERSPEC_STYLES = (
    ("PRT", 0),
    ("ART", 1),
    ("PR1", 2),
    ("PR2", 3),
    ("PR3", 4),
    ("PR4", 5),
    ("PR5", 6),
)

#: The reported shape: a fifth level under a 1) in PART 2, and a 1) sibling
#: after it that must stay the 1)'s sibling.
_SPRINKLER_OUTLINE = (
    ("GENERAL", "PRT"),
    ("SUMMARY", "ART"),
    ("Section includes wet-pipe sprinkler systems.", "PR1"),
    ("PRODUCTS", "PRT"),
    ("SPRINKLERS", "ART"),
    ("Sprinklers:", "PR1"),
    ("Quick-response, standard coverage:", "PR2"),
    ("Pendent:", "PR3"),
    ("Finish:", "PR4"),
    ("Chrome plated.", "PR5"),
    ("White polyester.", "PR5"),
    ("Brass.", "PR5"),
    ("Escutcheons: Recessed.", "PR4"),
    ("Spare sprinklers: Provide a cabinet.", "PR1"),
    ("EXECUTION", "PRT"),
    ("INSTALLATION", "ART"),
    ("Install sprinklers per NFPA 13.", "PR1"),
)

_TYPED_OUTLINE = (
    "A.\tSprinklers:",
    "1.\tQuick-response, standard coverage:",
    "a.\tPendent:",
    "1)\tFinish:",
    "a)\tChrome plated.",
    "b)\tWhite polyester.",
    "c)\tBrass.",
    "2)\tEscutcheons: Recessed.",
    "B.\tSpare sprinklers: Provide a cabinet.",
)


def _style_master(
    lines=_SPRINKLER_OUTLINE, *, numbered: bool = True, levels=None
) -> bytes:
    """The office-master shape: the outline lives on PRT/ART/PR1..PR5 styles
    (``numbered``), or on nothing but the style NAMES (not ``numbered``)."""
    document = Document()
    if numbered:
        _define_numbering(document, 70, levels or _MASTERSPEC_LEVELS)
    body = _add_style(document, "SpecBody", num_id=None, ilvl=None)
    styles = {
        name: _add_style(
            document,
            name,
            num_id=70 if numbered else None,
            ilvl=ilvl if numbered else None,
            based_on=body,
        )
        for name, ilvl in _MASTERSPEC_STYLES
    }
    for line in ("SECTION 21 13 13", "WET-PIPE SPRINKLER SYSTEMS"):
        document.add_paragraph(line)
    for text, style in lines:
        document.add_paragraph(text, style=styles[style])
    document.add_paragraph("END OF SECTION")
    return _save(document)


def _typed_master(lines=_TYPED_OUTLINE) -> bytes:
    """Typed labels with a TAB after each — no Word numbering anywhere."""
    document = Document()
    for line in (
        "SECTION 21 13 13",
        "WET-PIPE SPRINKLER SYSTEMS",
        "PART 2 - PRODUCTS",
        "2.1\tSPRINKLERS",
        *lines,
        "END OF SECTION",
    ):
        document.add_paragraph(line)
    return _save(document)


def _five_levels() -> list[dict]:
    """Ops building A. / 1. / a. / 1) / a) + b) under one article."""
    return [
        {"action": "add_article", "target_id": "pt2", "text": "SPRINKLERS"},
        {"action": "add_paragraph", "target_id": "pt2.a1", "text": "Sprinklers:"},
        {"action": "add_paragraph", "target_id": "pt2.a1.p1", "text": "Quick response:"},
        {"action": "add_paragraph", "target_id": "pt2.a1.p1.p1", "text": "Pendent:"},
        {"action": "add_paragraph", "target_id": "pt2.a1.p1.p1.p1", "text": "Finish:"},
        {
            "action": "add_paragraph",
            "target_id": "pt2.a1.p1.p1.p1.p1",
            "text": "Chrome plated.",
        },
        {
            "action": "add_paragraph",
            "target_id": "pt2.a1.p1.p1.p1.p1",
            "text": "White polyester.",
        },
    ]


def _refs(section: SpecSection) -> dict[str, str]:
    return {p.text: ref for _part, _article, p, _depth, ref in iter_paragraphs(section)}


def _depth_of(section: SpecSection) -> dict[str, int]:
    return {p.text: depth for _part, _article, p, depth, _ref in iter_paragraphs(section)}


def _body_children(payload: bytes) -> list:
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        xml = archive.read("word/document.xml")
    body = etree.fromstring(xml).find(qn("w:body"))
    return [child for child in body.iterchildren() if isinstance(child.tag, str)]


def _body_texts(payload: bytes) -> list[str]:
    return [
        paragraph.text
        for paragraph in Document(io.BytesIO(payload)).paragraphs
        if paragraph.text.strip()
    ]


# ---------------------------------------------------------------------------
# The model
# ---------------------------------------------------------------------------


def test_the_fifth_level_is_lettered_with_a_closing_parenthesis():
    section, _ = apply_edits(SpecSection.empty(), _five_levels())
    finish = section.to_dict()["parts"][1]["articles"][0]["paragraphs"][0][
        "children"
    ][0]["children"][0]["children"][0]
    assert finish["label"] == "1)"
    assert [child["label"] for child in finish["children"]] == ["a)", "b)"]
    assert _refs(section)["White polyester."] == "2.1.A.1.a.1.b"
    assert (
        "\n            a) (assumed) Chrome plated.  [id: pt2.a1.p1.p1.p1.p1.p1]\n"
        in outline(section)
    )


def test_the_fifth_levels_letters_run_past_z_like_the_third_levels():
    assert _paragraph_label(4, 0) == "a)"
    assert _paragraph_label(4, 25) == "z)"
    assert _paragraph_label(4, 26) == "aa)"
    # The third level's sequence, with its own closing punctuation.
    assert _paragraph_label(2, 26) == "aa."
    assert _paragraph_label(3, 1) == "2)"


def test_a_sixth_level_is_refused_by_name():
    section, _ = apply_edits(SpecSection.empty(), _five_levels())
    assert MAX_PARAGRAPH_DEPTH == 5
    with pytest.raises(SpecEditError, match=r"5 levels: A\. / 1\. / a\. / 1\) / a\)"):
        apply_edits(
            section,
            [
                {
                    "action": "add_paragraph",
                    "target_id": "pt2.a1.p1.p1.p1.p1.p1",
                    "text": "Too deep.",
                }
            ],
        )


def test_a_saved_five_level_document_loads_and_a_six_level_one_does_not():
    section, _ = apply_edits(SpecSection.empty(), _five_levels())
    store = DocumentStore()
    store.load({"versions": [section.to_dict()], "index": 0})
    assert _depth_of(store.doc)["White polyester."] == 4

    forged = section.to_dict()
    deepest = forged["parts"][1]["articles"][0]["paragraphs"][0]
    for _ in range(4):
        deepest = deepest["children"][0]
    deepest["children"] = [
        {
            "id": f"{deepest['id']}.p1",
            "label": "",
            "text": "A sixth level.",
            "status": "assumed",
            "children": [],
            "seq": 1,
        }
    ]
    deepest["seq"] = 2
    with pytest.raises(ValueError, match="exceeds 5 levels"):
        DocumentStore().load({"versions": [forged], "index": 0})


def test_the_edit_tool_states_the_models_own_limit():
    """The tool description is what the model reads; it must name the limit
    the model enforces, not a number that drifted from it."""
    description = APPLY_SPEC_EDITS_TOOL["description"]
    assert f"max {MAX_PARAGRAPH_DEPTH} levels" in description
    assert "1), a))" in description


# ---------------------------------------------------------------------------
# The importer
# ---------------------------------------------------------------------------


def test_a_pr5_level_nests_under_its_pr4_parent(tmp_path):
    """The reported master: numbering on PRT/ART/PR1..PR5 styles."""
    result = _parse(tmp_path, _style_master())
    assert not any("nesting deeper" in w for w in result.warnings), result.warnings
    article = result.section.parts[1].articles[0]
    pendent = article.paragraphs[0].children[0].children[0]
    assert [p.text for p in pendent.children] == [
        "Finish:",
        "Escutcheons: Recessed.",
    ]
    assert [p.text for p in pendent.children[0].children] == [
        "Chrome plated.",
        "White polyester.",
        "Brass.",
    ]
    refs = _refs(result.section)
    assert refs["Brass."] == "2.1.A.1.a.1.c"
    assert refs["Escutcheons: Recessed."] == "2.1.A.1.a.2"
    assert refs["Spare sprinklers: Provide a cabinet."] == "2.1.B"


def test_typed_fifth_level_labels_nest_and_are_stripped(tmp_path):
    """``a)`` is a label: stripped from the text, placed at the fifth level,
    and the provisions after it keep their own letters."""
    result = _parse(tmp_path, _typed_master())
    assert not any("jumped deeper" in w for w in result.warnings), result.warnings
    refs = _refs(result.section)
    assert refs["Chrome plated."] == "2.1.A.1.a.1.a"
    assert refs["Brass."] == "2.1.A.1.a.1.c"
    assert refs["Escutcheons: Recessed."] == "2.1.A.1.a.2"
    assert refs["Spare sprinklers: Provide a cabinet."] == "2.1.B"
    assert not any(text.startswith(("a)", "b)", "c)")) for text in refs)


def test_direct_word_numbering_reaches_the_fifth_level(tmp_path):
    document = Document()
    for line in (
        "SECTION 21 13 13",
        "WET-PIPE SPRINKLER SYSTEMS",
        "PART 2 - PRODUCTS",
        "2.1 SPRINKLERS",
    ):
        document.add_paragraph(line)
    _define_numbering(
        document,
        50,
        {
            0: ("upperLetter", "%1."),
            1: ("decimal", "%2."),
            2: ("lowerLetter", "%3."),
            3: ("decimal", "%4)"),
            4: ("lowerLetter", "%5)"),
        },
    )
    for ilvl, text in enumerate(
        ("Sprinklers:", "Quick response:", "Pendent:", "Finish:", "Chrome plated.")
    ):
        _numbered(document, text, ilvl, "50")
    _numbered(document, "White polyester.", 4, "50")
    result = _parse(tmp_path, _save(document))
    assert not any("nesting deeper" in w for w in result.warnings), result.warnings
    depths = _depth_of(result.section)
    assert depths["Chrome plated."] == depths["White polyester."] == 4
    assert _refs(result.section)["White polyester."] == "2.1.A.1.a.1.b"


def test_the_pr5_style_name_alone_places_the_fifth_level(tmp_path):
    """No numbering anywhere, no typed label: the CSI style NAME is the
    secondary signal, and PR5 is the fifth level — not a second PR4."""
    result = _parse(tmp_path, _style_master(numbered=False))
    depths = _depth_of(result.section)
    assert depths["Finish:"] == 3
    assert depths["Chrome plated."] == depths["Brass."] == 4
    assert _refs(result.section)["Escutcheons: Recessed."] == "2.1.A.1.a.2"


def test_a_sixth_level_still_clamps_to_the_fifth_and_says_where(tmp_path):
    """Past SectionFormat's five levels the importer still keeps everything
    and warns — the warning names the level it clamped to and where."""
    document = Document()
    for line in (
        "SECTION 21 13 13",
        "WET-PIPE SPRINKLER SYSTEMS",
        "PART 2 - PRODUCTS",
        "2.1 SPRINKLERS",
    ):
        document.add_paragraph(line)
    _define_numbering(
        document,
        50,
        {
            0: ("upperLetter", "%1."),
            1: ("decimal", "%2."),
            2: ("lowerLetter", "%3."),
            3: ("decimal", "%4)"),
            4: ("lowerLetter", "%5)"),
            5: ("lowerRoman", "(%6)"),
        },
    )
    for ilvl, text in enumerate(
        ("Sprinklers:", "Quick response:", "Pendent:", "Finish:", "Chrome plated.")
    ):
        _numbered(document, text, ilvl, "50")
    _numbered(document, "Mirror polished.", 5, "50")
    result = _parse(tmp_path, _save(document))
    clamps = [w for w in result.warnings if "nesting deeper" in w]
    assert len(clamps) == 1
    assert "nesting deeper than 5 levels — clamped to level 5." in clamps[0]
    assert "(at 2.1.A.1.a.1.b, id " in clamps[0]
    assert _depth_of(result.section)["Mirror polished."] == 4


# ---------------------------------------------------------------------------
# The normalized exports
# ---------------------------------------------------------------------------


def _five_level_store() -> DocumentStore:
    store = DocumentStore()
    store.begin_turn()
    store.apply_edits(
        [
            {
                "action": "replace",
                "target_id": "sec",
                "text": "WET-PIPE SPRINKLER SYSTEMS",
                "numbering": "21 13 13",
            },
            *_five_levels(),
        ]
    )
    store.commit_turn()
    return store


def test_the_normalized_export_numbers_the_fifth_level_and_reimports_it(tmp_path):
    store = _five_level_store()
    payload = build_docx(store.doc)
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        numbering = etree.fromstring(archive.read("word/numbering.xml"))
    (abstract,) = [
        a
        for a in numbering.iter(qn("w:abstractNum"))
        if (a.find(qn("w:name")) is not None)
        and a.find(qn("w:name")).get(qn("w:val"))
        == "Build-a-Spec SectionFormat Provisions"
    ]
    levels = {
        lvl.get(qn("w:ilvl")): (
            lvl.find(qn("w:numFmt")).get(qn("w:val")),
            lvl.find(qn("w:lvlText")).get(qn("w:val")),
        )
        for lvl in abstract.iter(qn("w:lvl"))
    }
    assert levels["3"] == ("decimal", "%4)")
    assert levels["4"] == ("lowerLetter", "%5)")
    ilvls = {
        p.text: p._p.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:ilvl')}").get(
            qn("w:val")
        )
        for p in Document(io.BytesIO(payload)).paragraphs
        if p.text in ("Finish:", "Chrome plated.", "White polyester.")
    }
    assert ilvls == {"Finish:": "3", "Chrome plated.": "4", "White polyester.": "4"}
    reread = _parse(tmp_path, payload, "clean.docx")
    assert _shape(reread.section) == _shape(store.doc)


def test_the_normalized_redline_letters_the_fifth_level_and_reimports_it(
    tmp_path,
):
    """The normalized redline writes literal labels, and the Accept-All
    reading has to find ``a)`` as one — or the fifth level comes back as
    top-level provisions with the letter left in their text."""
    store = _five_level_store()
    redline = build_docx(
        store.doc,
        redline=diff_sections(SpecSection.empty(), store.doc),
        redline_date="2026-09-23T12:00:00Z",
    )
    # Accept-All text: vs an empty base every word is inside a w:ins, which
    # python-docx's Paragraph.text does not read.
    texts = [
        _accept_all_paragraph_text(p._p)
        for p in Document(io.BytesIO(redline)).paragraphs
    ]
    assert "a)\tChrome plated." in texts and "b)\tWhite polyester." in texts
    reread = _parse(tmp_path, redline, "redline.docx")
    assert _shape(reread.section) == _shape(store.doc)


# ---------------------------------------------------------------------------
# Export Word (keeps your formatting)
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("master", ["style", "typed"])
def test_an_untouched_five_level_master_round_trips_element_for_element(
    tmp_path, master
):
    """For a typed master this is the bug's worst half: the unrecognized
    ``a)`` lines relettered every provision after them in the TREE, and the
    formatted export of an import nobody had touched wrote those letters
    into the file ("B." became "E.", "2)" became "1.")."""
    source = _style_master() if master == "style" else _typed_master()
    imported = _parse(tmp_path, source)
    stats: dict = {}
    exported = render_preserving_docx(
        source_bytes=source,
        format_map=imported.format_map,
        current=imported.section,
        stats=stats,
    )
    before = [etree.tostring(el) for el in _body_children(source)]
    after = [etree.tostring(el) for el in _body_children(exported)]
    assert before == after
    assert stats["spliced"] == 0 and not stats["fallback"]


def test_a_new_fifth_level_provision_in_a_typed_master_reletters_its_siblings(
    tmp_path,
):
    source = _typed_master()
    imported = _parse(tmp_path, source)
    finish = imported.section.parts[1].articles[0].paragraphs[0].children[0]
    finish = finish.children[0].children[0]
    section = _edit(
        imported.section,
        {
            "action": "add_paragraph",
            "target_id": finish.uid,
            "position": 0,
            "text": "Satin chrome.",
        },
    )
    exported = render_preserving_docx(
        source_bytes=source, format_map=imported.format_map, current=section
    )
    texts = _body_texts(exported)
    start = texts.index("a)\tSatin chrome.")
    assert texts[start : start + 5] == [
        "a)\tSatin chrome.",
        "b)\tChrome plated.",
        "c)\tWhite polyester.",
        "d)\tBrass.",
        "2)\tEscutcheons: Recessed.",
    ]
    reread = _parse(tmp_path, exported, "exported.docx")
    assert _shape(reread.section) == _shape(section)


def test_a_new_fifth_level_provision_takes_its_own_numbering_level(tmp_path):
    """The master's list defines the fifth level but its body never uses it,
    so the new ``a)`` is cloned from its PR4 parent and must take the level
    below — or Word draws it as a ``2)`` and it re-imports as a sibling."""
    lines = tuple(item for item in _SPRINKLER_OUTLINE if item[1] != "PR5")
    source = _style_master(lines)
    imported = _parse(tmp_path, source)
    finish = imported.section.parts[1].articles[0].paragraphs[0].children[0]
    finish = finish.children[0].children[0]
    assert finish.text == "Finish:" and not finish.children
    section = _edit(
        imported.section,
        {"action": "add_paragraph", "target_id": finish.uid, "text": "Chrome plated."},
    )
    stats: dict = {}
    exported = render_preserving_docx(
        source_bytes=source,
        format_map=imported.format_map,
        current=section,
        stats=stats,
    )
    assert (stats["level_offset"], stats["level_kept"]) == (1, 0)
    (added,) = [
        p._p
        for p in Document(io.BytesIO(exported)).paragraphs
        if p.text == "Chrome plated."
    ]
    ilvl = added.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:ilvl')}")
    assert ilvl is not None and ilvl.get(qn("w:val")) == "6"
    reread = _parse(tmp_path, exported, "exported.docx")
    assert _depth_of(reread.section)["Chrome plated."] == 4


@pytest.mark.parametrize("master", ["style", "typed"])
def test_the_redline_on_the_original_tracks_a_fifth_level_edit(tmp_path, master):
    """Edit one ``a)``, add another: Accept All is the formatted export,
    Reject All the upload, and the redline re-imports as the tree."""
    source = _style_master() if master == "style" else _typed_master()
    imported = _parse(tmp_path, source)
    refs = {
        p.text: p
        for _part, _article, p, _depth, _ref in iter_paragraphs(imported.section)
    }
    section = _edit(
        imported.section,
        {
            "action": "replace",
            "target_id": refs["White polyester."].uid,
            "text": "White polyester, factory applied.",
        },
        {
            "action": "add_paragraph",
            "target_id": refs["Finish:"].uid,
            "text": "Black.",
        },
    )
    redline, _stats = _verify(source, imported, section)
    reread = _parse(tmp_path, redline, "redline.docx")
    assert _shape(reread.section) == _shape(section)


# ---------------------------------------------------------------------------
# The import route
# ---------------------------------------------------------------------------


@pytest.fixture
def client():
    sessions.reset_session()
    with TestClient(create_app()) as api:
        yield api
    sessions.reset_session()


def test_an_imported_fifth_level_reaches_the_panel_with_its_letters(client):
    response = client.post(
        "/api/import/master",
        files={"file": ("21 13 13.docx", _style_master(), DOCX_MEDIA_TYPE)},
        data={"detach": "true"},
    )
    assert response.status_code == 200, response.text
    data = response.json()
    assert not any("nesting deeper" in w for w in data["warnings"])

    def labels(nodes):
        for node in nodes:
            yield node["label"], node["text"]
            yield from labels(node["children"])

    found = {
        text: label
        for part in data["doc"]["parts"]
        for article in part["articles"]
        for label, text in labels(article["paragraphs"])
    }
    assert found["Chrome plated."] == "a)"
    assert found["Brass."] == "c)"
    assert found["Escutcheons: Recessed."] == "2)"
