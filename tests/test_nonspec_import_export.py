"""A non-spec import exports as the file it was — no headings it never had.

An import with no spec structure (a memo) is still a SectionFormat tree: the
importer wraps its content in "PART 1 - GENERAL" and a synthetic
"1.1 IMPORTED CONTENT" article, which the panel shows as editor scaffolding
under a note that the file had no spec structure. *Export Word (keeps your
formatting)* used to print them — and so the redline on the original showed
them as insertions nobody made. While the section has no number or title,
they are now left out of both (``source_render._importer_placeholders``).
"""
from __future__ import annotations

import io
import zipfile

from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from lxml import etree

from backend import sessions
from backend.app import create_app
from backend.spec_doc.importer import parse_master_docx
from backend.spec_doc.model import SpecSection, apply_edits
from backend.spec_doc.revisions import accept_all, first_difference, reject_all
from backend.spec_doc.source_render import (
    render_preserving_docx,
    render_preserving_redline,
)
from tests.docx_fidelity_helpers import DOCX_MEDIA_TYPE

MEMO_LINES = (
    "Memo to file",
    "This memo describes the vibration isolation approach.",
    "Isolators are selected for the lowest operating speed.",
)
PLACEHOLDER_TEXT = ("PART 1 - GENERAL", "IMPORTED CONTENT")


def _memo_bytes() -> bytes:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "ACME ENGINEERING"
    for line in MEMO_LINES:
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _document_xml(payload: bytes) -> bytes:
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        return archive.read("word/document.xml")


def _body(payload: bytes):
    return etree.fromstring(_document_xml(payload)).find(qn("w:body"))


def _body_children(payload: bytes) -> list[bytes]:
    return [
        etree.tostring(child)
        for child in _body(payload).iterchildren()
        if isinstance(child.tag, str)
    ]


def _texts(payload: bytes) -> list[str]:
    return [
        p.text for p in Document(io.BytesIO(payload)).paragraphs if p.text.strip()
    ]


def _import(tmp_path):
    source = _memo_bytes()
    path = tmp_path / "memo.docx"
    path.write_bytes(source)
    imported = parse_master_docx(path)
    assert imported.spec_shape_detected is False
    return source, imported


def _render(source, imported, section, **flags) -> bytes:
    return render_preserving_docx(
        source_bytes=source,
        format_map=imported.format_map,
        current=section,
        **flags,
    )


def _article(section: SpecSection):
    return section.parts[0].articles[0]


# ---------------------------------------------------------------------------
# Through the app: the capture reads the condition, the render honours it
# ---------------------------------------------------------------------------


def _client_with_memo() -> TestClient:
    client = TestClient(create_app())
    response = client.post(
        "/api/import/master",
        files={"file": ("memo.docx", _memo_bytes(), DOCX_MEDIA_TYPE)},
        data={"detach": "true"},
    )
    assert response.status_code == 200, response.text
    assert sessions.get_session().import_is_unstructured() is True
    return client


def test_an_untouched_memo_exports_element_for_element():
    client = _client_with_memo()

    exported = client.get("/api/export/docx").content

    assert _body_children(exported) == _body_children(_memo_bytes())
    texts = _texts(exported)
    for heading in PLACEHOLDER_TEXT:
        assert not any(heading in text for text in texts), (heading, texts)


def test_the_redline_on_the_original_shows_no_invented_headings():
    client = _client_with_memo()
    article = _article(sessions.get_session().doc.doc)
    edit = client.post(
        "/api/doc/edit",
        json={
            "ops": [
                {
                    "action": "add_paragraph",
                    "target_id": article.uid,
                    "text": "Provide restraints at every isolator.",
                }
            ]
        },
    )
    assert edit.status_code == 200, edit.text

    response = client.get("/api/export/docx?redline=master&mode=preserved")

    # 200 means the redline passed its own D-7 check: Accept All is the
    # formatted export and Reject All is the upload.
    assert response.status_code == 200, response.text
    body = _body(response.content)
    inserted = "".join(
        text
        for ins in body.iter(qn("w:ins"))
        for text in ins.itertext()
    )
    assert "Provide restraints at every isolator." in inserted
    for heading in PLACEHOLDER_TEXT:
        assert heading not in inserted, heading
    formatted = client.get("/api/export/docx?mode=preserved").content
    assert first_difference(accept_all(body), _body(formatted)) is None
    assert first_difference(reject_all(body), _body(_memo_bytes())) is None


def test_the_export_capture_threads_the_condition_to_the_render(monkeypatch):
    """The render never reads the live session: the condition and the
    imported tree ride the captured inputs."""
    from backend import app as app_module

    client = _client_with_memo()
    seen: dict = {}
    real = app_module.render_preserving_docx

    def recording(**kwargs):
        seen.update(kwargs)
        return real(**kwargs)

    monkeypatch.setattr(app_module, "render_preserving_docx", recording)
    assert client.get("/api/export/docx").status_code == 200

    assert seen["unstructured_import"] is True
    baseline = seen["baseline"]
    session = sessions.get_session()
    assert isinstance(baseline, SpecSection)
    assert baseline.to_dict() == session.doc.versions[session.doc.baseline_index]
    # A detached copy, never a live object.
    assert baseline is not session.doc.doc


def test_naming_the_section_brings_the_headings_back():
    client = _client_with_memo()
    named = client.post(
        "/api/doc/edit",
        json={
            "ops": [
                {
                    "action": "replace",
                    "target_id": "sec",
                    "text": "VIBRATION CONTROLS",
                    "numbering": "23 05 48",
                }
            ]
        },
    )
    assert named.status_code == 200, named.text

    texts = _texts(client.get("/api/export/docx").content)

    assert "PART 1 - GENERAL" in texts
    assert any("IMPORTED CONTENT" in text for text in texts)


# ---------------------------------------------------------------------------
# The rule: what is the importer's, and what is somebody's heading
# ---------------------------------------------------------------------------


def test_a_heading_somebody_added_is_exported_with_its_part(tmp_path):
    source, imported = _import(tmp_path)
    section, _ = apply_edits(
        imported.section,
        [
            {"action": "add_article", "target_id": "pt1", "text": "SCOPE"},
            {
                "action": "add_paragraph",
                "target_id": "pt1.a2",
                "text": "Scope of the isolation work.",
            },
        ],
    )

    stats: dict = {}
    texts = _texts(
        _render(
            source,
            imported,
            section,
            unstructured_import=True,
            baseline=imported.section,
            stats=stats,
        )
    )

    # The part now holds somebody's article, so it keeps its heading; the
    # importer's container is still left out.
    assert "PART 1 - GENERAL" in texts
    assert any(text.endswith("SCOPE") for text in texts)
    assert not any("IMPORTED CONTENT" in text for text in texts)
    assert stats["placeholders_omitted"] == 1


def test_a_renamed_container_is_the_users_heading(tmp_path):
    source, imported = _import(tmp_path)
    article = _article(imported.section)
    section, _ = apply_edits(
        imported.section,
        [{"action": "replace", "target_id": article.uid, "text": "BACKGROUND"}],
    )

    texts = _texts(
        _render(
            source,
            imported,
            section,
            unstructured_import=True,
            baseline=imported.section,
        )
    )

    assert "PART 1 - GENERAL" in texts
    assert any(text.endswith("BACKGROUND") for text in texts)


def test_without_the_flag_or_the_baseline_every_heading_is_exported(tmp_path):
    """A structured import (or a caller that knows nothing of the import)
    renders exactly as before."""
    source, imported = _import(tmp_path)
    for flags in (
        {},
        {"unstructured_import": False, "baseline": imported.section},
        {"unstructured_import": True, "baseline": None},
    ):
        texts = _texts(_render(source, imported, imported.section, **flags))
        assert texts[:2] == ["PART 1 - GENERAL", "1.1 IMPORTED CONTENT"], flags


def test_the_redline_inherits_the_rule_from_the_shared_plan(tmp_path):
    source, imported = _import(tmp_path)
    article = _article(imported.section)
    section, _ = apply_edits(
        imported.section,
        [
            {
                "action": "add_paragraph",
                "target_id": article.uid,
                "text": "Provide restraints at every isolator.",
            }
        ],
    )

    stats: dict = {}
    redline = render_preserving_redline(
        source_bytes=source,
        format_map=imported.format_map,
        baseline=imported.section,
        current=section,
        author="Build-a-Spec",
        date="2026-09-23T12:00:00Z",
        stats=stats,
        unstructured_import=True,
    )

    assert stats["placeholders_omitted"] == 2
    assert stats["redline"]["inserted"] == 1
    clean = _render(
        source, imported, section, unstructured_import=True, baseline=imported.section
    )
    assert first_difference(accept_all(_body(redline)), _body(clean)) is None
    assert first_difference(reject_all(_body(redline)), _body(source)) is None


def test_a_heading_the_upload_really_has_is_never_a_placeholder(tmp_path):
    """The origin test: whatever else holds, a heading anchored to an element
    of the file is the file's content. An unstructured import anchors no
    heading, so this is the rule's floor rather than a case the app
    produces — pinned by handing a structured master the flag."""
    document = Document()
    for line in (
        "PART 1 - GENERAL",
        "1.1 SUMMARY",
        "A. Section includes vibration isolation.",
    ):
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    source = buffer.getvalue()
    path = tmp_path / "headings.docx"
    path.write_bytes(source)
    imported = parse_master_docx(path)
    assert not (imported.section.number or imported.section.title)

    exported = _render(
        source,
        imported,
        imported.section,
        unstructured_import=True,
        baseline=imported.section,
    )

    assert _body_children(exported) == _body_children(source)
