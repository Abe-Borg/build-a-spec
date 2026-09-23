"""A preserved block takes no letter anywhere a provision is named.

The panel, ``outline()`` and the formatting-preserving export number through
``model.labelled_paragraphs``: a preserved table between A and B is not B.
Every other place a provision is named reads ``iter_paragraphs`` refs — the
open-items list, lint issues, Final QC's ``reviewed_ref``, the Word export's
schedules, importer warnings and the diff — and those used to count the
table, so the provision the panel calls "B" was "1.1.C" in all of them.

These tests build ONE real imported master (a table between A and B) and ask
each of those surfaces what it calls the provision after the table, and what
it calls the table.
"""
from __future__ import annotations

import io

from docx import Document

from backend.qc.engine import _reviewed_location
from backend.spec_doc import open_questions
from backend.spec_doc.diffing import diff_sections
from backend.spec_doc.docx_export import build_docx
from backend.spec_doc.importer import parse_master_docx
from backend.spec_doc.linting import lint_document
from backend.spec_doc.model import SpecSection, apply_edits, iter_paragraphs, outline
from backend.spec_modules.registry import get_module

TABLE_REF = "1.1 [preserved table after A]"


def _master(tmp_path):
    document = Document()
    for line in (
        "SECTION 23 05 48",
        "VIBRATION AND SEISMIC CONTROLS",
        "PART 1 - GENERAL",
        "1.1 SUMMARY",
        "A. Section includes vibration isolation.",
    ):
        document.add_paragraph(line)
    table = document.add_table(rows=2, cols=2)
    for row, (left, right) in zip(
        table.rows, (("Equipment", "Deflection"), ("AHU-1", "1 inch"))
    ):
        row.cells[0].text = left
        row.cells[1].text = right
    for line in (
        "B. Provide isolators [TBD: deflection] [INSERT manufacturer].",
        # No "1." above it: the importer attaches it one level up and
        # warns, naming where it landed.
        "a. Restraint anchors.",
        "END OF SECTION",
    ):
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    path = tmp_path / "master.docx"
    path.write_bytes(buffer.getvalue())
    return parse_master_docx(path)


def _paragraphs(section: SpecSection):
    return section.parts[0].articles[0].paragraphs


def test_the_panel_and_every_ref_agree_after_a_preserved_table(tmp_path):
    imported = _master(tmp_path)
    section = imported.section
    first, table, second = _paragraphs(section)
    assert table.locked == "table"
    # The panel's letters (the serialized labels) and the model's outline.
    labels = [p["label"] for p in section.to_dict()["parts"][0]["articles"][0]["paragraphs"]]
    assert labels == ["A.", "", "B."]
    assert "B. (imported) Provide isolators" in outline(section)

    refs = {p.uid: ref for _pt, _a, p, _d, ref in iter_paragraphs(section)}
    assert refs[first.uid] == "1.1.A"
    assert refs[table.uid] == TABLE_REF
    assert refs[second.uid] == "1.1.B"
    assert refs[second.children[0].uid] == "1.1.B.1"
    assert "1.1.C" not in refs.values()


def test_importer_warnings_name_the_panels_ref(tmp_path):
    imported = _master(tmp_path)
    jumped = [w for w in imported.warnings if "jumped deeper" in w]
    assert jumped, imported.warnings
    child = _paragraphs(imported.section)[2].children[0]
    assert f"(at 1.1.B.1, id {child.uid})" in jumped[0]


def test_open_items_and_lint_name_the_panels_ref(tmp_path):
    section = _master(tmp_path).section
    items = open_questions(section)
    assert [item["ref"] for item in items if item["kind"] == "tbd"] == ["1.1.B"]

    issues = lint_document(section, get_module("generic"))
    placeholders = [i for i in issues if i["rule"] == "placeholder_marker"]
    assert [i["ref"] for i in placeholders] == ["1.1.B"]


def test_final_qc_reviewed_ref_names_the_panels_ref(tmp_path):
    section = _master(tmp_path).section
    _first, table, second = _paragraphs(section)
    assert _reviewed_location(section, second.uid)[0] == "1.1.B"
    assert _reviewed_location(section, table.uid)[0] == TABLE_REF


def test_the_word_schedules_name_the_panels_ref(tmp_path):
    section = _master(tmp_path).section
    exported = Document(io.BytesIO(build_docx(section)))
    cells = [
        (row.cells[0].text, row.cells[1].text)
        for table in exported.tables
        for row in table.rows
    ]
    # IMPORTED PROVISIONS NOT YET REVIEWED lists every imported block.
    refs = [ref for ref, _text in cells]
    assert "1.1.B" in refs
    assert TABLE_REF in refs
    assert "1.1.C" not in refs


def test_the_diff_names_the_panels_ref(tmp_path):
    base = _master(tmp_path).section
    second = _paragraphs(base)[2]
    current, _ = apply_edits(
        base,
        [{"action": "replace", "target_id": second.uid, "text": "Provide isolators."}],
    )
    diff = diff_sections(base, current)
    changed = next(e for e in diff.elements if e.uid == second.uid)
    assert changed.ref_base == "1.1.B"
    assert changed.ref_cur == "1.1.B"
