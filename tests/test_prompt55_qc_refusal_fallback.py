"""Final QC falls back when a streamed call is declined (5.5 prompting upgrade, P55-7).

Opus 5.5's safety classifiers can decline a benign-adjacent review — clean
agent chemistry, hazmat classification, a security discipline — and before
this session a declined lens failed (the report went partial), a declined
seat left its candidate inconclusive, and a declined grouping call fell back
to singletons. Now every request a STREAMED Final QC call sends carries
``fallbacks: "default"`` and its beta (``server-side-fallback-2026-07-01``),
so the API retries a declined request itself on the model it chooses and
returns that model's answer. The record the call produced names the model
(``served_by_model``), the report says so on that record and once in
Limitations, the rescued usage is estimated at the configured QC model's
rates (decision D6), and the continuation tail's value check (CT-2) never
measures a response another model answered.

What never carries it: the batched params (the Batches API rejects the
field), research and the chat (out of scope, D6). A request the provider
refuses BECAUSE of the parameter is sent once more without it and the
fallback switches off until the app restarts — CT-1's shape.
"""
from __future__ import annotations

import ast
import dataclasses
import json
import logging
import re
from pathlib import Path
from types import SimpleNamespace

import anthropic
import httpx
import pytest

from backend import cost_checks, settings
from backend.qc import engine as qc_engine
from backend.qc.engine import (
    QCConsolidation,
    QCFinding,
    QCLensStatus,
    QCResult,
    QCVerdict,
    run_final_qc,
)
from backend.qc.schema import QC_LENSES
from backend.research import engine as research_engine
from backend.research.schema import PRESERVED_THINKING_BETA, with_drop_block
from backend.spec_doc import docx_export
from backend.spec_modules import DEFAULT_MODULE
from backend.usage_ledger import estimate_usage_cost
from tests.fakes import (
    SequencedFakeClient,
    bad_request,
    fallback_served,
    pause_response,
    qc_consolidation_response,
    qc_findings_response,
    qc_verdict_response,
    text_block,
    thinking_block,
    tool_use_block,
)
from tests.test_qc_batch_verification import _finding, _scripts, _store
from tests.test_qc_warm_launch import _fixed_clock
from tests.test_retry_resume import _ResearchHarness

_REPO = Path(__file__).resolve().parents[1]
_BETA = qc_engine.REFUSAL_FALLBACK_BETA
_FALLBACK_MODEL = "claude-opus-5"
_WEB_LENS = next(lens.lens_id for lens in QC_LENSES if lens.web)
_DOC_LENS = next(lens.lens_id for lens in QC_LENSES if not lens.web)
_TITLE = "Fallback finding"
_REJECTED = (
    "fallbacks: Extra inputs are not permitted (server-side-fallback-2026-07-01)"
)


@pytest.fixture(autouse=True)
def _no_backoff(monkeypatch):
    monkeypatch.setattr(qc_engine.time, "sleep", lambda _seconds: None)


class _StreamLog:
    """The requests a client STREAMED — never a batched seat's params, which
    ``SequencedFakeClient.requests`` also holds (the batch fake resolves them
    through the same scripts)."""

    def __init__(self, client: SequencedFakeClient):
        self.requests: list[dict] = []
        original = client.stream

        def stream(**request):
            self.requests.append(dict(request))
            return original(**request)

        client.stream = stream


def _carries(request: dict) -> bool:
    body = request.get("extra_body") or {}
    betas = str((request.get("extra_headers") or {}).get("anthropic-beta", ""))
    return body.get("fallbacks") == "default" and _BETA in betas.split(",")


def _mentions_fallback(request: dict) -> bool:
    return "fallback" in json.dumps(request, sort_keys=True, default=repr)


def _run(client, *, batch: bool = True, refusal_fallback=True, lead: bool = False,
         store=None, events: list | None = None):
    store = store or _store()
    return run_final_qc(
        store.doc,
        None,
        DEFAULT_MODULE,
        client,
        model=settings.QC_MODEL,
        max_tokens=4096,
        version_index=store.index,
        started_at="2026-09-30T10:00:00-07:00",
        finished_at="2026-09-30T10:01:00-07:00",
        run_id="qc-refusal-fallback-test",
        batch_verification=batch,
        batch_warm_lead=lead,
        refusal_fallback=refusal_fallback,
        event_sink=(events.append if events is not None else (lambda _e: None)),
    )


def _one_finding(*, verdicts=None, lens_turns=None) -> dict:
    scripts = _scripts(
        **{
            _WEB_LENS: lens_turns
            or [qc_findings_response(_WEB_LENS, findings=[_finding(_TITLE)])]
        }
    )
    scripts[_TITLE] = list(
        verdicts
        if verdicts is not None
        else [qc_verdict_response(True), qc_verdict_response(True)]
    )
    return scripts


def _lens_status(result: QCResult, lens_id: str) -> QCLensStatus:
    return next(s for s in result.lens_statuses if s.lens_id == lens_id)


# ---------------------------------------------------------------------------
# P55-7.1 — streamed requests carry it; batched params never do
# ---------------------------------------------------------------------------


def test_every_streamed_request_carries_the_fallback_and_batched_params_never_do():
    client = SequencedFakeClient(_one_finding())
    log = _StreamLog(client)
    _run(client, batch=True)

    # Not vacuous: five lenses streamed, and the panel really rode a batch.
    assert len(log.requests) == len(QC_LENSES)
    assert all(_carries(request) for request in log.requests)
    batched = [entry["params"] for batch in client.batches.created for entry in batch]
    assert len(batched) == 2
    assert not any(_mentions_fallback(params) for params in batched)
    # The one request shape both transports build from never carries it.
    shape = qc_engine._qc_request_kwargs(
        system_prompt="s", tools=[{"name": "t"}], model=settings.QC_MODEL,
        max_tokens=10, effort="medium", cache_ttl="",
    )
    assert not _mentions_fallback(shape)


def test_streamed_seats_carry_it_too():
    client = SequencedFakeClient(_one_finding())
    log = _StreamLog(client)
    _run(client, batch=False)

    seats = [r for r in log.requests if "[[QC-VERIFY:" in json.dumps(r["messages"], default=repr)]
    assert len(seats) == 2
    assert all(_carries(request) for request in seats)
    assert not client.batches.created


def test_a_warm_lead_carries_it_and_its_batch_does_not(monkeypatch):
    monkeypatch.setattr(qc_engine, "_WARM_LEAD_MIN_SEATS_NO_WEB", 8)
    titles = [f"Doc lineage {number:02d}" for number in range(1, 6)]
    scripts = _scripts(
        **{
            _DOC_LENS: [
                qc_findings_response(
                    _DOC_LENS, findings=[_finding(title) for title in titles]
                )
            ]
        }
    )
    for title in titles:
        scripts[title] = [qc_verdict_response(True), qc_verdict_response(True)]
    client = SequencedFakeClient(scripts)
    log = _StreamLog(client)
    _run(client, batch=True, lead=True)

    seats = [r for r in log.requests if "[[QC-VERIFY:" in json.dumps(r["messages"], default=repr)]
    assert len(seats) == 1  # the lead, streamed
    assert _carries(seats[0])
    batched = [entry["params"] for batch in client.batches.created for entry in batch]
    assert len(batched) == 9
    assert not any(_mentions_fallback(params) for params in batched)


def test_the_switch_off_sends_neither():
    client = SequencedFakeClient(_one_finding())
    log = _StreamLog(client)
    _run(client, batch=False, refusal_fallback=False)

    assert len(log.requests) == len(QC_LENSES) + 2
    assert not any(_mentions_fallback(request) for request in log.requests)
    assert all("extra_body" not in request for request in log.requests)


def test_the_setting_reaches_a_run(monkeypatch):
    monkeypatch.setattr(settings, "QC_REFUSAL_FALLBACK", False)
    client = SequencedFakeClient(_one_finding())
    log = _StreamLog(client)
    _run(client, batch=True, refusal_fallback=None)
    assert not any(_mentions_fallback(request) for request in log.requests)

    monkeypatch.setattr(settings, "QC_REFUSAL_FALLBACK", True)
    client = SequencedFakeClient(_one_finding())
    log = _StreamLog(client)
    _run(client, batch=True, refusal_fallback=None)
    assert log.requests and all(_carries(request) for request in log.requests)


def test_the_fallback_ships_switched_on():
    """Read from the source, so no developer's environment decides it."""
    tree = ast.parse((_REPO / "backend" / "settings.py").read_text(encoding="utf-8"))
    for node in ast.walk(tree):
        if (
            isinstance(node, ast.Assign)
            and any(
                isinstance(t, ast.Name) and t.id == "QC_REFUSAL_FALLBACK"
                for t in node.targets
            )
        ):
            call = node.value
            assert isinstance(call, ast.Call)
            assert call.func.id == "_bool_env"
            assert call.args[0].value == "BUILD_A_SPEC_QC_REFUSAL_FALLBACK"
            assert call.args[1].value is True
            return
    raise AssertionError("QC_REFUSAL_FALLBACK is not assigned in settings.py")


def test_it_merges_with_the_preserved_thinking_beta_and_leaves_it_behind():
    edited = with_drop_block(
        {"model": "m", "thinking": {"type": "adaptive"}, "extra_body": {"x": 1}}
    )
    both = qc_engine._with_refusal_fallback(edited)
    betas = both["extra_headers"]["anthropic-beta"].split(",")
    assert betas == [PRESERVED_THINKING_BETA, _BETA]
    assert both["extra_body"] == {"x": 1, "fallbacks": "default"}
    # The copy is new: the request it was built from is untouched.
    assert "fallbacks" not in edited["extra_body"]

    back = qc_engine._without_refusal_fallback(both)
    assert back["extra_headers"] == {"anthropic-beta": PRESERVED_THINKING_BETA}
    assert back["extra_body"] == {"x": 1}
    alone = qc_engine._without_refusal_fallback(
        qc_engine._with_refusal_fallback({"model": "m"})
    )
    assert alone == {"model": "m"}


def test_research_never_carries_it():
    """Out of scope (D6): research is Sonnet-driven and its declines are
    already named as refusals. Pinned so a shared helper cannot leak it."""
    harness = _ResearchHarness()
    call = harness.run([harness.final()])
    assert call.requests
    assert not any(_mentions_fallback(request) for request in call.requests)


# ---------------------------------------------------------------------------
# P55-7.2 — a rescued call completes and says who answered it
# ---------------------------------------------------------------------------


def _rescued_lens(**signals) -> SimpleNamespace:
    return fallback_served(
        qc_findings_response(_WEB_LENS, findings=[_finding(_TITLE)]),
        to_model=_FALLBACK_MODEL,
        **signals,
    )


@pytest.mark.parametrize(
    "signals",
    [
        {},
        {"iteration": False, "model": False},  # the switch-point block alone
        {"block": False, "model": False},  # the fallback_message iteration alone
        {"block": False, "iteration": False},  # a sticky turn: the model alone
    ],
    ids=["all", "block", "iteration", "model"],
)
def test_a_rescued_lens_completes_and_records_who_answered(signals):
    result = _run(
        SequencedFakeClient(_one_finding(lens_turns=[_rescued_lens(**signals)])),
        batch=True,
    )
    status = _lens_status(result, _WEB_LENS)
    assert status.status == "completed"
    assert status.served_by_model == _FALLBACK_MODEL
    # The rescued lens's finding went on to its panel like any other.
    assert [f.title for f in result.findings] == [_TITLE]
    # And only that record says so.
    others = [s for s in result.lens_statuses if s.lens_id != _WEB_LENS]
    assert not any(s.served_by_model for s in others)


def test_another_model_name_alone_is_not_a_fallback_when_none_was_asked_for():
    """An alias or a dated snapshot echoed back is not a rescue: the model
    signal counts only for a request that carried the fallback."""
    turn = _rescued_lens(block=False, iteration=False)
    result = _run(
        SequencedFakeClient(_one_finding(lens_turns=[turn])),
        batch=True,
        refusal_fallback=False,
    )
    assert _lens_status(result, _WEB_LENS).served_by_model == ""


def test_a_decline_the_fallback_also_declines_stays_a_refusal():
    declined = fallback_served(
        qc_findings_response(
            _WEB_LENS, findings=None, stop_reason="refusal",
            refusal_category="cyber",
        ),
        to_model=_FALLBACK_MODEL,
    )
    result = _run(SequencedFakeClient(_one_finding(lens_turns=[declined])))
    status = _lens_status(result, _WEB_LENS)
    assert status.status == "failed"
    assert "safety classifier declined" in status.error
    assert status.served_by_model == ""


def test_a_rescued_seat_records_who_answered_it():
    seat = fallback_served(qc_verdict_response(True), to_model=_FALLBACK_MODEL)
    result = _run(
        SequencedFakeClient(_one_finding(verdicts=[seat, qc_verdict_response(True)])),
        batch=False,
    )
    verdicts = result.findings[0].verdicts
    assert sorted(v.served_by_model for v in verdicts) == ["", _FALLBACK_MODEL]
    assert all(v.status == "completed" for v in verdicts)


def test_the_payload_is_read_after_the_switch_point():
    """A mid-output decline keeps the declined partial in ``content``; its
    output-tool call may be cut short. The payload is the fallback model's."""
    partial = tool_use_block(
        "toolu_partial", "submit_qc_findings",
        {"summary": "", "reviewed_checks": [], "findings": [_finding("Cut short")]},
    )
    answer = qc_findings_response(_WEB_LENS, findings=None)
    answer.content.append(
        text_block(
            "<qc_json>"
            + json.dumps({"summary": "", "reviewed_checks": [], "findings": [_finding(_TITLE)]})
            + "</qc_json>"
        )
    )
    answer.stop_reason = "end_turn"
    turn = fallback_served(answer, to_model=_FALLBACK_MODEL, partial=[partial])
    result = _run(SequencedFakeClient(_one_finding(lens_turns=[turn])))
    assert [f.title for f in result.findings] == [_TITLE]


def test_a_paused_rescue_is_echoed_by_the_rule():
    """Before the final switch point only text, paired server-tool blocks and
    the switch marks are re-sent; thinking and a partial tool call are not.
    Everything after the boundary is re-sent as it came."""
    pending = pause_response(pending_query="nfpa 13 current edition")
    partial = [
        thinking_block("declined thinking"),
        tool_use_block("toolu_cut", "submit_qc_findings", {"summary": "x"}),
        text_block("A partial note."),
    ]
    turn = fallback_served(pending, to_model=_FALLBACK_MODEL, partial=partial)
    final = qc_findings_response(_WEB_LENS, findings=[])
    client = SequencedFakeClient(_one_finding(lens_turns=[turn, final]))
    log = _StreamLog(client)
    _run(client)

    lens = [r for r in log.requests if f"[[QC-LENS:{_WEB_LENS}]]" in json.dumps(r["messages"], default=repr)]
    assert len(lens) == 2
    echoed = lens[1]["messages"][-1]["content"]
    kinds = [getattr(block, "type", None) or block.get("type") for block in echoed]
    assert kinds == ["text", "fallback", "server_tool_use"]


def test_an_ordinary_response_is_re_sent_exactly_as_it_came():
    content = [text_block("a"), thinking_block("b")]
    assert qc_engine._fallback_echo_content(content) is content
    response = SimpleNamespace(content=content)
    assert qc_engine._payload_view(response) is response


def test_a_rescued_grouping_call_records_who_answered_it():
    titles = ["Grouped one", "Grouped two"]
    scripts = _scripts(
        **{
            _WEB_LENS: [
                qc_findings_response(
                    _WEB_LENS, findings=[_finding(title) for title in titles]
                )
            ]
        }
    )
    for title in titles:
        scripts[title] = [qc_verdict_response(True), qc_verdict_response(True)]
    singletons = [
        {
            "member_indexes": [index],
            "canonical_title": None,
            "canonical_issue": None,
            "canonical_rationale": None,
            "grouping_rationale": None,
            "reconciled_ops": None,
        }
        for index in (0, 1)
    ]
    scripts["[[QC-CONSOLIDATE:"] = [
        fallback_served(qc_consolidation_response(singletons), to_model=_FALLBACK_MODEL)
    ]
    client = SequencedFakeClient(scripts)
    log = _StreamLog(client)
    result = _run(client)
    grouping = [
        r for r in log.requests
        if "[[QC-CONSOLIDATE:" in json.dumps(r["messages"], default=repr)
    ]
    assert len(grouping) == 1 and _carries(grouping[0])
    assert result.consolidation.status == "complete"
    assert result.consolidation.served_by_model == _FALLBACK_MODEL
    payload = result.to_dict()
    assert payload["consolidation"]["served_by_model"] == _FALLBACK_MODEL
    assert QCResult.from_dict(json.loads(json.dumps(payload))).consolidation.served_by_model == _FALLBACK_MODEL
    assert docx_export.qc_refusal_fallback(payload)[0] == 1

    import io

    from docx import Document

    from backend.spec_doc.docx_export import build_qc_memo

    document = Document(
        io.BytesIO(build_qc_memo(payload, _store().doc, stale=False))
    )
    text = "\n".join(p.text for p in document.paragraphs)
    note = docx_export.QC_FALLBACK_RECORD_TEMPLATE.format(models=_FALLBACK_MODEL)
    assert "Refusal fallback: " + note in text


def test_a_rescued_lens_that_still_fails_says_who_answered():
    cut_off = fallback_served(
        qc_findings_response(_WEB_LENS, findings=None, stop_reason="max_tokens"),
        to_model=_FALLBACK_MODEL,
    )
    result = _run(SequencedFakeClient(_one_finding(lens_turns=[cut_off])))
    status = _lens_status(result, _WEB_LENS)
    assert status.status == "failed"
    assert status.served_by_model == _FALLBACK_MODEL


def test_a_rescued_seat_that_still_fails_says_who_answered():
    cut_off = fallback_served(
        SimpleNamespace(
            content=[text_block("Partial verdict")],
            stop_reason="max_tokens",
            usage=pause_response().usage,
        ),
        to_model=_FALLBACK_MODEL,
    )
    result = _run(
        SequencedFakeClient(_one_finding(verdicts=[cut_off, qc_verdict_response(True)])),
        batch=False,
    )
    verdicts = [v for c in [*result.findings, *result.inconclusive] for v in c.verdicts]
    failed = [v for v in verdicts if v.status == "failed"]
    assert len(failed) == 1
    assert failed[0].served_by_model == _FALLBACK_MODEL


def test_a_rescued_seat_with_a_malformed_verdict_still_says_who_answered():
    seat = fallback_served(qc_verdict_response(True), to_model=_FALLBACK_MODEL)
    seat.content[-1].input["upholds"] = "maybe"
    result = _run(
        SequencedFakeClient(_one_finding(verdicts=[seat, qc_verdict_response(True)])),
        batch=False,
    )
    candidates = [*result.findings, *result.inconclusive]
    verdicts = [v for c in candidates for v in c.verdicts]
    failed = [v for v in verdicts if v.status == "failed"]
    assert len(failed) == 1 and failed[0].error.startswith("Malformed QC verdict")
    assert failed[0].served_by_model == _FALLBACK_MODEL


def test_a_rescued_reply_without_its_tool_is_reminded_by_the_echo_rule():
    """P55-4's reminder re-sends the reply: a rescued one by the echo rule, and
    the reminder answers only what the fallback model actually said."""
    reply = SimpleNamespace(
        content=[text_block("A status note; the tool comes next.")],
        stop_reason="end_turn",
        usage=pause_response().usage,
    )
    partial = [
        thinking_block("declined thinking"),
        tool_use_block("toolu_invented", "not_a_real_tool", {}),
    ]
    turn = fallback_served(reply, to_model=_FALLBACK_MODEL, partial=partial)
    final = qc_findings_response(_WEB_LENS, findings=[])
    client = SequencedFakeClient(_one_finding(lens_turns=[turn, final]))
    log = _StreamLog(client)
    result = _run(client)

    assert _lens_status(result, _WEB_LENS).status == "completed"
    lens = [r for r in log.requests if f"[[QC-LENS:{_WEB_LENS}]]" in json.dumps(r["messages"], default=repr)]
    assert len(lens) == 2
    assistant, reminder = lens[1]["messages"][-2:]
    kinds = [getattr(block, "type", None) for block in assistant["content"]]
    assert kinds == ["fallback", "text"]
    # A plain text reminder: the invented call before the switch point was
    # the declined model's, and is not answered as if it were a request.
    assert [block["type"] for block in reminder["content"]] == ["text"]


# ---------------------------------------------------------------------------
# P55-7.2 — the record round-trips, and an older one still loads
# ---------------------------------------------------------------------------


def test_the_records_round_trip_and_stay_quiet_when_unset():
    lens = QCLensStatus(lens_id="x", title="X", status="completed", brief="b",
                        served_by_model=_FALLBACK_MODEL)
    assert QCLensStatus.from_dict(lens.to_dict()).served_by_model == _FALLBACK_MODEL
    assert "served_by_model" not in QCLensStatus(
        lens_id="x", title="X", status="completed", brief="b"
    ).to_dict()

    seat = QCVerdict(upholds=True, served_by_model=_FALLBACK_MODEL)
    assert (
        QCVerdict.from_dict(dataclasses.asdict(seat)).served_by_model
        == _FALLBACK_MODEL
    )
    # A seat serializes through its finding, which drops the unset key.
    finding = QCFinding(
        finding_id="qc-x", lens_id="x", title="t", severity="low", element_id="",
        issue="i", rationale="r", verdicts=[QCVerdict(upholds=True), seat],
    )
    verdicts = finding.to_dict()["verdicts"]
    assert "served_by_model" not in verdicts[0]
    assert verdicts[1]["served_by_model"] == _FALLBACK_MODEL

    step = QCConsolidation(status="complete", served_by_model=_FALLBACK_MODEL)
    assert "served_by_model" not in QCConsolidation(status="complete").to_dict()
    assert QCConsolidation.from_dict(step.to_dict()).served_by_model == _FALLBACK_MODEL


@pytest.mark.parametrize(
    "value", [None, 7, "", "   ", "not a model!", "a, b, c, d, e", ["m"]]
)
def test_a_record_never_fails_to_load_over_the_disclosure(value):
    raw = {"lens_id": "x", "title": "X", "status": "completed", "brief": "b"}
    if value is not None:
        raw["served_by_model"] = value
    assert QCLensStatus.from_dict(raw).served_by_model == ""


def test_a_whole_result_keeps_it_and_an_older_one_still_loads():
    seat = fallback_served(qc_verdict_response(True), to_model=_FALLBACK_MODEL)
    result = _run(
        SequencedFakeClient(
            _one_finding(
                lens_turns=[_rescued_lens()],
                verdicts=[seat, qc_verdict_response(True)],
            )
        ),
        batch=False,
    )
    payload = result.to_dict()
    reloaded = QCResult.from_dict(json.loads(json.dumps(payload)))
    assert reloaded is not None
    assert _lens_status(reloaded, _WEB_LENS).served_by_model == _FALLBACK_MODEL
    assert _FALLBACK_MODEL in {v.served_by_model for v in reloaded.findings[0].verdicts}

    # An older record: the key simply absent everywhere.
    for status in payload["lens_statuses"]:
        status.pop("served_by_model", None)
    for verdict in payload["findings"][0]["verdicts"]:
        verdict.pop("served_by_model", None)
    older = QCResult.from_dict(json.loads(json.dumps(payload)))
    assert older is not None
    assert not any(s.served_by_model for s in older.lens_statuses)


# ---------------------------------------------------------------------------
# P55-7.3 — the accounting reconciles, at the QC model's rates, disclosed
# ---------------------------------------------------------------------------


def test_rescued_usage_is_priced_at_the_qc_models_rates_and_reconciles():
    tokens = {"input": 1_000, "output": 2_000, "cache_read": 3_000}
    lens = fallback_served(
        qc_findings_response(_WEB_LENS, findings=[_finding(_TITLE)], tokens=tokens),
        to_model=_FALLBACK_MODEL,
    )
    result = _run(SequencedFakeClient(_one_finding(lens_turns=[lens])), batch=False)
    status = _lens_status(result, _WEB_LENS)
    assert status.estimated_cost_usd == estimate_usage_cost(
        settings.QC_MODEL, status.usage_totals
    )
    # The cost basis is the configured model's, untouched by the rescue.
    assert result.cost_basis["rate_model"] == settings.QC_MODEL
    # And the whole record reconciles on reload (a failed check discards it).
    assert QCResult.from_dict(json.loads(json.dumps(result.to_dict()))) is not None
    count, limitation = docx_export.qc_refusal_fallback(result.to_dict())
    assert count == 1
    assert f"estimated at {settings.QC_MODEL} rates" in limitation


# ---------------------------------------------------------------------------
# P55-7.4 — CT-2 never measures a response another model answered
# ---------------------------------------------------------------------------


def _measurable(**extra) -> SimpleNamespace:
    return SimpleNamespace(
        content=[text_block("done")],
        stop_reason="end_turn",
        usage=SimpleNamespace(
            input_tokens=10, output_tokens=5, cache_read_input_tokens=400,
            cache_creation_input_tokens=0,
            server_tool_use=SimpleNamespace(web_search_requests=0, web_fetch_requests=0),
        ),
        **extra,
    )


def _kinds() -> dict:
    tail = cost_checks.snapshot()["continuation_tail"][cost_checks.ENGINE_QC]
    return {key: tail[key] for key in ("exact", "bound", "unmeasured")}


def test_ct2_leaves_a_rescued_response_unmeasured():
    model = settings.QC_MODEL
    cost_checks.observe_continuation(
        cost_checks.ENGINE_QC, model=model, opening=None, response=_measurable()
    )
    assert _kinds() == {"exact": 0, "bound": 1, "unmeasured": 0}  # the control

    rescued = [
        # the engine's own signal (a sticky turn: another model named)
        dict(response=_measurable(), fallback_served=True),
        # a switch point in the continuation
        dict(response=fallback_served(_measurable(), iteration=False, model=False)),
        # a fallback_message iteration in the continuation
        dict(response=fallback_served(_measurable(), block=False, model=False)),
    ]
    for case in rescued:
        cost_checks.observe_continuation(
            cost_checks.ENGINE_QC, model=model, opening=None, **case
        )
    # A switch point beside a usage record that reports ONE ``message``
    # iteration: CT-2's iteration rules would measure it, so only the leaf's
    # own read of the block keeps it out. The leaf trusts the block, never
    # the iteration's label.
    labelled = fallback_served(_measurable(), iteration=False, model=False)
    labelled.usage.iterations = [
        {
            "type": "message",
            "input_tokens": 10,
            "output_tokens": 5,
            "cache_read_input_tokens": 400,
            "cache_creation_input_tokens": 0,
        }
    ]
    cost_checks.observe_continuation(
        cost_checks.ENGINE_QC, model=model, opening=None, response=labelled
    )
    # an opening another model answered, told by its switch point
    cost_checks.observe_continuation(
        cost_checks.ENGINE_QC, model=model,
        opening=fallback_served(_measurable(), iteration=False, model=False),
        response=_measurable(),
    )
    # an opening told only by its ``fallback_message`` iteration: the declined
    # attempt's ``message`` entry would otherwise be read as the opening's
    # prefix (nothing read, nothing written) and the saving computed from it
    cost_checks.observe_continuation(
        cost_checks.ENGINE_QC, model=model,
        opening=fallback_served(_measurable(), block=False, model=False),
        response=_measurable(),
    )
    assert _kinds() == {"exact": 0, "bound": 1, "unmeasured": 6}


def test_a_model_name_alone_is_not_read_by_the_leaf():
    """An alias must never switch the value check off: the leaf reads only
    the block and the iteration; the engine alone decides the model signal."""
    cost_checks.observe_continuation(
        cost_checks.ENGINE_QC, model=settings.QC_MODEL, opening=None,
        response=_measurable(model="some-alias-snapshot"),
    )
    assert _kinds()["bound"] == 1


def test_the_engine_tells_ct2_about_a_rescued_conversation(monkeypatch):
    """End to end: a paused lens a fallback model answered continues with
    the tail on, and its continuation is recorded unmeasured."""
    pending = fallback_served(
        pause_response(pending_query="nfpa 13"), to_model=_FALLBACK_MODEL,
        block=False, iteration=False,
    )
    final = _measurable_final()
    seen: list[dict] = []
    real = cost_checks.observe_continuation

    def spy(engine, **kwargs):
        seen.append(kwargs)
        return real(engine, **kwargs)

    monkeypatch.setattr(qc_engine.cost_checks, "observe_continuation", spy)
    store = _store()
    run_final_qc(
        store.doc, None, DEFAULT_MODULE,
        SequencedFakeClient(_one_finding(lens_turns=[pending, final])),
        model=settings.QC_MODEL, max_tokens=4096, version_index=store.index,
        started_at="2026-09-30T10:00:00-07:00",
        finished_at="2026-09-30T10:01:00-07:00",
        run_id="qc-fallback-ct2", batch_verification=True, batch_warm_lead=False,
        continuation_cache=True, refusal_fallback=True,
    )
    assert len(seen) == 1
    assert seen[0]["fallback_served"] is True
    assert _kinds()["unmeasured"] == 1


def test_a_restart_measures_its_new_conversation_again(monkeypatch):
    """A rescued conversation marks only itself: after the retry policy
    restarts the call (the final attempt always starts fresh), the new
    conversation's continuation is measured as usual."""
    reset = anthropic.APIConnectionError(
        message="connection reset",
        request=httpx.Request("POST", "https://api.anthropic.com/v1/messages"),
    )
    rescued = fallback_served(
        pause_response(pending_query="nfpa 13"), to_model=_FALLBACK_MODEL,
        block=False, iteration=False,
    )
    turns = [
        rescued,  # attempt 1: the opening, answered by another model
        reset,  # its continuation fails with progress: resume
        reset,  # the resend fails too, on the next-to-last attempt: restart
        pause_response(pending_query="nfpa 13"),  # attempt 3: a fresh opening
        _measurable_final(),  # its continuation, carrying the tail
    ]
    seen: list[dict] = []
    real = cost_checks.observe_continuation

    def spy(engine, **kwargs):
        seen.append(kwargs)
        return real(engine, **kwargs)

    monkeypatch.setattr(qc_engine.cost_checks, "observe_continuation", spy)
    store = _store()
    result = run_final_qc(
        store.doc, None, DEFAULT_MODULE,
        SequencedFakeClient(_one_finding(lens_turns=turns)),
        model=settings.QC_MODEL, max_tokens=4096, version_index=store.index,
        started_at="2026-09-30T10:00:00-07:00",
        finished_at="2026-09-30T10:01:00-07:00",
        run_id="qc-fallback-restart", batch_verification=True,
        batch_warm_lead=False, continuation_cache=True, refusal_fallback=True,
    )
    assert _lens_status(result, _WEB_LENS).status == "completed"
    # the abandoned conversation was still billed to another model, and says so
    assert _lens_status(result, _WEB_LENS).served_by_model == _FALLBACK_MODEL
    assert len(seen) == 1
    assert seen[0]["fallback_served"] is False
    assert _kinds()["unmeasured"] == 0


def _measurable_final() -> SimpleNamespace:
    final = qc_findings_response(_WEB_LENS, findings=[_finding(_TITLE)])
    final.usage.cache_read_input_tokens = 500
    return final


# ---------------------------------------------------------------------------
# The fallback's own refusal guard (CT-1's shape)
# ---------------------------------------------------------------------------


def test_a_refused_fallback_costs_one_request_and_switches_off(caplog):
    caplog.set_level(logging.WARNING, logger="buildaspec.qc")
    lens_turns = [bad_request(_REJECTED), qc_findings_response(_WEB_LENS, findings=[_finding(_TITLE)])]
    client = SequencedFakeClient(_one_finding(lens_turns=lens_turns))
    log = _StreamLog(client)
    result = _run(client, batch=True)

    status = _lens_status(result, _WEB_LENS)
    assert status.status == "completed"
    assert status.api_request_count == 2  # the refused request and its resend
    lens = [r for r in log.requests if f"[[QC-LENS:{_WEB_LENS}]]" in json.dumps(r["messages"], default=repr)]
    assert _carries(lens[0]) and not _mentions_fallback(lens[1])
    assert qc_engine.refusal_fallback_available() is False
    warnings = [r for r in caplog.records if "refusal fallback is switched off" in r.getMessage()]
    assert len(warnings) == 1

    # Every later request goes without it, until the app restarts.
    client = SequencedFakeClient(_one_finding())
    log = _StreamLog(client)
    _run(client, batch=True)
    assert log.requests and not any(_mentions_fallback(r) for r in log.requests)


def test_a_400_that_outlives_the_fallback_switches_nothing_off():
    lens_turns = [bad_request(_REJECTED), bad_request(_REJECTED)]
    result = _run(SequencedFakeClient(_one_finding(lens_turns=lens_turns)))
    status = _lens_status(result, _WEB_LENS)
    assert status.status == "failed"
    assert status.api_request_count == 2
    assert qc_engine.refusal_fallback_available() is True


def test_a_400_about_something_else_is_not_the_fallbacks():
    lens_turns = [bad_request("messages.0: invalid content")]
    result = _run(SequencedFakeClient(_one_finding(lens_turns=lens_turns)))
    status = _lens_status(result, _WEB_LENS)
    assert status.status == "failed"
    assert status.api_request_count == 1
    assert qc_engine.refusal_fallback_available() is True


def test_prompt_too_long_is_never_read_as_the_fallbacks():
    assert not qc_engine._is_fallback_rejection(
        bad_request("prompt is too long: fallbacks cannot help")
    )
    assert qc_engine._is_fallback_rejection(bad_request(_REJECTED))
    assert not qc_engine._is_fallback_rejection(RuntimeError(_REJECTED))


# ---------------------------------------------------------------------------
# P55-7.5 — both report projections state the same sentences
# ---------------------------------------------------------------------------


def _ts_literal(name: str) -> str:
    source = (_REPO / "frontend" / "src" / "lib" / "qcReport.ts").read_text(encoding="utf-8")
    match = re.search(rf'export const {name} =\s*"((?:[^"\\]|\\.)*)";', source)
    assert match, name
    return json.loads(f'"{match.group(1)}"')


def test_the_modal_states_the_memos_sentences_word_for_word():
    assert _ts_literal("QC_FALLBACK_RECORD_TEMPLATE") == docx_export.QC_FALLBACK_RECORD_TEMPLATE
    assert (
        _ts_literal("QC_FALLBACK_LIMITATION_TEMPLATE")
        == docx_export.QC_FALLBACK_LIMITATION_TEMPLATE
    )


def test_the_word_memo_discloses_every_rescued_record():
    from backend.spec_doc.docx_export import build_qc_memo

    seat = fallback_served(qc_verdict_response(True), to_model=_FALLBACK_MODEL)
    refuting = fallback_served(
        qc_verdict_response(False, note="Not a defect."), to_model=_FALLBACK_MODEL
    )
    lens = fallback_served(
        qc_findings_response(
            _WEB_LENS,
            findings=[_finding(_TITLE), {**_finding("Refuted one"), "element_id": "pt1.a1"}],
        ),
        to_model=_FALLBACK_MODEL,
    )
    scripts = _one_finding(lens_turns=[lens], verdicts=[seat, qc_verdict_response(True)])
    scripts["Refuted one"] = [refuting, qc_verdict_response(False, note="No.")]
    store = _store()
    result = _run(SequencedFakeClient(scripts), batch=False, store=store)
    assert [f.title for f in result.refuted] == ["Refuted one"]
    payload = result.to_dict()
    note = docx_export.QC_FALLBACK_RECORD_TEMPLATE.format(models=_FALLBACK_MODEL)
    rescued = next(s for s in payload["lens_statuses"] if s.get("served_by_model"))
    assert docx_export.qc_fallback_record_note(rescued) == note
    assert all(
        docx_export.qc_fallback_record_note(s) == ""
        for s in payload["lens_statuses"]
        if s is not rescued
    )
    count, limitation = docx_export.qc_refusal_fallback(payload)
    assert count == 3  # the lens, the surviving seat and the refuting seat
    assert limitation == docx_export.QC_FALLBACK_LIMITATION_TEMPLATE.format(
        count=3, models=_FALLBACK_MODEL, qc_model=settings.QC_MODEL
    )

    import io

    from docx import Document

    document = Document(io.BytesIO(build_qc_memo(payload, store.doc, stale=False)))
    text = "\n".join(p.text for p in document.paragraphs)
    cells = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    assert "Refusal fallback: " + note in text  # the lens record
    assert f"Refusal fallback (seat 1): {note}" in text  # the surviving seat
    assert f"Seat 1: {note}" in cells  # the refuted candidate's digest row
    assert limitation in text


def test_a_clean_run_discloses_nothing():
    result = _run(SequencedFakeClient(_one_finding()))
    assert docx_export.qc_refusal_fallback(result.to_dict()) == (0, "")


# ---------------------------------------------------------------------------
# F3 — a retained result stays current with the switch either way
# ---------------------------------------------------------------------------


def test_a_retained_result_stays_current_with_the_switch_either_way(monkeypatch):
    _fixed_clock(monkeypatch)
    monkeypatch.setattr(settings, "QC_BATCH_VERIFICATION", True)
    store = _store()
    on = _run(SequencedFakeClient(_one_finding()), refusal_fallback=True, store=store)
    off = _run(SequencedFakeClient(_one_finding()), refusal_fallback=False, store=store)
    assert on.input_fingerprint == off.input_fingerprint
    manifest = json.dumps(on.input_manifest, sort_keys=True, default=repr)
    assert "fallback" not in manifest.lower()
    for value in (True, False):
        monkeypatch.setattr(settings, "QC_REFUSAL_FALLBACK", value)
        assert on.matches_inputs(store.index, store.doc, None, DEFAULT_MODULE)
        assert off.matches_inputs(store.index, store.doc, None, DEFAULT_MODULE)


def test_research_engine_has_no_fallback_surface():
    assert not hasattr(research_engine, "REFUSAL_FALLBACK_BETA")
