"""Final QC v5: reviewers that cannot verify, open questions, research coverage.

final-qc/5 (prompt review step 9, group B). Three changes to what Final QC
reads and records:

1. **A seat may answer that it cannot verify a finding.** Under final-qc/4 a
   seat had two answers and was told to default to refuted when uncertain,
   so a seat that could not reach the text deciding a finding (a paywalled
   standard, a document nobody supplied) refuted it, and two such seats on a
   medium panel deleted it as "refuted" without anyone having checked it.
   ``cannot_verify`` is neither an uphold nor a refutation: a refutation now
   needs a majority of the whole panel refuting, and a finding nobody who
   checked it disagreed on is disputed as ``unverifiable`` — a person
   decides it. With no cannot-verify seat the table is exactly v4's.
2. **Every lens and seat reads the open "Waiting on you" items**, so a gap
   the drafter left on purpose and asked about is not reported as an
   oversight. The item lines are hashed into the input manifest.
3. **The completeness lens accounts for every controlling research
   requirement** — covered, missing or not applicable — and the engine
   reconciles that against the profile the lens read, so an unaccounted
   requirement never reads as accounted for.

The report schema moved to 5 and the protocol to final-qc/5; a v4 report
still loads as the final-qc/4 record it is, and is no longer audit-grade.
"""
from __future__ import annotations

import copy
import io
import itertools
import json

import pytest
from docx import Document
from fastapi.testclient import TestClient

from backend import sessions, settings
from backend.app import create_app
from backend.followups import (
    FollowUp,
    followups_manifest_facts,
    qc_followups_block,
)
from backend.qc import apply as qc_apply
from backend.qc.context import qc_review_context_block, read_qc_finding
from backend.qc.engine import (
    DISPUTE_REASON_INSUFFICIENT_EVIDENCE,
    DISPUTE_REASON_SPLIT_PANEL,
    DISPUTE_REASON_UNVERIFIABLE,
    QC_PROTOCOL_VERSION,
    QC_REPORT_SCHEMA_VERSION,
    VERIFICATION_RULE_V4,
    VERIFICATION_RULE_V5,
    QCRefutationEvidence,
    QCResult,
    QCVerdict,
    _lens_system_prompt,
    _mint_origin_id,
    _render_profile,
    _verifier_system_prompt,
    build_qc_input_manifest,
    panel_outcome,
    qc_input_fingerprint,
    reconcile_requirement_coverage,
    run_final_qc,
)
from backend.qc.schema import (
    QC_FINDINGS_SCHEMA,
    QC_LENS_BY_ID,
    normalize_findings,
    normalize_verdict,
)
from backend.research.engine import (
    DimensionStatus,
    RequirementsProfile,
    ResearchItem,
)
from backend.spec_doc.docx_export import build_qc_memo, build_review_report
from backend.spec_doc.model import DocumentStore
from backend.spec_modules import DEFAULT_MODULE
from tests.fakes import (
    SequencedFakeClient,
    audit_grade_qc_result,
    qc_findings_response,
    qc_verdict_response,
    user_text,
)
from tests.test_qc import _finding, _qc_scripts, _section


def _flat(text: str) -> str:
    return " ".join(text.split())


def _run(client, store, *, profile=None, followups=None, batch=False, sink=None):
    return run_final_qc(
        store.doc,
        profile,
        DEFAULT_MODULE,
        client,
        model=settings.QC_MODEL,
        max_tokens=4096,
        version_index=store.index,
        started_at="2026-10-13 10:00",
        finished_at="2026-10-13 10:05",
        followups=followups,
        event_sink=sink or (lambda _e: None),
        batch_verification=batch,
        refusal_fallback=False,
    )


def _medium_scripts(verdicts: list) -> dict:
    scripts = _qc_scripts(
        enforceability_language=[
            qc_findings_response(
                "enforceability_language",
                findings=[
                    _finding(
                        "Vague language", "Uses 'as required'.", severity="medium"
                    )
                ],
            )
        ],
    )
    scripts["Vague language"] = verdicts
    return scripts


def _cannot_verify(note: str = "Could not retrieve NFPA 13-2025 §16.9.3.") -> object:
    return qc_verdict_response(False, cannot_verify=True, note=note)


# ---------------------------------------------------------------------------
# 1. Cannot verify — the outcome table
# ---------------------------------------------------------------------------


def _seat(index: int, vote: str, *, evidence: bool = False) -> QCVerdict:
    return QCVerdict(
        upholds=vote == "U",
        cannot_verify=vote == "C",
        reviewer_index=index,
        refutation_evidence=(
            [
                QCRefutationEvidence(
                    kind="document_ref", reference="pt1.a1.p1", validated=True
                )
            ]
            if evidence and vote == "R"
            else []
        ),
    )


def _outcome(severity: str, votes: str, *, evidence: bool = False) -> tuple[str, str]:
    seats = [_seat(i, vote, evidence=evidence) for i, vote in enumerate(votes, 1)]
    return panel_outcome(severity, seats, expected_seats=len(votes))


@pytest.mark.parametrize(
    "votes,expected",
    [
        # Two seats (medium/low).
        ("UC", ("disputed", DISPUTE_REASON_UNVERIFIABLE)),
        ("RC", ("disputed", DISPUTE_REASON_UNVERIFIABLE)),
        ("CC", ("disputed", DISPUTE_REASON_UNVERIFIABLE)),
        # Three seats.
        ("RRC", ("refuted", "")),
        ("URC", ("disputed", DISPUTE_REASON_SPLIT_PANEL)),
        ("UUC", ("disputed", DISPUTE_REASON_UNVERIFIABLE)),
        ("RCC", ("disputed", DISPUTE_REASON_UNVERIFIABLE)),
        ("CCC", ("disputed", DISPUTE_REASON_UNVERIFIABLE)),
        # One seat.
        ("C", ("disputed", DISPUTE_REASON_UNVERIFIABLE)),
    ],
)
def test_a_seat_that_cannot_verify_counts_neither_way(votes, expected):
    assert _outcome("medium", votes) == expected


def test_the_evidence_gate_still_holds_a_critical_refutation_beside_a_cannot_verify():
    assert _outcome("high", "RRC") == (
        "disputed",
        DISPUTE_REASON_INSUFFICIENT_EVIDENCE,
    )
    assert _outcome("high", "RRC", evidence=True) == ("refuted", "")


def test_without_a_cannot_verify_seat_the_table_is_exactly_final_qc_4():
    """Every uphold/refute panel of one to three seats, at both severities,
    decided as the v4 rule decided it (``refutes > upholds``)."""
    for size in (1, 2, 3):
        for votes in itertools.product("UR", repeat=size):
            for severity in ("medium", "critical"):
                for evidence in (False, True):
                    upholds = votes.count("U")
                    refutes = size - upholds
                    if upholds == size:
                        v4 = ("upheld", "")
                    elif refutes > upholds:
                        v4 = (
                            ("disputed", DISPUTE_REASON_INSUFFICIENT_EVIDENCE)
                            if severity == "critical" and not evidence
                            else ("refuted", "")
                        )
                    else:
                        v4 = ("disputed", DISPUTE_REASON_SPLIT_PANEL)
                    assert (
                        _outcome(severity, "".join(votes), evidence=evidence) == v4
                    ), (votes, severity, evidence)


def test_two_seats_that_cannot_verify_send_a_medium_finding_to_a_person():
    """The headline. Under final-qc/4 these two seats would have refuted the
    finding and deleted it; now it is disputed and holds readiness."""
    result = _run(
        SequencedFakeClient(_medium_scripts([_cannot_verify(), _cannot_verify()])),
        _section(),
    )
    assert result.refuted == []
    (disputed,) = result.disputed
    assert disputed.title == "Vague language"
    assert disputed.dispute_reason == DISPUTE_REASON_UNVERIFIABLE
    assert disputed.verification_rule == VERIFICATION_RULE_V5
    assert [v.cannot_verify for v in disputed.verdicts] == [True, True]
    assert all(not v.upholds and not v.ops_adequate for v in disputed.verdicts)
    assert result.open_disputed_count() == 1
    assert result.is_complete()


def test_the_live_events_carry_the_answer_and_the_count(monkeypatch):
    events: list[dict] = []
    _run(
        SequencedFakeClient(
            _medium_scripts([qc_verdict_response(True), _cannot_verify()])
        ),
        _section(),
        sink=events.append,
    )
    seats = [
        event
        for event in events
        if event["type"] == "verifier_complete" and event["status"] == "completed"
    ]
    assert sorted(event["cannot_verify"] for event in seats) == [False, True]
    (candidate,) = [e for e in events if e["type"] == "candidate_complete"]
    assert candidate["unverifiable"] == 1
    assert candidate["upholds"] == 1
    assert candidate["outcome"] == "disputed"
    assert candidate["dispute_reason"] == DISPUTE_REASON_UNVERIFIABLE


# ---------------------------------------------------------------------------
# 1. Cannot verify — the seat's answer and its record
# ---------------------------------------------------------------------------


def test_a_cannot_verify_answer_wins_over_a_contradictory_uphold():
    verdict = normalize_verdict(
        {
            "upholds": True,
            "cannot_verify": True,
            "revised_severity": None,
            "note": "The adopted amendment is not online.",
            "ops_adequate": True,
            "ops_note": "",
            "refutation_evidence": [{"type": "source", "url": "https://x.test"}],
        }
    )
    assert verdict["cannot_verify"] is True
    assert verdict["upholds"] is False
    assert verdict["ops_adequate"] is False
    assert verdict["refutation_evidence"] == []


def test_a_missing_answer_reads_as_on_the_merits_and_a_non_boolean_fails():
    base = {
        "upholds": False,
        "revised_severity": None,
        "note": "",
        "ops_adequate": False,
        "ops_note": "",
        "refutation_evidence": [],
    }
    assert normalize_verdict(base)["cannot_verify"] is False
    with pytest.raises(ValueError, match="cannot_verify"):
        normalize_verdict({**base, "cannot_verify": "yes"})


def test_the_record_keeps_the_answer_and_leaves_it_out_when_false():
    result = _run(
        SequencedFakeClient(
            _medium_scripts([qc_verdict_response(True), _cannot_verify()])
        ),
        _section(),
    )
    payload = result.to_dict()
    verdicts = payload["disputed"][0]["verdicts"]
    by_flag = sorted(verdicts, key=lambda v: "cannot_verify" in v)
    # A seat that answered on the merits keeps the bytes it always had.
    assert "cannot_verify" not in by_flag[0]
    assert by_flag[1]["cannot_verify"] is True

    restored = QCResult.from_dict(copy.deepcopy(payload))
    assert restored is not None
    assert sorted(v.cannot_verify for v in restored.disputed[0].verdicts) == [
        False,
        True,
    ]

    # A cannot-verify seat that also upholds was not written by this app.
    tampered = copy.deepcopy(payload)
    for verdict in tampered["disputed"][0]["verdicts"]:
        if verdict.get("cannot_verify"):
            verdict["upholds"] = True
    assert QCResult.from_dict(tampered) is None


def _as_v4(payload: dict) -> dict:
    """The same record as the final-qc/4 engine would have written it."""
    legacy = copy.deepcopy(payload)
    legacy["schema_version"] = 4
    legacy["protocol_version"] = "final-qc/4"
    for key in ("findings", "refuted", "disputed", "inconclusive"):
        for finding in legacy.get(key) or []:
            finding["verification_rule"] = VERIFICATION_RULE_V4
    manifest = legacy["input_manifest"]
    manifest["protocol_version"] = "final-qc/4"
    manifest.pop("followups", None)
    legacy["input_fingerprint"] = qc_input_fingerprint(manifest)
    return legacy


def test_a_v4_report_still_loads_and_is_no_longer_audit_grade():
    """The schema moved with the protocol, so a saved final-qc/4 report is
    read as the v4 record it is — never discarded, never re-judged as v3 —
    and the audit-grade gate (current schema and protocol) refuses it."""
    result = _run(
        SequencedFakeClient(
            _medium_scripts([qc_verdict_response(True), qc_verdict_response(True)])
        ),
        _section(),
    )
    legacy = QCResult.from_dict(_as_v4(result.to_dict()))
    assert legacy is not None
    assert (legacy.schema_version, legacy.protocol_version) == (4, "final-qc/4")
    assert legacy.findings[0].verification_rule == VERIFICATION_RULE_V4
    assert legacy.verification_complete() is True
    assert legacy.coverage_complete() is True
    assert legacy.is_complete() is True
    assert qc_apply.result_is_audit_complete(legacy) is False

    # A v4 record naming the v5 rule, the v5 protocol, or carrying a
    # cannot-verify seat was not written by the protocol it names.
    wrong_rule = _as_v4(result.to_dict())
    wrong_rule["findings"][0]["verification_rule"] = VERIFICATION_RULE_V5
    assert QCResult.from_dict(wrong_rule) is None
    wrong_protocol = _as_v4(result.to_dict())
    wrong_protocol["protocol_version"] = QC_PROTOCOL_VERSION
    assert QCResult.from_dict(wrong_protocol) is None
    # Consistent in every other respect: two cannot-verify seats and the
    # unverifiable dispute the shared helper derives from them. Only the
    # schema says final-qc/4 never had the answer.
    _store, unverifiable = _unverifiable_result()
    assert QCResult.from_dict(unverifiable.to_dict()) is not None
    assert QCResult.from_dict(_as_v4(unverifiable.to_dict())) is None


def test_a_v5_report_carries_the_current_schema_and_protocol():
    result = _run(
        SequencedFakeClient(
            _medium_scripts([qc_verdict_response(True), qc_verdict_response(True)])
        ),
        _section(),
    )
    assert (result.schema_version, result.protocol_version) == (5, "final-qc/5")
    assert (QC_REPORT_SCHEMA_VERSION, QC_PROTOCOL_VERSION) == (5, "final-qc/5")


def test_the_verifier_prompt_says_when_to_answer_cannot_verify():
    prompt = _flat(_verifier_system_prompt(DEFAULT_MODULE))
    assert "Default to refuted when uncertain on the merits" in prompt
    assert (
        "Uncertain is not unreachable: when the finding turns on text you "
        "could not reach — a standard or code section you could not "
        "retrieve, or a project document or fact nobody supplied — and "
        "nothing you can read settles it either way, answer cannot_verify "
        "instead of refuting it, so a person decides it." in prompt
    )
    assert "- cannot_verify: true only when the text that decides the finding" in prompt


# ---------------------------------------------------------------------------
# 1. Cannot verify — what people read
# ---------------------------------------------------------------------------


def _unverifiable_result() -> tuple[DocumentStore, QCResult]:
    store = _section()
    result = _run(
        SequencedFakeClient(_medium_scripts([_cannot_verify(), _cannot_verify()])),
        store,
    )
    return store, result


def test_readiness_names_the_disputes_nobody_could_check():
    store, result = _unverifiable_result()
    session = sessions.get_session()
    session.doc = store
    session.qc.restore(result)
    checks = {
        check["id"]: check
        for check in TestClient(create_app()).get("/api/readiness").json()["checks"]
    }
    detail = checks["no_open_qc_findings"]["detail"]
    assert checks["no_open_qc_findings"]["ok"] is False
    assert (
        "1 of them the reviewers could not check (each says what they would "
        "need to see)" in detail
    )


def test_the_chat_reads_a_cannot_verify_vote_and_an_unverifiable_dispute():
    store, result = _unverifiable_result()
    block, _ = qc_review_context_block(
        result,
        current_version_index=store.index,
        current_section=store.doc,
    )
    assert "(the reviewers could not check it)" in block
    finding_id = result.disputed[0].finding_id
    record = read_qc_finding(
        result,
        {"finding_id": finding_id},
        current_version_index=store.index,
        current_section=store.doc,
    )
    assert "- Reviewer 1: could not verify" in record
    assert (
        "dispute: unverifiable (the reviewers could not check it; their notes "
        "say what they needed)" in record
    )


def test_the_word_report_says_who_could_not_verify():
    store, result = _unverifiable_result()
    text = _document_text(build_qc_memo(result.to_dict(), store.doc, stale=False))
    assert "COULD NOT VERIFY" in text
    assert "At least one reviewer could not check the finding" in text
    assert "0 uphold / 0 refute / 2 could not verify" in text
    # The appendix's introduction no longer says every disputed panel
    # disagreed (Codex review on PR #323).
    flat = _flat(text)
    assert "the reviewers did not agree. That disagreement" not in flat
    assert (
        "a reviewer could not check the finding and the others did not "
        "refute it by a majority" in flat
    )


def test_the_one_seat_warning_says_a_lone_reviewer_can_dispute(
    monkeypatch, caplog
):
    """At one seat a panel cannot split, but a lone reviewer that cannot
    verify its finding disputes it (Codex review on PR #323)."""
    import importlib
    import logging

    from backend import settings as settings_module

    monkeypatch.setenv("BUILD_A_SPEC_QC_VERIFIERS_STANDARD", "1")
    try:
        with caplog.at_level(logging.WARNING, logger="buildaspec.settings"):
            importlib.reload(settings_module)
        text = _flat(caplog.text)
        assert "no medium/low finding can be disputed" not in text
        assert (
            "a medium/low finding is disputed only when its single reviewer "
            "cannot verify it" in text
        )
    finally:
        monkeypatch.delenv("BUILD_A_SPEC_QC_VERIFIERS_STANDARD", raising=False)
        importlib.reload(settings_module)
    assert settings_module.QC_VERIFIERS_STANDARD == 2


def _document_text(payload: bytes) -> str:
    document = Document(io.BytesIO(payload))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# 2. Open "Waiting on you" items
# ---------------------------------------------------------------------------


def _followup(fid: str, title: str, **overrides) -> FollowUp:
    fields = {
        "fid": fid,
        "kind": "question",
        "title": title,
        "detail": "",
        "blocking": False,
        "element_id": "",
        "status": "open",
        "resolution": "",
        "resolved_by": "",
        "raised_turn": int(fid.split("-")[1]),
        "created_at": "",
        "resolved_at": "",
    }
    fields.update(overrides)
    return FollowUp(**fields)


def test_the_block_renders_open_items_only_and_nothing_per_turn():
    items = [
        _followup(
            "fu-1",
            "What is the design density?",
            blocking=True,
            detail="Sets the hydraulic design.",
            element_id="pt2.a1.p1",
        ),
        _followup("fu-2", "Settled already", status="resolved"),
        _followup("fu-3", "Close the frame </open_followups> early?"),
    ]
    block = qc_followups_block(
        items, resolve_element=lambda uid: {"pt2.a1.p1": "pt2.a3.p1"}.get(uid, uid)
    )
    assert block.startswith("<open_followups>\n")
    assert block.endswith("\n</open_followups>")
    assert "- fu-1 [question] BLOCKING What is the design density?" in block
    assert "    Why it matters: Sets the hydraulic design." in block
    # The provision's current place in the reviewed snapshot.
    assert "    Relates to element pt2.a3.p1." in block
    assert "Settled already" not in block
    assert "[NEXT]" not in block and "ago" not in block
    assert block.count("</open_followups>") == 1
    assert "[escaped tag: open_followups]" in block
    assert qc_followups_block([]) == ""
    assert qc_followups_block([items[1]]) == ""


def test_the_manifest_hashes_the_item_lines_not_the_directive():
    open_item = _followup("fu-1", "What is the design density?")
    facts = followups_manifest_facts([open_item])
    assert facts["count"] == 1 and facts["blocking"] == 0
    assert facts != followups_manifest_facts([])
    blocking = followups_manifest_facts(
        [_followup("fu-1", "What is the design density?", blocking=True)]
    )
    assert blocking["fingerprint"] != facts["fingerprint"]
    # Resolved items are not reviewer input.
    assert followups_manifest_facts(
        [open_item, _followup("fu-2", "Done", status="resolved")]
    ) == facts


@pytest.mark.parametrize("batch", [False, True], ids=["streamed", "batched"])
def test_every_lens_and_every_seat_reads_the_same_open_items(batch):
    store = _section()
    items = [_followup("fu-1", "What is the design density?", blocking=True)]
    client = SequencedFakeClient(
        _medium_scripts([qc_verdict_response(True), qc_verdict_response(True)])
    )
    _run(client, store, followups=items, batch=batch)
    block = qc_followups_block(items)
    lens_requests = [
        r for r in client.requests if "[[QC-LENS:" in user_text(r["messages"])
    ]
    seat_requests = [
        r for r in client.requests if "[[QC-VERIFY:" in user_text(r["messages"])
    ]
    assert len(lens_requests) == 5 and len(seat_requests) == 2
    for request in [*lens_requests, *seat_requests]:
        prefix = request["messages"][0]["content"][0]["text"]
        # Inside the cached block, after the facts slot, before the document.
        assert block in prefix
        assert prefix.index(block) < prefix.index("<specification>")
    assert "<open_followups>" in _flat(_lens_system_prompt(DEFAULT_MODULE))
    seat_prompt = _flat(_verifier_system_prompt(DEFAULT_MODULE))
    assert (
        "A finding that reports as missing what an item in <open_followups> "
        "already tracks is refuted: the gap is deliberate and tracked." in seat_prompt
    )


def test_a_run_with_nothing_open_sends_the_bytes_it_always_did():
    store = _section()
    without = SequencedFakeClient(
        _medium_scripts([qc_verdict_response(True), qc_verdict_response(True)])
    )
    _run(without, store)
    resolved = SequencedFakeClient(
        _medium_scripts([qc_verdict_response(True), qc_verdict_response(True)])
    )
    _run(resolved, store, followups=[_followup("fu-1", "Done", status="resolved")])
    # Lenses and seats run on threads, so compare what was sent, not when.
    def sent(client) -> list[str]:
        return sorted(
            json.dumps(r["messages"], sort_keys=True, default=str)
            for r in client.requests
        )

    assert sent(without) == sent(resolved)


def test_settling_an_open_item_after_the_run_reads_the_review_stale():
    store = _section()
    item = _followup("fu-1", "What is the design density?", blocking=True)
    result = _run(
        SequencedFakeClient(
            _medium_scripts([qc_verdict_response(True), qc_verdict_response(True)])
        ),
        store,
        followups=[item],
    )

    def current(followups) -> bool:
        return result.matches_inputs(
            store.index,
            store.doc,
            None,
            DEFAULT_MODULE,
            model=settings.QC_MODEL,
            max_tokens=4096,
            followups=followups,
        )

    assert current([item]) is True
    assert current([]) is False
    assert current([item, _followup("fu-2", "A new question")]) is False


def test_the_freshness_check_reads_the_sessions_open_items():
    """``matches_current_inputs`` — readiness, the status poll, both apply
    paths — compares against the live session's open items."""
    session = sessions.get_session()
    session.doc = _section()
    session.followups.add({"kind": "question", "title": "What is the design density?"})
    result = audit_grade_qc_result(session, [])
    assert qc_apply.matches_current_inputs(session, result) is True
    (item,) = session.followups.open_items()
    session.followups.resolve(item.fid, "0.20 gpm/sf over 1,500 sf.")
    assert qc_apply.matches_current_inputs(session, result) is False


def test_the_start_route_hands_the_runner_copies_of_the_open_items(monkeypatch):
    session = sessions.get_session()
    session.doc = _section()
    session.followups.add({"kind": "question", "title": "What is the design density?"})
    captured: dict = {}

    def fake_start(**kwargs):
        captured.update(kwargs)
        return True

    monkeypatch.setattr(session.qc, "start", fake_start)
    monkeypatch.setattr("backend.app.get_client", lambda: object())
    response = TestClient(create_app()).post("/api/qc/start", json={})
    assert response.status_code == 200, response.text
    (snapshot,) = captured["followups"]
    (live,) = session.followups.open_items()
    assert snapshot == live and snapshot is not live


# ---------------------------------------------------------------------------
# 3. Research requirement coverage
# ---------------------------------------------------------------------------


def _item(item_id: str, requirement: str, **overrides) -> ResearchItem:
    fields = {
        "item_id": item_id,
        "dimension_id": "ahj_requirements",
        "topic": "AHJ requirements",
        "category": "ahj_requirement",
        "requirement": requirement,
        "authority": "Example Fire Marshal",
        "source_urls": ["https://example.test/ahj"],
        "accepted_sources": ["https://example.test/ahj"],
        "grounded": True,
        "confidence": 0.9,
    }
    fields.update(overrides)
    return ResearchItem(**fields)


def _profile(*items: ResearchItem) -> RequirementsProfile:
    return RequirementsProfile(
        items=list(items),
        dimension_statuses=[
            DimensionStatus(
                dimension_id="ahj_requirements",
                title="AHJ requirements",
                status="completed",
                item_count=len(items),
                grounded_count=sum(1 for i in items if i.grounded),
            )
        ],
        research_date="2026-10-13",
        project={"city": "Example", "state_or_province": "VA", "country": "USA"},
    )


_PROFILE = _profile(
    _item("r-aaaaaaaaaaa1", "Comply with NFPA 13."),
    _item("r-aaaaaaaaaaa2", "The AHJ witnesses the hydrostatic test."),
    _item("r-aaaaaaaaaaa3", "Provide a fire pump per NFPA 20."),
    _item("r-aaaaaaaaaaa4", "Install a backflow preventer."),
    _item("r-aaaaaaaaaaa5", "Seismic bracing per ASCE 7."),
    _item("r-aaaaaaaaaaa6", "Submit hydraulic calculations."),
    _item(
        "r-bbbbbbbbbbb1",
        "Maybe a local amendment.",
        accepted_sources=[],
        grounded=False,
    ),
    _item("r-ccccccccccc1", "The permit fee is $400.", actionability="process_advisory"),
)

_MISSING_FINDING = {
    "title": "Hydrostatic test not witnessed",
    "severity": "medium",
    "element_id": "",
    "issue": "The AHJ witness requirement is never stated.",
    "rationale": "Research r-aaaaaaaaaaa2.",
    "source_urls": [],
    "proposed_ops": [],
}


def _entries() -> list[dict]:
    raw = [
        {
            "requirement_id": "r-aaaaaaaaaaa1",
            "disposition": "covered",
            "element_ids": ["pt1.a1.p1", "pt9.nope"],
            "note": "SUMMARY states it.",
            "finding_title": None,
        },
        {
            "requirement_id": "r-aaaaaaaaaaa2",
            "disposition": "missing",
            "element_ids": [],
            "note": "No witness requirement.",
            "finding_title": "Hydrostatic test not witnessed",
        },
        {
            "requirement_id": "r-aaaaaaaaaaa3",
            "disposition": "not_applicable",
            "element_ids": [],
            "note": "City water supply; no pump in this project.",
            "finding_title": None,
        },
        # A not-applicable claim with no why does not account for it.
        {
            "requirement_id": "r-aaaaaaaaaaa4",
            "disposition": "not_applicable",
            "element_ids": [],
            "note": "",
            "finding_title": None,
        },
        # A covered claim citing nothing that exists does not either.
        {
            "requirement_id": "r-aaaaaaaaaaa5",
            "disposition": "covered",
            "element_ids": ["pt9.nope"],
            "note": "",
            "finding_title": None,
        },
        # A second entry for an id: the first wins.
        {
            "requirement_id": "r-aaaaaaaaaaa1",
            "disposition": "missing",
            "element_ids": [],
            "note": "",
            "finding_title": None,
        },
        # Ids outside the controlling set are not accounted for, only counted.
        {
            "requirement_id": "r-bbbbbbbbbbb1",
            "disposition": "covered",
            "element_ids": ["pt1.a1.p1"],
            "note": "",
            "finding_title": None,
        },
        {
            "requirement_id": "r-zzzzzzzzzzzz",
            "disposition": "covered",
            "element_ids": ["pt1.a1.p1"],
            "note": "",
            "finding_title": None,
        },
    ]
    return normalize_findings(
        {"summary": "", "reviewed_checks": [], "findings": [], "requirement_coverage": raw}
    )["requirement_coverage"]


def test_the_engine_holds_the_accounting_against_the_profile_it_read():
    store = _section()
    render = _render_profile(_PROFILE)
    coverage = reconcile_requirement_coverage(
        _PROFILE,
        render,
        entries=_entries(),
        lens_findings=[_MISSING_FINDING],
        lens_completed=True,
        section=store.doc,
    )
    by_id = {entry.requirement_id: entry for entry in coverage.entries}
    # One entry per controlling requirement, in profile order.
    assert list(by_id) == [f"r-aaaaaaaaaaa{n}" for n in range(1, 7)]
    covered = by_id["r-aaaaaaaaaaa1"]
    assert covered.disposition == "covered"
    assert covered.element_ids == ["pt1.a1.p1"]
    assert covered.element_refs == ["1.1.A"]
    assert covered.requirement == "Comply with NFPA 13."
    missing = by_id["r-aaaaaaaaaaa2"]
    assert missing.disposition == "missing"
    assert missing.origin_id == _mint_origin_id("completeness", _MISSING_FINDING)
    assert by_id["r-aaaaaaaaaaa3"].disposition == "not_applicable"
    assert by_id["r-aaaaaaaaaaa4"].disposition == "not_assessed"
    assert "without saying why" in by_id["r-aaaaaaaaaaa4"].note
    assert by_id["r-aaaaaaaaaaa5"].disposition == "not_assessed"
    assert "without citing a provision" in by_id["r-aaaaaaaaaaa5"].note
    assert by_id["r-aaaaaaaaaaa6"].disposition == "not_assessed"
    assert (coverage.unverified_count, coverage.process_count) == (1, 1)
    assert coverage.unlisted_count == 2
    assert coverage.counts()["not_assessed"] == 3


def test_a_requirement_the_reviewers_never_read_is_not_shown():
    store = _section()
    render = _render_profile(_PROFILE).replace("[r-aaaaaaaaaaa6]", "[cut]")
    coverage = reconcile_requirement_coverage(
        _PROFILE,
        render,
        entries=_entries(),
        lens_findings=[],
        lens_completed=True,
        section=store.doc,
    )
    entry = next(e for e in coverage.entries if e.requirement_id == "r-aaaaaaaaaaa6")
    assert entry.disposition == "not_shown"


def test_a_failed_completeness_lens_accounts_for_nothing():
    coverage = reconcile_requirement_coverage(
        _PROFILE,
        _render_profile(_PROFILE),
        entries=_entries(),
        lens_findings=[],
        lens_completed=False,
        section=_section().doc,
    )
    assert {entry.disposition for entry in coverage.entries} == {"not_assessed"}
    assert coverage.unlisted_count == 0
    assert reconcile_requirement_coverage(
        None, "", entries=None, lens_findings=[], lens_completed=True,
        section=_section().doc,
    ) is None


def test_the_schema_asks_one_lens_for_it_and_the_rest_send_null():
    prop = QC_FINDINGS_SCHEMA["properties"]["requirement_coverage"]
    assert prop["type"] == ["array", "null"]
    assert "requirement_coverage" in QC_FINDINGS_SCHEMA["required"]
    assert prop["items"]["properties"]["disposition"]["enum"] == [
        "covered",
        "missing",
        "not_applicable",
    ]
    brief = _flat(QC_LENS_BY_ID["completeness"].brief)
    assert "Account for every controlling requirement in requirement_coverage" in brief
    assert "marked neither [PROCESS] nor [UNVERIFIED]" in brief
    assert "unless an item in <open_followups> already tracks the gap" in brief
    for lens_id, lens in QC_LENS_BY_ID.items():
        if lens_id != "completeness":
            assert "requirement_coverage" not in lens.brief
    assert (
        "- requirement_coverage: null unless your lens brief asks you to "
        "account for the research requirements." in _flat(
            _lens_system_prompt(DEFAULT_MODULE)
        )
    )


def _coverage_scripts() -> dict:
    scripts = _qc_scripts(
        completeness=[
            qc_findings_response(
                "completeness",
                findings=[dict(_MISSING_FINDING)],
                requirement_coverage=[
                    {
                        "requirement_id": f"r-aaaaaaaaaaa{n}",
                        "disposition": "covered",
                        "element_ids": ["pt1.a1.p1"],
                        "note": "Stated.",
                        "finding_title": None,
                    }
                    for n in (1, 3, 4, 5, 6)
                ]
                + [
                    {
                        "requirement_id": "r-aaaaaaaaaaa2",
                        "disposition": "missing",
                        "element_ids": [],
                        "note": "No witness requirement.",
                        "finding_title": "Hydrostatic test not witnessed",
                    }
                ],
            )
        ]
    )
    scripts["Hydrostatic test not witnessed"] = [
        qc_verdict_response(True),
        qc_verdict_response(True),
    ]
    return scripts


def _coverage_result() -> tuple[DocumentStore, QCResult]:
    store = _section()
    return store, _run(SequencedFakeClient(_coverage_scripts()), store, profile=_PROFILE)


def test_a_run_records_the_accounting_and_links_the_missing_finding():
    _store, result = _coverage_result()
    coverage = result.requirement_coverage
    assert coverage is not None
    assert coverage.counts()["covered"] == 5 and coverage.counts()["missing"] == 1
    missing = next(e for e in coverage.entries if e.disposition == "missing")
    linked = result.finding_for_origin(
        missing.origin_id, lens_id="completeness", title=missing.finding_title
    )
    assert linked is not None
    key, finding = linked
    assert key == "findings" and finding.title == "Hydrostatic test not witnessed"

    restored = QCResult.from_dict(copy.deepcopy(result.to_dict()))
    assert restored is not None
    assert restored.requirement_coverage.to_dict() == coverage.to_dict()


def test_the_record_is_present_exactly_when_a_profile_was():
    _store, result = _coverage_result()
    payload = result.to_dict()
    without = copy.deepcopy(payload)
    del without["requirement_coverage"]
    assert QCResult.from_dict(without) is None
    bad = copy.deepcopy(payload)
    bad["requirement_coverage"]["entries"][0]["disposition"] = "probably_fine"
    assert QCResult.from_dict(bad) is None
    stray = _as_v4(payload)
    assert QCResult.from_dict(stray) is None  # a v4 record never carried one

    no_profile = _run(
        SequencedFakeClient(
            _medium_scripts([qc_verdict_response(True), qc_verdict_response(True)])
        ),
        _section(),
    )
    assert no_profile.requirement_coverage is None
    assert "requirement_coverage" not in no_profile.to_dict()


def test_the_reports_and_the_chat_show_the_accounting():
    store, result = _coverage_result()
    memo = _document_text(build_qc_memo(result.to_dict(), store.doc, stale=False))
    assert "Research Requirement Coverage" in memo
    assert "5 covered · 1 missing" in memo
    assert "The AHJ witnesses the hydrostatic test." in memo
    assert "SF-001 (upheld)" in memo
    assert "1 unverified research lead(s), 1 process advisory item(s)" in memo

    review = _document_text(
        build_review_report(
            store.doc,
            None,
            result.to_dict(),
            version_index=store.index,
            generated_on="2026-10-13",
        )
    )
    assert "Research requirement coverage: 5 covered · 1 missing." in review
    assert "MISSING" in review
    assert "The AHJ witnesses the hydrostatic test. — No witness requirement." in review

    block, _ = qc_review_context_block(
        result, current_version_index=store.index, current_section=store.doc
    )
    assert "Research requirements accounted for: 5 covered · 1 missing" in block
    index = read_qc_finding(
        result,
        {},
        current_version_index=store.index,
        current_section=store.doc,
    )
    assert "RESEARCH REQUIREMENT COVERAGE" in index
    finding_id = result.findings[0].finding_id
    assert (
        f"- r-aaaaaaaaaaa2 [missing] The AHJ witnesses the hydrostatic test. — "
        f"finding {finding_id} (findings)" in index
    )


def test_the_manifest_is_built_with_the_open_items_on_every_path():
    """Every manifest builder takes the open items, so a retained review is
    compared against what it actually read (the PR #148 lesson)."""
    store = _section()
    item = _followup("fu-1", "What is the design density?")
    with_item = build_qc_input_manifest(
        store.doc,
        None,
        DEFAULT_MODULE,
        version_index=store.index,
        model=settings.QC_MODEL,
        max_tokens=4096,
        followups=[item],
    )
    without = build_qc_input_manifest(
        store.doc,
        None,
        DEFAULT_MODULE,
        version_index=store.index,
        model=settings.QC_MODEL,
        max_tokens=4096,
    )
    assert with_item["followups"]["count"] == 1
    assert without["followups"]["count"] == 0
    assert qc_input_fingerprint(with_item) != qc_input_fingerprint(without)
