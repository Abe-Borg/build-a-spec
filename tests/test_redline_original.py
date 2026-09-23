"""Redline on your original (Phase 1): the promise, row by row.

The export hands over a copy of the uploaded Word file with every change
since the import as a Word tracked change, and promises two things:

* **Accept All gives exactly what Export Word (keeps your formatting)
  produces** — checked here against ``render_preserving_docx`` itself, not
  against the redline builder's own idea of the clean body;
* **Reject All gives back the upload**, except that a moved provision's
  bookmarks stay with its new copy (D-6).

Every row of the plan's invariant matrix asserts both, with the oracles of
``backend.spec_doc.revisions`` — which share no code with the writer — and
the package-level invariants: every other member byte-identical, revision
ids unique and above the package's own, author and date on every change,
Word's schema order, and no tracked change ever deleting a paragraph mark
that holds a section break.
"""
from __future__ import annotations

import collections
import copy
import io
import random
import zipfile

import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from lxml import etree

from backend import sessions
from backend.app import create_app
from backend.spec_doc.docx_export import upload_redline_filename
from backend.spec_doc.importer import parse_master_docx
from backend.spec_doc.model import SpecEditError, SpecSection, apply_edits
from backend.spec_doc.revisions import (
    REVISION_TAGS,
    accept_all,
    first_difference,
    has_revisions,
    reject_all,
)
from backend.spec_doc.source_render import (
    SourceRedlineError,
    render_preserving_docx,
    render_preserving_redline,
)
from backend.spec_doc.source_splice import FALLBACK_REASONS
from tests.test_preserving_export import (
    HEADER_TEXT,
    _break_master,
    _hold_break,
    _import,
    _master_bytes,
    _not_used_master,
    _save,
    _typed_letter_master,
)

AUTHOR = "Build-a-Spec"
DATE = "2026-09-22T12:00:00Z"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


# ---------------------------------------------------------------------------
# The harness
# ---------------------------------------------------------------------------


def _parse(tmp_path, source: bytes, name: str = "master.docx"):
    path = tmp_path / name
    path.write_bytes(source)
    return parse_master_docx(path)


def _member(payload: bytes, name: str) -> bytes:
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        return archive.read(name)


def _body(payload: bytes):
    return etree.fromstring(_member(payload, "word/document.xml")).find(qn("w:body"))


def _members(payload: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _bookmarks(body) -> set[str]:
    return {b.get(qn("w:name")) for b in body.iter(qn("w:bookmarkStart"))}


def _highest_id(source: bytes) -> int:
    highest = 0
    for name, data in _members(source).items():
        if not (name.startswith("word/") and name.endswith(".xml")):
            continue
        for element in etree.fromstring(data).iter():
            value = element.get(qn("w:id")) if isinstance(element.tag, str) else None
            if value is not None and value.isdigit():
                highest = max(highest, int(value))
    return highest


def _revisions(body) -> list:
    return [
        element
        for element in body.iter()
        if isinstance(element.tag, str)
        and element.tag
        in (qn("w:ins"), qn("w:del"), qn("w:pPrChange"), qn("w:rPrChange"))
    ]


def _assert_schema_order(body) -> None:
    """What Word is strict about (D-6)."""
    flags = (qn("w:ins"), qn("w:del"))
    for properties in body.iter(qn("w:pPr")):
        tags = [c.tag for c in properties if isinstance(c.tag, str)]
        if qn("w:rPr") in tags:
            position = tags.index(qn("w:rPr"))
            assert all(
                t in (qn("w:sectPr"), qn("w:pPrChange")) for t in tags[position + 1 :]
            ), tags
        if qn("w:pPrChange") in tags:
            assert tags[-1] == qn("w:pPrChange"), tags
        mark = properties.find(qn("w:rPr"))
        if mark is not None:
            mark_tags = [c.tag for c in mark if isinstance(c.tag, str)]
            flagged = [i for i, t in enumerate(mark_tags) if t in flags]
            assert flagged in ([], [0]), mark_tags  # the flag leads CT_ParaRPr
            if qn("w:rPrChange") in mark_tags:  # ...and its change ends it
                assert mark_tags[-1] == qn("w:rPrChange"), mark_tags
    for properties in body.iter(qn("w:trPr")):
        tags = [c.tag for c in properties if isinstance(c.tag, str)]
        flagged = [i for i, t in enumerate(tags) if t in flags]
        others = [i for i, t in enumerate(tags) if t not in flags and t != qn("w:trPrChange")]
        assert all(f > o for f in flagged for o in others), tags


def _inserted_bookmarks(body) -> set[str]:
    """Bookmarks that sit in content a ``w:ins`` wraps — on a moved copy."""
    names = set()
    for start in body.iter(qn("w:bookmarkStart")):
        if any(
            a.tag == qn("w:ins") and a.getparent().tag != qn("w:rPr")
            for a in start.iterancestors()
        ):
            names.add(start.get(qn("w:name")))
    return names


def _verify(
    source: bytes,
    imported,
    section,
    *,
    moved_bookmarks=frozenset(),
    allow_moved_bookmarks: bool = False,
):
    """Render the redline and assert the whole promise. Returns the redline
    bytes and the export stats.

    ``moved_bookmarks`` names the bookmarks a test expects to travel with a
    moved copy (the one Reject-All limit, D-6); ``allow_moved_bookmarks``
    instead admits whichever bookmarks the redline itself carries on an
    inserted copy — for sweeps whose edits are not written by hand."""
    stats: dict = {}
    redline = render_preserving_redline(
        source_bytes=source,
        format_map=imported.format_map,
        baseline=imported.section,
        current=section,
        author=AUTHOR,
        date=DATE,
        stats=stats,
    )
    clean = render_preserving_docx(
        source_bytes=source, format_map=imported.format_map, current=section
    )
    r_body, c_body, u_body = _body(redline), _body(clean), _body(source)
    if allow_moved_bookmarks:
        moved_bookmarks = set(moved_bookmarks) | _inserted_bookmarks(r_body)

    # Accept All == the formatted export; Reject All == the upload.
    assert first_difference(accept_all(r_body), c_body) is None, first_difference(
        accept_all(r_body), c_body
    )
    rejected = reject_all(r_body)
    lost = _bookmarks(u_body) - _bookmarks(rejected)
    assert lost <= set(moved_bookmarks), lost
    assert first_difference(rejected, u_body, exclude_bookmarks=lost) is None, (
        first_difference(rejected, u_body, exclude_bookmarks=lost)
    )
    assert not has_revisions(accept_all(r_body))
    assert not has_revisions(rejected)

    # Every member but the body is the upload's, byte for byte.
    before, after = _members(source), _members(redline)
    assert list(before) == list(after)
    assert [n for n in before if before[n] != after[n]] in (
        [],
        ["word/document.xml"],
    )

    # Revision metadata (D-6).
    revisions = _revisions(r_body)
    ids = [int(r.get(qn("w:id"))) for r in revisions]
    assert len(ids) == len(set(ids))
    assert all(i > _highest_id(source) for i in ids)
    assert all(r.get(qn("w:author")) == AUTHOR for r in revisions)
    assert all(r.get(qn("w:date")) == DATE for r in revisions)
    _assert_schema_order(r_body)

    # No tracked change ever deletes a mark holding a section break.
    for paragraph in r_body.iter(qn("w:p")):
        properties = paragraph.find(qn("w:pPr"))
        if properties is None or properties.find(qn("w:sectPr")) is None:
            continue
        mark = properties.find(qn("w:rPr"))
        assert mark is None or mark.find(qn("w:del")) is None
    return redline, stats


def _tracked(payload: bytes, tag: str) -> list[str]:
    """The text inside every ``w:ins``/``w:del`` of the redline, in order."""
    body = _body(payload)
    text_tags = (qn("w:t"), qn("w:delText"))
    return [
        "".join(t.text or "" for t in wrapper.iter(*text_tags))
        for wrapper in body.iter(qn(tag))
        if wrapper.getparent().tag != qn("w:rPr")
    ]


def _edit(section, *ops):
    edited, _ = apply_edits(section, list(ops))
    return edited


# ---------------------------------------------------------------------------
# The invariant matrix
# ---------------------------------------------------------------------------


def test_no_edits_means_no_tracked_changes(tmp_path):
    source = _master_bytes()
    imported = _parse(tmp_path, source)
    redline, stats = _verify(source, imported, imported.section)
    assert _revisions(_body(redline)) == []
    assert stats["redline"]["revisions"] == 0
    assert stats["redline"]["last_mark_untracked"] == 0


def test_a_one_word_edit(tmp_path):
    source = _master_bytes()
    imported = _parse(tmp_path, source)
    first = imported.section.parts[0].articles[0].paragraphs[0]
    section = _edit(
        imported.section,
        {
            "action": "replace",
            "target_id": first.uid,
            "text": "Section includes seismic isolation for mechanical equipment.",
        },
    )
    redline, stats = _verify(source, imported, section)
    assert _tracked(redline, "w:del") == ["vibration"]
    assert _tracked(redline, "w:ins") == ["seismic"]
    assert stats["redline"]["spliced"] == 1


def test_an_edit_inside_a_bold_phrase(tmp_path):
    source = _typed_letter_master()
    imported = _parse(tmp_path, source)
    first = imported.section.parts[0].articles[0].paragraphs[0]
    section = _edit(
        imported.section,
        {
            "action": "replace",
            "target_id": first.uid,
            "text": "Section includes seismic isolation for mechanical equipment.",
        },
    )
    redline, _ = _verify(source, imported, section)
    (inserted,) = [
        w for w in _body(redline).iter(qn("w:ins")) if w.getparent().tag != qn("w:rPr")
    ]
    assert inserted.find(f".//{qn('w:b')}") is not None  # typed over bold


def test_a_deletion_at_a_run_boundary(tmp_path):
    """The bold phrase is its own run; deleting it deletes that run."""
    source = _typed_letter_master()
    imported = _parse(tmp_path, source)
    first = imported.section.parts[0].articles[0].paragraphs[0]
    section = _edit(
        imported.section,
        {
            "action": "replace",
            "target_id": first.uid,
            "text": "Section includes for mechanical equipment.",
        },
    )
    redline, _ = _verify(source, imported, section)
    assert _tracked(redline, "w:del") == ["vibration isolation "]


@pytest.mark.parametrize(
    "text",
    [
        "Also Section includes vibration isolation for mechanical equipment.",
        "Section includes vibration isolation for mechanical equipment. Comply.",
    ],
    ids=["at-the-start", "at-the-end"],
)
def test_an_insertion_at_the_start_or_end(tmp_path, text):
    source = _master_bytes()
    imported = _parse(tmp_path, source)
    first = imported.section.parts[0].articles[0].paragraphs[0]
    section = _edit(
        imported.section, {"action": "replace", "target_id": first.uid, "text": text}
    )
    redline, _ = _verify(source, imported, section)
    assert _tracked(redline, "w:del") == []
    assert len(_tracked(redline, "w:ins")) == 1


def test_relettering_in_a_typed_letter_master_tracks_the_letters(tmp_path):
    """Decision 2: a relettered provision carries a tracked letter change —
    the only way Reject All gives back your letters."""
    source = _typed_letter_master()
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    section = _edit(
        imported.section,
        {
            "action": "add_paragraph",
            "target_id": article.uid,
            "position": 0,
            "text": "Provide seismic restraints.",
        },
    )
    redline, _ = _verify(source, imported, section)
    assert _tracked(redline, "w:del") == ["A.", "B."]
    inserted = _tracked(redline, "w:ins")
    assert "B." in inserted and "C." in inserted
    assert any("Provide seismic restraints." in text for text in inserted)


def _numbered_master() -> bytes:
    from tests.test_importer import _define_numbering, _numbered

    document = Document()
    document.add_paragraph("SECTION 23 05 48")
    document.add_paragraph("VIBRATION CONTROLS")
    document.add_paragraph("PART 1 - GENERAL")
    _define_numbering(
        document, 50, {1: ("decimal", "%1.%2"), 2: ("upperLetter", "%3.")}
    )
    _numbered(document, "SUMMARY", 1, "50")
    _numbered(document, "Section includes vibration isolation.", 2, "50")
    _numbered(document, "Related requirements.", 2, "50")
    _numbered(document, "Provide isolators as scheduled.", 2, "50")
    document.add_paragraph("END OF SECTION")
    return _save(document)


def test_insert_and_delete_in_a_word_numbered_master(tmp_path):
    """Word numbers these itself: no letter is tracked, only the words."""
    source = _numbered_master()
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    section = _edit(
        imported.section,
        {"action": "delete", "target_id": article.paragraphs[1].uid},
        {
            "action": "add_paragraph",
            "target_id": article.uid,
            "position": 0,
            "text": "Provide seismic restraints.",
        },
    )
    redline, _ = _verify(source, imported, section)
    assert _tracked(redline, "w:del") == ["Related requirements."]
    assert _tracked(redline, "w:ins") == ["Provide seismic restraints."]


def _nested_master() -> bytes:
    document = Document()
    for line in (
        "SECTION 23 05 48",
        "VIBRATION CONTROLS",
        "PART 1 - GENERAL",
        "1.1 SUMMARY",
        "A. Section includes:",
        "1. Spring isolators.",
        "2. Seismic restraints.",
        "B. Related requirements.",
        "1.2 SUBMITTALS",
        "A. Product data.",
        "END OF SECTION",
    ):
        document.add_paragraph(line)
    return _save(document)


def test_deleting_a_paragraph_that_has_children(tmp_path):
    source = _nested_master()
    imported = _parse(tmp_path, source)
    parent = imported.section.parts[0].articles[0].paragraphs[0]
    assert len(parent.children) == 2
    section = _edit(imported.section, {"action": "delete", "target_id": parent.uid})
    redline, _ = _verify(source, imported, section)
    deleted = _tracked(redline, "w:del")
    assert deleted[:3] == [
        "A. Section includes:",
        "1. Spring isolators.",
        "2. Seismic restraints.",
    ]


def test_deleting_an_article(tmp_path):
    source = _nested_master()
    imported = _parse(tmp_path, source)
    submittals = imported.section.parts[0].articles[1]
    section = _edit(imported.section, {"action": "delete", "target_id": submittals.uid})
    redline, _ = _verify(source, imported, section)
    assert _tracked(redline, "w:del") == ["1.2 SUBMITTALS", "A. Product data."]


def test_deleting_a_table(tmp_path):
    source = _master_bytes()
    imported = _parse(tmp_path, source)
    schedule = imported.section.parts[0].articles[1]
    table = next(p for p in schedule.paragraphs if p.locked == "table")
    section = _edit(imported.section, {"action": "delete", "target_id": table.uid})
    redline, _ = _verify(source, imported, section)
    rows = list(_body(redline).iter(qn("w:tr")))
    assert rows and all(
        r.find(f"{qn('w:trPr')}/{qn('w:del')}") is not None for r in rows
    )


def test_a_pure_reorder_of_provisions_is_tracked_so_reject_restores_the_order(
    tmp_path,
):
    """The case the diff alone cannot see: every text unchanged."""
    source = _nested_master()
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    section = _edit(
        imported.section,
        {"action": "move", "target_id": article.paragraphs[1].uid, "position": 0},
    )
    redline, stats = _verify(source, imported, section)
    assert stats["redline"]["moved"] >= 1
    assert "A. Related requirements." in _tracked(redline, "w:ins")
    assert "B. Related requirements." in _tracked(redline, "w:del")


def _heavy_first_master() -> bytes:
    """A provision with two children, then two leaf siblings."""
    document = Document()
    for line in (
        "SECTION 23 05 48",
        "VIBRATION CONTROLS",
        "PART 1 - GENERAL",
        "1.1 SUMMARY",
        "A. Section includes:",
        "1. Spring isolators.",
        "2. Seismic restraints.",
        "B. Related requirements.",
        "C. Quality assurance.",
        "END OF SECTION",
    ):
        document.add_paragraph(line)
    return _save(document)


def test_the_redline_moves_what_the_diff_says_moved(tmp_path):
    """Moving the first provision (with its two children) below its two
    siblings: the diff keeps the LONGER run of siblings in place and reports
    the parent moved, subtree and all. The raw body-level chain would rather
    keep the parent's three paragraphs and move the two siblings — the
    redline follows the diff wherever no section break forbids it."""
    source = _heavy_first_master()
    imported = _parse(tmp_path, source)
    parent = imported.section.parts[0].articles[0].paragraphs[0]
    section = _edit(imported.section, {"action": "move", "target_id": parent.uid, "position": 2})
    redline, stats = _verify(source, imported, section)
    assert stats["redline"]["moved"] == 3  # the parent and both children
    assert stats["redline"]["moves_added"] == 0
    deleted, inserted = _tracked(redline, "w:del"), _tracked(redline, "w:ins")
    assert "A. Section includes:" in deleted
    assert any(text.endswith("Section includes:") for text in inserted)
    assert "1. Spring isolators." in deleted and "1. Spring isolators." in inserted


def test_a_section_break_stays_in_place_even_when_a_longer_chain_skips_it():
    """``_keep_in_place``: a record pinned around a break outweighs any
    number of others — here two elements that came from below the break
    would make a longer chain without it."""
    from backend.spec_doc.source_render import _keep_in_place

    sources = [5, 6, 3, 7]  # upload positions in clean order; 3 is a break
    assert _keep_in_place(sources, [False, False, True, False], [False] * 4) == [2, 3]


def test_the_diffs_stayers_outweigh_any_number_of_its_movers():
    """``_keep_in_place``: two siblings the diff kept outweigh the moved
    parent's three paragraphs, although the latter make the longer chain."""
    from backend.spec_doc.source_render import _keep_in_place

    sources = [3, 4, 0, 1, 2]
    movers = [False, False, True, True, True]
    assert _keep_in_place(sources, [False] * 5, movers) == [0, 1]
    # Without the diff's say, the longest chain wins.
    assert _keep_in_place(sources, [False] * 5, [False] * 5) == [2, 3, 4]


def test_the_writer_never_deletes_a_mark_that_holds_a_section_break():
    """Defence in depth under the builder's own refusal: Accept All of a
    deleted mark holding ``w:sectPr`` would merge two Word sections."""
    from backend.spec_doc.revision_marks import RevisionMarks, mark_paragraph

    paragraph = etree.fromstring(
        f'<w:p xmlns:w="{W}"><w:pPr><w:sectPr/></w:pPr><w:r><w:t>x</w:t></w:r></w:p>'
    )
    marks = RevisionMarks(author=AUTHOR, date=DATE, first_id=10)
    with pytest.raises(AssertionError):
        mark_paragraph(paragraph, "w:del", marks)


def test_neutralizing_the_last_paragraph_leaves_a_section_break_where_it_is():
    """``neutralize_last_paragraph`` moves formatting, never a break: the
    numbering goes into the ``w:pPrChange`` and the ``w:sectPr`` stays in
    the current properties, ahead of it — and only a paragraph-mark
    revision may ask. A paragraph that sets nothing is left alone and
    spends no revision id."""
    from backend.spec_doc.revision_marks import (
        RevisionMarks,
        neutralize_last_paragraph,
    )

    paragraph = etree.fromstring(
        f'<w:p xmlns:w="{W}"><w:pPr><w:numPr><w:ilvl w:val="0"/>'
        '<w:numId w:val="3"/></w:numPr><w:sectPr/></w:pPr>'
        "<w:r><w:t>x</w:t></w:r></w:p>"
    )
    marks = RevisionMarks(author=AUTHOR, date=DATE, first_id=10)
    neutralize_last_paragraph(paragraph, "w:del", marks)
    properties = paragraph.find(qn("w:pPr"))
    assert [etree.QName(c).localname for c in properties] == ["sectPr", "pPrChange"]
    recorded = properties.find(f"{qn('w:pPrChange')}/{qn('w:pPr')}")
    assert [etree.QName(c).localname for c in recorded] == ["numPr"]
    with pytest.raises(ValueError):
        neutralize_last_paragraph(paragraph, "w:moveTo", marks)
    before = marks.count
    for shape in ("", "<w:pPr><w:rPr/></w:pPr>"):
        for tag in ("w:del", "w:ins"):
            plain = etree.fromstring(
                f'<w:p xmlns:w="{W}">{shape}<w:r><w:t>x</w:t></w:r></w:p>'
            )
            unchanged = etree.tostring(plain)
            neutralize_last_paragraph(plain, tag, marks)
            assert etree.tostring(plain) == unchanged
    assert marks.count == before
    assert paragraph.find(f"{qn('w:pPr')}/{qn('w:rPr')}") is None


def test_a_deleted_row_keeps_its_flag_after_its_other_properties(tmp_path):
    """A row that already carries properties (cannot split, repeats as a
    header) is flagged after them — Word's schema order for ``w:trPr``."""
    from docx.oxml import OxmlElement

    document = Document()
    for line in (
        "SECTION 23 05 48",
        "VIBRATION CONTROLS",
        "PART 1 - GENERAL",
        "1.1 SCHEDULE",
        "A. Provide isolators as scheduled.",
    ):
        document.add_paragraph(line)
    table = document.add_table(rows=2, cols=2)
    for row, (left, right) in zip(table.rows, (("Equipment", "Deflection"), ("AHU-1", "1 inch"))):
        row.cells[0].text, row.cells[1].text = left, right
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        properties.append(OxmlElement("w:tblHeader"))
    document.add_paragraph("B. Coordinate with the schedule.")
    document.add_paragraph("END OF SECTION")
    source = _save(document)
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    block = next(p for p in article.paragraphs if p.locked == "table")
    section = _edit(imported.section, {"action": "delete", "target_id": block.uid})
    redline, _ = _verify(source, imported, section)  # asserts the order
    rows = list(_body(redline).iter(qn("w:tr")))
    assert rows
    for row in rows:
        tags = [etree.QName(c).localname for c in row.find(qn("w:trPr"))]
        assert tags == ["cantSplit", "tblHeader", "del"]


def test_a_deleted_mark_with_its_own_formatting_is_flagged_first(tmp_path):
    """A paragraph mark that already carries run properties (a bold mark):
    the deletion flag is the FIRST child of ``w:pPr/w:rPr``."""
    from docx.oxml import OxmlElement

    document = Document()
    for line in (
        "SECTION 23 05 48",
        "VIBRATION CONTROLS",
        "PART 1 - GENERAL",
        "1.1 SUMMARY",
        "A. Section includes vibration isolation.",
    ):
        document.add_paragraph(line)
    bold_mark = document.add_paragraph("B. Related requirements.")
    mark = OxmlElement("w:rPr")
    mark.append(OxmlElement("w:b"))
    bold_mark._p.get_or_add_pPr().append(mark)
    document.add_paragraph("C. Quality assurance.")
    document.add_paragraph("END OF SECTION")
    source = _save(document)
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    section = _edit(imported.section, {"action": "delete", "target_id": article.paragraphs[1].uid})
    redline, _ = _verify(source, imported, section)  # asserts the order
    flagged = [
        p
        for p in _body(redline).iter(qn("w:p"))
        if p.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:del')}") is not None
    ]
    assert len(flagged) == 1
    tags = [etree.QName(c).localname for c in flagged[0].find(f"{qn('w:pPr')}/{qn('w:rPr')}")]
    assert tags == ["del", "b"]


def test_a_pure_reorder_of_an_article_carries_its_children(tmp_path):
    source = _nested_master()
    imported = _parse(tmp_path, source)
    submittals = imported.section.parts[0].articles[1]
    section = _edit(
        imported.section, {"action": "move", "target_id": submittals.uid, "position": 0}
    )
    redline, _ = _verify(source, imported, section)
    moved = _tracked(redline, "w:ins") + _tracked(redline, "w:del")
    assert any("Product data." in text for text in moved)


def test_a_reorder_that_also_edits_the_moved_provision(tmp_path):
    source = _nested_master()
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    moved = article.paragraphs[1]
    section = _edit(
        imported.section,
        {"action": "move", "target_id": moved.uid, "position": 0},
        {
            "action": "replace",
            "target_id": moved.uid,
            "text": "Related requirements are in Division 22.",
        },
    )
    redline, _ = _verify(source, imported, section)
    inserted = _tracked(redline, "w:ins")
    assert any("Division 22." in text for text in inserted)
    assert "B. Related requirements." in _tracked(redline, "w:del")


def test_a_moved_provision_keeps_its_bookmarks_on_the_new_copy(tmp_path):
    """D-6: the copy that survives Accept All keeps the bookmark; Reject All
    restores the text at the old place, but not the bookmark — the one
    documented limit. No name ever appears twice."""
    document = Document()
    for line in (
        "SECTION 23 05 48",
        "VIBRATION CONTROLS",
        "PART 1 - GENERAL",
        "1.1 SUMMARY",
        "A. Section includes vibration isolation.",
    ):
        document.add_paragraph(line)
    marked = document.add_paragraph()
    start = etree.SubElement(marked._p, qn("w:bookmarkStart"))
    start.set(qn("w:id"), "77")
    start.set(qn("w:name"), "_Ref77")
    marked.add_run("B. Related requirements.")
    end = etree.SubElement(marked._p, qn("w:bookmarkEnd"))
    end.set(qn("w:id"), "77")
    document.add_paragraph("END OF SECTION")
    source = _save(document)
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    section = _edit(
        imported.section,
        {"action": "move", "target_id": article.paragraphs[1].uid, "position": 0},
    )
    redline, _ = _verify(source, imported, section, moved_bookmarks={"_Ref77"})
    names = [
        b.get(qn("w:name")) for b in _body(redline).iter(qn("w:bookmarkStart"))
    ]
    assert names == ["_Ref77"]
    # Revision ids start above the bookmark's id.
    assert min(int(r.get(qn("w:id"))) for r in _revisions(_body(redline))) > 77


def test_filling_a_not_used_part(tmp_path):
    source = _not_used_master()
    imported = _parse(tmp_path, source)
    section = _edit(
        imported.section,
        {"action": "add_article", "target_id": "pt2", "text": "ISOLATORS"},
    )
    redline, stats = _verify(source, imported, section)
    assert "(Not used.)" in _tracked(redline, "w:del")
    assert any("ISOLATORS" in text for text in _tracked(redline, "w:ins"))
    assert stats["redline"]["dropped"] == 1


def test_deleting_the_provision_under_a_section_break(tmp_path):
    source = _break_master(held=False)
    imported = _parse(tmp_path, source)
    schedule = imported.section.parts[0].articles[1]
    section = _edit(imported.section, {"action": "delete", "target_id": schedule.uid})
    redline, _ = _verify(source, imported, section)
    breaks = [
        p for p in _body(redline).iter(qn("w:p"))
        if p.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None
    ]
    assert len(breaks) == 1


def test_deleting_a_break_holder_keeps_its_mark_and_leaves_an_empty_break(tmp_path):
    source = _break_master(held=True)
    imported = _parse(tmp_path, source)
    holder = imported.section.parts[0].articles[0].paragraphs[1]
    section = _edit(imported.section, {"action": "delete", "target_id": holder.uid})
    redline, stats = _verify(source, imported, section)
    (holder_copy,) = [
        p for p in _body(redline).iter(qn("w:p"))
        if p.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None
    ]
    assert holder_copy.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:del')}") is None
    assert "".join(t.text for t in holder_copy.iter(qn("w:delText"))) == (
        "B. Related requirements."
    )
    assert stats["redline"]["leftovers"] == 1


def test_moving_a_break_holder_leaves_the_break_where_it_was(tmp_path):
    source = _break_master(held=True)
    imported = _parse(tmp_path, source)
    holder = imported.section.parts[0].articles[0].paragraphs[1]
    section = _edit(
        imported.section, {"action": "move", "target_id": holder.uid, "position": 0}
    )
    redline, _ = _verify(source, imported, section)
    breaks = [
        p for p in _body(redline).iter(qn("w:p"))
        if p.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None
    ]
    assert len(breaks) == 1
    assert breaks[0].find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:del')}") is None


def test_a_deleted_word_numbered_break_holder_records_its_numbering_cancel(tmp_path):
    """Deviation 6: the clean export's leftover takes ``w:numId 0``, so the
    redline's emptied holder records that as a paragraph-property change —
    Accept All cancels the number, Reject All restores it."""
    from tests.test_importer import _define_numbering, _numbered

    document = Document()
    document.add_paragraph("SECTION 23 05 48")
    document.add_paragraph("VIBRATION CONTROLS")
    document.add_paragraph("PART 1 - GENERAL")
    _define_numbering(
        document, 50, {1: ("decimal", "%1.%2"), 2: ("upperLetter", "%3.")}
    )
    _numbered(document, "SUMMARY", 1, "50")
    _numbered(document, "Section includes vibration isolation.", 2, "50")
    holder = _numbered(document, "Related requirements.", 2, "50")
    _hold_break(holder, document)
    _numbered(document, "SCHEDULE", 1, "50")
    _numbered(document, "Provide isolators as scheduled.", 2, "50")
    source = _save(document)
    imported = _parse(tmp_path, source)
    summary = imported.section.parts[0].articles[0]
    section = _edit(
        imported.section,
        {"action": "delete", "target_id": summary.paragraphs[1].uid},
    )
    redline, _ = _verify(source, imported, section)
    (change,) = list(_body(redline).iter(qn("w:pPrChange")))
    original = change.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}")
    assert original.get(qn("w:val")) == "50"


def test_a_section_renumber_on_a_header_line(tmp_path):
    source = _master_bytes()
    imported = _parse(tmp_path, source)
    section = _edit(
        imported.section,
        {
            "action": "replace",
            "target_id": "sec",
            "text": "VIBRATION AND SEISMIC CONTROLS",
            "numbering": "23 05 49",
        },
    )
    redline, _ = _verify(source, imported, section)
    assert "48" in _tracked(redline, "w:del")
    assert "49" in _tracked(redline, "w:ins")


def test_a_fallback_paragraph_deletes_all_and_inserts_the_new_text(tmp_path):
    from tests.docx_fidelity_helpers import _append_hyperlink

    document = Document()
    for line in (
        "SECTION 23 05 48",
        "VIBRATION CONTROLS",
        "PART 1 - GENERAL",
        "1.1 SUMMARY",
    ):
        document.add_paragraph(line)
    linked = document.add_paragraph("A. See ")
    _append_hyperlink(linked, "the client standard", "https://example.com/std")
    linked.add_run(" for isolators.")
    document.add_paragraph("END OF SECTION")
    source = _save(document)
    imported = _parse(tmp_path, source)
    first = imported.section.parts[0].articles[0].paragraphs[0]
    section = _edit(
        imported.section,
        {"action": "replace", "target_id": first.uid, "text": "See Section 23 05 00."},
    )
    redline, stats = _verify(source, imported, section)
    assert stats["redline"]["fallback"] == 1
    assert stats["fallback"] == {"hyperlink": 1}
    # The hyperlink's runs are deleted where they sit (a hyperlink cannot be
    # wrapped in w:del).
    hyperlink = _body(redline).find(f".//{qn('w:hyperlink')}")
    assert hyperlink.find(f"{qn('w:del')}/{qn('w:r')}") is not None


def test_a_complex_field_in_a_rewritten_provision_is_deleted_as_runs(tmp_path):
    from tests.docx_fidelity_helpers import _append_page_field

    document = Document()
    for line in (
        "SECTION 23 05 48",
        "VIBRATION CONTROLS",
        "PART 1 - GENERAL",
        "1.1 SUMMARY",
    ):
        document.add_paragraph(line)
    fielded = document.add_paragraph("A. See page ")
    _append_page_field(fielded)
    document.add_paragraph("END OF SECTION")
    source = _save(document)
    imported = _parse(tmp_path, source)
    first = imported.section.parts[0].articles[0].paragraphs[0]
    section = _edit(
        imported.section,
        {"action": "replace", "target_id": first.uid, "text": "See Part 3."},
    )
    redline, _ = _verify(source, imported, section)
    assert _body(redline).find(f".//{qn('w:delInstrText')}") is not None


def test_spacers_travel_with_their_provision(tmp_path):
    """The master's blank line above B moves with it, as a tracked move."""
    source = _master_bytes()
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    section = _edit(
        imported.section,
        {"action": "move", "target_id": article.paragraphs[1].uid, "position": 0},
    )
    _verify(source, imported, section)


def test_deleting_a_provision_deletes_its_spacers_with_it(tmp_path):
    source = _master_bytes()
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    section = _edit(
        imported.section, {"action": "delete", "target_id": article.paragraphs[1].uid}
    )
    redline, stats = _verify(source, imported, section)
    assert stats["redline"]["dropped"] == 1  # the blank line above B


def test_front_matter_is_untouched(tmp_path):
    from tests.test_import_office_master import _office_master

    source = _office_master()
    imported = _parse(tmp_path, source)
    target = next(
        p
        for part in imported.section.parts
        for article in part.articles
        for p in article.paragraphs
        if p.text == "Install per NFPA 13."
    )
    section = _edit(
        imported.section,
        {"action": "replace", "target_id": target.uid, "text": "Install per NFPA 13-2025."},
    )
    redline, _ = _verify(source, imported, section)
    before = _body(source)
    after = _body(redline)
    front = len(imported.front_matter) if hasattr(imported, "front_matter") else 0
    assert front > 0
    for index in range(front):
        assert etree.tostring(before[index]) == etree.tostring(after[index])


# ---------------------------------------------------------------------------
# Refusals
# ---------------------------------------------------------------------------


def test_a_master_with_pending_revisions_is_refused_by_name(tmp_path):
    """D-5: Reject All would also reject someone else's pending changes."""
    from tests.docx_fidelity_helpers import rewrite_zip_members

    source = _master_bytes()
    document_xml = _member(source, "word/document.xml").decode("utf-8")
    revised = document_xml.replace(
        "<w:t>END OF SECTION</w:t>",
        "<w:t>END OF SECTION</w:t></w:r><w:ins w:id=\"900\" w:author=\"Someone\" "
        "w:date=\"2026-01-01T00:00:00Z\"><w:r><w:t> (rev)</w:t></w:r></w:ins><w:r>",
        1,
    )
    assert revised != document_xml
    pending = rewrite_zip_members(
        source, replacements={"word/document.xml": revised.encode("utf-8")}
    )
    imported = _parse(tmp_path, pending)
    with pytest.raises(SourceRedlineError) as caught:
        render_preserving_redline(
            source_bytes=pending,
            format_map=imported.format_map,
            baseline=imported.section,
            current=imported.section,
            author=AUTHOR,
            date=DATE,
        )
    assert caught.value.reason == "pending_revisions"
    message = str(caught.value)
    assert "Accept or reject" in message and "import the file again" in message
    assert "extracted provisions" in message


def test_track_changes_merely_switched_on_is_not_a_refusal(tmp_path):
    """``w:trackRevisions`` with nothing pending does not break Reject All."""
    from tests.docx_fidelity_helpers import rewrite_zip_members

    source = _master_bytes()
    settings = _member(source, "word/settings.xml").decode("utf-8")
    switched = settings.replace("<w:zoom", "<w:trackRevisions/><w:zoom", 1)
    assert switched != settings
    tracked = rewrite_zip_members(
        source, replacements={"word/settings.xml": switched.encode("utf-8")}
    )
    imported = _parse(tmp_path, tracked)
    first = imported.section.parts[0].articles[0].paragraphs[0]
    section = _edit(
        imported.section, {"action": "replace", "target_id": first.uid, "text": "Changed."}
    )
    redline, _ = _verify(tracked, imported, section)
    # The switch stays as the user left it: settings.xml is untouched.
    assert _member(redline, "word/settings.xml") == _member(tracked, "word/settings.xml")


def test_deleting_a_block_content_control_is_refused_by_name(tmp_path):
    from tests.docx_fidelity_helpers import _append_opaque_sdt

    document = Document()
    for line in (
        "SECTION 23 05 48",
        "VIBRATION CONTROLS",
        "PART 1 - GENERAL",
        "1.1 SUMMARY",
        "A. Section includes vibration isolation.",
    ):
        document.add_paragraph(line)
    _append_opaque_sdt(document)
    document.add_paragraph("B. Related requirements.")
    document.add_paragraph("END OF SECTION")
    source = _save(document)
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    control = next(p for p in article.paragraphs if p.locked == "content_control")
    section = _edit(imported.section, {"action": "delete", "target_id": control.uid})
    with pytest.raises(SourceRedlineError) as caught:
        render_preserving_redline(
            source_bytes=source,
            format_map=imported.format_map,
            baseline=imported.section,
            current=section,
            author=AUTHOR,
            date=DATE,
        )
    assert caught.value.reason == "block_content_control"
    assert "extracted provisions still works" in str(caught.value)


def test_the_upload_is_never_modified(tmp_path):
    source = _typed_letter_master()
    snapshot = bytes(source)
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    section = _edit(
        imported.section, {"action": "delete", "target_id": article.paragraphs[0].uid}
    )
    _verify(source, imported, section)
    assert source == snapshot


def test_a_deleted_last_paragraph_keeps_words_marks_and_mark(tmp_path):
    """Word cannot track a document's last paragraph mark: its words are
    deleted and the mark is left alone."""
    document = Document()
    for line in (
        "SECTION 23 05 48",
        "VIBRATION CONTROLS",
        "PART 1 - GENERAL",
        "1.1 SUMMARY",
        "A. Section includes vibration isolation.",
        "B. Related requirements.",
    ):
        document.add_paragraph(line)
    source = _save(document)
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    section = _edit(
        imported.section, {"action": "delete", "target_id": article.paragraphs[1].uid}
    )
    redline, _ = _verify(source, imported, section)
    last = [c for c in _body(redline) if c.tag == qn("w:p")][-1]
    assert last.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:del')}") is None
    assert last.find(f".//{qn('w:delText')}") is not None
    # A plain paragraph has no formatting to record.
    assert last.find(f".//{qn('w:pPrChange')}") is None
    assert last.find(f".//{qn('w:rPrChange')}") is None


def _formatted_to_the_end_master(shape: str) -> bytes:
    """A master with no END OF SECTION, so its last body paragraph is a
    provision — one carrying formatting that an EMPTY paragraph still shows:
    Word numbering (an empty numbered paragraph prints its number) or a page
    break before it (a blank page), plus run formatting on its mark (the
    height of the empty line)."""
    from docx.oxml import OxmlElement

    from tests.test_importer import _define_numbering, _numbered

    document = Document()
    for line in ("SECTION 23 05 48", "VIBRATION CONTROLS", "PART 1 - GENERAL"):
        document.add_paragraph(line)
    if shape == "numbered":
        _define_numbering(
            document, 50, {1: ("decimal", "%1.%2"), 2: ("upperLetter", "%3.")}
        )
        _numbered(document, "SUMMARY", 1, "50")
        _numbered(document, "Section includes vibration isolation.", 2, "50")
        last = _numbered(document, "Provide isolators as scheduled.", 2, "50")
    else:
        for line in ("1.1 SUMMARY", "A. Section includes vibration isolation."):
            document.add_paragraph(line)
        last = document.add_paragraph("B. Provide isolators as scheduled.")
        last.paragraph_format.page_break_before = True
    mark = OxmlElement("w:rPr")
    mark.append(OxmlElement("w:b"))
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "28")
    mark.append(size)
    last._p.get_or_add_pPr().append(mark)
    return _save(document)


_SHOWN = {"numbered": "numPr", "page-break": "pageBreakBefore"}


def _tail(body):
    """The body's last paragraph."""
    return [c for c in body if c.tag == qn("w:p")][-1]


def _words(paragraph) -> str:
    return "".join(
        t.text or "" for t in paragraph.iter(qn("w:t"), qn("w:delText"))
    )


def _formatting(paragraph) -> list[str]:
    """What a resolved paragraph still sets, by local name: its own
    paragraph properties, then its mark's run formatting."""
    properties = paragraph.find(qn("w:pPr"))
    if properties is None:
        return []
    names = [
        etree.QName(c).localname
        for c in properties
        if isinstance(c.tag, str) and c.tag != qn("w:rPr")
    ]
    mark = properties.find(qn("w:rPr"))
    if mark is not None:
        names += [etree.QName(c).localname for c in mark if isinstance(c.tag, str)]
    return names


@pytest.mark.parametrize("shape", ["numbered", "page-break"])
def test_a_deleted_last_paragraph_accepts_to_a_plain_empty_one(tmp_path, shape):
    """Codex, PR #187: Word cannot track the last paragraph mark, so Accept
    All of a deleted last provision leaves its paragraph behind, empty — and
    an empty paragraph that keeps its number or its page break is not
    invisible. The redline records that formatting as a tracked change whose
    current side is plain: Accept All leaves a plain empty paragraph, Reject
    All the provision exactly as it was."""
    source = _formatted_to_the_end_master(shape)
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    section = _edit(
        imported.section,
        {"action": "delete", "target_id": article.paragraphs[-1].uid},
    )
    redline, stats = _verify(source, imported, section)
    assert stats["redline"]["last_mark_untracked"] == 1
    r_body = _body(redline)
    shown = _SHOWN[shape]
    tail = _tail(r_body)
    assert tail.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:del')}") is None
    assert tail.find(f"{qn('w:pPr')}/{qn('w:' + shown)}") is None
    recorded = tail.find(f"{qn('w:pPr')}/{qn('w:pPrChange')}/{qn('w:pPr')}")
    assert recorded is not None and recorded.find(qn("w:" + shown)) is not None
    leftover = _tail(accept_all(r_body))
    assert _words(leftover) == ""
    assert _formatting(leftover) == []
    restored = _tail(reject_all(r_body))
    assert _words(restored).endswith("Provide isolators as scheduled.")
    assert {shown, "b", "sz"} <= set(_formatting(restored))


@pytest.mark.parametrize("shape", ["numbered", "page-break"])
def test_an_appended_last_paragraph_rejects_to_a_plain_empty_one(tmp_path, shape):
    """The mirror image: a provision added after the last one is the
    document's new last paragraph, so Reject All leaves it behind, empty.
    The formatting it took from its neighbour is recorded as a change FROM
    nothing: Reject All leaves a plain empty paragraph, Accept All the
    provision exactly as exported."""
    source = _formatted_to_the_end_master(shape)
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    section = _edit(
        imported.section,
        {
            "action": "add_paragraph",
            "target_id": article.uid,
            "position": len(article.paragraphs),
            "text": "Provide seismic restraints.",
        },
    )
    redline, stats = _verify(source, imported, section)
    assert stats["redline"]["last_mark_untracked"] == 1
    r_body = _body(redline)
    shown = _SHOWN[shape]
    tail = _tail(r_body)
    assert tail.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:ins')}") is None
    recorded = tail.find(f"{qn('w:pPr')}/{qn('w:pPrChange')}/{qn('w:pPr')}")
    assert recorded is not None and len(recorded) == 0
    leftover = _tail(reject_all(r_body))
    assert _words(leftover) == ""
    assert _formatting(leftover) == []
    exported = _tail(accept_all(r_body))
    assert _words(exported).endswith("Provide seismic restraints.")
    assert {shown, "b", "sz"} <= set(_formatting(exported))


def test_revision_tags_are_the_oracle_vocabulary():
    """The writer only ever emits what the oracle resolves."""
    assert {
        qn("w:ins"),
        qn("w:del"),
        qn("w:pPrChange"),
        qn("w:rPrChange"),
    } <= REVISION_TAGS


def _break_page_master() -> bytes:
    document = Document()
    for line in (
        "SECTION 23 05 48",
        "VIBRATION CONTROLS",
        "PART 1 - GENERAL",
        "1.1 SUMMARY",
    ):
        document.add_paragraph(line)
    leading = document.add_paragraph()
    leading.add_run().add_break(WD_BREAK.PAGE)
    leading.add_run("A. Section includes vibration isolation.")
    document.add_paragraph("END OF SECTION")
    return _save(document)


def test_words_prepended_after_a_leading_page_break_track_after_it(tmp_path):
    """Deviation 8 in the redline: the inserted words sit after the break,
    exactly where the clean export puts them."""
    source = _break_page_master()
    imported = _parse(tmp_path, source)
    first = imported.section.parts[0].articles[0].paragraphs[0]
    section = _edit(
        imported.section,
        {
            "action": "replace",
            "target_id": first.uid,
            "text": "Also Section includes vibration isolation.",
        },
    )
    redline, _ = _verify(source, imported, section)
    paragraph = next(
        p for p in _body(redline).iter(qn("w:p")) if p.find(f".//{qn('w:ins')}") is not None
    )
    order = [
        "break" if el.tag == qn("w:br") else "ins"
        for el in paragraph.iter(qn("w:br"), qn("w:ins"))
    ]
    assert order.index("break") < order.index("ins")


def test_copies_are_independent_of_the_upload(tmp_path):
    """Rendering twice from the same inputs gives the same body (ids and
    all): nothing is mutated along the way."""
    source = _nested_master()
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    section = _edit(
        imported.section,
        {"action": "move", "target_id": article.paragraphs[1].uid, "position": 0},
    )
    first, _ = _verify(source, imported, copy.deepcopy(section))
    second, _ = _verify(source, imported, copy.deepcopy(section))
    assert _member(first, "word/document.xml") == _member(second, "word/document.xml")


# ---------------------------------------------------------------------------
# The corners: section breaks under reorders, tables, pictures, fields
# ---------------------------------------------------------------------------


def _two_break_master() -> bytes:
    """Three provisions with an empty section-break paragraph between each."""
    document = Document()
    for line in (
        "SECTION 23 05 48",
        "VIBRATION CONTROLS",
        "PART 1 - GENERAL",
        "1.1 SUMMARY",
        "A. First provision.",
    ):
        document.add_paragraph(line)
    _hold_break(document.add_paragraph(), document)
    document.add_paragraph("B. Second provision.")
    _hold_break(document.add_paragraph(), document)
    document.add_paragraph("C. Third provision.")
    document.add_paragraph("END OF SECTION")
    return _save(document)


def test_a_reversal_across_two_section_breaks_stays_exact(tmp_path):
    """Deviation 4: breaks go where ``_place`` puts them in the clean export,
    and the redline follows. Reversing three provisions across two breaks is
    where the diff's own choice of what moved (it keeps the last one) cannot
    sit between the breaks ``_place`` placed; the redline keeps the breaks
    fixed and reports the fewest moves around them instead."""
    source = _two_break_master()
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    first, second, third = (p.uid for p in article.paragraphs)
    section = _edit(
        imported.section,
        {"action": "move", "target_id": third, "position": 0},
        {"action": "move", "target_id": first, "position": 2},
    )
    assert [p.uid for p in section.parts[0].articles[0].paragraphs] == [
        third,
        second,
        first,
    ]
    redline, stats = _verify(source, imported, section)
    assert stats["redline"]["moves_added"] >= 1
    assert stats["redline"]["moved"] == 2  # two moves, not three
    breaks = [
        p for p in _body(redline).iter(qn("w:p"))
        if p.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None
    ]
    assert len(breaks) == 2


def test_a_moved_table_is_deleted_there_and_inserted_here(tmp_path):
    """Two provisions follow the schedule; moving the schedule below both
    keeps them in place and moves the table (the longest order kept)."""
    document = Document()
    for line in (
        "SECTION 23 05 48",
        "VIBRATION CONTROLS",
        "PART 1 - GENERAL",
        "1.1 SCHEDULE",
    ):
        document.add_paragraph(line)
    table = document.add_table(rows=2, cols=2)
    for row, (left, right) in zip(table.rows, [("AHU-1", "1 inch"), ("P-1", "0.75 inch")]):
        row.cells[0].text = left
        row.cells[1].text = right
    document.add_paragraph("A. Provide isolators as scheduled.")
    document.add_paragraph("B. Coordinate with the structural engineer.")
    document.add_paragraph("END OF SECTION")
    source = _save(document)
    imported = _parse(tmp_path, source)
    schedule = imported.section.parts[0].articles[0]
    block = next(p for p in schedule.paragraphs if p.locked == "table")
    section = _edit(
        imported.section,
        {"action": "move", "target_id": block.uid, "position": len(schedule.paragraphs) - 1},
    )
    redline, stats = _verify(source, imported, section)
    assert stats["redline"]["moved"] == 1
    tables = list(_body(redline).iter(qn("w:tbl")))
    assert len(tables) == 2
    flags = []
    for table_el in tables:
        row_flags = set()
        for row in table_el.iter(qn("w:tr")):
            properties = row.find(qn("w:trPr"))
            assert properties is not None
            row_flags |= {etree.QName(f).localname for f in properties}
        flags.append(row_flags)
    assert sorted(flags, key=sorted) == [{"del"}, {"ins"}]


def test_deleting_a_picture_block(tmp_path):
    """A provision holding a picture is a preserved block; deleting it is a
    tracked deletion of the paragraph, drawing and all."""
    from tests.docx_fidelity_helpers import _png_bytes

    document = Document()
    for line in (
        "SECTION 23 05 48",
        "VIBRATION CONTROLS",
        "PART 1 - GENERAL",
        "1.1 SUMMARY",
        "A. Section includes vibration isolation.",
    ):
        document.add_paragraph(line)
    figure = document.add_paragraph("B. Isolator detail: ")
    figure.add_run().add_picture(io.BytesIO(_png_bytes()))
    document.add_paragraph("C. See the figure above.")
    document.add_paragraph("END OF SECTION")
    source = _save(document)
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    picture = next(p for p in article.paragraphs if p.locked == "image")
    section = _edit(imported.section, {"action": "delete", "target_id": picture.uid})
    redline, _ = _verify(source, imported, section)
    drawing = _body(redline).find(f".//{qn('w:drawing')}")
    assert any(a.tag == qn("w:del") for a in drawing.iterancestors())


def test_deleting_part_of_a_table_of_contents_is_refused_by_name(tmp_path):
    """A cached TOC is one Word field across many paragraphs; deleting one
    of them cannot be tracked without breaking the field."""
    from tests.test_import_office_master import _toc_paragraphs

    document = Document()
    for line in (
        "SECTION 23 05 48",
        "VIBRATION CONTROLS",
        "PART 1 - GENERAL",
        "1.1 SUMMARY",
        "A. Section includes vibration isolation.",
    ):
        document.add_paragraph(line)
    _toc_paragraphs(document, ["1.1 SUMMARY 3", "1.2 SUBMITTALS 3"])
    document.add_paragraph("B. Related requirements.")
    document.add_paragraph("END OF SECTION")
    source = _save(document)
    imported = _parse(tmp_path, source)
    field = next(
        p
        for part in imported.section.parts
        for article in part.articles
        for p in article.paragraphs
        if p.locked == "field"
    )
    section = _edit(imported.section, {"action": "delete", "target_id": field.uid})
    with pytest.raises(SourceRedlineError) as caught:
        render_preserving_redline(
            source_bytes=source,
            format_map=imported.format_map,
            baseline=imported.section,
            current=section,
            author=AUTHOR,
            date=DATE,
        )
    assert caught.value.reason == "field_block"


def test_a_moved_provision_with_a_comment_is_refused_by_name(tmp_path):
    """One comment cannot be anchored in two places: moving the provision it
    annotates cannot be shown as a tracked delete-and-insert."""
    document = Document()
    for line in (
        "SECTION 23 05 48",
        "VIBRATION CONTROLS",
        "PART 1 - GENERAL",
        "1.1 SUMMARY",
        "A. Section includes vibration isolation.",
    ):
        document.add_paragraph(line)
    annotated = document.add_paragraph()
    etree.SubElement(annotated._p, qn("w:commentRangeStart")).set(qn("w:id"), "3")
    annotated.add_run("B. Related requirements.")
    etree.SubElement(annotated._p, qn("w:commentRangeEnd")).set(qn("w:id"), "3")
    document.add_paragraph("END OF SECTION")
    source = _save(document)
    imported = _parse(tmp_path, source)
    article = imported.section.parts[0].articles[0]
    section = _edit(
        imported.section,
        {"action": "move", "target_id": article.paragraphs[1].uid, "position": 0},
    )
    with pytest.raises(SourceRedlineError) as caught:
        render_preserving_redline(
            source_bytes=source,
            format_map=imported.format_map,
            baseline=imported.section,
            current=section,
            author=AUTHOR,
            date=DATE,
        )
    assert caught.value.reason == "moved_annotation"


# ---------------------------------------------------------------------------
# Sweeps: re-import, and every corpus master under a mix of edits
# ---------------------------------------------------------------------------

_MIX_WORDS = ("seismic", "isolation", "Provide", "restraints", "per", "spring")

#: Named refusals a real master can legitimately earn. A failed self-check
#: (accept / reject / bookmarks / package) or unaccounted content is never
#: one of them: those mean the writer broke its own promise.
_STRUCTURAL_REFUSALS = frozenset(
    {
        "pending_revisions",
        "revision_scan_unavailable",
        "section_break_reorder",
        "moved_annotation",
        "simple_field",
        "block_content_control",
        "field_block",
        "pending_revisions_in_body",
        "untrackable_markup",
    }
)


def _paragraph_lists(section):
    """(parent uid, sibling list) for every paragraph sibling list."""
    lists = []

    def walk(parent, paragraphs):
        lists.append((parent, paragraphs))
        for paragraph in paragraphs:
            walk(paragraph.uid, paragraph.children)

    for part in section.parts:
        for article in part.articles:
            walk(article.uid, article.paragraphs)
    return lists


def _scripted_edit(section, rng):
    """One edit of the kind a user makes: reword, delete, add, move — and
    the same for articles. ``None`` when the tree offers nothing to do."""
    lists = _paragraph_lists(section)
    paragraphs = [p for _parent, ps in lists for p in ps]
    kind = rng.choice(
        ("reword", "delete", "add", "move", "move_article", "add_article", "delete_article")
    )
    if kind == "reword":
        editable = [p for p in paragraphs if not p.locked]
        if not editable:
            return None
        paragraph = rng.choice(editable)
        words = paragraph.text.split()
        if words and rng.random() < 0.7:
            words[rng.randrange(len(words))] = rng.choice(_MIX_WORDS)
        else:
            words.insert(rng.randint(0, len(words)), rng.choice(_MIX_WORDS))
        return {"action": "replace", "target_id": paragraph.uid, "text": " ".join(words)}
    if kind == "delete":
        return (
            {"action": "delete", "target_id": rng.choice(paragraphs).uid}
            if paragraphs
            else None
        )
    if kind == "add":
        if not lists:
            return None
        parent, siblings = rng.choice(lists)
        return {
            "action": "add_paragraph",
            "target_id": parent,
            "position": rng.randint(0, len(siblings)),
            "text": f"Provide {rng.choice(_MIX_WORDS)} restraints.",
        }
    if kind == "move":
        movable = [(parent, ps) for parent, ps in lists if len(ps) >= 2]
        if not movable:
            return None
        _parent, siblings = rng.choice(movable)
        return {
            "action": "move",
            "target_id": rng.choice(siblings).uid,
            "position": rng.randrange(len(siblings)),
        }
    articles = [(part, a) for part in section.parts for a in part.articles]
    if kind == "move_article":
        parts = [part for part in section.parts if len(part.articles) >= 2]
        if not parts:
            return None
        part = rng.choice(parts)
        return {
            "action": "move",
            "target_id": rng.choice(part.articles).uid,
            "position": rng.randrange(len(part.articles)),
        }
    if kind == "add_article":
        part = rng.choice(section.parts)
        return {
            "action": "add_article",
            "target_id": part.uid,
            "position": rng.randint(0, len(part.articles)),
            "text": "ISOLATORS",
        }
    if not articles:
        return None
    return {"action": "delete", "target_id": rng.choice(articles)[1].uid}


def _edit_mix(section, seed: int, count: int):
    """``count`` scripted edits, deterministic per ``seed``. An edit the
    model refuses (a retyped preserved block, say) is simply skipped."""
    rng = random.Random(seed)
    for _ in range(count):
        op = _scripted_edit(section, rng)
        if op is None:
            continue
        try:
            section, _ = apply_edits(section, [op])
        except SpecEditError:
            continue
    return section


def _shape(section):
    """What the importer can see of a tree: the identity, and per PART each
    article's title and its provisions' depth and text, in order."""

    def paragraphs(nodes, depth):
        out = []
        for node in nodes:
            out.append((depth, node.text))
            out.extend(paragraphs(node.children, depth + 1))
        return out

    return (
        section.number,
        section.title,
        [
            (part.uid, [(a.title, paragraphs(a.paragraphs, 0)) for a in part.articles])
            for part in section.parts
        ],
    )


@pytest.mark.parametrize("seed", range(10))
def test_reimporting_the_redline_reproduces_the_current_tree(tmp_path, seed):
    """The Batch 5 invariant, now on your original: the app's own Accept-All
    reader, given the redline, reads back the tree it was exported from.

    Typed letters here, because a new provision nested deeper than any
    provision a Word-numbered master already has takes its kin's
    ``w:ilvl`` in today's formatted export (a Phase 0 limit, recorded in the
    plan) — so for those the redline re-imports exactly as the formatted
    export does, which is the promise, but not as the tree."""
    source = _typed_letter_master()
    imported = _parse(tmp_path, source)
    section = _edit_mix(imported.section, seed, 1 + seed % 6)
    redline, _ = _verify(source, imported, section, allow_moved_bookmarks=True)
    (tmp_path / "redline.docx").write_bytes(redline)
    reread = parse_master_docx(tmp_path / "redline.docx").section
    assert _shape(reread) == _shape(section)


def test_reimporting_the_redline_matches_reimporting_the_formatted_export(tmp_path):
    """Whatever the formatted export does, the redline's Accept All does too
    — through the importer, for masters of every labelling kind."""
    masters = {
        "typed": _typed_letter_master(),
        "numbered": _numbered_master(),
        "nested": _nested_master(),
        "held_break": _break_master(held=True),
    }
    for name, source in masters.items():
        imported = _parse(tmp_path, source, f"{name}.docx")
        for seed in range(4):
            section = _edit_mix(imported.section, seed * 13 + len(name), 5)
            redline, _ = _verify(source, imported, section, allow_moved_bookmarks=True)
            clean = render_preserving_docx(
                source_bytes=source, format_map=imported.format_map, current=section
            )
            (tmp_path / "r.docx").write_bytes(redline)
            (tmp_path / "c.docx").write_bytes(clean)
            assert _shape(parse_master_docx(tmp_path / "r.docx").section) == _shape(
                parse_master_docx(tmp_path / "c.docx").section
            ), (name, seed)


def test_every_corpus_master_keeps_the_promise_under_a_mix_of_edits(tmp_path):
    """Word-saved, LibreOffice-saved and hand-made masters alike: every one,
    run through a scripted mix of edits, keeps both halves of the promise —
    or is refused by name for a structural reason. Never by a failed
    self-check, which would mean the writer broke its own promise.

    The fallback count is the evidence D-2's eligibility is widened from;
    every fallback carries a reason from the closed vocabulary."""
    from tests.docx_corpus import build_case, corpus_cases

    fallbacks: collections.Counter = collections.Counter()
    produced = 0
    for case in corpus_cases():
        source = build_case(case, tmp_path)
        imported = _parse(tmp_path, source, f"{case.case_id}.docx")
        for seed in range(3):
            section = _edit_mix(imported.section, seed * 31 + len(case.case_id), 6)
            try:
                _redline, stats = _verify(
                    source, imported, section, allow_moved_bookmarks=True
                )
            except SourceRedlineError as exc:
                assert exc.reason in _STRUCTURAL_REFUSALS, (
                    case.case_id,
                    seed,
                    exc.reason,
                    exc.detail,
                )
                continue
            produced += 1
            fallbacks.update(stats.get("fallback", {}))
    assert produced
    assert set(fallbacks) <= set(FALLBACK_REASONS)


# ---------------------------------------------------------------------------
# Through the API (D-8, backend half)
# ---------------------------------------------------------------------------

_FIXED_DATE = "2026-09-22T12:00:00Z"


@pytest.fixture
def client(monkeypatch):
    """An app whose revision timestamp is pinned, so two exports of one
    document are byte-comparable."""
    import backend.app as app_module

    monkeypatch.setattr(app_module, "_revision_timestamp", lambda: _FIXED_DATE)
    return TestClient(create_app())


def _reword_first(client) -> None:
    doc = client.get("/api/doc").json()["doc"]
    first = doc["parts"][0]["articles"][0]["paragraphs"][0]["id"]
    edit = client.post(
        "/api/doc/edit",
        json={
            "ops": [
                {
                    "action": "replace",
                    "target_id": first,
                    "text": "Section includes seismic restraint and isolation.",
                }
            ]
        },
    )
    assert edit.status_code == 200, edit.text


def _export(client, **params):
    return client.get("/api/export/docx", params=params)


def _filename(response) -> str:
    disposition = response.headers["content-disposition"]
    return disposition.split('filename="', 1)[1].split('"', 1)[0]


def test_the_route_hands_over_the_redline_on_the_original(client):
    source = _master_bytes()
    payload = _import(client, source)
    assert payload["preserved_redline_available"] is True
    assert payload["preserved_redline_reason"] is None
    _reword_first(client)

    response = _export(client, redline="master", mode="preserved")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    # Named after the upload, so replacing the master is a rename (D-6).
    assert _filename(response) == "office-master - REDLINE.docx"
    redline = response.content

    # Accept All is the route's own formatted export; Reject All the upload.
    clean = _export(client, mode="preserved")
    assert clean.status_code == 200
    r_body = _body(redline)
    assert first_difference(accept_all(r_body), _body(clean.content)) is None
    assert first_difference(reject_all(r_body), _body(source)) is None
    before, after = _members(source), _members(redline)
    assert [n for n in before if before[n] != after[n]] == ["word/document.xml"]
    changes = _revisions(r_body)
    assert changes
    assert {c.get(qn("w:author")) for c in changes} == {"Build-a-Spec"}
    assert {c.get(qn("w:date")) for c in changes} == {_FIXED_DATE}
    # The firm's header travelled with the file untouched.
    assert Document(io.BytesIO(redline)).sections[0].header.paragraphs[0].text == (
        HEADER_TEXT
    )


def test_a_redline_names_its_mode_or_takes_the_default_for_its_kind(client):
    """``redline=master`` with no mode is the redline on the original when
    one is available; ``redline=version`` with no mode stays extracted
    provisions (the preserved redline is master-only in Phase 1)."""
    _import(client, _master_bytes())
    _reword_first(client)

    bare = _export(client, redline="master")
    explicit = _export(client, redline="master", mode="preserved")
    assert bare.status_code == explicit.status_code == 200
    assert bare.content == explicit.content
    assert _filename(bare) == "office-master - REDLINE.docx"

    extracted = _export(client, redline="master", mode="normalized")
    assert extracted.status_code == 200
    assert _filename(extracted).startswith("SECTION 23 05 48")
    # Build-a-Spec's own document: the firm's header does not travel.
    assert (
        Document(io.BytesIO(extracted.content)).sections[0].header.paragraphs[0].text
        != HEADER_TEXT
    )

    version = _export(client, redline="version", base=0)
    assert version.status_code == 200
    assert _filename(version).startswith("SECTION 23 05 48")
    assert (
        Document(io.BytesIO(version.content)).sections[0].header.paragraphs[0].text
        != HEADER_TEXT
    )


def test_a_redline_against_a_version_on_the_original_is_a_400(client):
    _import(client, _master_bytes())
    _reword_first(client)
    response = _export(client, redline="version", base=0, mode="preserved")
    assert response.status_code == 400
    assert "imported master only" in response.json()["error"]


def test_without_an_imported_master_there_is_nothing_to_compare(client):
    """A document written from scratch: the redline on the original is not
    offered, and asking for it is the existing no-master 400."""
    added = client.post(
        "/api/doc/edit",
        json={"ops": [{"action": "add_article", "target_id": "pt1", "text": "SUMMARY"}]},
    )
    assert added.status_code == 200
    payload = client.get("/api/doc").json()
    assert payload["preserved_redline_available"] is False
    assert payload["preserved_redline_reason"]["code"] == "no_baseline"
    response = _export(client, redline="master", mode="preserved")
    assert response.status_code == 400
    assert "no imported master" in response.json()["error"]


def test_a_project_without_its_format_map_names_what_is_missing(client):
    """The upload without the map built from it (a project imported before
    1.14.0): the redline on the original is not offered, an explicit request
    is refused by name, and the default falls back to extracted provisions."""
    _import(client, _master_bytes())
    _reword_first(client)
    sessions.get_session().source_format_map = None

    payload = client.get("/api/doc").json()
    assert payload["preserved_redline_available"] is False
    reason = payload["preserved_redline_reason"]
    assert reason["code"] == "no_original"
    assert "importing the file again" in reason["message"]

    refused = _export(client, redline="master", mode="preserved")
    assert refused.status_code == 409
    assert refused.json()["code"] == "no_original"
    assert refused.json()["error"] == reason["message"]

    fallback = _export(client, redline="master")
    assert fallback.status_code == 200
    assert _filename(fallback).startswith("SECTION 23 05 48")


def _pending_master() -> bytes:
    from tests.docx_fidelity_helpers import rewrite_zip_members

    source = _master_bytes()
    document_xml = _member(source, "word/document.xml").decode("utf-8")
    revised = document_xml.replace(
        "<w:t>END OF SECTION</w:t>",
        "<w:t>END OF SECTION</w:t></w:r><w:ins w:id=\"900\" w:author=\"Someone\" "
        "w:date=\"2026-01-01T00:00:00Z\"><w:r><w:t> (rev)</w:t></w:r></w:ins><w:r>",
        1,
    )
    assert revised != document_xml
    return rewrite_zip_members(
        source, replacements={"word/document.xml": revised.encode("utf-8")}
    )


def test_pending_revisions_are_refused_by_the_route_and_named_in_the_payload(client):
    """D-5, end to end: the payload says why before the click, an explicit
    request is a 409 naming the fix, and the default falls back to the
    redline of extracted provisions — which still works."""
    _import(client, _pending_master())
    _reword_first(client)

    payload = client.get("/api/doc").json()
    assert payload["preserved_redline_available"] is False
    reason = payload["preserved_redline_reason"]
    assert reason["code"] == "pending_revisions"
    assert "Accept or reject those changes in Word" in reason["message"]
    assert "extracted provisions still works" in reason["message"]

    refused = _export(client, redline="master", mode="preserved")
    assert refused.status_code == 409
    assert refused.json()["code"] == "pending_revisions"
    assert refused.json()["error"] == reason["message"]

    fallback = _export(client, redline="master")
    assert fallback.status_code == 200
    assert _filename(fallback).startswith("SECTION 23 05 48")


def test_track_changes_switched_on_is_offered_through_the_route(client):
    """``w:trackRevisions`` alone is not a pending change: the package-wide
    scan reports both, and only the pending half refuses."""
    from tests.docx_fidelity_helpers import rewrite_zip_members

    source = _master_bytes()
    settings_xml = _member(source, "word/settings.xml").decode("utf-8")
    switched = settings_xml.replace("<w:zoom", "<w:trackRevisions/><w:zoom", 1)
    tracked = rewrite_zip_members(
        source, replacements={"word/settings.xml": switched.encode("utf-8")}
    )
    assert _import(client, tracked)["preserved_redline_available"] is True
    _reword_first(client)
    response = _export(client, redline="master", mode="preserved")
    assert response.status_code == 200
    assert _member(response.content, "word/settings.xml") == switched.encode("utf-8")


def test_a_failed_self_check_is_a_409_naming_the_check(client, monkeypatch):
    """D-7: a redline that cannot prove its own promise is never handed
    over. Forced here by making the comparator report a difference."""
    from backend.spec_doc import revisions

    _import(client, _master_bytes())
    _reword_first(client)
    monkeypatch.setattr(
        revisions,
        "first_difference",
        lambda *_a, **_k: revisions.Difference(index=4, left="p", right="p", path="p/r/t"),
    )
    response = _export(client, redline="master", mode="preserved")
    assert response.status_code == 409
    body = response.json()
    assert body["code"] == "accept_check_failed"
    assert "failed its own check" in body["error"]
    assert "extracted provisions still works" in body["error"]


def test_the_payload_and_the_route_answer_from_one_derivation(client, monkeypatch):
    """The Export menu must never offer a redline the route refuses: both
    read ``_preserved_redline_availability``."""
    import backend.app as app_module

    _import(client, _master_bytes())
    _reword_first(client)
    monkeypatch.setattr(
        app_module,
        "_preserved_redline_availability",
        lambda _session: (False, "revision_scan_unavailable"),
    )
    payload = client.get("/api/doc").json()
    assert payload["preserved_redline_available"] is False
    assert payload["preserved_redline_reason"]["code"] == "revision_scan_unavailable"
    refused = _export(client, redline="master", mode="preserved")
    assert refused.status_code == 409
    assert refused.json()["code"] == "revision_scan_unavailable"
    assert _filename(_export(client, redline="master")).startswith("SECTION 23 05 48")


def test_the_revision_scan_runs_once_per_upload(client, monkeypatch):
    """The payload is built on every edit, undo and poll, and the pending-
    revisions scan reads every story part of the package — so it is cached
    per upload, keyed on the upload's content.

    Counted relatively: the session object outlives a test (``reset()`` is
    in place), so an earlier test that imported this very master may have
    left its answer cached — a correct hit, since the key is the content."""
    import backend.app as app_module

    calls = []
    real = app_module.detect_pending_revisions

    def counting(source):
        calls.append(len(source))
        return real(source)

    monkeypatch.setattr(app_module, "detect_pending_revisions", counting)
    _import(client, _master_bytes())
    settled = len(calls)
    assert settled <= 1
    for _ in range(3):
        assert client.get("/api/doc").json()["preserved_redline_available"] is True
    _reword_first(client)
    assert _export(client, redline="master", mode="preserved").status_code == 200
    assert len(calls) == settled  # no rescans for the same upload

    # A different upload is a different answer: it is scanned once.
    assert client.post("/api/session/reset").status_code == 200
    _import(client, _typed_letter_master())
    assert len(calls) == settled + 1
    for _ in range(2):
        client.get("/api/doc")
    assert len(calls) == settled + 1


def test_the_redline_filename_comes_from_the_upload():
    section = SpecSection()
    section.number, section.title = "23 05 48", "VIBRATION CONTROLS"
    assert upload_redline_filename("office-master.docx", section) == (
        "office-master - REDLINE.docx"
    )
    assert upload_redline_filename("Fire Pump.DOCX", section) == (
        "Fire Pump - REDLINE.docx"
    )
    assert upload_redline_filename('a:b*?"<>|.docx', section) == "ab - REDLINE.docx"
    assert upload_redline_filename("bad\x0bname.docx", section) == (
        "badu000Bname - REDLINE.docx"
    )
    # Unknown or scrubbed away: the section-derived redline name.
    for unknown in ("", ".docx", ":*?.docx", None):
        assert upload_redline_filename(unknown, section) == (
            "SECTION 23 05 48 - VIBRATION CONTROLS - REDLINE.docx"
        )
