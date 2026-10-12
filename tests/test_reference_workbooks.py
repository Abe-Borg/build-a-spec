"""Excel workbooks attached as reference documents.

An equipment schedule or a calc workbook is the most common background a
designer has that is neither prose nor a PDF. The reader's contract, pinned
here:

1. every sheet is read, in workbook order, under a marker naming it (hidden
   sheets too, marked hidden), its rows as CSV;
2. a cell reads the way Excel last showed it — a formula's saved result, a
   date as YYYY-MM-DD, a percentage with its % — and a formula with no saved
   result is read as its formula and reported, never left silently blank;
3. what cannot be read is refused with the reason and the way out (legacy
   formats, a password-protected workbook, junk), and the reader never parses
   an entity declaration (defusedxml) or more than the ZIP bounds allow.
"""
from __future__ import annotations

import csv
import datetime
import io
import re
import zipfile
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from backend import reference_extract, sessions
from backend.app import create_app
from backend.llm.conversation import _run_tool
from backend.reference_extract import (
    ReferenceExtractError,
    extract_reference_document,
)
from backend.spec_doc.project import load_project

REPO_ROOT = Path(__file__).resolve().parents[1]


def _xlsx(build) -> bytes:
    """A workbook from ``build(workbook)``; the first sheet is ``Pumps``."""
    workbook = Workbook()
    workbook.active.title = "Pumps"
    build(workbook)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _schedule(workbook: Workbook) -> None:
    pumps = workbook["Pumps"]
    pumps.append(["Tag", "Rated flow (gpm)", "Rated pressure (psi)", "Driver"])
    pumps.append(["FP-1", 2500, 125, "Diesel, UL listed"])
    pumps.append(["FP-2", 2500, 125, "Electric"])
    valves = workbook.create_sheet("Valves")
    valves.append(["Tag", "Size (in)", "Note"])
    valves.append(["V-1", 8, "Chilled water 44 degrees F at design"])


def _rewrite(payload: bytes, part: str, transform) -> bytes:
    """``payload`` with one part's XML passed through ``transform``."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(payload)) as source, zipfile.ZipFile(
        out, "w", zipfile.ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            body = source.read(info.filename)
            if info.filename == part:
                body = transform(body.decode("utf-8")).encode("utf-8")
            target.writestr(info.filename, body)
    return out.getvalue()


def _with_cell(payload: bytes, ref: str, cell_xml: str) -> bytes:
    """Replace one cell element of the first sheet with ``cell_xml``."""

    def transform(xml: str) -> str:
        replaced, count = re.subn(
            rf'<c r="{ref}"[^>]*?(?:/>|>.*?</c>)', cell_xml, xml, count=1
        )
        assert count == 1, xml
        return replaced

    return _rewrite(payload, "xl/worksheets/sheet1.xml", transform)


def _upload(client: TestClient, name: str, payload: bytes):
    return client.post(
        "/api/reference/upload",
        files={"file": (name, payload, "application/octet-stream")},
    )


def _read_tool(session, ref_id: str) -> str:
    return _run_tool(
        session,
        {"name": "read_reference_doc", "id": "tu1", "input": {"ref_id": ref_id}},
    )[0]["content"]


def _sheet_rows(text: str, marker: str) -> list[list[str]]:
    """The CSV rows under ``marker``, up to the next marker (or the tool
    result's closing frame)."""
    lines = text.split("\n")
    start = lines.index(marker) + 1
    end = next(
        (i for i in range(start, len(lines)) if lines[i].startswith(("[", "</"))),
        len(lines),
    )
    return list(csv.reader(lines[start:end]))


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------


def test_a_workbook_attaches_as_excel_and_reads_sheet_by_sheet():
    client = TestClient(create_app())

    response = _upload(client, "Pump schedule.xlsx", _xlsx(_schedule))

    assert response.status_code == 200, response.json()
    doc = response.json()["reference_doc"]
    assert doc["kind"] == "xlsx"
    assert doc["kind_label"] == "Excel"
    assert doc["filename"] == "Pump schedule.xlsx"
    assert doc["block_count"] == 5
    assert response.json()["warnings"] == []
    session = sessions.get_session()
    text = _read_tool(session, "ref-1")
    assert text.index("[workbook sheets: 1. Pumps; 2. Valves]") < text.index(
        "[sheet 1 of 2: Pumps]"
    ) < text.index("[sheet 2 of 2: Valves]")
    assert "(Excel," in text
    assert _sheet_rows(text, "[sheet 1 of 2: Pumps]") == [
        ["Tag", "Rated flow (gpm)", "Rated pressure (psi)", "Driver"],
        ["FP-1", "2500", "125", "Diesel, UL listed"],
        ["FP-2", "2500", "125", "Electric"],
    ]
    assert _sheet_rows(text, "[sheet 2 of 2: Valves]") == [
        ["Tag", "Size (in)", "Note"],
        ["V-1", "8", "Chilled water 44 degrees F at design"],
    ]
    stubs = session.references.context_stubs()
    assert "Excel" in stubs
    assert "sheet markers [sheet N of M: Name]" in stubs


def test_a_single_sheet_workbook_has_no_contents_line():
    text = extract_reference_document(
        _xlsx(lambda wb: wb["Pumps"].append(["FP-1", 2500])),
        filename="one.xlsx",
    ).text

    assert text == "[sheet 1 of 1: Pumps]\nFP-1,2500"


def test_cells_read_the_way_excel_shows_them():
    def build(workbook: Workbook) -> None:
        sheet = workbook["Pumps"]
        sheet.append(
            [
                datetime.date(2026, 10, 12),
                datetime.datetime(2026, 10, 12, 14, 30),
                datetime.datetime(2026, 10, 12, 14, 30, 15),
                datetime.time(7, 45),
            ]
        )
        sheet.append([0.65, 0.07, 5, 5])
        sheet["A2"].number_format = "0%"
        sheet["B2"].number_format = "0.0%"
        # A % inside a quoted literal or escaped is decoration, not a percent.
        sheet["C2"].number_format = '0" %"'
        sheet["D2"].number_format = "0\\%"
        sheet.append([True, False, 0.1 + 0.2, 2500.0, 125.5, 1234567.891])
        sheet.append(['He said "go", twice', "line one\nline two", "#N/A"])

    text = extract_reference_document(_xlsx(build), filename="cells.xlsx").text

    assert _sheet_rows(text, "[sheet 1 of 1: Pumps]") == [
        ["2026-10-12", "2026-10-12 14:30", "2026-10-12 14:30:15", "07:45"],
        ["65%", "7%", "5", "5"],
        ["TRUE", "FALSE", "0.3", "2500", "125.5", "1234567.891"],
        ['He said "go", twice', "line one line two", "#N/A"],
    ]


def test_a_saved_formula_result_is_read_not_the_formula():
    payload = _xlsx(lambda wb: wb["Pumps"].append(["FP-1", 2500, 125, "=B1*C1"]))
    saved = _with_cell(payload, "D1", '<c r="D1"><f>B1*C1</f><v>312500</v></c>')

    extraction = extract_reference_document(saved, filename="calc.xlsx")

    assert extraction.text.endswith("FP-1,2500,125,312500")
    assert extraction.warnings == ()


def test_a_formula_whose_saved_result_is_empty_reads_blank():
    """``=IF(A1="","",…)`` is how schedules hide unused rows. Excel saves its
    empty result as a ``str`` value, and that cell IS blank — it must not be
    mistaken for a formula Excel never calculated."""
    payload = _xlsx(lambda wb: wb["Pumps"].append(["FP-1", 2500, "=B1"]))
    for cell in (
        '<c r="C1" t="str"><f>IF(B1="","",B1)</f><v></v></c>',
        '<c r="C1" t="str"><f>IF(B1="","",B1)</f><v/></c>',
    ):
        extraction = extract_reference_document(
            _with_cell(payload, "C1", cell), filename="calc.xlsx"
        )
        assert extraction.text.endswith("\nFP-1,2500"), cell
        assert extraction.warnings == ()


def test_a_formula_with_no_saved_result_is_read_as_its_formula_and_reported():
    """A workbook written by a script carries formulas Excel never ran.
    Leaving those cells blank would tell the model a schedule has no totals."""
    client = TestClient(create_app())
    payload = _xlsx(
        lambda wb: wb["Pumps"].append(["FP-1", 2500, 125, "=B1*C1", "=SUM(B1:C1)"])
    )

    body = _upload(client, "scripted.xlsx", payload).json()

    assert any(
        "2 formula cell(s) had no saved result" in warning
        for warning in body["warnings"]
    )
    text = _read_tool(sessions.get_session(), "ref-1")
    assert _sheet_rows(text, "[sheet 1 of 1: Pumps]") == [
        ["FP-1", "2500", "125", "=B1*C1", "=SUM(B1:C1)"]
    ]


def test_spacer_columns_and_runs_of_blank_rows_are_dropped():
    """A formatted schedule's empty margin and spacer columns, and its blank
    rows, are layout. One blank row still separates two groups."""

    def build(workbook: Workbook) -> None:
        sheet = workbook["Pumps"]
        sheet["B2"] = "Tag"
        sheet["D2"] = "Flow"
        sheet["B3"] = "FP-1"
        sheet["D3"] = 2500
        sheet["B7"] = "V-1"
        sheet["E7"] = "8 in"
        sheet["B12"] = "End"

    text = extract_reference_document(_xlsx(build), filename="laid-out.xlsx").text

    assert text == (
        "[sheet 1 of 1: Pumps]\n"
        "Tag,Flow\n"
        "FP-1,2500\n"
        "\n"
        "V-1,,8 in\n"
        "\n"
        "End"
    )


def test_hidden_sheets_are_read_and_marked():
    def build(workbook: Workbook) -> None:
        workbook["Pumps"].append(["FP-1", 2500])
        lookup = workbook.create_sheet("Lookup")
        lookup.append(["Diesel", "UL listed"])
        lookup.sheet_state = "hidden"

    extraction = extract_reference_document(_xlsx(build), filename="h.xlsx")

    assert "[workbook sheets: 1. Pumps; 2. Lookup (hidden)]" in extraction.text
    assert "[hidden sheet 2 of 2: Lookup]\nDiesel,UL listed" in extraction.text
    assert any("1 hidden sheet(s) were read" in w for w in extraction.warnings)


def test_a_wrong_recorded_dimension_does_not_hide_rows():
    """Read-only mode trusts the sheet's recorded ``<dimension>``, and some
    writers record ``A1`` for a full sheet."""
    payload = _xlsx(
        lambda wb: [wb["Pumps"].append([f"FP-{n}", 2500]) for n in range(1, 4)]
    )
    lying = _rewrite(
        payload,
        "xl/worksheets/sheet1.xml",
        lambda xml: re.sub(r'<dimension ref="[^"]*"\s*/>', '<dimension ref="A1"/>', xml),
    )

    text = extract_reference_document(lying, filename="dim.xlsx").text

    assert text.endswith("FP-1,2500\nFP-2,2500\nFP-3,2500")


def test_chart_sheets_are_skipped_and_said():
    from openpyxl.chart import BarChart, Reference

    def build(workbook: Workbook) -> None:
        sheet = workbook["Pumps"]
        sheet.append(["FP-1", 2500])
        sheet.append(["FP-2", 1500])
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2))
        workbook.create_chartsheet("Flow chart").add_chart(chart)

    extraction = extract_reference_document(_xlsx(build), filename="c.xlsx")

    assert "Flow chart" not in extraction.text
    assert any("1 chart sheet(s)" in w for w in extraction.warnings)


def test_a_macro_enabled_workbook_keeps_its_name_and_reads_as_excel():
    client = TestClient(create_app())

    body = _upload(client, "Pump schedule.xlsm", _xlsx(_schedule)).json()

    assert body["reference_doc"]["filename"] == "Pump schedule.xlsm"
    assert body["reference_doc"]["kind"] == "xlsx"
    assert reference_extract.sanitize_reference_filename(
        "calc.xlsm", kind="xlsx"
    ) == "calc.xlsm"


def test_the_type_survives_a_project_round_trip():
    client = TestClient(create_app())
    _upload(client, "schedule.xlsx", _xlsx(_schedule))
    session = sessions.get_session()

    load_project(sessions.project_payload(session), session)

    assert session.references.docs[0].kind == "xlsx"
    assert session.references.docs[0].kind_label() == "Excel"


# ---------------------------------------------------------------------------
# Refusals — each with the reason and the way out
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "name, described",
    [
        ("old.xls", "Excel 97-2003 workbook (.xls)"),
        ("fast.xlsb", "Excel binary workbook (.xlsb)"),
        ("libre.ods", "OpenDocument spreadsheet (.ods)"),
        ("mac.numbers", "Numbers spreadsheet (.numbers)"),
    ],
)
def test_a_spreadsheet_format_no_reader_handles_names_the_way_out(name, described):
    client = TestClient(create_app())

    refused = _upload(client, name, b"anything")

    assert refused.status_code == 400
    error = refused.json()["error"]
    assert described in error
    assert "Excel Workbook (.xlsx)" in error
    assert ".csv" in error


def test_a_password_protected_workbook_is_refused_with_the_reason():
    """Excel encrypts a workbook that needs a password to open into an OLE
    compound file — not a ZIP at all."""
    client = TestClient(create_app())
    encrypted = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512

    refused = _upload(client, "locked.xlsx", encrypted)

    assert refused.status_code == 400
    assert "password-protected" in refused.json()["error"]


@pytest.mark.parametrize(
    "payload",
    [
        b"not a zip at all",
        # A ZIP, but a Word package rather than a workbook.
        pytest.param("docx", id="word-package"),
    ],
)
def test_a_file_that_is_not_a_workbook_is_refused(payload):
    if payload == "docx":
        from docx import Document

        buffer = io.BytesIO()
        Document().save(buffer)
        payload = buffer.getvalue()
    client = TestClient(create_app())

    refused = _upload(client, "renamed.xlsx", payload)

    assert refused.status_code == 400
    assert "could not be read as an Excel workbook" in refused.json()["error"]


def test_a_workbook_with_no_values_is_refused():
    with pytest.raises(ReferenceExtractError, match="No cell values"):
        extract_reference_document(_xlsx(lambda wb: None), filename="blank.xlsx")


def test_an_entity_declaration_is_refused_not_expanded():
    """defusedxml is what makes openpyxl refuse this. Without it the standard
    parser would expand the entity and read "44 degrees F"."""
    payload = _xlsx(lambda wb: wb["Pumps"].append(["placeholder"]))
    declared = _rewrite(
        _with_cell(
            payload, "A1", '<c r="A1" t="inlineStr"><is><t>&chw;</t></is></c>'
        ),
        "xl/worksheets/sheet1.xml",
        lambda xml: xml.replace(
            "<worksheet",
            '<!DOCTYPE worksheet [<!ENTITY chw "44 degrees F">]><worksheet',
            1,
        ),
    )

    with pytest.raises(ReferenceExtractError, match="Excel workbook"):
        extract_reference_document(declared, filename="entity.xlsx")


def test_an_entity_expansion_bomb_is_refused():
    payload = _xlsx(lambda wb: wb["Pumps"].append(["placeholder"]))
    entities = '<!ENTITY lol0 "lollollollollollollollollollol">' + "".join(
        f'<!ENTITY lol{i} "' + f"&lol{i - 1};" * 10 + '">' for i in range(1, 10)
    )
    bomb = _rewrite(
        _with_cell(
            payload, "A1", '<c r="A1" t="inlineStr"><is><t>&lol9;</t></is></c>'
        ),
        "xl/worksheets/sheet1.xml",
        lambda xml: xml.replace(
            "<worksheet", f"<!DOCTYPE worksheet [{entities}]><worksheet", 1
        ),
    )

    with pytest.raises(ReferenceExtractError):
        extract_reference_document(bomb, filename="bomb.xlsx")


def test_the_reader_fails_closed_without_its_xml_guard(monkeypatch):
    monkeypatch.setattr(openpyxl, "DEFUSEDXML", False)

    with pytest.raises(ReferenceExtractError, match="XML guard is missing"):
        extract_reference_document(_xlsx(_schedule), filename="s.xlsx")


def test_the_reader_and_its_guard_are_dependencies():
    """The guard is not optional, so it is a declared dependency, not an
    accident of whatever else is installed."""
    requirements = (REPO_ROOT / "requirements.txt").read_text(encoding="utf-8")
    assert re.search(r"^openpyxl>=", requirements, re.MULTILINE)
    assert re.search(r"^defusedxml>=", requirements, re.MULTILINE)
    assert openpyxl.DEFUSEDXML is True
    spec = (REPO_ROOT / "packaging" / "windows" / "build-a-spec.spec").read_text(
        encoding="utf-8"
    )
    assert '"defusedxml.ElementTree"' in spec


def test_an_oversized_member_is_refused_before_the_reader(monkeypatch):
    monkeypatch.setattr(reference_extract, "MAX_ZIP_MEMBER_BYTES", 64)

    with pytest.raises(ReferenceExtractError, match="unpacks to more data"):
        extract_reference_document(_xlsx(_schedule), filename="big.xlsx")


def test_too_many_members_are_refused_before_the_reader(monkeypatch):
    monkeypatch.setattr(reference_extract, "MAX_ZIP_MEMBERS", 3)

    with pytest.raises(ReferenceExtractError, match="too many parts"):
        extract_reference_document(_xlsx(_schedule), filename="many.xlsx")


def test_an_encrypted_member_is_refused():
    payload = bytearray(_xlsx(_schedule))
    # Set the encryption bit in the first central-directory entry's flags.
    central = payload.index(b"PK\x01\x02")
    payload[central + 8] |= 0x1

    with pytest.raises(ReferenceExtractError, match="encrypted"):
        extract_reference_document(bytes(payload), filename="enc.xlsx")


# ---------------------------------------------------------------------------
# The runaway breaker
# ---------------------------------------------------------------------------


def test_a_huge_workbook_is_bounded_and_says_so(monkeypatch):
    monkeypatch.setattr(reference_extract, "MAX_WORKBOOK_TEXT_CHARS", 40)

    def build(workbook: Workbook) -> None:
        for n in range(1, 11):
            workbook["Pumps"].append([f"FP-{n}", 2500])
        workbook.create_sheet("Valves").append(["V-1", 8])
        workbook.create_sheet("Notes").append(["note"])

    extraction = extract_reference_document(_xlsx(build), filename="huge.xlsx")

    assert "FP-1,2500" in extraction.text
    assert "FP-10" not in extraction.text
    assert "V-1" not in extraction.text
    assert extraction.text.endswith(
        "[... reading stopped at row 5 of sheet 'Pumps' (limit 40 "
        "characters); the rest of that sheet and 2 later sheet(s) were not "
        "read.]"
    )
    assert any("Reading stopped at row 5" in w for w in extraction.warnings)


def test_the_breaker_sits_above_what_a_reference_stores():
    """The stored truncation, which tells the model the tail was not read,
    is what normally cuts a long workbook; the breaker only bounds the CPU."""
    from backend.reference_docs import MAX_TEXT_CHARS

    assert reference_extract.MAX_WORKBOOK_TEXT_CHARS > MAX_TEXT_CHARS


def test_the_model_is_told_how_a_workbook_reads():
    """The markers, the hidden-sheet marker and the unsaved-formula reading
    are the app's conventions, not Excel's, so the stable policy says them."""
    from backend.llm.prompts import _REFERENCE_DOC_POLICY

    assert "Markdown, or an Excel workbook;" in _REFERENCE_DOC_POLICY
    assert "[sheet N of M: Name]" in _REFERENCE_DOC_POLICY
    assert "[hidden sheet …]" in _REFERENCE_DOC_POLICY
    assert "a cell that reads as a formula (=…) had no saved result" in (
        _REFERENCE_DOC_POLICY
    )


def _with_far_row(payload: bytes, row: int) -> bytes:
    """The first sheet's row 3 moved to ``row``, with no recorded
    ``<dimension>`` — what a crafted (or merely stray) far-out cell looks
    like to the reader."""

    def transform(xml: str) -> str:
        xml = re.sub(r'<dimension ref="[^"]*"\s*/>', "", xml)
        xml = xml.replace('<row r="3"', f'<row r="{row}"', 1)
        return xml.replace('r="A3"', f'r="A{row}"', 1)

    return _rewrite(payload, "xl/worksheets/sheet1.xml", transform)


def _far_build(workbook: Workbook) -> None:
    workbook["Pumps"].append(["FP-1", 2500])
    workbook["Pumps"]["A3"] = "far away"


def test_blank_rows_count_toward_the_breaker(monkeypatch):
    """openpyxl yields every blank row between two with values, and blank
    rows add no text, so the character breaker alone never stops a crafted
    row index far past Excel's last row (Codex review on PR #304). Row
    3,000,000 rather than a billion keeps a reverted breaker a failure, not
    a hang."""
    monkeypatch.setattr(reference_extract, "MAX_WORKBOOK_POSITIONS", 1_000)
    payload = _with_far_row(_xlsx(_far_build), 3_000_000)

    extraction = extract_reference_document(payload, filename="far.xlsx")

    assert "FP-1,2500" in extraction.text
    assert "far away" not in extraction.text
    assert re.search(
        r"\[\.\.\. reading stopped at row [\d,]+ of sheet 'Pumps' \(limit "
        r"1,000 rows and cells, blank ones included\); the rest of that "
        r"sheet were not read\.\]$",
        extraction.text,
    ), extraction.text[-200:]
    assert any("Reading stopped at row" in w for w in extraction.warnings)


def test_cell_positions_count_and_the_budget_spans_the_workbook(monkeypatch):
    """A wide row costs every position up to its last cell, and the budget is
    the workbook's, not each sheet's."""
    monkeypatch.setattr(reference_extract, "MAX_WORKBOOK_POSITIONS", 30)

    def build(workbook: Workbook) -> None:
        for n in range(1, 4):
            workbook["Pumps"].append([f"FP-{n}", 2500, 125])  # 4 positions a row
        workbook.create_sheet("Valves")["Z1"] = "wide"  # 27 positions

    extraction = extract_reference_document(_xlsx(build), filename="wide.xlsx")

    assert "FP-3,2500,125" in extraction.text
    assert "wide" not in extraction.text
    assert "row 1 of sheet 'Valves' (limit 30 rows and cells" in extraction.text


def test_a_stray_cell_at_excels_last_row_is_still_read():
    """The breaker is sized past a real sheet: one value at row 1,048,576 —
    a million blank rows away — is read under the default limit."""
    assert reference_extract.MAX_WORKBOOK_POSITIONS > 2 * 1_048_576
    payload = _with_far_row(_xlsx(_far_build), 1_048_576)

    extraction = extract_reference_document(payload, filename="last-row.xlsx")

    assert extraction.text.endswith("FP-1,2500\n\nfar away")
    assert extraction.warnings == ()
