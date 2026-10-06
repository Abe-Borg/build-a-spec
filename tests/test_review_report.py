"""The specification ends at END OF SECTION; the review trail is its own file.

Owner rule (2026-10-06): the specification carries specification text and
nothing else. The assumptions schedule, the unreviewed imported provisions,
the open items and the Final QC / compliance-audit closing used to follow END
OF SECTION in every Word export; they are now the review report, downloaded
only on request. A blank section number or title prints blank, never a
``[TBD]``. These tests pin both documents and the boundary between them.
"""
from __future__ import annotations

import io

from docx import Document
from fastapi.testclient import TestClient

from backend import sessions
from backend.app import create_app
from backend.spec_doc.docx_export import (
    build_docx,
    build_review_report,
    review_report_filename,
)
from backend.spec_doc.model import SpecSection, apply_edits
from tests.fakes import audit_grade_qc_result

_REVIEW_HEADINGS = (
    "ASSUMPTIONS SCHEDULE",
    "IMPORTED PROVISIONS NOT YET REVIEWED",
    "OPEN ITEMS",
    "FINAL QC SUMMARY",
    "COMPLIANCE AUDIT SUMMARY",
)


def _texts(payload: bytes) -> list[str]:
    return [p.text for p in Document(io.BytesIO(payload)).paragraphs]


def _section(**header) -> SpecSection:
    """Every kind of reviewer material: assumed, imported, needs-input, TBD."""
    ops = []
    if header:
        ops.append({"action": "replace", "target_id": "sec", **header})
    ops += [
        {"action": "add_article", "target_id": "pt1", "text": "SUMMARY"},
        {"action": "add_paragraph", "target_id": "pt1.a1", "text": "Provide sprinklers.", "status": "assumed"},
        {"action": "add_paragraph", "target_id": "pt1.a1", "text": "Coordinate risers.", "status": "confirmed"},
        {"action": "add_paragraph", "target_id": "pt1.a1", "text": "Density: [TBD: density].", "status": "needs_input"},
    ]
    section, _ = apply_edits(SpecSection.empty(), ops)
    section.parts[0].articles[0].paragraphs[1].status = "imported"
    return section


def test_the_specification_ends_at_end_of_section():
    section = _section(text="WET-PIPE SPRINKLER SYSTEMS", numbering="21 13 13")
    texts = _texts(build_docx(section))
    assert [t for t in texts if t.strip()][-1] == "END OF SECTION 21 13 13"
    for heading in _REVIEW_HEADINGS:
        assert heading not in texts
    assert Document(io.BytesIO(build_docx(section))).tables == []


def test_a_blank_header_prints_blank_never_tbd():
    section = _section()
    texts = _texts(build_docx(section))
    assert texts[:2] == ["SECTION", ""]
    assert [t for t in texts if t.strip()][-1] == "END OF SECTION"
    # The only TBD left is the provision's own text, which the export copies
    # as written; the export adds none.
    assert sum("[TBD" in t for t in texts) == 1


def test_the_review_report_carries_everything_written_for_the_reviewer():
    section = _section(text="WET-PIPE SPRINKLER SYSTEMS", numbering="21 13 13")
    audit = {
        "audited_at": "2026-10-06",
        "summary": "Two requirements covered.",
        "coverage": [{"status": "covered", "requirement_id": "r-1", "note": "Density per profile."}],
        "findings": [],
    }
    payload = build_review_report(
        section, audit_result=audit, version_index=4, generated_on="2026-10-06"
    )
    texts = _texts(payload)
    assert texts[:3] == [
        "REVIEW REPORT",
        "SECTION 21 13 13 - WET-PIPE SPRINKLER SYSTEMS",
        # The panel's v5 is stored index 4 (PR #278 review): both, the QC
        # report's wording.
        "Document version v5 (stored index 4) | Generated 2026-10-06",
    ]
    for heading in (
        "ASSUMPTIONS SCHEDULE",
        "IMPORTED PROVISIONS NOT YET REVIEWED",
        "OPEN ITEMS",
        "COMPLIANCE AUDIT SUMMARY",
    ):
        assert heading in texts
    cells = [
        (row.cells[0].text, row.cells[1].text)
        for table in Document(io.BytesIO(payload)).tables
        for row in table.rows[1:]
    ]
    assert ("1.1.A", "Provide sprinklers.") in cells
    assert ("1.1.B", "Coordinate risers.") in cells
    assert any(text.startswith("[NEEDS INPUT]") for _ref, text in cells)
    assert ("1.1.C", "[TBD] density") in cells
    # Nothing of the specification body: no PART heading, no END OF SECTION.
    assert "PART 1 - GENERAL" not in texts
    assert not any(t.startswith("END OF SECTION") for t in texts)


def test_a_blank_header_leaves_the_report_heading_out():
    texts = _texts(build_review_report(_section(), generated_on="2026-10-06"))
    assert texts[:2] == ["REVIEW REPORT", "Generated 2026-10-06"]


def test_the_report_is_named_after_the_section():
    section = _section(text="WET-PIPE SPRINKLER SYSTEMS", numbering="21 13 13")
    assert (
        review_report_filename(section)
        == "SECTION 21 13 13 - WET-PIPE SPRINKLER SYSTEMS - REVIEW REPORT.docx"
    )
    assert review_report_filename(_section()) == "DRAFT SECTION - REVIEW REPORT.docx"


def test_the_endpoint_serves_a_current_qc_closing_and_the_spec_export_does_not():
    client = TestClient(create_app())
    session = sessions.get_session()
    session.doc.begin_turn()
    session.doc.apply_edits(
        [
            {"action": "replace", "target_id": "sec", "text": "WET-PIPE SPRINKLER SYSTEMS", "numbering": "21 13 13"},
            {"action": "add_article", "target_id": "pt1", "text": "SUMMARY"},
            {"action": "add_paragraph", "target_id": "pt1.a1", "text": "Provide sprinklers.", "status": "confirmed"},
        ]
    )
    session.doc.commit_turn()
    session.qc.restore(audit_grade_qc_result(session, []))

    report = client.get("/api/export/review-report")
    assert report.status_code == 200
    assert report.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml"
    )
    assert "REVIEW REPORT.docx" in report.headers["content-disposition"]
    assert "FINAL QC SUMMARY" in _texts(report.content)

    spec = client.get("/api/export/docx")
    assert spec.status_code == 200
    assert "FINAL QC SUMMARY" not in _texts(spec.content)


def test_an_empty_document_still_gets_a_report():
    response = TestClient(create_app()).get("/api/export/review-report")
    assert response.status_code == 200
    texts = _texts(response.content)
    assert texts[0] == "REVIEW REPORT"
    assert "None — every provision is confirmed." in texts
