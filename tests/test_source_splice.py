"""The word-level splice (Redline on your original, D-2; built in Phase 0).

``plan_splice`` is pure — source text and target text in, an edit script
out — so its contract is pinned directly: the keep and delete steps
partition the source in order (Reject All would give back the original),
keep + insert reproduce the target's words (Accept All gives the edit), and
whitespace follows the rules the module states. The rendering half is
pinned on real paragraphs: unchanged words keep their runs, markers keep
their places, and anything the splice cannot map is refused with a named
reason rather than guessed at.
"""
from __future__ import annotations

import random

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from lxml import etree

from backend.spec_doc.importer import _accept_all_paragraph_text
from backend.spec_doc.source_splice import (
    OP_DELETE,
    OP_INSERT,
    OP_KEEP,
    map_paragraph,
    plan_splice,
    splice_paragraph,
    word_spans,
)
from tests.docx_fidelity_helpers import _append_hyperlink, _append_page_field


def _views(source: str, ops) -> tuple[str, str]:
    """(accept, reject): the texts Accept All and Reject All would give."""
    accept = []
    reject = []
    for op in ops:
        if op.op == OP_INSERT:
            accept.append(op.text)
        elif op.op == OP_KEEP:
            accept.append(source[op.start : op.end])
            reject.append(source[op.start : op.end])
        else:
            reject.append(source[op.start : op.end])
    return "".join(accept), "".join(reject)


def _assert_script(source: str, target: str) -> str:
    ops = plan_splice(source, target)
    cursor = 0
    for op in ops:
        if op.op in (OP_KEEP, OP_DELETE):
            assert op.start == cursor, (source, target, ops)
            assert op.end > op.start
            cursor = op.end
        else:
            assert op.text
    assert cursor == len(source), (source, target, ops)
    accept, reject = _views(source, ops)
    assert reject == source
    assert accept.split() == target.split(), (source, target, accept)
    return accept


def test_word_spans_follow_str_split():
    for text in ("", "   ", "a", " a  b\tc\n", "A. Nbsp", "x y"):
        assert [text[a:b] for a, b in word_spans(text)] == text.split()


def test_relettering_keeps_the_tab_after_the_letter():
    assert _assert_script(
        "A.\tSection includes vibration isolation.",
        "B. Section includes vibration isolation.",
    ) == "B.\tSection includes vibration isolation."


def test_an_insertion_after_the_label_keeps_the_tab_before_it():
    assert _assert_script("A.\tSection includes.", "A. New Section includes.") == (
        "A.\tNew Section includes."
    )


def test_a_deleted_word_takes_the_gap_after_it():
    assert _assert_script(
        "A.\tSection includes vibration.", "A. includes vibration."
    ) == "A.\tincludes vibration."
    assert _assert_script("Foo bar baz", "bar baz") == "bar baz"


def test_a_deletion_at_the_end_leaves_no_trailing_gap():
    assert _assert_script("Provide spring isolators", "Provide spring") == (
        "Provide spring"
    )


def test_survivors_keep_the_source_spacing_between_them():
    assert _assert_script(
        "Provide isolators.  Comply.", "Provide isolators. Comply. Now."
    ) == "Provide isolators.  Comply. Now."


def test_an_insertion_at_the_end_is_separated_from_the_last_word():
    assert _assert_script("Provide isolators.", "Provide isolators. Comply.") == (
        "Provide isolators. Comply."
    )


def test_edge_scripts():
    assert _assert_script("", "New text") == "New text"
    assert _assert_script("   ", "New") == "   New"
    assert _assert_script("Old words", "Brand new") == "Brand new"
    assert _assert_script("Some words", "") == ""
    assert _assert_script("Same words", "Same words") == "Same words"


def test_new_text_takes_the_formatting_word_would_give_it():
    """Typed over a selection: the first replaced character's formatting.
    Inserted: the character before it; at the very start, the first."""
    source = "Section includes vibration isolation."
    (replace,) = [
        op for op in plan_splice(source, "Section includes seismic isolation.")
        if op.op == OP_INSERT
    ]
    assert replace.style_at == source.index("vibration")
    (insert,) = [
        op for op in plan_splice(source, "Section includes spring vibration isolation.")
        if op.op == OP_INSERT
    ]
    assert insert.style_at == source.index("vibration") - 1
    (first,) = [
        op for op in plan_splice(source, "Also Section includes vibration isolation.")
        if op.op == OP_INSERT
    ]
    assert first.style_at == 0
    (last,) = [
        op for op in plan_splice(source, "Section includes vibration isolation. Done.")
        if op.op == OP_INSERT
    ]
    assert last.style_at == len(source) - 1


def test_the_script_invariants_hold_across_many_edits():
    """Random words, random whitespace (tabs, double spaces, a leading and
    trailing gap), random insertions, deletions and replacements."""
    rng = random.Random(20260922)
    vocabulary = [
        "A.", "B.", "1.", "Section", "includes", "vibration", "isolation",
        "for", "mechanical", "equipment.", "NFPA", "13-2025.", "Provide",
        "spring", "isolators", "and", "restraints", "(TBD)",
    ]
    gaps = [" ", " ", " ", "  ", "\t", "\n", " \t"]
    for _ in range(600):
        words = [rng.choice(vocabulary) for _ in range(rng.randint(0, 9))]
        source = rng.choice(["", " ", "\t"])
        for index, word in enumerate(words):
            if index:
                source += rng.choice(gaps)
            source += word
        source += rng.choice(["", "", " "])
        edited = list(words)
        for _ in range(rng.randint(0, 4)):
            action = rng.choice(["insert", "delete", "replace"])
            position = rng.randint(0, len(edited))
            if action == "insert":
                edited.insert(position, rng.choice(vocabulary))
            elif edited and position < len(edited):
                if action == "delete":
                    del edited[position]
                else:
                    edited[position] = rng.choice(vocabulary)
        _assert_script(source, " ".join(edited))


# ---------------------------------------------------------------------------
# The character map and the rendering
# ---------------------------------------------------------------------------


def _paragraph():
    return Document().add_paragraph()


def test_the_map_reads_what_the_importer_reads():
    paragraph = _paragraph()
    paragraph.add_run("A.\tSection ")
    run = paragraph.add_run("vibration")
    run.bold = True
    run.add_break()  # a text-wrapping break reads as "\n"
    paragraph.add_run("isolation").add_break(WD_BREAK.PAGE)  # reads as ""
    etree.SubElement(paragraph.runs[0]._r, qn("w:lastRenderedPageBreak"))
    pmap, reason = map_paragraph(paragraph._p)
    assert reason == ""
    assert pmap.text == _accept_all_paragraph_text(paragraph._p)
    assert pmap.text == "A.\tSection vibration\nisolation"


def test_unmappable_paragraphs_are_refused_by_name():
    # A paragraph's own hyperlink is mapped (the hyperlink block below); what
    # a link may not hold is refused by name, like the same thing outside it.
    for inner, reason in (
        ("w:hyperlink", "hyperlink"),  # a link nested in a link
        ("w:fldSimple", "field"),
        ("w:sdt", "content_control"),
        ("w:commentRangeStart", "comment"),
        ("w:ins", "revisions"),
        ("w:smartTag", "other_markup"),
    ):
        linked = _paragraph()
        linked.add_run("See ")
        _append_hyperlink(linked, "the client standard", "https://example.com")
        linked._p.find(qn("w:hyperlink")).append(etree.Element(qn(inner)))
        assert map_paragraph(linked._p) == (None, reason), inner
    run_field = _paragraph()
    _append_hyperlink(run_field, "page", "https://example.com")
    run = run_field._p.find(f"{qn('w:hyperlink')}/{qn('w:r')}")
    etree.SubElement(run, qn("w:fldChar")).set(qn("w:fldCharType"), "begin")
    assert map_paragraph(run_field._p) == (None, "field")

    field = _paragraph()
    field.add_run("Page ")
    _append_page_field(field)
    assert map_paragraph(field._p) == (None, "field")

    revised = _paragraph()
    inserted = etree.SubElement(revised._p, qn("w:ins"))
    etree.SubElement(etree.SubElement(inserted, qn("w:r")), qn("w:t")).text = "new"
    assert map_paragraph(revised._p) == (None, "revisions")

    symbol = _paragraph()
    etree.SubElement(symbol.add_run("x")._r, qn("w:sym"))
    assert map_paragraph(symbol._p) == (None, "symbol")

    plain = _paragraph()
    plain.add_run("Plain.")
    assert map_paragraph(plain._p, expected_text="Something else") == (
        None,
        "text_mismatch",
    )


def test_custom_xml_never_reaches_the_word_level_redline():
    """Word will not load inline custom XML inside ``w:ins``/``w:del``
    ([MS-OI29500] §2.1.188(a)). ``render_redline`` wraps only the runs it
    maps, so the guarantee is that it never maps one: a paragraph holding
    custom XML — at its own level or inside a hyperlink — is refused
    (``other_markup``), and the fallback marks it through
    ``revision_marks``, which keeps the element outside every wrapper."""
    for inside_link in (False, True):
        paragraph = _paragraph()
        paragraph.add_run("See ")
        if inside_link:
            _append_hyperlink(paragraph, "the client standard", "https://example.com")
            container = paragraph._p.find(qn("w:hyperlink"))
        else:
            container = paragraph._p
        custom = etree.SubElement(container, qn("w:customXml"))
        etree.SubElement(custom, qn("w:customXmlPr"))
        etree.SubElement(etree.SubElement(custom, qn("w:r")), qn("w:t")).text = "term"
        assert map_paragraph(paragraph._p) == (None, "other_markup"), inside_link


def _runs(element) -> list[tuple[bool, str]]:
    """(bold, text) per run, tabs shown as \\t."""
    runs = []
    for run in element.iter(qn("w:r")):
        properties = run.find(qn("w:rPr"))
        bold = properties is not None and properties.find(qn("w:b")) is not None
        text = "".join(
            "\t" if node.tag == qn("w:tab") else (node.text or "")
            for node in run
            if node.tag in (qn("w:t"), qn("w:tab"))
        )
        runs.append((bold, text))
    return runs


def test_unchanged_words_keep_their_own_runs():
    paragraph = _paragraph()
    paragraph.add_run("A.\tSection includes ")
    paragraph.add_run("vibration isolation").bold = True
    paragraph.add_run(" for mechanical equipment.")

    spliced, reason = splice_paragraph(
        paragraph._p, "B. Section includes vibration isolation for HVAC equipment."
    )

    assert reason == ""
    assert "".join(text for _bold, text in _runs(spliced)) == (
        "B.\tSection includes vibration isolation for HVAC equipment."
    )
    assert [text for bold, text in _runs(spliced) if bold] == ["vibration isolation"]
    # The source element itself is never mutated.
    assert _accept_all_paragraph_text(paragraph._p) == (
        "A.\tSection includes vibration isolation for mechanical equipment."
    )


def test_markers_keep_their_places_and_a_leading_page_break_survives():
    """A bookmark around the text still wraps it after the letter changes,
    and a page break before a relettered label is not deleted with the old
    letter (it sits on the edge of the deletion, not inside it)."""
    paragraph = _paragraph()
    start = etree.SubElement(paragraph._p, qn("w:bookmarkStart"))
    start.set(qn("w:id"), "1")
    start.set(qn("w:name"), "_Ref1")
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    paragraph.add_run("A.\tProvide isolators.")
    end = etree.SubElement(paragraph._p, qn("w:bookmarkEnd"))
    end.set(qn("w:id"), "1")

    spliced, reason = splice_paragraph(paragraph._p, "B. Provide isolators.")

    assert reason == ""
    tags = [etree.QName(child).localname for child in spliced if child.tag != qn("w:pPr")]
    assert tags[0] == "bookmarkStart" and tags[-1] == "bookmarkEnd"
    breaks = [
        br for br in spliced.iter(qn("w:br")) if br.get(qn("w:type")) == "page"
    ]
    assert len(breaks) == 1
    assert _accept_all_paragraph_text(spliced) == "B.\tProvide isolators."


def _layout(element) -> list[str]:
    """The paragraph's content in document order: ``<page>`` / ``<column>``
    for a break, ``[name]`` / ``[/name]`` for a bookmark, text otherwise."""
    names = {}
    tokens: list[str] = []
    for node in element.iter():
        if node.tag == qn("w:bookmarkStart"):
            names[node.get(qn("w:id"))] = node.get(qn("w:name"))
            tokens.append(f"[{node.get(qn('w:name'))}]")
        elif node.tag == qn("w:bookmarkEnd"):
            tokens.append(f"[/{names.get(node.get(qn('w:id')), '?')}]")
        elif node.tag == qn("w:br") and node.get(qn("w:type")) in ("page", "column"):
            tokens.append(f"<{node.get(qn('w:type'))}>")
        elif node.tag == qn("w:tab"):
            tokens.append("\t")
        elif node.tag == qn("w:t") and node.text:
            tokens.append(node.text)
    # Adjacent text tokens are one run of text to the reader.
    merged: list[str] = []
    for token in tokens:
        plain = not token.startswith(("<", "["))
        if plain and merged and not merged[-1].startswith(("<", "[")):
            merged[-1] += token
        else:
            merged.append(token)
    return merged


def test_words_prepended_to_a_paragraph_join_it_after_its_leading_break():
    """A page or column break in front of a paragraph's first word starts the
    paragraph on a new page or column. Words prepended to the paragraph join
    it there; they are not stranded on the page before the break (Codex
    review on PR #184). The rule is about LEADING content only: a break at
    the end of a paragraph stays after words appended to it."""
    own_run = _paragraph()
    own_run.add_run().add_break(WD_BREAK.PAGE)
    own_run.add_run("Provide isolators.")
    spliced, reason = splice_paragraph(own_run._p, "Also Provide isolators.")
    assert reason == ""
    assert _layout(spliced) == ["<page>", "Also Provide isolators."]

    same_run = _paragraph()
    run = same_run.add_run()
    run.add_break(WD_BREAK.COLUMN)
    etree.SubElement(run._r, qn("w:t")).text = "Provide isolators."
    spliced, reason = splice_paragraph(same_run._p, "Also Provide isolators.")
    assert reason == ""
    assert _layout(spliced) == ["<column>", "Also Provide isolators."]

    # "Before the first word", not "at offset 0": leading whitespace is kept
    # where it was, and the break still leads the words.
    indented = _paragraph()
    indented.add_run("\t")
    indented.add_run().add_break(WD_BREAK.PAGE)
    indented.add_run("Provide isolators.")
    spliced, reason = splice_paragraph(indented._p, "Also Provide isolators.")
    assert reason == ""
    assert _layout(spliced) == ["\t", "<page>", "Also Provide isolators."]

    # A bookmark opening in front of the first word still wraps all of it.
    bookmarked = _paragraph()
    start = etree.SubElement(bookmarked._p, qn("w:bookmarkStart"))
    start.set(qn("w:id"), "7")
    start.set(qn("w:name"), "_Toc7")
    bookmarked.add_run("SUMMARY")
    end = etree.SubElement(bookmarked._p, qn("w:bookmarkEnd"))
    end.set(qn("w:id"), "7")
    spliced, reason = splice_paragraph(bookmarked._p, "GENERAL SUMMARY")
    assert reason == ""
    assert _layout(spliced) == ["[_Toc7]", "GENERAL SUMMARY", "[/_Toc7]"]

    # The end is not the start: a trailing break stays after appended words.
    trailing = _paragraph()
    trailing.add_run("Provide isolators.")
    trailing.add_run().add_break(WD_BREAK.PAGE)
    spliced, reason = splice_paragraph(trailing._p, "Provide isolators. Comply.")
    assert reason == ""
    assert _layout(spliced) == ["Provide isolators. Comply.", "<page>"]

    # And a break in front of a later word stays in front of that word.
    middle = _paragraph()
    middle.add_run("Provide ")
    middle.add_run().add_break(WD_BREAK.PAGE)
    middle.add_run("isolators.")
    spliced, reason = splice_paragraph(middle._p, "Provide spring isolators.")
    assert reason == ""
    assert _layout(spliced) == ["Provide spring ", "<page>", "isolators."]


def test_a_zero_width_node_inside_a_deleted_word_goes_with_it():
    paragraph = _paragraph()
    run = paragraph.add_run("Provide spring")
    etree.SubElement(run._r, qn("w:softHyphen"))
    etree.SubElement(run._r, qn("w:t")).text = "loaded isolators."

    spliced, reason = splice_paragraph(paragraph._p, "Provide isolators.")

    assert reason == ""
    assert spliced.find(f".//{qn('w:softHyphen')}") is None
    assert _accept_all_paragraph_text(spliced) == "Provide isolators."


# ---------------------------------------------------------------------------
# The redline rendering (Phase 1, D-2): the same pieces, tracked
# ---------------------------------------------------------------------------


def _marks():
    from backend.spec_doc.revision_marks import RevisionMarks

    return RevisionMarks(author="Build-a-Spec", date="2026-09-22T00:00:00Z", first_id=100)


def _in_body(*paragraphs):
    from copy import deepcopy

    body = etree.Element(qn("w:body"))
    for paragraph in paragraphs:
        body.append(deepcopy(paragraph))
    return body


def _rich_paragraph():
    """Typed letter + tab, a bold phrase, a bookmark over a word, a page
    break before a word and Word's layout cache — every kind of atom."""
    paragraph = _paragraph()
    paragraph.add_run("A.\tSection includes ")
    paragraph.add_run("vibration isolation").bold = True
    start = etree.SubElement(paragraph._p, qn("w:bookmarkStart"))
    start.set(qn("w:id"), "5")
    start.set(qn("w:name"), "_Ref5")
    paragraph.add_run(" for mechanical")
    end = etree.SubElement(paragraph._p, qn("w:bookmarkEnd"))
    end.set(qn("w:id"), "5")
    tail = paragraph.add_run(" equipment and ")
    etree.SubElement(tail._r, qn("w:lastRenderedPageBreak"))
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    paragraph.add_run("piping.")
    return paragraph._p


def test_the_redline_accepts_to_the_clean_render_and_rejects_to_the_source():
    """D-2's promise on one paragraph, over 400 seeded edits: Accept All of
    the tracked rendering IS the clean export's rendering, and Reject All
    IS the source paragraph — every run, bookmark and break in place."""
    from backend.spec_doc.revisions import accept_all, first_difference, reject_all
    from backend.spec_doc.source_splice import render_clean, render_redline

    source = _rich_paragraph()
    pmap, reason = map_paragraph(source, expected_text=_accept_all_paragraph_text(source))
    assert reason == ""
    words = pmap.text.split()
    vocabulary = ["seismic", "Provide", "restraints", "B.", "per", "NFPA", "13."]
    rng = random.Random(20260923)
    for _ in range(400):
        edited = list(words)
        for _ in range(rng.randint(1, 4)):
            action = rng.choice(["insert", "delete", "replace"])
            position = rng.randint(0, len(edited))
            if action == "insert":
                edited.insert(position, rng.choice(vocabulary))
            elif edited and position < len(edited):
                if action == "delete":
                    del edited[position]
                else:
                    edited[position] = rng.choice(vocabulary)
        ops = plan_splice(pmap.text, " ".join(edited))
        redline = _in_body(render_redline(pmap, ops, _marks()))
        clean = _in_body(render_clean(pmap, ops))
        assert first_difference(accept_all(redline), clean) is None, edited
        assert first_difference(reject_all(redline), _in_body(source)) is None, edited


def test_deleted_words_are_deleted_text_inside_their_own_runs():
    from backend.spec_doc.source_splice import render_redline

    paragraph = _paragraph()
    paragraph.add_run("Section includes ")
    paragraph.add_run("vibration isolation").bold = True
    pmap, _ = map_paragraph(paragraph._p)
    redline = render_redline(
        pmap, plan_splice(pmap.text, "Section includes seismic isolation"), _marks()
    )
    (deleted,) = redline.findall(qn("w:del"))
    (inserted,) = redline.findall(qn("w:ins"))
    assert [t.text for t in deleted.iter(qn("w:delText"))] == ["vibration"]
    assert deleted.find(f".//{qn('w:t')}") is None
    # The replaced word was bold, and the word typed over it takes that.
    assert inserted.find(f"{qn('w:r')}/{qn('w:rPr')}/{qn('w:b')}") is not None
    for wrapper in (deleted, inserted):
        assert wrapper.get(qn("w:author")) == "Build-a-Spec"
        assert wrapper.get(qn("w:date")) == "2026-09-22T00:00:00Z"
        assert int(wrapper.get(qn("w:id"))) >= 100


# ---------------------------------------------------------------------------
# Hyperlinks: mapped like the paragraph around them, and never grown or split
# ---------------------------------------------------------------------------


def _link(paragraph, *content, anchor="_Target"):
    """Append a ``w:hyperlink`` to an internal ``anchor``: each item of
    ``content`` is a run's text (``"*text"``: bold), a ``("[", name, id)`` /
    ``("]", id)`` bookmark start / end, or ``"proofErr"``. Its runs carry the
    Hyperlink character style, which is what a word added beside the link
    must never borrow."""
    link = etree.SubElement(paragraph._p, qn("w:hyperlink"))
    link.set(qn("w:anchor"), anchor)
    for item in content:
        if isinstance(item, tuple):
            if item[0] == "[":
                marker = etree.SubElement(link, qn("w:bookmarkStart"))
                marker.set(qn("w:id"), item[2])
                marker.set(qn("w:name"), item[1])
            else:
                etree.SubElement(link, qn("w:bookmarkEnd")).set(qn("w:id"), item[1])
            continue
        if item == "proofErr":
            etree.SubElement(link, qn("w:proofErr")).set(qn("w:type"), "spellStart")
            continue
        run = etree.SubElement(link, qn("w:r"))
        properties = etree.SubElement(run, qn("w:rPr"))
        etree.SubElement(properties, qn("w:rStyle")).set(qn("w:val"), "Hyperlink")
        if item.startswith("*"):
            etree.SubElement(properties, qn("w:b"))
            item = item[1:]
        node = etree.SubElement(run, qn("w:t"))
        node.text = item
        node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return link


def _shape(element) -> list[str]:
    """The paragraph as a reader sees it: ``LINK[...]`` around a link's
    content, ``~`` in front of a run that carries the Hyperlink style, and
    ``<name>`` for a marker."""
    def run_text(run):
        styled = run.find(f"{qn('w:rPr')}/{qn('w:rStyle')}") is not None
        text = "".join(
            "\t" if node.tag == qn("w:tab") else (node.text or "")
            for node in run
            if node.tag in (qn("w:t"), qn("w:tab"))
        )
        return ("~" if styled else "") + text

    shape = []
    for child in element:
        if child.tag == qn("w:pPr"):
            continue
        if child.tag == qn("w:hyperlink"):
            inner = [
                run_text(c) if c.tag == qn("w:r") else f"<{etree.QName(c).localname}>"
                for c in child
            ]
            shape.append("LINK[" + "|".join(inner) + "]")
        elif child.tag == qn("w:r"):
            shape.append(run_text(child))
        else:
            shape.append(f"<{etree.QName(child).localname}>")
    return shape


def _see_the_standard():
    paragraph = _paragraph()
    paragraph.add_run("See ")
    _link(paragraph, "the client standard")
    paragraph.add_run(" for isolators.")
    return paragraph


def _clean(paragraph, target):
    spliced, reason = splice_paragraph(
        paragraph._p, target, expected_text=_accept_all_paragraph_text(paragraph._p)
    )
    assert reason == ""
    assert _accept_all_paragraph_text(spliced) == target
    return _shape(spliced)


def test_the_map_reads_a_hyperlink_the_way_the_importer_does():
    """The importer reads a link as the text of its own runs (python-docx
    ``CT_Hyperlink.text``); the map reads the same characters, and every
    atom remembers the link it sits in."""
    paragraph = _paragraph()
    paragraph.add_run("See ")
    _link(paragraph, ("[", "_Ref2", "2"), "the ", "proofErr", "*client", " standard", ("]", "2"))
    paragraph.add_run(".")
    pmap, reason = map_paragraph(
        paragraph._p, expected_text=_accept_all_paragraph_text(paragraph._p)
    )
    assert reason == ""
    assert pmap.text == "See the client standard."
    assert len(pmap.links) == 1
    assert [pmap.link_at(i) for i in range(len(pmap.text))] == (
        [-1] * 4 + [0] * len("the client standard") + [-1]
    )
    assert pmap.run_links == (-1, 0, 0, 0, -1)
    markers = [a for a in pmap.atoms if a.run < 0]
    assert [etree.QName(a.node).localname for a in markers] == [
        "bookmarkStart",
        "proofErr",
        "bookmarkEnd",
    ]
    assert {a.link for a in markers} == {0}


def test_an_unchanged_link_keeps_its_runs_target_and_formatting():
    """The reported shape: relettering a provision that holds a link keeps
    the link (its target and its runs, bold included) and the tab after the
    letter — it used to take the fallback and become plain text."""
    paragraph = _paragraph()
    paragraph.add_run("A.\tSection includes ")
    paragraph.add_run("vibration isolation").bold = True
    paragraph.add_run(" per ")
    link = _link(paragraph, "the ", "*client", " standard", anchor="_Std")
    paragraph.add_run(" for equipment.")
    spliced, reason = splice_paragraph(
        paragraph._p,
        "B. Section includes vibration isolation per the client standard for equipment.",
    )
    assert reason == ""
    assert _shape(spliced) == [
        "B.",
        "\tSection includes ",
        "vibration isolation",
        " per ",
        "LINK[~the |~client|~ standard]",
        " for equipment.",
    ]
    (copied,) = spliced.findall(qn("w:hyperlink"))
    assert copied.get(qn("w:anchor")) == link.get(qn("w:anchor"))
    assert [t.text for t in copied.iter(qn("w:t"))] == ["the ", "client", " standard"]
    assert copied.find(f"{qn('w:r')}[2]/{qn('w:rPr')}/{qn('w:b')}") is not None


def test_words_typed_over_a_links_own_words_stay_in_the_link():
    """The link's display text changed; its target did not."""
    assert _clean(_see_the_standard(), "See the owner standard for isolators.") == [
        "See ",
        "LINK[~the |~owner|~ standard]",
        " for isolators.",
    ]
    assert _clean(_see_the_standard(), "See the NFPA code for isolators.") == [
        "See ",
        "LINK[~the |~NFPA code]",
        " for isolators.",
    ]


def test_words_inserted_between_two_characters_of_a_link_go_in_it():
    assert _clean(_see_the_standard(), "See the client design standard for isolators.") == [
        "See ",
        "LINK[~the client |~design |~standard]",
        " for isolators.",
    ]


def test_a_link_never_grows_to_cover_words_added_beside_it():
    """At either edge of a link — and after a link that holds its own
    trailing space — new words go outside it, and do not borrow its look."""
    assert _clean(_see_the_standard(), "See the client standard now for isolators.") == [
        "See ",
        "LINK[~the client standard]",
        " ",
        "now ",
        "for isolators.",
    ]
    assert _clean(_see_the_standard(), "See also the client standard for isolators.") == [
        "See ",
        "also ",
        "LINK[~the client standard]",
        " for isolators.",
    ]
    ending = _paragraph()
    ending.add_run("B. Unrelated linked provision: ")
    _link(ending, "client requirements")
    assert _clean(ending, "B. Unrelated linked provision: client requirements apply.") == [
        "B. Unrelated linked provision: ",
        "LINK[~client requirements]",
        " apply.",
    ]
    opening = _paragraph()
    _link(opening, "NFPA 13")
    opening.add_run(" applies.")
    assert _clean(opening, "Comply with NFPA 13 applies.") == [
        "Comply with ",
        "LINK[~NFPA 13]",
        " applies.",
    ]
    spaced = _paragraph()
    spaced.add_run("See ")
    _link(spaced, "NFPA 13 ")
    spaced.add_run("for design.")
    assert _clean(spaced, "See NFPA 13 current for design.") == [
        "See ",
        "LINK[~NFPA 13 ]",
        "current ",
        "for design.",
    ]


def test_new_words_beside_a_link_take_the_nearest_formatting_outside_it():
    """Word's rule, refined at a link: typed over a link's first word but
    going outside it, the new words take the nearest character outside the
    link — here a bold one — and a paragraph that is nothing but a link has
    no such character, so they take the first run's outside it: none."""
    paragraph = _paragraph()
    paragraph.add_run("See ")
    _link(paragraph, "the client standard")
    paragraph.add_run(" for").bold = True
    paragraph.add_run(" isolators.")
    spliced, reason = splice_paragraph(paragraph._p, "See the client rules isolators.")
    assert reason == ""
    new = [r for r in spliced.findall(qn("w:r")) if "rules" in "".join(r.itertext())]
    (run,) = new
    assert run.find(f"{qn('w:rPr')}/{qn('w:b')}") is not None
    assert run.find(f"{qn('w:rPr')}/{qn('w:rStyle')}") is None

    only = _paragraph()
    _link(only, "NFPA 13")
    spliced, reason = splice_paragraph(only._p, "Comply with NFPA 13")
    assert reason == ""
    assert _shape(spliced) == ["Comply with ", "LINK[~NFPA 13]"]
    assert spliced.find(f"{qn('w:r')}/{qn('w:rPr')}") is None


def test_a_replacement_running_into_a_link_goes_in_front_of_it():
    """Words typed over words on both sides of a link's edge go outside the
    link: in front of it when the replacement runs INTO the link (written
    after the deletion, they would cut the link in two), after it when it
    runs out."""
    assert _clean(_see_the_standard(), "Under client standard for isolators.") == [
        "Under",
        "LINK[~ client standard]",
        " for isolators.",
    ]
    assert _clean(_see_the_standard(), "See the client rules isolators.") == [
        "See ",
        "LINK[~the client ]",
        "rules",
        " isolators.",
    ]


def test_a_link_whose_words_are_all_deleted_is_gone_unless_a_marker_keeps_it():
    assert _clean(_see_the_standard(), "See for isolators.") == [
        "See ",
        "for isolators.",
    ]
    marked = _paragraph()
    marked.add_run("See ")
    _link(marked, "the standard", ("]", "3"))
    marked.add_run(" now.")
    assert _clean(marked, "See now.") == ["See ", "LINK[<bookmarkEnd>]", "now."]


def test_a_links_zero_width_content_stays_inside_it():
    """A bookmark closing inside a link stays inside it when new words land
    right after it, and a bookmark opening a link stays with it when words
    are prepended in front of it — neither splits the link."""
    closing = _paragraph()
    closing.add_run("See ")
    _link(closing, "the standard", ("]", "3"))
    closing.add_run(" now.")
    assert _clean(closing, "Rules now.") == ["LINK[<bookmarkEnd>]", "Rules", " now."]

    opening = _paragraph()
    _link(opening, ("[", "_Ref4", "4"), "NFPA 13")
    opening.add_run(" applies.")
    assert _clean(opening, "Comply with NFPA 13 applies.") == [
        "Comply with ",
        "LINK[<bookmarkStart>|~NFPA 13]",
        " applies.",
    ]


def test_an_empty_hyperlink_is_carried_where_it_was():
    paragraph = _paragraph()
    paragraph.add_run("Provide ")
    etree.SubElement(paragraph._p, qn("w:hyperlink")).set(qn("w:anchor"), "_Empty")
    paragraph.add_run("isolators.")
    pmap, reason = map_paragraph(paragraph._p)
    assert reason == "" and pmap.links == ()
    spliced, _ = splice_paragraph(paragraph._p, "Provide spring isolators.")
    assert _shape(spliced) == ["Provide ", "spring ", "LINK[]", "isolators."]


def test_a_paragraph_without_links_keeps_the_scripts_own_choices():
    """Every link rule is a no-op without a link: the placement of new words
    is exactly the edit script's (``style_at``, outside every link), which
    is what keeps a link-free paragraph's rendering byte for byte what it
    was before links were mapped."""
    from backend.spec_doc.source_splice import _placements

    source = _rich_paragraph()
    pmap, _ = map_paragraph(source)
    ops = plan_splice(pmap.text, "B. Also Section seismic isolation for piping.")
    placements = _placements(pmap, ops)
    assert placements
    for index, placement in placements.items():
        assert (placement.link, placement.style, placement.split) == (
            -1,
            ops[index].style_at,
            -1,
        )


def _linked_paragraph():
    """Every link shape at once: a link opening the paragraph, a bold word
    beside one, a link holding a bookmark and a spelling marker and a bold
    run, a link holding its trailing space, and one ending the paragraph."""
    paragraph = _paragraph()
    _link(paragraph, "Scope", anchor="_Top")
    paragraph.add_run(":\tSection includes ")
    paragraph.add_run("vibration isolation").bold = True
    paragraph.add_run(" per ")
    _link(
        paragraph,
        ("[", "_Ref9", "9"),
        "the client ",
        "*standard",
        ("]", "9"),
        "proofErr",
        anchor="_Std",
    )
    paragraph.add_run(" and ")
    _link(paragraph, "NFPA 13 ", anchor="_Nfpa")
    paragraph.add_run("for equipment, see ")
    _link(paragraph, "Section 23 05 00.", anchor="_Sec")
    return paragraph._p


def test_the_redline_keeps_its_promise_on_paragraphs_with_links():
    """D-2's promise over 600 seeded edits of a paragraph holding every link
    shape: Accept All of the tracked rendering IS the clean rendering,
    Reject All IS the source paragraph, the importer reads the clean
    rendering as the edit, and no link is ever split in two."""
    from backend.spec_doc.revisions import accept_all, first_difference, reject_all
    from backend.spec_doc.source_splice import render_clean, render_redline

    source = _linked_paragraph()
    pmap, reason = map_paragraph(source, expected_text=_accept_all_paragraph_text(source))
    assert reason == ""
    links = len(source.findall(qn("w:hyperlink")))
    words = pmap.text.split()
    vocabulary = ["seismic", "Provide", "restraints", "per", "NFPA", "client", "Section"]
    rng = random.Random(20260923)
    for _ in range(600):
        edited = list(words)
        for _ in range(rng.randint(1, 5)):
            action = rng.choice(["insert", "delete", "replace"])
            position = rng.randint(0, len(edited))
            if action == "insert":
                edited.insert(position, rng.choice(vocabulary))
            elif edited and position < len(edited):
                if action == "delete":
                    del edited[position]
                else:
                    edited[position] = rng.choice(vocabulary)
        ops = plan_splice(pmap.text, " ".join(edited))
        redline = render_redline(pmap, ops, _marks())
        clean = render_clean(pmap, ops)
        assert first_difference(accept_all(_in_body(redline)), _in_body(clean)) is None, edited
        assert first_difference(reject_all(_in_body(redline)), _in_body(source)) is None, edited
        assert _accept_all_paragraph_text(clean).split() == edited
        assert len(redline.findall(qn("w:hyperlink"))) <= links, edited
        assert len(clean.findall(qn("w:hyperlink"))) <= links, edited


def test_a_tracked_edit_inside_a_link_is_marked_inside_it():
    """A hyperlink cannot sit inside ``w:del``/``w:ins``; its runs can. The
    edit of a link's word is tracked inside the one copy of the link."""
    from backend.spec_doc.source_splice import render_redline

    paragraph = _see_the_standard()
    pmap, _ = map_paragraph(paragraph._p)
    redline = render_redline(
        pmap, plan_splice(pmap.text, "See the owner standard for isolators."), _marks()
    )
    (link,) = redline.findall(qn("w:hyperlink"))
    (deleted,) = link.findall(qn("w:del"))
    (inserted,) = link.findall(qn("w:ins"))
    assert [t.text for t in deleted.iter(qn("w:delText"))] == ["client"]
    assert [t.text for t in inserted.iter(qn("w:t"))] == ["owner"]
    assert inserted.find(f".//{qn('w:rStyle')}") is not None  # typed over the link
    assert redline.find(qn("w:del")) is None and redline.find(qn("w:ins")) is None
