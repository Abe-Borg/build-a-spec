"""The primary Word export keeps formatting AND native, reviewable revisions."""
from __future__ import annotations

import io
import zipfile

import pytest
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from lxml import etree

from backend import sessions
from backend.app import create_app
from backend.spec_doc.revisions import accept_all, first_difference, has_revisions, reject_all
from backend.spec_doc.tracked_export import accepted_revision_baseline, enable_track_changes
from tests.test_preserving_export import _import, _master_bytes, _save
from tests.test_redline_original import _body, _members, _pending_master, _reword_first


@pytest.fixture
def client(monkeypatch):
    from backend import settings
    import backend.app as app_module

    monkeypatch.setattr(app_module, "_revision_timestamp", lambda: "2026-10-01T12:00:00Z")
    monkeypatch.setattr(settings, "REDLINE_COMMENTS", False)
    return TestClient(create_app())


def _export(client, **params):
    return client.get("/api/export/docx", params={"mode": "preserved", "track_changes": True, **params})


def _tracking(payload: bytes):
    return etree.fromstring(_members(payload)["word/settings.xml"]).find(qn("w:trackRevisions"))


def test_formatted_export_tracks_existing_edits_and_enables_future_edits(client):
    source = _master_bytes()
    _import(client, source)
    _reword_first(client)
    response = _export(client)
    assert response.status_code == 200, response.text
    assert " - REDLINE.docx" not in response.headers["content-disposition"]
    assert _tracking(response.content).get(qn("w:val")) == "true"
    body = _body(response.content)
    assert has_revisions(body)
    assert list(body.iter(qn("w:ins"))) and list(body.iter(qn("w:del")))
    clean = client.get("/api/export/docx", params={"mode": "preserved"})
    assert clean.status_code == 200
    assert first_difference(accept_all(body), _body(clean.content)) is None
    assert first_difference(reject_all(body), _body(source)) is None
    before, after = _members(source), _members(response.content)
    assert before.keys() == after.keys()
    assert {name for name in before if before[name] != after[name]} == {
        "word/document.xml", "word/settings.xml",
    }
    assert sessions.get_session().source_docx_bytes == source
    assert client.get("/api/import/original").content == source
    # The separate original-format redline keeps its settings and body.
    redline = client.get("/api/export/docx", params={"mode": "preserved", "redline": "master"})
    assert redline.status_code == 200
    assert _members(redline.content)["word/settings.xml"] == before["word/settings.xml"]
    assert _members(redline.content)["word/document.xml"] == after["word/document.xml"]


@pytest.mark.parametrize("params", [
    {},
    {"track_changes": False, "redline": "master"},
    {"track_changes": False, "redline": "master", "mode": "normalized"},
])
@pytest.mark.parametrize("author,initials", [(None, "AB"), ("Renée Garcia", "RG")])
def test_exports_use_the_personal_author_and_keep_other_reviewers(client, monkeypatch, params, author, initials):
    from docx import Document
    from backend import settings
    from tests.test_redline_comments import _install_fix

    monkeypatch.setattr(settings, "REDLINE_COMMENTS", True)
    if author is not None:
        monkeypatch.setattr(settings, "WORD_AUTHOR", author)
    expected_author = author or "Abraham Borg"
    document = Document(io.BytesIO(_master_bytes()))
    document.add_comment(document.paragraphs[4].runs, text="Confirm scope.", author="Reviewer", initials="RV")
    source = _save(document)
    _import(client, source)
    target = sessions.get_session().doc.doc.parts[0].articles[0].paragraphs[0].uid
    _install_fix(client, target, [{
        "action": "replace", "target_id": target,
        "text": "Section includes seismic isolation for mechanical equipment.",
    }])
    result = _export(client, **params)
    assert result.status_code == 200, result.text
    changes = [element for element in _body(result.content).iter()
               if element.tag in {qn("w:ins"), qn("w:del")}]
    assert changes and {element.get(qn("w:author")) for element in changes} == {expected_author}
    if params.get("mode") != "normalized":
        before = etree.fromstring(_members(source)["word/comments.xml"])
        after = etree.fromstring(_members(result.content)["word/comments.xml"])
        original = {comment.get(qn("w:id")): comment for comment in before}
        added = [comment for comment in after if comment.get(qn("w:id")) not in original]
        assert added
        assert {comment.get(qn("w:author")) for comment in added} == {expected_author}
        assert {comment.get(qn("w:initials")) for comment in added} == {initials}
        for comment in after:
            if comment.get(qn("w:id")) in original:
                assert first_difference(original[comment.get(qn("w:id"))], comment) is None
    assert sessions.get_session().source_docx_bytes == source


def test_unedited_export_has_tracking_on_without_inventing_revisions(client):
    source = _master_bytes()
    _import(client, source)
    response = _export(client)
    assert response.status_code == 200, response.text
    assert not has_revisions(_body(response.content))
    assert _tracking(response.content).get(qn("w:val")) == "true"
    assert first_difference(_body(response.content), _body(source)) is None


def test_saved_baspec_opens_through_the_app_and_exports_tracked_changes(client):
    source = _master_bytes()
    _import(client, source)
    _reword_first(client)
    saved = client.get("/api/project/save")
    assert saved.status_code == 200
    sessions.reset_session()
    loaded = client.post("/api/project/load-file", files={
        "file": ("section.baspec", saved.content, "application/octet-stream"),
    })
    assert loaded.status_code == 200, loaded.text
    assert loaded.json()["preserved_redline_available"] is True
    result = _export(client)
    assert result.status_code == 200, result.text
    assert has_revisions(_body(result.content))
    assert first_difference(reject_all(_body(result.content)), _body(source)) is None
    assert _tracking(result.content).get(qn("w:val")) == "true"


def test_revision_bearing_baspec_tracks_app_edits_against_the_accepted_import_view(client):
    source = _pending_master()
    _import(client, source)
    _reword_first(client)
    saved = client.get("/api/project/save").content
    sessions.reset_session()
    loaded = client.post("/api/project/load-file", files={
        "file": ("revision-bearing.baspec", saved, "application/octet-stream"),
    })
    assert loaded.status_code == 200
    assert loaded.json()["tracked_export_available"] is True
    assert loaded.json()["tracked_export_reason"] is None
    assert loaded.json()["preserved_redline_available"] is False
    session = sessions.get_session()
    accepted, _ = accepted_revision_baseline(source, session.source_format_map)
    result = _export(client)
    assert result.status_code == 200, result.text
    assert has_revisions(_body(result.content))
    assert first_difference(reject_all(_body(result.content)), _body(accepted)) is None
    assert first_difference(accept_all(_body(source)), _body(accepted)) is None
    assert _tracking(result.content).get(qn("w:val")) == "true"
    assert session.source_docx_bytes == source
    # The stricter, separate redline still makes its original promise.
    original_redline = client.get("/api/export/docx", params={"mode": "preserved", "redline": "master"})
    assert original_redline.status_code == 409
    assert original_redline.json()["code"] == "pending_revisions"


def test_pending_paragraph_mark_remaps_origins_without_losing_the_original(client):
    from docx import Document
    from backend.spec_doc.source_render import render_preserving_docx

    document = Document(io.BytesIO(_master_bytes()))
    paragraph = document.paragraphs[4]._p
    properties = paragraph.get_or_add_pPr()
    mark = etree.SubElement(properties, qn("w:rPr"))
    deletion = etree.SubElement(mark, qn("w:del"))
    deletion.set(qn("w:id"), "800")
    deletion.set(qn("w:author"), "Previous reviewer")
    source = _save(document)
    _import(client, source)
    _reword_first(client)
    session = sessions.get_session()
    accepted, mapped = accepted_revision_baseline(source, session.source_format_map)
    assert any(anchor.origin_index < 0 for anchor in mapped.anchors)
    result = _export(client)
    assert result.status_code == 200, result.text
    body = _body(result.content)
    assert first_difference(reject_all(body), _body(accepted)) is None
    clean = render_preserving_docx(source_bytes=accepted, format_map=mapped, current=session.doc.doc)
    assert first_difference(accept_all(body), _body(clean)) is None
    assert session.source_docx_bytes == source
    assert b"basOriginIndex" not in _members(result.content)["word/document.xml"]


def test_prior_header_and_footer_revisions_keep_their_accepted_appearance(client):
    from docx import Document

    document = Document(io.BytesIO(_master_bytes()))
    for paragraph in (document.sections[0].header.paragraphs[0], document.sections[0].footer.paragraphs[0]):
        inserted = etree.SubElement(paragraph._p, qn("w:ins"))
        inserted.set(qn("w:id"), "801")
        inserted.set(qn("w:author"), "Previous reviewer")
        run = etree.SubElement(inserted, qn("w:r"))
        etree.SubElement(run, qn("w:t")).text = " REVISED"
    source = _save(document)
    _import(client, source)
    _reword_first(client)
    result = _export(client)
    assert result.status_code == 200, result.text
    before, after = _members(source), _members(result.content)
    for name in ("word/header1.xml", "word/footer1.xml"):
        expected = accept_all(etree.fromstring(before[name]))
        actual = etree.fromstring(after[name])
        assert first_difference(expected, actual) is None
        assert not has_revisions(actual)
    assert sessions.get_session().source_docx_bytes == source


@pytest.mark.parametrize("header_part", [
    "word/review-header.dat",
    "client-data/review-header.xml",
    "client-data/review-header.dat",
])
def test_relocated_revision_part_exports_with_tracking_and_preserves_formatting(client, header_part):
    from tests.test_source_global_blockers import _with_relocated_revision_header

    source = _with_relocated_revision_header(_master_bytes(), new_part=header_part)
    imported = _import(client, source)
    assert imported["tracked_export_available"] is True
    _reword_first(client)
    session = sessions.get_session()
    original_map = session.source_format_map
    result = _export(client)
    assert result.status_code == 200, result.text
    assert _tracking(result.content).get(qn("w:val")) == "true"
    before, after = _members(source), _members(result.content)
    expected_header = accept_all(etree.fromstring(before[header_part]))
    exported_header = etree.fromstring(after[header_part])
    assert first_difference(expected_header, exported_header) is None
    assert not has_revisions(exported_header)
    assert has_revisions(_body(result.content))
    assert first_difference(reject_all(_body(result.content)), _body(source)) is None
    clean = client.get("/api/export/docx", params={"mode": "preserved"})
    assert clean.status_code == 200
    assert first_difference(accept_all(_body(result.content)), _body(clean.content)) is None
    assert before.keys() == after.keys()
    assert {name for name in before if before[name] != after[name]} == {
        header_part, "word/document.xml", "word/settings.xml",
    }
    assert session.source_docx_bytes == source
    assert session.source_format_map is original_map
    assert client.get("/api/import/original").content == source


@pytest.mark.parametrize("params", [
    {"mode": "normalized"}, {"mode": "source"}, {"mode": None},
    {"redline": "version", "base": 0}, {"redline": "unknown"},
])
def test_tracking_cannot_silently_use_another_export_mode(client, params):
    _import(client, _master_bytes())
    assert _export(client, **params).status_code == 400


def test_unavailable_original_refuses_instead_of_delivering_a_clean_file(client):
    _import(client, _master_bytes())
    sessions.get_session().source_format_map = None
    response = _export(client)
    assert response.status_code == 409
    assert response.json()["code"] == "no_original"


def _rewrite(payload, replacements=None, removed=()):
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(payload)) as source, zipfile.ZipFile(result, "w") as target:
        for entry in source.infolist():
            if entry.filename not in removed:
                target.writestr(entry, (replacements or {}).get(entry.filename, source.read(entry.filename)))
    return result.getvalue()


@pytest.mark.parametrize("value", [None, "false", "0", "true"])
def test_tracking_switch_keeps_all_other_settings_and_uses_schema_order(value):
    source = _master_bytes()
    settings = etree.fromstring(_members(source)["word/settings.xml"])
    view = etree.Element(qn("w:revisionView"))
    position = settings.index(settings.find(qn("w:defaultTabStop")))
    settings.insert(position, view)
    if value is not None:
        tracking = etree.Element(qn("w:trackRevisions"))
        tracking.set(qn("w:val"), value)
        settings.insert(position + 1, tracking)
    payload = _rewrite(source, {"word/settings.xml": etree.tostring(settings)})
    result = enable_track_changes(payload)
    after = _members(result)
    original = _members(payload)
    assert all(after[name] == data for name, data in original.items() if name != "word/settings.xml")
    updated = etree.fromstring(after["word/settings.xml"])
    assert [child.tag for child in updated][position:position + 3] == [
        qn("w:revisionView"), qn("w:trackRevisions"), qn("w:defaultTabStop"),
    ]
    switches = updated.findall(qn("w:trackRevisions"))
    assert len(switches) == 1 and switches[0].get(qn("w:val")) == "true"
    updated.remove(switches[0])
    previous = settings.find(qn("w:trackRevisions"))
    if previous is not None:
        settings.remove(previous)
    assert etree.tostring(updated) == etree.tostring(settings)


def test_document_without_settings_gets_only_the_tracking_switch():
    source = _master_bytes()
    before = _members(source)
    rels = etree.fromstring(before["word/_rels/document.xml.rels"])
    for rel in list(rels):
        if rel.get("Type", "").endswith("/settings"):
            rels.remove(rel)
    types = etree.fromstring(before["[Content_Types].xml"])
    for child in list(types):
        if child.get("PartName") == "/word/settings.xml":
            types.remove(child)
    payload = _rewrite(source, {
        "word/_rels/document.xml.rels": etree.tostring(rels),
        "[Content_Types].xml": etree.tostring(types),
    }, removed=("word/settings.xml",))
    result = enable_track_changes(payload)
    after = _members(result)
    settings = etree.fromstring(after["word/settings.xml"])
    assert [child.tag for child in settings] == [qn("w:trackRevisions")]
    assert all(after[name] == data for name, data in _members(payload).items()
               if name not in ("[Content_Types].xml", "word/_rels/document.xml.rels"))
    from docx import Document
    assert Document(io.BytesIO(result)).settings.element.find(qn("w:trackRevisions")) is not None


def test_nonstandard_settings_part_is_resolved_through_its_relationship():
    source = _master_bytes()
    before = _members(source)
    rels = etree.fromstring(before["word/_rels/document.xml.rels"])
    for rel in rels:
        if rel.get("Type", "").endswith("/settings"):
            rel.set("Target", "custom-settings.xml")
    result = io.BytesIO()
    with zipfile.ZipFile(result, "w") as target:
        for name, data in before.items():
            if name == "word/_rels/document.xml.rels":
                data = etree.tostring(rels)
            target.writestr("word/custom-settings.xml" if name == "word/settings.xml" else name, data)
    after = _members(enable_track_changes(result.getvalue()))
    assert "word/settings.xml" not in after
    assert etree.fromstring(after["word/custom-settings.xml"]).find(qn("w:trackRevisions")).get(qn("w:val")) == "true"
