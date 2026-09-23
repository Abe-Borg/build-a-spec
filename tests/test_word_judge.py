"""The real-Word judge, proved without Word.

``tests/word_judge.py`` compares what real Word saved after resolving a
redline with what the same Word saved for the reference. Everything but the
Word call is pure, so it is proved here on any platform: each tolerance for
what a Word save writes on its own removes exactly that and nothing more,
the app's own self-check applies none of them, the targeted cases really
cover the markup shapes they claim, and the whole pipeline — driven by fake
Words — passes a faithful Word (noisy or not) and fails an unfaithful one,
saying where.
"""
from __future__ import annotations

import json
import random
from pathlib import Path

import pytest
from docx.oxml.ns import qn
from lxml import etree

from backend.spec_doc.revisions import accept_all, first_difference, reject_all
from tests import word_judge as judge
from tools.render_docx_word import ResolvedDocx, WordResolution

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _body(inner: str):
    return etree.fromstring(f'<w:body xmlns:w="{W}">{inner}</w:body>')


def _run(text: str, *, bold: bool = False, attrs: str = "") -> str:
    properties = "<w:rPr><w:b/></w:rPr>" if bold else ""
    return f'<w:r{attrs}>{properties}<w:t xml:space="preserve">{text}</w:t></w:r>'


def _paragraph(*runs: str, attrs: str = "") -> str:
    return f"<w:p{attrs}>{''.join(runs)}</w:p>"


def _difference(left, right):
    """The judge's own comparison — never a copy of it, so each tolerance
    test proves the verdicts the judge actually reaches."""
    return judge.body_difference(left, right)


# ---------------------------------------------------------------------------
# Each tolerance removes what a Word save writes, and nothing more
# ---------------------------------------------------------------------------


def test_the_tolerances_are_exactly_the_four_named():
    """A new tolerance is a decision: it needs evidence from a real run, a
    reason in the table, and a test here."""
    assert list(judge.WORD_SAVE_TOLERANCES) == [
        "rsid attributes",
        "w:proofErr",
        "w:lastRenderedPageBreak",
        "the _GoBack bookmark",
    ]
    assert all(len(reason) > 80 for reason in judge.WORD_SAVE_TOLERANCES.values())


def test_rsid_attributes_are_a_save_stamp_not_content():
    stamped = _body(
        _paragraph(
            _run("Provide isolators.", attrs=' w:rsidR="00A1B2C3" w:rsidRPr="00D4E5F6"'),
            attrs=' w:rsidR="00112233" w:rsidRDefault="00445566" w:rsidP="00778899"',
        )
        + '<w:sectPr w:rsidR="00ABCDEF" w:rsidSect="00FEDCBA"/>'
    )
    plain = _body(_paragraph(_run("Provide isolators.")) + "<w:sectPr/>")
    assert _difference(stamped, plain) is None
    assert _difference(plain, stamped) is None  # both sides
    # Only the rsids go: any other attribute, on the same elements, stays.
    styled = _body(
        _paragraph(_run("Provide isolators."), attrs=' w:rsidR="00112233" w14:paraId="1" xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"')
        .replace("<w:r>", '<w:pPr><w:jc w:val="center"/></w:pPr><w:r>', 1)
        + "<w:sectPr/>"
    )
    assert _difference(styled, plain) is not None


def test_a_proofing_mark_is_not_content_but_its_words_are():
    """Word splits a run around a flagged word; stripped of the marks, the
    split runs merge back — and a real word change beside them still
    shows."""
    flagged = _body(
        _paragraph(
            _run("Provide "),
            '<w:proofErr w:type="spellStart"/>',
            _run("isolaters"),
            '<w:proofErr w:type="spellEnd"/>',
            _run("."),
        )
    )
    plain = _body(_paragraph(_run("Provide isolaters.")))
    assert _difference(flagged, plain) is None
    changed = _body(_paragraph(_run("Provide isolators.")))
    assert _difference(flagged, changed) is not None


def test_a_rendered_page_break_is_a_layout_cache_but_a_real_break_is_not():
    cached = _body(
        _paragraph('<w:r><w:lastRenderedPageBreak/><w:t>Provide isolators.</w:t></w:r>')
    )
    plain = _body(_paragraph(_run("Provide isolators.")))
    assert _difference(cached, plain) is None
    real = _body(
        _paragraph('<w:r><w:br w:type="page"/><w:t>Provide isolators.</w:t></w:r>')
    )
    assert _difference(real, plain) is not None


def test_the_go_back_bookmark_is_navigation_state_but_every_other_is_compared():
    def marked(name: str, where: int) -> object:
        runs = [_run("Provide "), _run("isolators.")]
        marker = (
            f'<w:bookmarkStart w:id="9" w:name="{name}"/><w:bookmarkEnd w:id="9"/>'
        )
        runs.insert(where, marker)
        return _body(_paragraph(*runs))

    plain = _body(_paragraph(_run("Provide isolators.")))
    assert _difference(marked("_GoBack", 0), plain) is None
    assert _difference(marked("_GoBack", 1), marked("_GoBack", 2)) is None
    for name in ("_Hlk154", "_Toc12", "_Ref77", "Anchor"):
        assert _difference(marked(name, 1), plain) is not None, name


def test_the_apps_own_self_check_applies_none_of_the_tolerances():
    """The tolerances exist only because both files the judge compares were
    written by Word. The self-check compares the app's XML with the app's
    XML, and must stay as strict as it is."""
    plain = _body(_paragraph(_run("Provide isolators.")))
    word_saved = (
        _body(_paragraph(_run("Provide isolators."), attrs=' w:rsidR="00112233"')),
        _body(
            _paragraph(
                '<w:proofErr w:type="spellStart"/>',
                _run("Provide isolators."),
                '<w:proofErr w:type="spellEnd"/>',
            )
        ),
        _body(
            _paragraph('<w:r><w:lastRenderedPageBreak/><w:t>Provide isolators.</w:t></w:r>')
        ),
        _body(
            _paragraph(
                '<w:bookmarkStart w:id="0" w:name="_GoBack"/><w:bookmarkEnd w:id="0"/>',
                _run("Provide isolators."),
            )
        ),
    )
    for body in word_saved:
        assert first_difference(body, plain) is not None
        assert _difference(body, plain) is None


def test_a_difference_is_reported_with_both_sides_rendered(tmp_path):
    left = judge.with_body(_package(), _body(_paragraph(_run("Provide isolators."))))
    right = judge.with_body(_package(), _body(_paragraph(_run("Provide restraints."))))
    difference = judge.word_difference(left, right)
    assert difference is not None and difference.index == 0
    described = judge.describe_difference(left, right, difference)
    assert "'Provide isolators.'" in described["resolved_by_word"]
    assert "'Provide restraints.'" in described["reference_resaved_by_word"]
    assert described["path"].startswith("p/r")


def _package() -> bytes:
    from tests.test_preserving_export import _master_bytes

    return _master_bytes()


# ---------------------------------------------------------------------------
# The cases
# ---------------------------------------------------------------------------


def test_the_judge_is_off_unless_asked():
    assert not judge.judge_is_configured({})
    assert not judge.judge_is_configured({judge.JUDGE_ENV: "0"})
    assert not judge.judge_is_configured({judge.JUDGE_ENV: "  "})
    assert judge.judge_is_configured({judge.JUDGE_ENV: "1"})
    assert judge.judge_is_configured({judge.JUDGE_ENV: "Yes"})


def test_the_judge_replays_every_corpus_master_under_the_sweeps_own_edits():
    from tests.docx_corpus import corpus_cases
    from tests.test_redline_original import CORPUS_SWEEP_SEEDS

    groups = judge.judge_groups()
    ids = [group.group_id for group in groups]
    assert len(ids) == len(set(ids))
    assert len({group.slug for group in groups}) == len(groups)
    corpus = [group for group in groups if group.group_id.startswith("corpus/")]
    assert [group.group_id for group in corpus] == [
        f"corpus/{case.case_id}" for case in corpus_cases()
    ]
    for group in corpus:
        assert [name for name, _edit in group.edits] == [
            f"seed-{seed}" for seed in CORPUS_SWEEP_SEEDS
        ]


@pytest.fixture(scope="module")
def targeted(tmp_path_factory):
    workspace = tmp_path_factory.mktemp("judge-targeted")
    return [
        judge.build_judge_cases(group, workspace / group.slug)
        for group in judge.judge_groups()
        if group.group_id.startswith("targeted/")
    ]


def test_every_targeted_case_is_a_real_redline(targeted):
    """None is refused, and every one carries tracked changes — the app's own
    self-check already passed for each, inside the render."""
    for batch in targeted:
        for case in batch.cases:
            assert isinstance(case, judge.JudgeCase), (batch.group.group_id, case)
            assert case.stats["redline"]["revisions"] > 0, (
                batch.group.group_id,
                case.name,
            )


def test_the_targeted_cases_cover_every_shape_the_writer_emits(targeted):
    """What the judge exists to put in front of Word: each markup shape the
    redline writer can emit appears in at least one targeted case."""
    totals: dict[str, int] = {}
    markup: set[str] = set()
    moved: set[str] = set()
    for batch in targeted:
        for case in batch.rendered:
            for key, value in case.stats["redline"].items():
                totals[key] = totals.get(key, 0) + value
            body = judge.word_body(case.redline)
            moved |= case.moved_bookmarks
            for element in body.iter():
                if not isinstance(element.tag, str):
                    continue
                name = etree.QName(element).localname
                parent = element.getparent()
                if name in ("ins", "del") and parent is not None:
                    where = etree.QName(parent).localname
                    markup.add(f"{name}@{where}")
                elif name in ("pPrChange", "rPrChange", "delText", "delInstrText"):
                    markup.add(name)
                elif name == "drawing" and any(
                    a.tag == qn("w:del") for a in element.iterancestors()
                ):
                    markup.add("deleted-drawing")
    for kind in (
        "spliced",
        "fallback",
        "inserted",
        "deleted",
        "dropped",
        "moved",
        "moves_added",
        "leftovers",
        "last_mark_untracked",
    ):
        assert totals.get(kind, 0) > 0, kind
    assert {
        "ins@p",  # inserted runs
        "del@p",  # deleted runs
        "ins@rPr",  # an inserted paragraph mark
        "del@rPr",  # a deleted paragraph mark
        "ins@trPr",  # an inserted table row (a moved table)
        "del@trPr",  # a deleted table row
        "ins@hyperlink",  # changed words inside a link
        "del@hyperlink",
        "pPrChange",  # a numbering cancel, a neutralized last paragraph
        "rPrChange",  # a neutralized last paragraph mark
        "delText",
        "delInstrText",  # a deleted field instruction
        "deleted-drawing",
    } <= markup, markup
    assert moved == {"_Ref77"}


# ---------------------------------------------------------------------------
# The pipeline, driven by fake Words
# ---------------------------------------------------------------------------

_COUNTED = {qn(f"w:{name}") for name in ("ins", "del", "moveFrom", "moveTo", "pPrChange", "rPrChange")}


def _revision_count(body) -> int:
    return sum(1 for element in body.iter() if element.tag in _COUNTED)


def _rewrite_like_word(body):
    """What a Word save does to EVERY file the same way. Not a tolerance —
    this is content the comparison keeps — so only comparing Word's save
    with Word's save cancels it: every hyperlink gains a history flag and
    every run states its language."""
    for link in body.iter(qn("w:hyperlink")):
        link.set(qn("w:history"), "1")
    for run in body.iter(qn("w:r")):
        properties = run.find(qn("w:rPr"))
        if properties is None:
            properties = etree.Element(qn("w:rPr"))
            run.insert(0, properties)
        if properties.find(qn("w:lang")) is None:
            etree.SubElement(properties, qn("w:lang")).set(qn("w:val"), "en-US")
    return body


def _add_word_noise(body, seed: str):
    """What a Word save writes: the same rewrite in every file
    (:func:`_rewrite_like_word`), plus what it writes on its own — rsids
    everywhere, a flagged word splitting a run, a rendered page break, and a
    _GoBack bookmark — placed differently in every file."""
    _rewrite_like_word(body)
    rng = random.Random(seed)
    for element in body.iter(qn("w:p"), qn("w:r"), qn("w:tr"), qn("w:sectPr")):
        element.set(qn("w:rsidR"), f"{rng.getrandbits(32):08X}")
    # Word splits a run around a flagged word into runs of the same
    # formatting, each holding its share of the text. Only a run holding
    # nothing but one w:t is split here, so the copy duplicates nothing.
    runs = [
        run
        for run in body.iter(qn("w:r"))
        if [child.tag for child in run if child.tag != qn("w:rPr")] == [qn("w:t")]
        and len(run.find(qn("w:t")).text or "") > 3
    ]
    if runs:
        run = rng.choice(runs)
        text = run.find(qn("w:t"))
        cut = rng.randint(1, len(text.text) - 1)
        tail = etree.fromstring(etree.tostring(run))
        tail.find(qn("w:t")).text = text.text[cut:]
        tail.find(qn("w:t")).set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        text.text = text.text[:cut]
        text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        parent = run.getparent()
        index = parent.index(run)
        start = etree.Element(qn("w:proofErr"))
        start.set(qn("w:type"), "spellStart")
        end = etree.Element(qn("w:proofErr"))
        end.set(qn("w:type"), "spellEnd")
        parent.insert(index + 1, start)
        parent.insert(index + 2, tail)
        parent.insert(index + 3, end)
        first = rng.choice(runs)
        properties = first.find(qn("w:rPr"))
        first.insert(0 if properties is None else 1, etree.Element(qn("w:lastRenderedPageBreak")))
    paragraphs = list(body.iter(qn("w:p")))
    if paragraphs:
        paragraph = rng.choice(paragraphs)
        properties = paragraph.find(qn("w:pPr"))
        position = 0 if properties is None else 1
        # Word numbers its _GoBack bookmark like any other: an id no other
        # bookmark in the document holds.
        taken = [
            int(marker.get(qn("w:id")))
            for marker in body.iter(qn("w:bookmarkStart"), qn("w:bookmarkEnd"))
            if (marker.get(qn("w:id")) or "").isdigit()
        ]
        go_back_id = str(max(taken, default=-1) + 1)
        start = etree.Element(qn("w:bookmarkStart"))
        start.set(qn("w:id"), go_back_id)
        start.set(qn("w:name"), "_GoBack")
        end = etree.Element(qn("w:bookmarkEnd"))
        end.set(qn("w:id"), go_back_id)
        paragraph.insert(position, start)
        paragraph.insert(position + 1, end)
    return body


def _fake_word(
    *,
    accept=accept_all,
    reject=reject_all,
    noise: bool = False,
    authors=(judge.JUDGE_AUTHOR,),
    left_behind: int = 0,
    fail: str = "",
):
    """A Word that resolves with the given functions (the app's own oracle
    by default) and saves what it resolved."""

    def resolve(jobs, **_kwargs):
        documents = []
        for job in jobs:
            payload = job.input_path.read_bytes()
            body = judge.word_body(payload)
            count = _revision_count(body)
            if fail and job.input_path.name.endswith(fail):
                documents.append(ResolvedDocx(job, False, "Word could not read it.", -1, -1, ()))
                continue
            if job.action == "accept":
                body = accept(body)
            elif job.action == "reject":
                body = reject(body)
            if noise:
                body = _add_word_noise(body, job.output_path.name)
            job.output_path.write_bytes(judge.with_body(payload, body))
            documents.append(
                ResolvedDocx(
                    job,
                    True,
                    "",
                    count,
                    left_behind if job.action != "resave" else 0,
                    tuple(authors) if count else (),
                )
            )
        return WordResolution("16.0", "16.0.fake", tuple(documents))

    return resolve


def _group(group_id: str) -> judge.JudgeGroup:
    return next(group for group in judge.judge_groups() if group.group_id == group_id)


@pytest.fixture(scope="module")
def batches(tmp_path_factory):
    workspace = tmp_path_factory.mktemp("judge-batches")
    return {
        group_id: judge.build_judge_cases(_group(group_id), workspace / group_id.replace("/", "-"))
        for group_id in (
            "targeted/nested",
            "targeted/bookmarked",
            "targeted/last-numbered",
            "targeted/break-holder",
        )
    }


@pytest.mark.parametrize("noise", [False, True], ids=["plain", "with-word-save-noise"])
def test_a_faithful_word_passes_every_targeted_case(targeted, tmp_path, noise):
    """Every targeted shape, resolved by a Word that resolves the way the
    oracle does — and, noisy, saves the way Word does. The noise lands
    somewhere different in every file, so this is also what shows the
    tolerances swallow Word's save markup without swallowing content."""
    for batch in targeted:
        verdict = judge.judge_batch(
            batch, tmp_path / f"{noise}" / batch.group.slug, resolve=_fake_word(noise=noise)
        )
        assert verdict.failures == [], verdict.failures
        assert {case["status"] for case in verdict.cases} == {"pass"}, batch.group.group_id
        assert verdict.word_build == "16.0.fake"


def test_a_faithful_noisy_word_passes_every_corpus_mix(tmp_path):
    """The exact batches the Windows run hands real Word — every corpus
    master under the sweep's edit mixes — pass a Word that resolves as the
    oracle does and saves the way Word does. So a failure on Windows is real
    Word disagreeing, never an artifact of the harness."""
    for group in judge.judge_groups():
        if not group.group_id.startswith("corpus/"):
            continue
        batch = judge.build_judge_cases(group, tmp_path / "cases" / group.slug)
        verdict = judge.judge_batch(
            batch, tmp_path / "word" / group.slug, resolve=_fake_word(noise=True)
        )
        assert verdict.failures == [], verdict.failures
        assert {case["status"] for case in verdict.cases} <= {"pass", "refused"}


def test_only_a_symmetric_comparison_survives_what_word_rewrites(batches, tmp_path):
    """A Word save rewrites content in every file the same way. Compared with
    Word's own save of the reference, that cancels; compared with the app's
    files as written, every case would read as a difference — which is why
    the judge has the same Word re-save the formatted export and the upload
    instead of tolerating what a save rewrites."""
    batch = batches["targeted/nested"]
    verdict = judge.judge_batch(batch, tmp_path / "judged", resolve=_fake_word(noise=True))
    assert verdict.failures == []
    for case in batch.rendered:
        accepted = (tmp_path / "judged" / f"{case.name}.accepted.word.docx").read_bytes()
        rejected = (tmp_path / "judged" / f"{case.name}.rejected.word.docx").read_bytes()
        assert judge.word_difference(accepted, case.clean) is not None, case.name
        assert judge.word_difference(rejected, batch.upload) is not None, case.name


def test_the_moved_bookmark_is_the_one_loss_reject_may_have(batches, tmp_path):
    verdict = judge.judge_batch(
        batches["targeted/bookmarked"], tmp_path / "bookmarked", resolve=_fake_word()
    )
    (entry,) = verdict.cases
    assert entry["status"] == "pass"
    assert entry["lost_bookmarks"] == ["_Ref77"] == entry["moved_bookmarks"]


def test_an_unfaithful_word_fails_and_says_where(batches, tmp_path):
    """A Word whose Accept All leaves a stray formatted paragraph at the top
    of the body: the judge names the case, the resolution and the body
    child, and renders what Word saved there."""

    def leaves_a_stray_paragraph(body):
        accepted = accept_all(body)
        properties = etree.Element(qn("w:pPr"))
        etree.SubElement(properties, qn("w:jc")).set(qn("w:val"), "center")
        stray = etree.Element(qn("w:p"))
        stray.append(properties)
        accepted.insert(0, stray)
        return accepted

    verdict = judge.judge_batch(
        batches["targeted/nested"],
        tmp_path / "unfaithful",
        resolve=_fake_word(accept=leaves_a_stray_paragraph),
    )
    assert verdict.failures
    assert all(
        failure.startswith("targeted/nested ") and "Accept All in Word differs" in failure
        for failure in verdict.failures
    ), verdict.failures
    failed = [case for case in verdict.cases if case["status"] == "fail"]
    assert failed and all("accept" in case and "reject" not in case for case in failed)
    assert failed[0]["accept"]["index"] == 0
    assert "<jc" in failed[0]["accept"]["resolved_by_word"]


def test_a_swapped_word_fails_both_ways(batches, tmp_path):
    verdict = judge.judge_batch(
        batches["targeted/break-holder"],
        tmp_path / "swapped",
        resolve=_fake_word(accept=reject_all, reject=accept_all),
    )
    for case in verdict.cases:
        assert case["status"] == "fail"
        assert "accept" in case and "reject" in case


@pytest.mark.parametrize(
    ("fake", "message"),
    [
        (_fake_word(authors=("Someone Else",)), "authored by ['Someone Else']"),
        (_fake_word(left_behind=2), "2 tracked change(s) were still there"),
        (
            _fake_word(accept=lambda body: body),
            "still carries revision markup",
        ),
    ],
    ids=["another-author", "changes-left", "markup-left"],
)
def test_what_word_read_and_left_is_judged_too(batches, tmp_path, fake, message):
    verdict = judge.judge_batch(
        batches["targeted/last-numbered"], tmp_path / "read", resolve=fake
    )
    assert verdict.failures
    assert any(message in failure for failure in verdict.failures), verdict.failures


def test_a_bookmark_lost_by_no_move_is_a_failure(batches, tmp_path):
    def reject_losing_bookmarks(body):
        rejected = reject_all(body)
        for marker in list(rejected.iter(qn("w:bookmarkStart"), qn("w:bookmarkEnd"))):
            marker.getparent().remove(marker)
        return rejected

    batch = batches["targeted/bookmarked"]
    unmoved = judge.JudgeBatch(
        batch.group,
        batch.upload,
        tuple(
            judge.JudgeCase(c.name, c.redline, c.clean, frozenset(), c.stats)
            for c in batch.rendered
        ),
    )
    verdict = judge.judge_batch(
        unmoved, tmp_path / "lost", resolve=_fake_word(reject=reject_losing_bookmarks)
    )
    assert any("lost bookmark(s) ['_Ref77']" in f for f in verdict.failures), verdict.failures


def test_a_file_word_cannot_open_is_an_error_not_a_pass(batches, tmp_path):
    verdict = judge.judge_batch(
        batches["targeted/nested"],
        tmp_path / "unreadable",
        resolve=_fake_word(fail="reorder-article.redline.docx"),
    )
    by_case = {case["case"]: case for case in verdict.cases}
    assert by_case["reorder-article"]["status"] == "error"
    assert any("Word could not accept" in f for f in verdict.failures)
    assert by_case["delete-article"]["status"] == "pass"


def test_refused_mixes_are_recorded_not_sent_to_word(tmp_path):
    group = judge.JudgeGroup("targeted/refused", lambda _w: b"", ())
    batch = judge.JudgeBatch(group, b"", (judge.RefusedCase("seed-0", "moved_annotation"),))
    called = []
    verdict = judge.judge_batch(batch, tmp_path, resolve=lambda jobs: called.append(jobs))
    assert called == []
    assert verdict.cases == [
        {"case": "seed-0", "status": "refused", "refusal": "moved_annotation"}
    ]
    assert verdict.failures == []


def test_a_refusal_that_is_not_structural_is_a_failure(tmp_path):
    """The judge and the corpus sweep agree on what a legitimate refusal is:
    a failed self-check is the app breaking its promise before Word sees
    the file, never something to skip quietly."""
    from backend.spec_doc.source_render import REDLINE_ACCEPT_CHECK

    assert REDLINE_ACCEPT_CHECK not in judge.structural_refusals()
    assert "moved_annotation" in judge.structural_refusals()
    group = judge.JudgeGroup("corpus/refused", lambda _w: b"", ())
    batch = judge.JudgeBatch(group, b"", (judge.RefusedCase("seed-1", REDLINE_ACCEPT_CHECK),))
    verdict = judge.judge_batch(batch, tmp_path, resolve=lambda jobs: None)
    (entry,) = verdict.cases
    assert entry["status"] == "error"
    assert verdict.failures == [
        f"corpus/refused seed-1: the app refused it for '{REDLINE_ACCEPT_CHECK}', not "
        "a structural reason: its own self-check failed before Word saw the file"
    ]


def test_the_report_names_files_never_paths(batches, tmp_path):
    report = judge.JudgeReport(tmp_path / "judge")
    verdict = judge.judge_batch(
        batches["targeted/bookmarked"], report.run_dir / "bookmarked", resolve=_fake_word()
    )
    report.record_group(verdict)
    text = (tmp_path / "judge" / "report.json").read_text(encoding="utf-8")
    assert str(tmp_path) not in text
    data = json.loads(text)
    assert data["word"] == {"version": "16.0", "build": "16.0.fake"}
    assert data["summary"]["pass"] == 1
    assert data["tolerances"] == dict(judge.WORD_SAVE_TOLERANCES)
    assert (report.run_dir / "report.json").read_text(encoding="utf-8") == text
    second = judge.JudgeReport(tmp_path / "judge")
    assert second.run_dir != report.run_dir  # a rerun never reuses a folder


def test_the_report_folder_defaults_under_gitignored_artifacts(monkeypatch):
    monkeypatch.delenv(judge.JUDGE_DIR_ENV, raising=False)
    repo = Path(__file__).resolve().parent.parent
    assert judge.DEFAULT_JUDGE_DIR == repo / "artifacts" / "word-judge"
    ignored = (repo / ".gitignore").read_text(encoding="utf-8").splitlines()
    assert "/artifacts/" in ignored


# ---------------------------------------------------------------------------
# Word's own tracked moves (the Word-saved sample)
# ---------------------------------------------------------------------------


def _word_shaped_move_sample() -> bytes:
    """A hand-built stand-in for the Word-saved sample: one paragraph moved
    below another, its bookmark on the moved-to copy."""
    from tests.test_preserving_export import _save

    from docx import Document

    document = Document()
    document.add_paragraph("A. Placeholder provision alpha.")
    document.add_paragraph("B. Placeholder provision bravo.")
    payload = _save(document)
    author = f'w:author="{judge.TRACKED_MOVE_AUTHOR}" w:date="2026-09-23T10:00:00Z"'
    moved_from = (
        f'<w:p><w:pPr><w:rPr><w:moveFrom w:id="1" {author}/></w:rPr></w:pPr>'
        f'<w:moveFromRangeStart w:id="2" {author} w:name="move1"/>'
        f'<w:moveFrom w:id="3" {author}><w:r><w:t>A. Placeholder provision alpha.</w:t></w:r></w:moveFrom>'
        '<w:moveFromRangeEnd w:id="2"/></w:p>'
    )
    moved_to = (
        f'<w:p><w:pPr><w:rPr><w:moveTo w:id="4" {author}/></w:rPr></w:pPr>'
        f'<w:moveToRangeStart w:id="5" {author} w:name="move1"/>'
        f'<w:moveTo w:id="6" {author}><w:bookmarkStart w:id="7" w:name="{judge.TRACKED_MOVE_BOOKMARK}"/>'
        '<w:r><w:t>A. Placeholder provision alpha.</w:t></w:r><w:bookmarkEnd w:id="7"/></w:moveTo>'
        '<w:moveToRangeEnd w:id="5"/></w:p>'
    )
    stay = "<w:p><w:r><w:t>B. Placeholder provision bravo.</w:t></w:r></w:p>"
    body = _body(moved_from + stay + moved_to + "<w:sectPr/>")
    return judge.with_body(payload, body)


def test_word_and_the_oracle_resolve_a_move_sample_alike(tmp_path):
    report = judge.judge_tracked_move_sample(
        _word_shaped_move_sample(),
        tmp_path,
        resolve=_fake_word(authors=(judge.TRACKED_MOVE_AUTHOR,), noise=True),
    )
    assert report["problems"] == []
    places = report["bookmark"]
    assert places["in_sample"] == [
        {
            "paragraph_text": "A. Placeholder provision alpha.",
            "inside": ["moveTo"],
            "paragraph_mark": "moveTo",
        }
    ]
    assert [p["inside"] for p in places["after_word_accept"]] == [[]]
    assert places["after_word_reject"] == []


@pytest.mark.parametrize(
    ("fake", "message"),
    [
        (_fake_word(authors=("Someone Else",)), "Word read the sample's authors as"),
        (
            _fake_word(authors=(judge.TRACKED_MOVE_AUTHOR,), left_behind=1),
            "1 tracked change(s) left after 'accept'",
        ),
        (
            _fake_word(authors=(judge.TRACKED_MOVE_AUTHOR,), fail="sample.docx"),
            "Word could not handle accept",
        ),
    ],
    ids=["another-author", "changes-left", "unreadable"],
)
def test_the_sample_is_judged_for_what_word_read_and_left(tmp_path, fake, message):
    report = judge.judge_tracked_move_sample(
        _word_shaped_move_sample(), tmp_path, resolve=fake
    )
    assert any(message in problem for problem in report["problems"]), report["problems"]


def test_a_word_that_resolves_moves_its_own_way_is_caught(tmp_path):
    report = judge.judge_tracked_move_sample(
        _word_shaped_move_sample(),
        tmp_path,
        resolve=_fake_word(accept=_unwrap_moves, authors=(judge.TRACKED_MOVE_AUTHOR,)),
    )
    assert any("Word's accept of its own moves differs" in p for p in report["problems"])


def _unwrap_moves(body):
    """Keep BOTH copies of every move — a resolution the oracle would not
    give."""
    copy = etree.fromstring(etree.tostring(body))
    for tag in ("w:moveFrom", "w:moveTo"):
        for wrapper in list(copy.iter(qn(tag))):
            parent = wrapper.getparent()
            if parent.tag == qn("w:rPr"):
                parent.remove(wrapper)
                continue
            index = parent.index(wrapper)
            for offset, child in enumerate(list(wrapper)):
                parent.insert(index + offset, child)
            parent.remove(wrapper)
    for tag in ("moveFromRangeStart", "moveFromRangeEnd", "moveToRangeStart", "moveToRangeEnd"):
        for marker in list(copy.iter(qn(f"w:{tag}"))):
            marker.getparent().remove(marker)
    return copy
